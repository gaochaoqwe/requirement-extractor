"""
FPGA静态分析增强服务 - 处理确定/不确定规则分流和AI审查
"""
import uuid
import asyncio
from typing import Dict, Any, List, Optional, Tuple
from pathlib import Path
from datetime import datetime
import logging

from app.staticAnalysis.filterFpga.services.fpga_excel_processor import FpgaExcelProcessor
from app.staticAnalysis.filterFpga.database.postgresql_database import get_db_session
from app.staticAnalysis.filterFpga.database.postgresql_models import (
    FPGARule, StaticAnalysisSession, StaticAnalysisIssue
)
from app.staticAnalysis.filterFpga.utils.logger import setup_module_logger
from app.staticAnalysis.filterFpga.services.fpga_ai_agent import FpgaAIAgent
from app.staticAnalysis.filter.services.ai_client import StaticAnalysisAIClient
from app.staticAnalysis.filter.services.code_analysis_agent_v2 import CodeAnalysisAgentV2
from app.staticAnalysis.filterFpga.services.fpga_output_generator import FpgaOutputGenerator
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

logger = setup_module_logger(__name__, "fpga_enhanced_service.log")


class FpgaEnhancedService:
    """FPGA静态分析增强服务"""
    
    def __init__(self, session_id: str = None, task_id: int = None):
        self.session_id = session_id or str(uuid.uuid4())
        self.task_id = task_id
        
        # 日志器
        self.log_dir = None
        self.logger = None
        if task_id:
            self._init_logger_with_task_id(task_id)
        else:
            self._init_temp_logger()
        
        # Excel处理器
        self.excel_processor = FpgaExcelProcessor(self.session_id)
        
        # 进度跟踪
        self.total_issues = 0
        self.accurate_issues = 0
        self.uncertain_issues = 0
        self.processed_issues = 0
        self.filtered_issues = 0
        self.current_status = "pending"
        self.current_analyzing_issue = ""
        self.tool_calls_count = 0
        self.accurate_ready = False  # 确定规则是否已就绪
        
        # 文件信息
        self.original_filename = ""
        self.output_filename = ""
        self.codebase_path = None
        self.selected_rules = []
        # 操作人员姓名（由API注入）
        self.operator_name: Optional[str] = None
        
        # ===== 预生成文件缓存 =====
        self.cached_filename = None  # FPGA预生成的四分分类文件名
    
    def _init_logger_with_task_id(self, task_id: int):
        """使用task_id初始化日志器"""
        log_folder = str(task_id)
        self.log_dir = Path(__file__).parent.parent / "logs" / log_folder
        self.log_dir.mkdir(parents=True, exist_ok=True)
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        log_file = self.log_dir / f"fpga_analysis_{self.session_id}_{timestamp}.log"
        
        self.logger = logging.getLogger(f"FPGA_{self.session_id}")
        self.logger.setLevel(logging.DEBUG)
        self.logger.handlers.clear()
        self.logger.propagate = False
        
        file_handler = logging.FileHandler(str(log_file), encoding='utf-8')
        file_handler.setLevel(logging.DEBUG)
        formatter = logging.Formatter('%(asctime)s - %(name)s - %(levelname)s - %(message)s')
        file_handler.setFormatter(formatter)
        self.logger.addHandler(file_handler)
        
        self.logger.info(f"=" * 100)
        self.logger.info(f"FPGA规则静态分析日志 - 会话ID: {self.session_id}")
        self.logger.info(f"=" * 100)
        # 切换到task_id后记录操作人员
        try:
            if getattr(self, 'operator_name', None):
                self.logger.info(f"操作人员: {self.operator_name}")
        except Exception:
            pass
    
    def _init_temp_logger(self):
        """创建临时日志器"""
        # 使用模块级日志器，写入系统日志文件，避免丢日志
        from app.staticAnalysis.filterFpga.utils.logger import setup_module_logger
        self.logger = setup_module_logger(__name__, "fpga_enhanced_service.log")
    
    async def process_files(self, excel_path: str, codebase_zip_path: str, 
                           selected_rules: List[str], skip_ai: bool = False) -> Dict[str, Any]:
        """
        处理上传的文件
        
        Args:
            excel_path: Excel文件路径
            codebase_zip_path: 代码库zip文件路径
            selected_rules: 选中的规则列表 (如 ["A-1-5", "R-2-1"])
        
        Returns:
            处理结果
        """
        try:
            self.selected_rules = selected_rules
            # 保存路径信息，便于后续导出阶段使用
            self.codebase_path = codebase_zip_path
            self.excel_path = excel_path
            try:
                self.original_filename = Path(excel_path).name
            except Exception:
                self.original_filename = ""
            self.current_status = "processing"
            
            self.logger.info(f"开始处理FPGA静态分析 - 选中规则数: {len(selected_rules)}")
            
            # 1. 解析Excel并提取代码
            parse_result = self.excel_processor.process_excel_file(excel_path, codebase_zip_path)
            if not parse_result["success"]:
                return parse_result
            
            all_issues = parse_result["issues"]
            self.logger.info(f"解析得到 {len(all_issues)} 条问题记录")
            
            # 2. 全量入库，不在此处按选中规则过滤（与C版一致，以支持未勾选Sheet导出）
            self.total_issues = len(all_issues)
            if self.total_issues == 0:
                # 无匹配项时，结束流程，避免前端一直等待
                self.current_status = "completed"
                self._update_session_status("completed")
                self.logger.info("所选规则未匹配到任何问题，流程结束")
                return {
                    "success": True,
                    "session_id": self.session_id,
                    "total_issues": 0,
                    "accurate_issues": 0,
                    "uncertain_issues": 0,
                    "filtered_issues": 0,
                    "message": "所选规则未匹配到任何问题"
                }
            
            # 3. 保存到数据库（全量）
            await self._save_issues_to_db(all_issues)
            
            # 4. 分类：确定规则 vs 不确定规则（对全量issues分类）
            accurate_issues, uncertain_issues = self._classify_issues(all_issues)
            
            self.accurate_issues = len(accurate_issues)
            self.uncertain_issues = len(uncertain_issues)
            
            self.logger.info(f"分类完成: 确定规则 {self.accurate_issues} 条, 不确定规则 {self.uncertain_issues} 条")
            
            # 5. 处理确定规则（直接生成问题单）
            if accurate_issues:
                self.logger.info("开始处理确定规则...")
                await self._process_accurate_issues(accurate_issues)
                # [FIX] 不在这里设置 accurate_ready，等预生成文件完成后再设置
                self.logger.info(f"确定规则处理完成，共 {len(accurate_issues)} 条")
            
            # [FIX] 数据入库和确定规则处理完成后，预生成文件缓存
            self._pre_generate_cache_file()
            
            # [FIX] 预生成完成后，设置 accurate_ready = True
            if self.total_issues > 0:
                self.accurate_ready = True
                self.logger.info(f"文件缓存已预生成，设置 accurate_ready = True")
            
            # [FIX] 提前计算"勾选规则的不确定问题数"，用于AI审查进度
            selected_set = set(self.selected_rules or [])
            uncertain_selected = [u for u in uncertain_issues if u.get("rule_code") in selected_set]
            # uncertain_issues 保持为所有不确定问题数（勾选+未勾选），用于前端显示
            # uncertain_selected 是勾选的不确定问题，用于AI审查
            self.logger.info(f"[初始化] 不确定问题总数: {len(uncertain_issues)} 条, 其中勾选: {len(uncertain_selected)} 条")
            
            # 6. 可选：跳过不确定规则AI审查（并发受限时）
            if skip_ai:
                self.logger.warning("并发限制触发：跳过不确定规则AI审查，仅提供未经AI过滤的问题单")
                # 将状态置为完成，前端可直接下载未经AI过滤
                self.current_status = "completed"
                self._update_session_status("completed")
                return {
                    "success": True,
                    "session_id": self.session_id,
                    "total_issues": self.total_issues,
                    "accurate_issues": self.accurate_issues,
                    "uncertain_issues": self.uncertain_issues,
                    "filtered_issues": self.filtered_issues,
                    "ai_disabled": True
                }

            # 7. 处理不确定规则（AI审查，仅对勾选规则）
            if uncertain_selected:
                self.logger.info("开始AI审查不确定规则...")
                await self._process_uncertain_issues(uncertain_selected, codebase_zip_path)
                self.logger.info(f"不确定规则审查完成，共 {len(uncertain_selected)} 条")
            
            self.current_status = "completed"
            self._update_session_status("completed")
            
            return {
                "success": True,
                "session_id": self.session_id,
                "total_issues": self.total_issues,
                "accurate_issues": self.accurate_issues,
                "uncertain_issues": self.uncertain_issues,
                "filtered_issues": self.filtered_issues
            }
            
        except Exception as e:
            self.logger.error(f"处理文件失败: {str(e)}", exc_info=True)
            self.current_status = "failed"
            self._update_session_status("error")
            return {"success": False, "error": str(e)}
    
    def _classify_issues(self, issues: List[Dict[str, Any]]) -> Tuple[List[Dict], List[Dict]]:
        """
        分类问题：确定规则 vs 不确定规则
        
        Returns:
            (确定规则问题列表, 不确定规则问题列表)
        """
        session = get_db_session()
        try:
            # 获取所有规则的准确性类型
            rule_codes = list(set(issue["rule_code"] for issue in issues))
            rules = session.query(FPGARule).filter(
                FPGARule.rule_code.in_(rule_codes)
            ).all()
            
            # 创建规则编号到准确性类型的映射
            rule_accuracy_map = {rule.rule_code: rule.accuracy_type for rule in rules}
            
            accurate = []
            uncertain = []
            
            for issue in issues:
                rule_code = issue["rule_code"]
                accuracy = rule_accuracy_map.get(rule_code, "uncertain")
                bucket = "accurate" if accuracy == "accurate" else "uncertain"
                # 记录分类日志
                self.logger.info(
                    f"[CLASSIFY] 规则={rule_code}, 文件={issue.get('file_name','')}, 行={issue.get('line_number','')}, accuracy={accuracy} -> {bucket}"
                )
                if accuracy == "accurate":
                    accurate.append(issue)
                else:
                    uncertain.append(issue)
            
            return accurate, uncertain
            
        finally:
            session.close()
    
    async def _save_issues_to_db(self, issues: List[Dict[str, Any]]):
        """保存问题到数据库"""
        session = get_db_session()
        try:
            # 创建会话记录
            db_session = StaticAnalysisSession(
                session_id=self.session_id,
                original_filename=self.original_filename,
                file_path=getattr(self, 'excel_path', ''),  # Excel文件路径
                codebase_path=getattr(self, 'codebase_path', None),  # 代码库路径
                total_issues=len(issues),
                status="processing"
            )
            session.add(db_session)

            # 在 add 之后、commit 之前设置操作人员姓名，便于SQLAlchemy跟踪并落库
            if getattr(self, 'operator_name', None):
                db_session.operator_name = self.operator_name
                self.logger.info(f"设置操作人员: {self.operator_name}")
            session.flush()
            
            self.task_id = db_session.id
            # 拿到 task_id 后，切换日志到该任务专属目录 filterFpga/logs/<task_id>
            try:
                self._init_logger_with_task_id(self.task_id)
                self.logger.info(f"已创建任务日志目录并切换日志输出至该目录: task_id={self.task_id}")
            except Exception as _log_e:
                # 即使日志切换失败，也不影响主流程
                logger = setup_module_logger(__name__, "fpga_enhanced_service.log")
                logger.warning(f"切换任务日志目录失败: {str(_log_e)}")
            
            # 批量插入问题
            db_issues = []
            # 预先根据规则表获取准确性映射，确保确定/不确定与规则库一致
            rule_codes = list(set(issue["rule_code"] for issue in issues))
            rules = session.query(FPGARule).filter(FPGARule.rule_code.in_(rule_codes)).all()
            rule_accuracy_map = {r.rule_code: r.accuracy_type for r in rules}
            for idx, issue in enumerate(issues, start=1):
                rule_code = issue["rule_code"]
                # 优先使用解析阶段传入的accuracy_type，否则回退到规则库中的定义，最后默认uncertain
                accuracy_type = issue.get("accuracy_type") or rule_accuracy_map.get(rule_code, "uncertain")
                # 记录入库前的识别信息（含提取的单行代码预览）
                preview = (issue.get("source_code") or "")[:120]
                self.logger.info(
                    f"[SAVE] 序号={idx}, 规则={rule_code}, 文件={issue.get('file_name','')}, 行={issue.get('line_number','')}, accuracy={accuracy_type}, 代码='{preview}'"
                )
                db_issue = StaticAnalysisIssue(
                    task_id=self.task_id,  # 使用task_id而不session_id
                    row_index=idx,  # 行索引
                    rule_code=rule_code,
                    rule_type="fpga",
                    rule_accuracy_type=accuracy_type,  # 与规则库保持一致的准确性
                    issue_name=f"{issue['rule_code']} - {issue.get('rule_description', '')[:50]}",
                    issue_file=issue["file_name"],
                    issue_line=str(issue["line_number"]),  # 转为字符串
                    issue_description=issue.get("rule_description", ""),
                    ai_analyzed=False,  # 字段名修正
                    is_filtered=False
                )
                db_issues.append(db_issue)
            
            session.bulk_save_objects(db_issues)
            session.commit()
            
            self.logger.info(f"保存了 {len(db_issues)} 条问题到数据库")
            
        except Exception as e:
            session.rollback()
            self.logger.error(f"保存问题到数据库失败: {str(e)}", exc_info=True)
            raise
        finally:
            session.close()
    
    def _pre_generate_cache_file(self):
        """
        预生成文件缓存，避免用户点击下载时文件还未准备好
        在数据入库完成后立即调用，生成FPGA的四分分类下载文件
        """
        try:
            self.logger.info("[预生成] 开始预生成FPGA四分分类文件缓存")
            
            # 预生成FPGA四分分类文件（如果有问题）
            if self.total_issues > 0:
                try:
                    self.logger.info(f"[预生成] 开始生成FPGA缓存文件，问题数: {self.total_issues}")
                    result = self.generate_split_output_file(highlight_ai=False)
                    if result.get("success"):
                        self.cached_filename = result.get("filename")
                        self.logger.info(f"[预生成] FPGA缓存文件生成成功: {self.cached_filename}")
                    else:
                        self.logger.error(f"[预生成] FPGA缓存文件生成失败: {result.get('error')}")
                except Exception as e:
                    self.logger.error(f"[预生成] 生成FPGA缓存文件异常: {str(e)}")
                    import traceback
                    self.logger.error(traceback.format_exc())
            else:
                self.logger.info("[预生成] FPGA问题为空，跳过文件生成")
            
            self.logger.info("[预生成] 文件缓存预生成完成")
            
        except Exception as e:
            self.logger.error(f"[预生成] 预生成文件缓存时出错: {str(e)}")
            import traceback
            self.logger.error(traceback.format_exc())
            # 不抛出异常，允许流程继续
    
    async def _process_accurate_issues(self, issues: List[Dict[str, Any]]):
        """
        处理确定规则（直接标记为保留，不需要AI审查）
        """
        session = get_db_session()
        try:
            for issue in issues:
                # 更新数据库：标记为已审查且保留
                session.query(StaticAnalysisIssue).filter(
                    StaticAnalysisIssue.task_id == self.task_id,
                    StaticAnalysisIssue.rule_code == issue["rule_code"],
                    StaticAnalysisIssue.issue_file == issue["file_name"],
                    StaticAnalysisIssue.issue_line == str(issue["line_number"])
                ).update({
                    "ai_analyzed": True,
                    "is_filtered": False,  # 保留
                    "ai_analysis_result": "确定规则-自动保留"
                })
                
                self.processed_issues += 1
            
            session.commit()
            
        except Exception as e:
            session.rollback()
            self.logger.error(f"处理确定规则失败: {str(e)}", exc_info=True)
            raise
        finally:
            session.close()
    
    async def _process_uncertain_issues(self, issues: List[Dict[str, Any]], codebase_path: str):
        """
        处理不确定规则（需要AI审查） - 使用Agent工具调用机制
        """
        # 创建Agent（使用CodeAnalysisAgentV2支持工具调用），传入任务专属的文件logger
        agent = CodeAnalysisAgentV2(self.session_id, codebase_path, logger=self.logger)
        
        session = get_db_session()
        try:
            for idx, issue in enumerate(issues, 1):
                self.current_analyzing_issue = f"{issue['rule_code']} - {issue['file_name']}:{issue['line_number']}"
                self.logger.info(f"[{idx}/{len(issues)}] AI审查(Agent模式): {self.current_analyzing_issue}")
                
                # 获取规则详情
                rule = session.query(FPGARule).filter(
                    FPGARule.rule_code == issue["rule_code"]
                ).first()
                
                if not rule:
                    self.logger.warning(f"未找到规则: {issue['rule_code']}")
                    continue
                
                # 准备AI审查上下文
                context = {
                    "rule_code": issue["rule_code"],
                    "rule_content": rule.rule_content,
                    "file_name": issue["file_name"],
                    "line_number": issue["line_number"],
                    "source_code": issue.get("source_code", ""),
                    "code_context": issue.get("code_context", ""),
                    "description": issue.get("rule_description", "")
                }
                
                # 调用AI审查 - 使用Agent工具调用机制
                try:
                    ai_result = await FpgaAIAgent.analyze_fpga_issue_with_agent_v2(
                        agent, context, self.logger
                    )
                    
                    # 更新工具调用统计
                    self.tool_calls_count = agent.tool_call_count
                    
                    # 更新数据库
                    is_filtered = not ai_result.get("should_retain", True)
                    
                    session.query(StaticAnalysisIssue).filter(
                        StaticAnalysisIssue.task_id == self.task_id,
                        StaticAnalysisIssue.rule_code == issue["rule_code"],
                        StaticAnalysisIssue.issue_file == issue["file_name"],
                        StaticAnalysisIssue.issue_line == str(issue["line_number"])
                    ).update({
                        "ai_analyzed": True,
                        "is_filtered": is_filtered,
                        "ai_is_real_issue": ai_result.get("should_retain", True),
                        "ai_analysis_result": ai_result.get("reason", "")
                    })
                    
                    if is_filtered:
                        self.filtered_issues += 1
                    
                except Exception as e:
                    self.logger.error(f"AI审查失败: {str(e)}")
                    # 失败时默认保留
                    session.query(StaticAnalysisIssue).filter(
                        StaticAnalysisIssue.task_id == self.task_id,
                        StaticAnalysisIssue.rule_code == issue["rule_code"],
                        StaticAnalysisIssue.issue_file == issue["file_name"],
                        StaticAnalysisIssue.issue_line == str(issue["line_number"])
                    ).update({
                        "ai_analyzed": True,
                        "is_filtered": False,
                        "ai_is_real_issue": True,
                        "ai_analysis_result": f"AI审查失败，默认保留: {str(e)}"
                    })
                
                self.processed_issues += 1
                session.commit()
            
        except Exception as e:
            session.rollback()
            self.logger.error(f"处理不确定规则失败: {str(e)}", exc_info=True)
            raise
        finally:
            session.close()
    
    def get_progress(self) -> Dict[str, Any]:
        """获取当前进度"""
        progress_pct = 0
        if self.total_issues > 0:
            progress_pct = int((self.processed_issues / self.total_issues) * 100)
        
        # [FIX] 从数据库或内存获取统计数据
        total_issues_db = self.total_issues
        accurate_issues_db = self.accurate_issues
        uncertain_issues_db = self.uncertain_issues
        ai_total = 0  # 仅勾选的不确定问题数，用于AI审查进度
        
        if self.task_id and (total_issues_db == 0 or uncertain_issues_db == 0):
            # 当内存值为0时，尝试从数据库查询（可能是刚开始处理，数据还在加载中）
            try:
                session = get_db_session()
                try:
                    # 查询总问题数
                    if total_issues_db == 0:
                        total_issues_db = session.query(StaticAnalysisIssue).filter(
                            StaticAnalysisIssue.task_id == self.task_id
                        ).count()
                    
                    # 查询确定规则数
                    if accurate_issues_db == 0:
                        accurate_issues_db = session.query(StaticAnalysisIssue).filter(
                            StaticAnalysisIssue.task_id == self.task_id,
                            StaticAnalysisIssue.rule_accuracy_type == 'accurate'
                        ).count()
                    
                    # 查询所有不确定规则数
                    if uncertain_issues_db == 0:
                        uncertain_issues_db = session.query(StaticAnalysisIssue).filter(
                            StaticAnalysisIssue.task_id == self.task_id,
                            StaticAnalysisIssue.rule_accuracy_type == 'uncertain'
                        ).count()
                finally:
                    session.close()
            except Exception:
                # 查询失败时保持为0
                pass
        
        # 查询仅勾选的不确定问题数（用于AI审查进度）
        if self.task_id and self.selected_rules:
            try:
                session = get_db_session()
                try:
                    ai_total = session.query(StaticAnalysisIssue).filter(
                        StaticAnalysisIssue.task_id == self.task_id,
                        StaticAnalysisIssue.rule_accuracy_type == 'uncertain',
                        StaticAnalysisIssue.rule_code.in_(list(set(self.selected_rules)))
                    ).count()
                finally:
                    session.close()
            except Exception:
                # 查询失败时保持为0
                pass

        return {
            "status": self.current_status,
            "progress": progress_pct,
            "total_issues": total_issues_db,
            "accurate_issues": accurate_issues_db,
            "uncertain_issues": uncertain_issues_db,  # 所有不确定问题数（勾选+未勾选）
            "processed_issues": self.processed_issues,
            "filtered_issues": self.filtered_issues,
            "current_analyzing_issue": self.current_analyzing_issue,
            "tool_calls_count": self.tool_calls_count,
            "accurate_ready": self.accurate_ready,
            # 新增：仅不确定且勾选的总数（用于AI审查进度）
            "uncertain_selected_issues": ai_total,
            # 兼容前端字段：AI处理总量
            "ai_processing_total": ai_total
        }

    def generate_split_output_file(self, highlight_ai: bool = False) -> Dict[str, Any]:
        """生成四分分类（确定/不确定 × 勾选/未勾选）Excel，支持可选AI高亮。
        仅当 highlight_ai=True 时，在“不确定规则（勾选）”sheet 中对 ai_is_real_issue=True 的行使用浅绿色高亮。
        """
        try:
            session = get_db_session()
            try:
                # 通过会话ID获取当前任务ID，确保限定范围
                db_sess = session.query(StaticAnalysisSession).filter(
                    StaticAnalysisSession.session_id == self.session_id
                ).first()
                if not db_sess:
                    return {"success": False, "error": "会话不存在或已过期"}

                task_id = db_sess.id

                # 取出当前会话的全部问题（不过滤 is_filtered）
                db_issues = session.query(StaticAnalysisIssue).filter(
                    StaticAnalysisIssue.task_id == task_id
                ).all()

                if not db_issues:
                    return {"success": False, "error": "当前会话下没有问题数据"}

                # 预取规则信息映射
                rule_codes = list({i.rule_code for i in db_issues})
                rules = session.query(FPGARule).filter(FPGARule.rule_code.in_(rule_codes)).all()
                rule_map = {r.rule_code: r for r in rules}

                # 代码库路径用于缺失源码时的回填
                codebase_path = db_sess.codebase_path

                # [FIX] 勾选集合：优先使用内存中的 selected_rules，否则从数据库读取
                if self.selected_rules:
                    # 使用内存中的勾选状态（处理过程中）
                    selected_set = set(self.selected_rules)
                    self.logger.info(f"[SPLIT-EXPORT] 使用内存勾选状态: {len(selected_set)} 条规则")
                else:
                    # 服务重新创建时，从数据库读取勾选状态
                    db_selected_rules = session.query(FPGARule.rule_code).filter(FPGARule.is_selected == True).all()
                    selected_set = set([r[0] for r in db_selected_rules])
                    self.logger.info(f"[SPLIT-EXPORT] 从数据库读取勾选状态: {len(selected_set)} 条规则")

                # 四个桶
                buckets: Dict[str, List[StaticAnalysisIssue]] = {
                    "确定规则（勾选)": [],
                    "确定规则（未勾选)": [],
                    "不确定规则（勾选)": [],
                    "不确定规则（未勾选)": [],
                }

                for issue in db_issues:
                    acc = issue.rule_accuracy_type or "uncertain"
                    # 优先使用规则的 is_selected 字段，否则回退到 selected_set
                    rule = rule_map.get(issue.rule_code)
                    if rule and rule.is_selected is not None:
                        selected = rule.is_selected
                    else:
                        selected = issue.rule_code in selected_set
                    if acc == "accurate":
                        key = "确定规则（勾选)" if selected else "确定规则（未勾选)"
                    else:
                        key = "不确定规则（勾选)" if selected else "不确定规则（未勾选)"
                    buckets[key].append(issue)

                # 构建生成器，复用列头/格式与八列内容生成逻辑
                generator = FpgaOutputGenerator(self.session_id)

                # 创建工作簿
                wb = openpyxl.Workbook()
                # 先清空默认sheet
                default_ws = wb.active
                wb.remove(default_ws)

                # 统一样式
                header_font = Font(bold=True, size=11)
                header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
                ai_confirm_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                normal_align = Alignment(vertical='center', wrap_text=True)

                def write_sheet(title: str, issues_list: List[StaticAnalysisIssue], with_ownership: bool = False, ownership_value: str = ""):
                    """写sheet，支持可选的"问题归属"列"""
                    ws = wb.create_sheet(title=title)
                    # 表头
                    start_col = 1
                    if with_ownership:
                        # 最前面增加"问题归属"列
                        cell = ws.cell(row=1, column=1, value="问题归属")
                        cell.font = header_font
                        cell.alignment = header_align
                        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                        start_col = 2
                    
                    for col_idx, col_name in enumerate(generator.OUTPUT_COLUMNS, start=start_col):
                        cell = ws.cell(row=1, column=col_idx, value=col_name)
                        cell.font = header_font
                        cell.alignment = header_align
                        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

                    # 行数据
                    row_num = 2
                    for issue in issues_list:
                        rule = rule_map.get(issue.rule_code)
                        issue_dict = {
                            "rule_code": issue.rule_code,
                            "rule_content": rule.rule_content if rule else "",
                            "language": (rule.language if rule else "") or "",
                            "issue_title": rule.issue_title if rule else "",
                            "impact_analysis": rule.impact_analysis if rule else "",
                            "fix_suggestion": rule.fix_suggestion if rule else "",
                            "issue_file": issue.issue_file,
                            "issue_line": issue.issue_line,
                            "issue_description": issue.issue_description or "",
                            # 模型无source_code字段，保持与输出生成器一致
                            "source_code": getattr(issue, 'source_code', '') or "",
                            "ai_result": getattr(issue, 'ai_analysis_result', '') or "",
                            "codebase_path": codebase_path,
                        }

                        # 生成8列值
                        row_values = generator._analyze_issue_for_output(issue_dict)
                        
                        col_offset = 0
                        if with_ownership:
                            # 写入"问题归属"列
                            ws.cell(row=row_num, column=1, value=ownership_value)
                            col_offset = 1
                        
                        for col_idx, value in enumerate(row_values, start=1+col_offset):
                            cell = ws.cell(row=row_num, column=col_idx, value=value)
                            cell.alignment = normal_align

                        # 仅在不确定(勾选)且需要高亮时对 AI 确认的问题高亮绿色
                        if highlight_ai and title == "不确定规则（勾选)" and getattr(issue, 'ai_is_real_issue', None) is True:
                            for cell in ws[row_num]:
                                cell.fill = ai_confirm_fill

                        row_num += 1

                    # 列宽、行高统一设置
                    from openpyxl.utils import get_column_letter
                    for row in range(1, ws.max_row + 1):
                        ws.row_dimensions[row].height = 40
                    
                    total_cols = 9 if with_ownership else 8
                    for col in range(1, total_cols + 1):
                        col_letter = get_column_letter(col)
                        if col == 1 and with_ownership:
                            ws.column_dimensions[col_letter].width = 25  # "问题归属"列宽度
                        elif col <= 6:
                            ws.column_dimensions[col_letter].width = 20
                        else:
                            ws.column_dimensions[col_letter].width = 60

                # [NEW] 先生成"过滤后问题"sheet（放在第一个位置）
                # 包含：确定规则（勾选）+ 不确定规则（勾选）中AI确认的问题
                filtered_issues = []
                # 1. 确定规则（勾选）全部加入
                for issue in buckets.get("确定规则（勾选)", []):
                    filtered_issues.append((issue, "确定勾选"))
                # 2. 不确定规则（勾选）中AI确认的加入
                for issue in buckets.get("不确定规则（勾选)", []):
                    if getattr(issue, 'ai_is_real_issue', None) is True:
                        filtered_issues.append((issue, "不确定勾选（AI已确认）"))
                
                # 写"过滤后问题"sheet（带问题归属列）
                if filtered_issues:
                    ws_filtered = wb.create_sheet(title="过滤后问题", index=0)  # index=0确保是第一个
                    # 表头
                    cell = ws_filtered.cell(row=1, column=1, value="问题归属")
                    cell.font = header_font
                    cell.alignment = header_align
                    cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                    for col_idx, col_name in enumerate(generator.OUTPUT_COLUMNS, start=2):
                        cell = ws_filtered.cell(row=1, column=col_idx, value=col_name)
                        cell.font = header_font
                        cell.alignment = header_align
                        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                    
                    # 行数据
                    row_num = 2
                    for issue, ownership in filtered_issues:
                        rule = rule_map.get(issue.rule_code)
                        issue_dict = {
                            "rule_code": issue.rule_code,
                            "rule_content": rule.rule_content if rule else "",
                            "language": (rule.language if rule else "") or "",
                            "issue_title": rule.issue_title if rule else "",
                            "impact_analysis": rule.impact_analysis if rule else "",
                            "fix_suggestion": rule.fix_suggestion if rule else "",
                            "issue_file": issue.issue_file,
                            "issue_line": issue.issue_line,
                            "issue_description": issue.issue_description or "",
                            "source_code": getattr(issue, 'source_code', '') or "",
                            "ai_result": getattr(issue, 'ai_analysis_result', '') or "",
                            "codebase_path": codebase_path,
                        }
                        row_values = generator._analyze_issue_for_output(issue_dict)
                        
                        # 写入"问题归属"列
                        ws_filtered.cell(row=row_num, column=1, value=ownership)
                        for col_idx, value in enumerate(row_values, start=2):
                            cell = ws_filtered.cell(row=row_num, column=col_idx, value=value)
                            cell.alignment = normal_align
                        
                        row_num += 1
                    
                    # 设置列宽和行高
                    from openpyxl.utils import get_column_letter
                    for row in range(1, ws_filtered.max_row + 1):
                        ws_filtered.row_dimensions[row].height = 40
                    ws_filtered.column_dimensions['A'].width = 25  # 问题归属列
                    for col in range(2, 10):
                        col_letter = get_column_letter(col)
                        if col <= 7:
                            ws_filtered.column_dimensions[col_letter].width = 20
                        else:
                            ws_filtered.column_dimensions[col_letter].width = 60
                    
                    self.logger.info(f"[SPLIT-EXPORT] 生成'过滤后问题'sheet，共 {len(filtered_issues)} 条问题")

                # 按固定顺序写四个Sheet
                sheet_order = [
                    "确定规则（勾选)",
                    "确定规则（未勾选)",
                    "不确定规则（勾选)",
                    "不确定规则（未勾选)",
                ]
                for k in sheet_order:
                    write_sheet(k, buckets.get(k, []))

                # 输出目录与文件名（简化版）
                from app.staticAnalysis.filter.config.config import UPLOAD_DIR
                output_dir = Path(UPLOAD_DIR) / "staticAnalysis" / "fpga" / "output"
                output_dir.mkdir(parents=True, exist_ok=True)
                filename = "FPGA_四分类问题单（AI过滤后）.xlsx" if highlight_ai else "FPGA_四分类问题单（未经AI过滤）.xlsx"
                output_path = output_dir / filename

                wb.save(str(output_path))
                wb.close()

                self.logger.info(f"[SPLIT-EXPORT] 生成成功: {output_path}")
                return {"success": True, "output_path": str(output_path), "filename": filename}

            finally:
                session.close()

        except Exception as e:
            self.logger.error(f"生成四分分类问题单失败: {str(e)}", exc_info=True)
            return {"success": False, "error": str(e)}
    
    def _update_session_status(self, status: str):
        """更新会话状态到数据库"""
        db = get_db_session()
        try:
            session = db.query(StaticAnalysisSession).filter(
                StaticAnalysisSession.session_id == self.session_id
            ).first()
            
            if session:
                session.status = status
                session.filtered_issues = self.filtered_issues
                session.ai_analyzed_issues = self.processed_issues
                session.accurate_issues = self.accurate_issues
                session.uncertain_issues = self.uncertain_issues
                db.commit()
                self.logger.info(f"[DB] 更新会话状态: {status}")
        except Exception as e:
            self.logger.error(f"[DB] 更新会话状态失败: {e}")
            db.rollback()
        finally:
            db.close()

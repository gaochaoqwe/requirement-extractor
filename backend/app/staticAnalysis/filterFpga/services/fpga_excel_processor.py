"""
FPGA静态分析Excel处理器
解析FPGA静态分析结果Excel文件（"结果1" sheet）
"""
import openpyxl
import re
import zipfile
from pathlib import Path
from typing import Dict, List, Any, Tuple, Optional

from app.staticAnalysis.filterFpga.utils.logger import setup_module_logger

logger = setup_module_logger(__name__, "fpga_excel_processor.log")


class FpgaExcelProcessor:
    """FPGA Excel解析器"""
    
    def __init__(self, session_id: str):
        self.session_id = session_id
        self.logger = logger
        
    def process_excel_file(self, excel_path: str, codebase_zip_path: Optional[str] = None) -> Dict[str, Any]:
        """
        处理FPGA静态分析Excel文件
        
        Args:
            excel_path: Excel文件路径
            codebase_zip_path: 代码库zip文件路径（可选）
        
        Returns:
            解析结果字典
        """
        try:
            self.logger.info(f"开始处理FPGA Excel文件: {excel_path}")
            
            # 1. 验证Excel文件
            if not self._validate_excel(excel_path):
                return {"success": False, "error": "Excel文件验证失败"}
            
            # 2. 读取Excel数据
            issues = self._read_excel_data(excel_path)
            if not issues:
                return {"success": False, "error": "未读取到有效数据"}
            
            self.logger.info(f"从Excel读取了 {len(issues)} 条问题记录")
            
            # 3. 如果提供了代码库，提取代码内容
            if codebase_zip_path:
                self.logger.info("开始从代码库提取源代码...")
                issues = self._extract_code_from_codebase(issues, codebase_zip_path)
            
            return {
                "success": True,
                "issues": issues,
                "total_count": len(issues)
            }
            
        except Exception as e:
            error_msg = f"处理Excel文件失败: {str(e)}"
            self.logger.error(error_msg, exc_info=True)
            return {"success": False, "error": error_msg}
    
    def _validate_excel(self, excel_path: str) -> bool:
        """验证Excel文件格式"""
        try:
            workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
            
            # 检查是否有"结果1"工作表
            if "结果1" not in workbook.sheetnames:
                self.logger.error(f"Excel文件中未找到'结果1'工作表，现有工作表: {workbook.sheetnames}")
                workbook.close()
                return False
            
            worksheet = workbook["结果1"]
            
            # 检查是否有足够的列（至少需要A-H列）
            if worksheet.max_column < 8:
                self.logger.error(f"'结果1'工作表列数不足，需要至少8列(A-H)，当前: {worksheet.max_column}")
                workbook.close()
                return False
            
            # 检查是否有数据行
            if worksheet.max_row < 2:
                self.logger.error("'结果1'工作表为空，没有数据行")
                workbook.close()
                return False
            
            workbook.close()
            self.logger.info("Excel文件验证通过")
            return True
            
        except Exception as e:
            self.logger.error(f"验证Excel文件时出错: {str(e)}")
            return False
    
    def _read_excel_data(self, excel_path: str) -> List[Dict[str, Any]]:
        """
        读取Excel数据
        
        从"结果1" sheet读取数据：
        A列：序号
        B列：规则编号（格式: "5.1.1.5 建议准则A-1-5"）
        C列：规则描述
        D列：文件（格式: "E:/xx/xx/xx/new/top.v"）
        E列：行
        F列：告警类型
        G列：确认状态
        H列：备注
        """
        issues = []
        
        try:
            workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
            worksheet = workbook["结果1"]
            
            # 逐行读取数据（跳过表头）
            for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    # 检查行是否有足够的数据
                    if not row or len(row) < 5:
                        continue
                    
                    # 提取基本信息
                    sequence = str(row[0]) if row[0] is not None else ""  # A列：序号
                    rule_text = str(row[1]) if row[1] is not None else ""  # B列：规则编号文本
                    rule_description = str(row[2]) if row[2] is not None else ""  # C列：规则描述
                    file_path = str(row[3]) if row[3] is not None else ""  # D列：文件路径
                    line_number = str(row[4]) if row[4] is not None else ""  # E列：行号
                    alert_type = str(row[5]) if row[5] is not None and len(row) > 5 else ""  # F列：告警类型
                    confirm_status = str(row[6]) if row[6] is not None and len(row) > 6 else ""  # G列：确认状态
                    remark = str(row[7]) if row[7] is not None and len(row) > 7 else ""  # H列：备注
                    
                    # 提取规则编号（从B列文本中提取，如 "A-1-5"）
                    rule_code = self._extract_fpga_rule_code(rule_text)
                    
                    # 提取文件名（从完整路径中提取）
                    file_name = self._extract_filename(file_path)
                    
                    # 跳过无效数据
                    if not rule_code or not file_name:
                        self.logger.warning(f"第{row_idx}行数据无效，跳过: rule_code={rule_code}, file_name={file_name}")
                        continue
                    
                    issues.append({
                        "row_index": row_idx,
                        "sequence": sequence,
                        "rule_text": rule_text,
                        "rule_code": rule_code,
                        "rule_description": rule_description,
                        "file_path": file_path,
                        "file_name": file_name,
                        "line_number": line_number,
                        "alert_type": alert_type,
                        "confirm_status": confirm_status,
                        "remark": remark,
                        "source_code": ""  # 稍后从代码库提取
                    })
                    
                except Exception as e:
                    self.logger.warning(f"处理第{row_idx}行数据时出错: {str(e)}")
                    continue
            
            workbook.close()
            self.logger.info(f"成功读取 {len(issues)} 条有效问题记录")
            return issues
            
        except Exception as e:
            self.logger.error(f"读取Excel数据时出错: {str(e)}")
            return []
    
    def _extract_fpga_rule_code(self, rule_text: str) -> str:
        """
        从规则文本中提取规则编号
        
        支持格式示例：
        - "5.1.1.5 建议准则A-1-5"
        - "5.2.5.2 强制准则5.3.1"
        - "5.9.1.a 一般准则"
        
        提取优先级：
        1. 若“建议准则/强制准则”后面明确跟着规则编号，则优先取该编号
        2. 否则返回文本中最后出现的一个受支持规则编号
        """
        if not rule_text:
            return ""
        
        text = str(rule_text).strip()

        # 0) 优先匹配 “建议准则/强制准则” 后跟随的规则编号（取最后一个）
        label_pat = re.compile(
            r'(建议准则|强制准则)\s*([AR]-\d+-\d+|A-5\.\d+\.\d+(?:\.\d+)?|5\.\d+\.\d+\.[A-Za-z]|5\.\d+\.\d+(?:\.\d+)?(?!\.[A-Za-z]))',
            re.IGNORECASE,
        )
        label_matches = list(label_pat.finditer(text))
        if label_matches:
            return label_matches[-1].group(2)

        # 1) 收集所有支持的编号，返回最后出现的那个
        patterns = [
            re.compile(r'\b([AR]-\d+-\d+)\b'),                # Verilog A-#-# / R-#-#
            re.compile(r'\b(A-5\.\d+\.\d+(?:\.\d+)?)\b'), # VHDL建议 A-5.x.x[.x]
            re.compile(r'\b(5\.\d+\.\d+\.[A-Za-z])\b'),    # Verilog新增一般准则 5.9.1.a
            re.compile(r'\b(5\.\d+\.\d+(?:\.\d+)?)(?!\.[A-Za-z])\b'),   # VHDL数字 5.x.x[.x]，避免吞掉 5.9.1.a 这类后缀规则
        ]
        candidates: List[Tuple[int, str]] = []
        for pat in patterns:
            for m in pat.finditer(text):
                candidates.append((m.start(), m.group(1)))
        if candidates:
            candidates.sort(key=lambda x: x[0])
            return candidates[-1][1]

        self.logger.warning(f"无法从规则文本中提取规则编号: {rule_text}")
        return ""
    
    def _extract_filename(self, file_path: str) -> str:
        """
        从文件路径中提取文件名
        
        输入: "E:/xx/xx/xx/new/top.v" 或 "E:\\xx\\xx\\xx\\new\\top.v"
        输出: "top.v"
        """
        if not file_path:
            return ""
        
        # 处理Windows和Unix风格的路径
        file_path = file_path.replace('\\', '/')
        
        # 提取最后一个斜杠后的部分
        parts = file_path.split('/')
        if parts:
            return parts[-1]
        
        return ""
    
    def _extract_code_from_codebase(self, issues: List[Dict[str, Any]], 
                                   codebase_zip_path: str) -> List[Dict[str, Any]]:
        """
        从代码库zip文件中提取源代码
        """
        try:
            # 1) 文件名 -> issues 映射
            file_to_issues: Dict[str, List[Dict[str, Any]]] = {}
            for it in issues:
                file_to_issues.setdefault(it["file_name"], []).append(it)

            self.logger.info(f"需要从代码库提取 {len(file_to_issues)} 个文件的代码")

            # 2) 打开zip并选择最佳条目
            with zipfile.ZipFile(codebase_zip_path, "r") as zf:
                entries = zf.namelist()
                entry_map: Dict[str, Optional[str]] = {}
                for fname, lst in file_to_issues.items():
                    hint = lst[0].get("file_path", "") if lst else ""
                    best = self._find_best_zip_entry(entries, fname, hint)
                    if not best:
                        self.logger.warning(f"未在代码库中找到文件: {fname}")
                    entry_map[fname] = best

                # 3) 提取源码与上下文
                for fname, lst in file_to_issues.items():
                    entry = entry_map.get(fname)
                    if not entry:
                        continue
                    try:
                        content = zf.read(entry).decode("utf-8", errors="ignore")
                        lines = content.split("\n")
                        self.logger.info(f"正在处理文件: {Path(entry).name} (共{len(lines)}行)")

                        for issue in lst:
                            issue["resolved_zip_entry"] = entry
                            # 行号解析
                            try:
                                ln = int(str(issue["line_number"]).strip())
                            except Exception:
                                self.logger.warning(f"无效的行号: {issue['line_number']}")
                                continue

                            if 1 <= ln <= len(lines):
                                extracted = lines[ln - 1].rstrip("\r").strip()
                                if not extracted:
                                    # 邻域回退（±3）
                                    for i in range(max(1, ln - 3), min(len(lines), ln + 3) + 1):
                                        cand = lines[i - 1].rstrip("\r").strip()
                                        if cand:
                                            self.logger.warning(f"[EXTRACT] 目标行为空，使用邻近非空行: 原行={ln} 替换为行={i}")
                                            ln = i
                                            extracted = cand
                                            break
                                issue["source_code"] = extracted
                                # 上下文窗口
                                c0, c1 = max(0, ln - 3), min(len(lines), ln + 2)
                                issue["code_context"] = "\n".join(x.rstrip("\r") for x in lines[c0:c1])
                                self.logger.info(
                                    f"[EXTRACT] 文件={Path(entry).name}, 行={ln}, 规则={issue.get('rule_code','')}, 代码='{issue['source_code'][:120]}'"
                                )
                            else:
                                self.logger.warning(f"行号 {ln} 超出文件 {Path(entry).name} 范围(共{len(lines)}行)")
                    except Exception as er:
                        self.logger.warning(f"读取文件 {entry} 时出错: {str(er)}")

            # 4) 统计
            success_cnt = sum(1 for it in issues if it.get("source_code"))
            self.logger.info(f"成功为 {success_cnt}/{len(issues)} 条问题提取了源代码")
            
            return issues
            
        except Exception as e:
            self.logger.error(f"从代码库提取源代码时出错: {str(e)}", exc_info=True)
            return issues

    def _find_best_zip_entry(self, entries: List[str], filename: str, original_path_hint: str = "") -> Optional[str]:
        """在zip条目中查找与给定文件名最匹配的条目。
        策略：
        - 文件名大小写不敏感相等为候选；
        - 使用原始路径的后缀段匹配进行打分（越多匹配分越高）；
        - 再按路径长度（较短优先），最后按字母序稳定选择。
        """
        if not filename:
            return None
        from pathlib import PurePosixPath
        target_name = filename.lower()
        candidates = [e for e in entries if PurePosixPath(e).name.lower() == Path(filename).name.lower()]
        if not candidates:
            return None
        # 准备原始路径后缀段
        hint = original_path_hint.replace('\\', '/').strip('/')
        hint_parts = [p.lower() for p in hint.split('/') if p]
        def score(entry: str) -> tuple:
            parts = [p.lower() for p in entry.split('/') if p]
            # 后缀从末尾开始对齐计数
            suffix_match = 0
            for a, b in zip(reversed(parts), reversed(hint_parts)):
                if a == b:
                    suffix_match += 1
                else:
                    break
            # 路径越短越好
            path_len = len(parts)
            return (suffix_match, -path_len, entry.lower())
        candidates.sort(key=score, reverse=True)
        return candidates[0]

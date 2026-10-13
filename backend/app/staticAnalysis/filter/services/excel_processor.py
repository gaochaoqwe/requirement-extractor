"""
静态分析问题过滤Excel处理器
"""
import openpyxl
import re
from pathlib import Path
from typing import Dict, List, Any, Tuple, Optional
import uuid
from datetime import datetime

from app.staticAnalysis.filter.config.config import REQUIRED_SHEET_NAME, REQUIRED_COLUMNS, OUTPUT_COLUMNS, FILTER_CONFIG, BATCH_PROCESSING_CONFIG, TARGET_SHEET_PATTERN
from app.staticAnalysis.filter.models.models import StaticAnalysisSession, StaticAnalysisIssue
from app.staticAnalysis.filter.database.database import get_db_session
from app.staticAnalysis.filter.utils.logger import get_session_logger

class ExcelProcessor:
    def __init__(self, session_id: str):
        self.session_id = session_id
        self.logger = get_session_logger(session_id)
        self.required_sheet = REQUIRED_SHEET_NAME
        self.required_columns = REQUIRED_COLUMNS
        
    def process_uploaded_file(self, file_path: str, original_filename: str, filter_keywords: List[str]) -> Tuple[str, Dict[str, Any]]:
        """
        处理上传的静态分析Excel文件
        """
        try:
            self.logger.info(f"开始处理静态分析Excel文件: {original_filename}")
            
            # 验证Excel文件
            validation_result = self._validate_excel_file(file_path)
            if not validation_result["success"]:
                self.logger.error(f"Excel文件验证失败: {validation_result['error']}")
                return self.session_id, validation_result
            
            # 读取Excel数据
            read_result = self._read_static_analysis_data(file_path)
            if not read_result["success"]:
                self.logger.error(f"Excel数据读取失败: {read_result['error']}")
                return self.session_id, read_result
            
            # 保存到数据库
            save_result = self._save_to_database(
                self.session_id, 
                file_path, 
                original_filename, 
                read_result["data"],
                filter_keywords
            )
            
            if save_result["success"]:
                self.logger.info(f"静态分析文件处理完成 - 会话ID: {self.session_id}, 问题数量: {len(read_result['data'])}")
            
            return self.session_id, save_result
            
        except Exception as e:
            error_msg = f"处理静态分析Excel文件时出错: {str(e)}"
            self.logger.error(error_msg)
            return self.session_id, {"success": False, "error": error_msg}
    
    def _validate_excel_file(self, file_path: str) -> Dict[str, Any]:
        """验证Excel文件格式"""
        try:
            self.logger.info("开始验证Excel文件")
            
            # 尝试打开Excel文件
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            
            # 获取所有sheet名称
            all_sheets = workbook.sheetnames
            self.logger.info(f"发现的工作表: {', '.join(all_sheets)}")
            
            # 找到需要处理的sheet（优先处理GJB_8114相关的sheet）
            from ..config.config import SKIP_SHEETS
            
            # 首先查找GJB_8114相关的sheet
            target_sheets = [sheet for sheet in all_sheets if TARGET_SHEET_PATTERN in sheet]
            
            if target_sheets:
                process_sheets = target_sheets
                self.logger.info(f"找到GJB_8114相关工作表: {', '.join(target_sheets)}")
            else:
                # 如果没有找到GJB_8114相关sheet，处理其他sheet（排除跳过的sheet）
                process_sheets = [sheet for sheet in all_sheets if sheet not in SKIP_SHEETS]
                self.logger.info(f"未找到GJB_8114工作表，将处理其他工作表: {', '.join(process_sheets)}")
            
            if not process_sheets:
                return {
                    "success": False, 
                    "error": f"没有找到需要处理的工作表。所有工作表: {', '.join(all_sheets)}"
                }
            
            # 验证每个需要处理的sheet
            for sheet_name in process_sheets:
                worksheet = workbook[sheet_name]
                
                # 检查文件是否为空
                if worksheet.max_row < 2:
                    self.logger.warning(f"工作表 '{sheet_name}' 为空或只有表头，将跳过")
                    continue
                
                # 检查必要的列数
                if worksheet.max_column < 13:  # 至少需要A-M列
                    return {
                        "success": False, 
                        "error": f"工作表 '{sheet_name}' 列数不足，需要至少13列(A-M)"
                    }
            
            workbook.close()
            
            self.logger.info("Excel文件验证成功")
            return {"success": True, "process_sheets": process_sheets}
            
        except Exception as e:
            return {"success": False, "error": f"验证Excel文件时出错: {str(e)}"}
    
    def _read_static_analysis_data(self, file_path: str) -> Dict[str, Any]:
        """读取静态分析数据"""
        try:
            self.logger.info("开始读取静态分析数据")
            
            # 使用openpyxl读取Excel文件
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            
            # 获取需要处理的sheet列表
            from ..config.config import SKIP_SHEETS
            all_sheets = workbook.sheetnames
            
            # 首先查找GJB_8114相关的sheet
            target_sheets = [sheet for sheet in all_sheets if TARGET_SHEET_PATTERN in sheet]
            
            if target_sheets:
                process_sheets = target_sheets
                self.logger.info(f"找到GJB_8114相关工作表: {', '.join(target_sheets)}")
            else:
                # 如果没有找到GJB_8114相关sheet，处理其他sheet（排除跳过的sheet）
                process_sheets = [sheet for sheet in all_sheets if sheet not in SKIP_SHEETS]
                self.logger.info(f"未找到GJB_8114工作表，将处理其他工作表: {', '.join(process_sheets)}")
            
            all_issues = []
            
            # 处理每个sheet
            for sheet_name in process_sheets:
                self.logger.info(f"正在处理工作表: {sheet_name}")
                worksheet = workbook[sheet_name]
                
                # 检查工作表是否为空
                if worksheet.max_row < 2:
                    self.logger.warning(f"工作表 '{sheet_name}' 为空，跳过")
                    continue
                
                # 读取该sheet的数据
                sheet_issues = self._read_sheet_data(worksheet, sheet_name)
                all_issues.extend(sheet_issues)
                
                self.logger.info(f"从工作表 '{sheet_name}' 读取了 {len(sheet_issues)} 个问题")
            
            workbook.close()
            
            self.logger.info(f"静态分析数据读取完成，总共 {len(all_issues)} 个问题")
            return {"success": True, "data": all_issues}
            
        except Exception as e:
            error_msg = f"读取静态分析数据时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg}
    
    def _read_sheet_data(self, worksheet, sheet_name: str) -> List[Dict[str, Any]]:
        """读取单个工作表的数据"""
        issues = []
        
        # 检查必要的列是否存在
        max_col_index = max(self.required_columns.values())
        if worksheet.max_column <= max_col_index:
            self.logger.warning(f"工作表 '{sheet_name}' 列数不足，需要至少{max_col_index + 1}列")
            return []
        
        valid_count = 0
        
        # 逐行处理数据（跳过表头行）
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            try:
                # 检查行是否有足够的列
                if len(row) <= max_col_index:
                    continue
                
                # 提取问题信息
                issue_id = str(row[self.required_columns["编号"]]) if row[self.required_columns["编号"]] is not None else ""
                defect_name = str(row[self.required_columns["缺陷名称"]]) if row[self.required_columns["缺陷名称"]] is not None else ""
                issue_file = str(row[self.required_columns["问题文件"]]) if row[self.required_columns["问题文件"]] is not None else ""
                issue_line = str(row[self.required_columns["问题行数"]]) if row[self.required_columns["问题行数"]] is not None else ""
                issue_description = str(row[self.required_columns["问题描述"]]) if row[self.required_columns["问题描述"]] is not None else ""
                
                # 使用缺陷名称作为问题名称，并提取GJB规则编号
                issue_name = defect_name
                gjb_rule_code = self._extract_gjb_rule_code(defect_name)
                
                # 清理问题描述中的括号（参考原有逻辑）
                issue_description = self._clean_issue_description(issue_description)
                
                # 跳过空行或无效数据
                if not issue_name or len(issue_name.strip()) == 0:
                    continue
                
                issues.append({
                    "row_index": row_idx,  # Excel行号
                    "sheet_name": sheet_name,  # 工作表名称
                    "issue_id": issue_id,  # 编号
                    "issue_name": issue_name,  # 缺陷名称
                    "gjb_rule_code": gjb_rule_code,  # 提取的GJB规则编号
                    "issue_file": issue_file,
                    "issue_line": issue_line,
                    "issue_description": issue_description
                })
                
                valid_count += 1
                
            except Exception as e:
                self.logger.warning(f"处理工作表 '{sheet_name}' 第{row_idx}行数据时出错: {str(e)}")
                continue
        
        return issues
    
    def _analyze_issue_for_output(self, issue) -> List[str]:
        """
        分析问题并生成12列输出数据
        根据原始项目的问题分析逻辑
        """
        # 基本信息
        wenti_mingcheng = issue.issue_name
        wenti_laiyuan = f"{issue.issue_file}文件{issue.issue_line}行" if issue.issue_file else "源代码"
        wenti_miaoshu = issue.issue_description or ""
        
        # 问题等级分析
        if '越界' in wenti_mingcheng or '章节缺失' in wenti_mingcheng or '失效' in wenti_mingcheng:
            wentidengji = '重要问题'
        else:
            wentidengji = '一般问题'
        
        # 问题数目分析
        if '、' in wenti_laiyuan:
            comma_count = wenti_laiyuan.count('、') + 1
            wenti_shumu = str(comma_count)
        else:
            wenti_shumu = '1'
        
        # 根据问题来源判断问题类型和测试活动
        if '.c' in wenti_laiyuan or '.h' in wenti_laiyuan or '.cpp' in wenti_laiyuan or '源代码' in wenti_laiyuan:
            wentileixing = '程序问题'
            if '编码规范' in wenti_mingcheng:
                ceshi_huodong = '静态分析'
            else:
                ceshi_huodong = '静态分析'
            ceping_danwei_yijian = f"请开发人员针对{wenti_mingcheng}问题对代码进行修改，保证代码实现正确。"
        elif '文档' in wenti_laiyuan:
            ceshi_huodong = '静态分析'
            wentileixing = '软件需求问题'
            ceping_danwei_yijian = '请开发人员对文档问题进行修改,确保文档描述的正确且测试人员能够理解测试需求。'
        elif '任务书' in wenti_laiyuan:
            ceshi_huodong = '静态分析'
            wentileixing = '用户需求问题'
            ceping_danwei_yijian = '请开发人员对文档问题进行修改,确保文档描述的正确且测试人员能够理解测试需求。'
        elif '规格说明' in wenti_laiyuan or '设计说明' in wenti_laiyuan or '影响域分析报告' in wenti_laiyuan or '接口' in wenti_laiyuan:
            ceshi_huodong = '静态分析'
            wentileixing = '软件需求问题'
            ceping_danwei_yijian = '请开发人员对文档问题进行修改,确保文档描述的正确且测试人员能够理解测试需求。'
        else:
            ceshi_huodong = '静态分析'
            wentileixing = '程序问题'
            ceping_danwei_yijian = f"请开发人员针对{wenti_mingcheng}问题对代码进行修改，保证代码实现正确。"
        
        # 问题描述处理（将问题名称转换为需求描述）
        wenti_xuqiu_miaoshu = self._convert_problem_to_requirement(wenti_mingcheng)
        
        # 生成四段式问题描述
        wenti_siduan_miaoshu = self._generate_four_part_description(
            wenti_mingcheng, wenti_laiyuan, wenti_miaoshu, wenti_xuqiu_miaoshu, ceshi_huodong
        )
        
        # 返回12列数据
        return [
            wenti_mingcheng,           # 第1列：问题名称
            wenti_laiyuan,             # 第2列：问题来源
            wenti_miaoshu,             # 第3列：问题复现
            wenti_xuqiu_miaoshu,       # 第4列：问题解释
            ceshi_huodong,             # 第5列：测试活动
            wentileixing,              # 第6列：问题类型
            wentidengji,               # 第7列：问题等级
            wenti_shumu,               # 第8列：问题数目
            wenti_mingcheng,           # 第9列：问题名称（重复）
            wenti_laiyuan,             # 第10列：问题来源（重复）
            wenti_siduan_miaoshu,      # 第11列：问题描述（四段式）
            ceping_danwei_yijian       # 第12列：测评单位意见
        ]
    
    def _convert_problem_to_requirement(self, wenti_mingcheng: str) -> str:
        """
        将问题名称转换为需求描述
        根据原始项目的转换逻辑
        """
        import re
        
        # 检查并替换关键词
        if '不规范' in wenti_mingcheng:
            return re.sub(r'不规范', '规范', wenti_mingcheng)
        elif '不清晰' in wenti_mingcheng:
            return re.sub(r'不清晰', '清晰', wenti_mingcheng)
        elif '不符合' in wenti_mingcheng:
            return re.sub(r'不符合', '符合', wenti_mingcheng)
        elif '不充分' in wenti_mingcheng:
            return re.sub(r'不充分', '充分', wenti_mingcheng)
        elif '不一致' in wenti_mingcheng:
            return re.sub(r'不一致', '一致', wenti_mingcheng)
        elif '混乱' in wenti_mingcheng:
            return re.sub(r'混乱', '清晰', wenti_mingcheng)
        elif '不对' in wenti_mingcheng:
            return re.sub(r'不对', '正确', wenti_mingcheng)
        elif '缺失' in wenti_mingcheng:
            return re.sub(r'缺失', '完整', wenti_mingcheng)
        elif '缺少' in wenti_mingcheng:
            return re.sub(r'缺少', '需要补充', wenti_mingcheng)
        elif '错误' in wenti_mingcheng:
            return re.sub(r'错误', '正确', wenti_mingcheng)
        elif '没有' in wenti_mingcheng:
            return re.sub(r'没有', '需要补充', wenti_mingcheng)
        elif '不全' in wenti_mingcheng:
            return re.sub(r'不全', '需要进行详细补充', wenti_mingcheng)
        elif '同名' in wenti_mingcheng:
            return re.sub(r'同名', '不能同名，具有各自独立的名称', wenti_mingcheng)
        elif '存在未初始化变量' in wenti_mingcheng:
            return re.sub(r'存在未初始化变量', '需要对变量进行初始化', wenti_mingcheng)
        elif '存在死循环' in wenti_mingcheng:
            return re.sub(r'存在死循环', '不存在死循环', wenti_mingcheng)
        elif '存在空语句' in wenti_mingcheng:
            return re.sub(r'存在空语句', '不存在空语句', wenti_mingcheng)
        elif '数组边界没有界定' in wenti_mingcheng:
            return re.sub(r'数组边界没有界定', '需要对数组边界进行界定', wenti_mingcheng)
        elif '存在数组越界' in wenti_mingcheng:
            return re.sub(r'存在数组越界', '不存在数组越界', wenti_mingcheng)
        elif '重定义' in wenti_mingcheng:
            return re.sub(r'重定义', '不能重定义，需要使用不同的名称', wenti_mingcheng)
        elif '不匹配' in wenti_mingcheng:
            return re.sub(r'不匹配', '匹配', wenti_mingcheng)
        elif '无法' in wenti_mingcheng:
            return re.sub(r'无法', '能够', wenti_mingcheng)
        elif '未' in wenti_mingcheng:
            return re.sub(r'未', '已', wenti_mingcheng)
        elif '没' in wenti_mingcheng:
            return re.sub(r'没', '有', wenti_mingcheng)
        elif '数组越界' in wenti_mingcheng:
            return re.sub(r'数组越界', '数组不越界', wenti_mingcheng)
        else:
            return wenti_mingcheng
    
    def _generate_four_part_description(self, wenti_mingcheng: str, wenti_laiyuan: str, 
                                      wenti_miaoshu: str, wenti_xuqiu_miaoshu: str, 
                                      ceshi_huodong: str) -> str:
        """
        生成四段式问题描述：需求描述、问题描述、问题定位、影响分析
        根据原始项目的格式
        """
        # 影响分析根据问题类型生成
        if ceshi_huodong == '文档审查':
            wenti_yingxiangfenxi = "不利于代码维护以及测试人员进行测试需求的提取以及测试用例的设计"
        else:
            # 根据问题名称生成影响分析
            wenti_yingxiangfenxi = self._generate_impact_analysis(wenti_mingcheng)
        
        # 构建四段式描述
        siduan_miaoshu = (
            f"需求描述:{wenti_xuqiu_miaoshu}。\n"
            f"问题描述:{wenti_mingcheng}。\n"
            f"问题定位: {wenti_laiyuan},具体内容如下：\"{wenti_miaoshu}\",如上述内容所示,{wenti_mingcheng}。\n"
            f"影响分析:{wenti_yingxiangfenxi}"
        )
        
        # 统一的格式处理
        siduan_miaoshu = siduan_miaoshu.replace('。。', '。')
        
        return siduan_miaoshu
    
    def _generate_impact_analysis(self, wenti_mingcheng: str) -> str:
        """
        根据问题名称生成影响分析
        """
        if '越界' in wenti_mingcheng:
            return "可能导致程序崩溃或产生不可预期的行为，严重影响软件稳定性"
        elif '内存' in wenti_mingcheng and '泄漏' in wenti_mingcheng:
            return "可能导致系统内存不足，长时间运行后影响系统性能"
        elif '死循环' in wenti_mingcheng:
            return "可能导致程序无响应，严重影响软件可用性"
        elif '未初始化' in wenti_mingcheng:
            return "可能导致程序行为不确定，影响软件可靠性"
        elif '编码规范' in wenti_mingcheng or '命名' in wenti_mingcheng:
            return "影响代码可读性和可维护性，不利于后续开发和维护工作"
        elif '注释' in wenti_mingcheng:
            return "影响代码可理解性，不利于代码维护和团队协作"
        elif '格式' in wenti_mingcheng:
            return "影响代码整体质量，不利于代码审查和维护"
        else:
            return f"{wenti_mingcheng},在某些特定场景下容易引发代码逻辑混乱，进而引发软件故障。"
    
    def _set_excel_formatting(self, worksheet):
        """
        设置Excel格式：行高、列宽、自动换行
        """
        # 设置行高为40
        for row in range(1, worksheet.max_row + 1):
            worksheet.row_dimensions[row].height = 40
        
        # 设置列宽：前10列为20，后2列为60
        from openpyxl.utils import get_column_letter
        
        for col in range(1, 13):  # 12列
            col_letter = get_column_letter(col)
            if col <= 10:
                worksheet.column_dimensions[col_letter].width = 20
            else:
                worksheet.column_dimensions[col_letter].width = 60
    
    def _extract_gjb_rule_code(self, defect_name: str) -> str:
        """从缺陷名称中提取GJB规则编号"""
        if not defect_name:
            return ""
        
        # 匹配GJB_R_XX_XX_XX格式的规则编号
        pattern = re.compile(r"GJB_R_(\d{2})_(\d{2})_(\d{2})")
        match = pattern.search(defect_name)
        
        if match:
            return f"GJB_R_{match.group(1)}_{match.group(2)}_{match.group(3)}"
        
        return ""
    
    def _clean_issue_name(self, issue_name: str) -> str:
        """清理问题名称中的异常字符"""
        if not issue_name:
            return ""
        
        # 去掉GJB规范编号（参考原有逻辑）
        pattern1 = re.compile(r"GJB_R_\d{2}_\d{2}_\d{2}")
        pattern2 = re.compile(r"S_\d{2}_\d{2}")
        
        cleaned = pattern1.sub("", issue_name)
        cleaned = pattern2.sub("", cleaned)
        cleaned = cleaned.replace(' ', '')
        
        return cleaned.strip()
    
    def _clean_issue_description(self, issue_description: str) -> str:
        """清理问题描述中的括号"""
        if not issue_description:
            return ""
        
        # 对科代分析工具中的括号和大括号进行去除（参考原有逻辑）
        cleaned = issue_description.replace('在(', "在").replace(')文件', '文件')
        cleaned = cleaned.replace('第(', "第").replace(')行', '行')
        cleaned = cleaned.replace('数字[', "数字").replace('件[', "件").replace(']件', "件")
        
        return cleaned.strip()
    
    def _save_to_database(self, session_id: str, file_path: str, original_filename: str, 
                         issues_data: List[Dict], filter_keywords: List[str]) -> Dict[str, Any]:
        """保存数据到数据库 - 采用分批次插入优化性能"""
        try:
            self.logger.info(f"开始保存数据到数据库 - 会话ID: {session_id}, 问题数: {len(issues_data)}")
            
            db = get_db_session()
            
            try:
                # 创建会话记录
                session = StaticAnalysisSession(
                    session_id=session_id,
                    original_filename=original_filename,
                    file_path=file_path,
                    status="uploading",  # 改为uploading状态，避免过早显示processing
                    total_issues=len(issues_data),
                    filter_keywords=",".join(filter_keywords) if filter_keywords else ""
                )
                db.add(session)
                db.commit()  # 立即提交会话记录
                db.refresh(session)  # 刷新以获取自增ID
                
                # 获取任务ID
                task_id = session.id
                self.logger.info(f"会话已创建 - Session ID: {session_id}, Task ID: {task_id}")
                
                # 分批次插入问题记录，避免单个事务过大
                BATCH_SIZE = BATCH_PROCESSING_CONFIG["db_insert_batch_size"]  # 从配置获取批次大小
                total_batches = (len(issues_data) + BATCH_SIZE - 1) // BATCH_SIZE
                
                self.logger.info(f"将分{total_batches}个批次插入数据")
                
                for batch_idx in range(0, len(issues_data), BATCH_SIZE):
                    batch_data = issues_data[batch_idx:batch_idx + BATCH_SIZE]
                    current_batch = (batch_idx // BATCH_SIZE) + 1
                    
                    self.logger.info(f"插入第{current_batch}/{total_batches}批次，{len(batch_data)}条记录")
                    
                    # 创建当前批次的问题记录
                    batch_issues = []
                    for issue_data in batch_data:
                        self.logger.debug(f"保存问题: {issue_data['issue_name']}, GJB编号: {issue_data.get('gjb_rule_code', '无')}")
                        issue = StaticAnalysisIssue(
                            task_id=task_id,
                            row_index=issue_data["row_index"],
                            issue_name=issue_data["issue_name"],
                            issue_file=issue_data["issue_file"],
                            issue_line=issue_data["issue_line"],
                            issue_description=issue_data["issue_description"],
                            gjb_rule_code=issue_data.get("gjb_rule_code", ""),
                            sheet_name=issue_data.get("sheet_name", "")
                        )
                        batch_issues.append(issue)
                    
                    # 批量插入当前批次
                    db.add_all(batch_issues)
                    db.commit()  # 每批次都提交，避免长事务
                    
                    # 释放内存
                    batch_issues.clear()
                
                # 更新会话状态为已上传
                session.status = "uploaded"
                db.commit()
                
                self.logger.info(f"数据保存成功 - 会话: {session_id}, 总问题数: {len(issues_data)}个")
                
                return {
                    "success": True,
                    "total_issues": len(issues_data),
                    "session_id": session_id
                }
                
            finally:
                db.close()
                
        except Exception as e:
            error_msg = f"保存数据到数据库时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg}
    
    def generate_filtered_excel(self, session_id: str) -> Dict[str, Any]:
        """
        生成过滤后的Excel文件
        """
        try:
            self.logger.info(f"开始生成过滤后的Excel文件")
            
            db = get_db_session()
            
            try:
                # 获取会话信息
                session = db.query(StaticAnalysisSession).filter(
                    StaticAnalysisSession.session_id == session_id
                ).first()
                
                if not session:
                    return {"success": False, "error": "会话不存在"}
                
                # 获取未被过滤的问题项
                task_id = session.id
                retained_issues = db.query(StaticAnalysisIssue).filter(
                    StaticAnalysisIssue.task_id == task_id,
                    StaticAnalysisIssue.is_filtered == False
                ).order_by(StaticAnalysisIssue.row_index).all()
                
                if not retained_issues:
                    return {"success": False, "error": "没有保留的问题项"}
                
                # 生成输出文件路径
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                output_filename = f"filtered_issues_{session_id}_{timestamp}.xlsx"
                from ..config.config import UPLOAD_DIR
                output_path = UPLOAD_DIR / "outputs" / output_filename
                output_path.parent.mkdir(parents=True, exist_ok=True)
                
                # 创建新的Excel文件
                workbook = openpyxl.Workbook()
                worksheet = workbook.active
                worksheet.title = "过滤后的问题"
                
                # 设置样式
                from openpyxl.styles import Alignment
                
                # 写入表头（12列格式）
                from ..config.config import OUTPUT_COLUMNS
                for col_idx, header in enumerate(OUTPUT_COLUMNS, 1):
                    cell = worksheet.cell(row=1, column=col_idx, value=header)
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                
                # 写入数据
                for row_idx, issue in enumerate(retained_issues, 2):
                    # 分析问题并生成12列数据
                    issue_data = self._analyze_issue_for_output(issue)
                    
                    for col_idx, value in enumerate(issue_data, 1):
                        cell = worksheet.cell(row=row_idx, column=col_idx, value=value)
                        # 设置单元格对齐方式
                        if col_idx <= 10:  # 前10列
                            cell.alignment = Alignment(horizontal='left', vertical='center')
                        else:  # 后2列设置自动换行
                            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
                
                # 设置行高和列宽
                self._set_excel_formatting(worksheet)
                
                # 保存文件
                workbook.save(output_path)
                workbook.close()
                
                # 更新会话状态
                session.output_filename = output_filename
                session.status = "completed"
                db.commit()
                
                self.logger.info(f"过滤后的Excel文件生成成功: {output_path}, 保留了{len(retained_issues)}个问题")
                
                return {
                    "success": True,
                    "output_path": str(output_path),
                    "total_retained": len(retained_issues),
                    "filename": output_filename
                }
                
            finally:
                db.close()
                
        except Exception as e:
            error_msg = f"生成过滤后的Excel文件时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg} 
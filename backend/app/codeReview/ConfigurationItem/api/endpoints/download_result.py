"""
下载分析结果的API端点 - Excel格式
"""
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from io import BytesIO
import openpyxl
from openpyxl.styles import Alignment
import json
import os
import re
import glob
import zipfile
from urllib.parse import quote
from ...config import ALL_SUPPORTED_EXTENSIONS
from ...config_database import get_db  # 函数审查数据库切换模块
from ...config_database_chain import get_chain_db  # 链路审查数据库切换模块
from ...config_database_feature import get_feature_db  # 功能审查数据库切换模块
from ...db.deps import get_db
from ...db.repository.problem_repo import ProblemRepository
from ...db.repository.task_repo import TaskRepository
from ...db.repository.file_repo import FileRepository
from ...db.repository.function_repo import FunctionRepository
from ...db.repository.chain_problem_repo import ChainProblemRepository
from ...db.repository.file_problem_repo import FileProblemRepository
from ...db.repository.feature_repo import FeatureReviewProblemRepository

# 导入统一的日志系统
from ...utils.logger import log_debug, log_error

router = APIRouter()

# 函数级审查列
FUNCTION_COLUMNS = [
    "文件名称",
    "源代码",
    "问题来源",
    "违反规则",
    "问题等级",
    "问题名称",
    "问题位置",
    "问题描述"
]

# 函数审查代码视图标记问题单列
FUNCTION_MARKED_PROBLEM_COLUMNS = [
    "源代码",
    "测试活动",
    "问题类型",
    "问题等级",
    "问题数目",
    "问题名称",
    "问题来源",
    "问题描述"
]

# 业务级审查列
BUSINESS_COLUMNS = [
    "链路名称/文件名称",
    "链路功能/文件功能",
    "涉及函数/源代码",
    "问题原因",
    "问题等级", 
    "问题名称",
    "问题位置",
    "问题描述"
]


def _extract_file_level_location(problem: dict) -> str:
    """从问题位置中提取文件级位置，不包含函数名。"""
    issue_location = str(problem.get("issue_location") or "").strip()
    file_name = str(problem.get("file_name") or "").strip()

    if not issue_location:
        return file_name

    for separator in ["::", "#", "函数", "方法", "行号", "第"]:
        if separator in issue_location:
            issue_location = issue_location.split(separator, 1)[0].strip()

    if ":" in issue_location:
        prefix, suffix = issue_location.rsplit(":", 1)
        if suffix.strip().isdigit():
            issue_location = prefix.strip()

    return issue_location or file_name


def _format_worksheet(ws):
    """格式化工作表：设置列宽和自动换行"""
    wrap_alignment = Alignment(wrapText=True, vertical="top")
    for column_cells in ws.columns:
        # 计算该列最大字符长度
        max_length = 0
        for cell in column_cells:
            cell.alignment = wrap_alignment  # 自动换行
            if cell.value:
                cell_length = len(str(cell.value))
                if cell_length > max_length:
                    max_length = cell_length
        # 基于最大长度设置列宽，限制在 20~60 之间
        adjusted_width = min(max(max_length + 2, 20), 60)
        column_letter = column_cells[0].column_letter
        ws.column_dimensions[column_letter].width = adjusted_width


def _build_commented_file_content(original_source: str, functions: list) -> str:
    """将原始文件内容与函数的注释后代码合并，生成完整的注释版本文件内容。

    规则（与前端 CodeMapSourceView 展示逻辑保持一致）：
    - 已生成注释的函数用 commented_code 替换原始代码段；
    - 未生成注释的函数、函数之间的代码（全局变量、结构体等）、
      文件头部与尾部均保留原始代码。

    Args:
        original_source: 文件原始内容
        functions: 该文件下的函数列表，每项为 dict，包含
                   start_line / end_line / commented_code / source_code

    Returns:
        合并后的完整注释版本文件内容
    """
    if not original_source:
        return original_source or ""

    file_lines = original_source.split('\n')
    max_line = len(file_lines)

    # 仅保留具备有效行号信息的函数，并按起始行排序
    valid_funcs = []
    for func in functions or []:
        try:
            start_line = int(func.get('start_line') or 0)
        except (TypeError, ValueError):
            start_line = 0
        if start_line > 0:
            valid_funcs.append((start_line, func))
    valid_funcs.sort(key=lambda item: item[0])

    if not valid_funcs:
        return original_source

    parts = []

    # 插入第一个函数之前的内容（文件头、include 等）
    first_start = valid_funcs[0][0] - 1
    if first_start > 0:
        header = '\n'.join(file_lines[:first_start]).rstrip()
        if header.strip():
            parts.append(header)
            parts.append('')

    prev_end = first_start

    for idx, (start_line, func) in enumerate(valid_funcs):
        start = max(start_line - 1, 0)
        try:
            end = int(func.get('end_line') or 0)
        except (TypeError, ValueError):
            end = 0
        if end < start_line:
            end = start_line
        end = min(end, max_line)

        # 插入函数之间的非函数代码
        if idx > 0 and start > prev_end:
            between = '\n'.join(file_lines[prev_end:start]).strip('\n')
            if between.strip():
                parts.append(between)
                parts.append('')

        # 优先使用注释后代码，其次原函数代码，最后回退到文件原始行
        body = func.get('commented_code') or func.get('source_code')
        if not body:
            body = '\n'.join(file_lines[start:end])
        parts.append(body)
        parts.append('')

        prev_end = max(prev_end, end)

    # 插入最后一个函数之后的内容
    if max_line > prev_end:
        footer = '\n'.join(file_lines[prev_end:]).strip('\n')
        if footer.strip():
            parts.append(footer)

    return '\n'.join(parts)


@router.get("/download_result/{task_id}")
async def download_result(task_id: int, db: Session = Depends(get_db)):
    """
    下载分析结果为Excel文件
    
    Args:
        task_id: 任务ID
        db: 数据库会话
        
    Returns:
        Excel文件下载响应
    """
    try:
        task_repo = TaskRepository(db)
        
        # 检查任务是否存在
        task = task_repo.get_task(task_id)
        if not task:
            raise HTTPException(status_code=404, detail="任务不存在")
        
        # 判断审查类型
        analysis_type = 'function'  # 默认函数审查
        
        if task.task_name:
            if '链路审查' in task.task_name:
                analysis_type = 'chain'
            elif '功能审查' in task.task_name or '业务审查' in task.task_name:
                analysis_type = 'feature'
        
        # 或者根据task_description判断
        if task.task_description:
            if '链路审查' in task.task_description or '链路功能' in task.task_description:
                analysis_type = 'chain'
            elif '项目功能级别' in task.task_description:
                analysis_type = 'feature'
        
        log_debug(f"任务{task_id}的审查类型: {analysis_type}", "download_result")
        
        # 初始化变量
        function_problems = []
        chain_problems = []
        file_problems = []
        feature_problems = []
        
        # 根据审查类型获取对应的问题
        if analysis_type == 'function':
            # 函数审查：只查询函数级审查问题
            problem_repo = ProblemRepository(db)
            function_problems = problem_repo.get_problems_by_task(task_id)
        elif analysis_type == 'chain':
            # 链路审查：使用链路审查数据库
            chain_db = next(get_chain_db())
            try:
                chain_problem_repo = ChainProblemRepository(chain_db)
                chain_problems_raw = chain_problem_repo.get_problems_by_task(task_id)
                chain_problems = [problem.to_dict() for problem in chain_problems_raw]
            finally:
                chain_db.close()
        elif analysis_type == 'feature':
            # 功能审查：使用功能审查数据库
            feature_db = next(get_feature_db())
            try:
                feature_problem_repo = FeatureReviewProblemRepository(feature_db)
                feature_problems_raw = feature_problem_repo.get_problems_by_task(task_id)
                feature_problems = [problem.to_dict() for problem in feature_problems_raw]
            finally:
                feature_db.close()
        
        # 检查是否有任何问题
        if not function_problems and not chain_problems and not file_problems and not feature_problems:
            raise HTTPException(status_code=404, detail="该任务没有发现问题")
        
        # 创建Excel工作簿
        wb = openpyxl.Workbook()
        
        # 创建函数级审查sheet
        if function_problems:
            ws_function = wb.active
            ws_function.title = "函数级审查结果"
            
            # 添加表头
            ws_function.append(FUNCTION_COLUMNS)
            
            # 添加数据行
            for problem in function_problems:
                # 合并四段式描述
                description = "\n".join([
                    f"需求描述：{str(problem.get('description_requirement', ''))}",
                    f"问题描述：{str(problem.get('description_issue', ''))}",
                    f"问题定位：{str(problem.get('description_location', ''))}",
                    f"影响分析：{str(problem.get('description_impact', ''))}"
                ])
                
                row = [
                    problem.get("file_name", ""),
                    problem.get("source_code", ""),
                    problem.get("problem_source", "main_review"),
                    problem.get("problem_reason", ""),
                    problem.get("issue_level", ""),
                    problem.get("issue_name", ""),
                    problem.get("issue_location", ""),
                    description
                ]
                ws_function.append(row)
            
            # 设置函数级审查sheet的格式
            _format_worksheet(ws_function)
            log_debug(f"添加函数级审查问题: {len(function_problems)} 个", "download_result")
        
        # 创建链路审查sheet
        if chain_problems:
            if function_problems:
                ws_chain = wb.create_sheet(title="链路审查结果")
            else:
                ws_chain = wb.active
                ws_chain.title = "链路审查结果"
            
            # 添加表头
            ws_chain.append(BUSINESS_COLUMNS)
            
            # 添加数据行
            for problem in chain_problems:
                try:
                    affected_functions = json.loads(problem.get("affected_functions", "[]"))
                    if isinstance(affected_functions, list):
                        affected_functions_str = ", ".join(affected_functions)
                    else:
                        affected_functions_str = str(affected_functions)
                except:
                    affected_functions_str = problem.get("affected_functions", "")
                
                # 合并四段式描述为详细描述
                detailed_description = "\n".join([
                    f"需求描述：{str(problem.get('description_requirement', ''))}",
                    f"问题描述：{str(problem.get('description_issue', ''))}",
                    f"问题定位：{str(problem.get('description_location', ''))}",
                    f"影响分析：{str(problem.get('description_impact', ''))}"
                ])
                
                row = [
                    problem.get("chain_name", ""),
                    problem.get("chain_function_description", ""),
                    affected_functions_str,
                    problem.get("problem_reason", ""),
                    problem.get("issue_level", ""),
                    problem.get("issue_name", ""),
                    problem.get("issue_location", ""),
                    detailed_description
                ]
                ws_chain.append(row)
            
            # 设置链路审查sheet的格式
            _format_worksheet(ws_chain)
            log_debug(f"添加链路审查问题: {len(chain_problems)} 个", "download_result")
        
        # 创建功能审查sheet
        if feature_problems:
            sheet_count = sum([bool(function_problems), bool(chain_problems)])
            if sheet_count > 0:
                ws_feature = wb.create_sheet(title="功能审查结果")
            else:
                ws_feature = wb.active
                ws_feature.title = "功能审查结果"
            
            # 添加表头
            ws_feature.append(BUSINESS_COLUMNS)
            
            # 添加数据行
            for problem in feature_problems:
                # 获取实现路径（相关函数列表）
                try:
                    implementation_path = json.loads(problem.get("implementation_path", "[]"))
                    if isinstance(implementation_path, list):
                        implementation_path_str = ", ".join(implementation_path)
                    else:
                        implementation_path_str = str(implementation_path)
                except:
                    implementation_path_str = problem.get("implementation_path", "")
                
                # 合并四段式描述为详细描述
                detailed_description = "\n".join([
                    f"需求描述：{str(problem.get('description_requirement', ''))}",
                    f"问题描述：{str(problem.get('description_issue', ''))}",
                    f"问题定位：{str(problem.get('description_location', ''))}",
                    f"影响分析：{str(problem.get('description_impact', ''))}"
                ])
                
                row = [
                    problem.get("feature_name", ""),  # 功能名称
                    problem.get("feature_description", ""),  # 功能描述
                    implementation_path_str,  # 实现路径（涉及函数）
                    problem.get("problem_reason", ""),
                    problem.get("issue_level", ""),
                    problem.get("issue_name", ""),
                    problem.get("issue_location", ""),
                    detailed_description
                ]
                ws_feature.append(row)
            
            # 设置功能审查sheet的格式
            _format_worksheet(ws_feature)
            log_debug(f"添加功能审查问题: {len(feature_problems)} 个", "download_result")
        
        # 导出为字节流
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # 生成文件名
        filename = f"code_review_result_task_{task_id}.xlsx"
        headers = {"Content-Disposition": f"attachment; filename={filename}"}
        
        return StreamingResponse(
            output, 
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
            headers=headers
        )
        
    except Exception as e:
        log_error(f"下载结果失败: task_id={task_id}", "download_result", e)
        raise HTTPException(status_code=500, detail=f"下载失败: {str(e)}")


@router.post("/export_marked_problems/{task_id}")
async def export_marked_problems(task_id: int, request_body: dict, db: Session = Depends(get_db)):
    """
    导出用户在代码视图中标记的问题为Excel文件
    
    Args:
        task_id: 任务ID
        request_body: 包含已标记问题列表的请求体 {"problems": [...]}
        db: 数据库会话
        
    Returns:
        Excel文件下载响应
    """
    try:
        task_repo = TaskRepository(db)
        
        # 检查任务是否存在
        task = task_repo.get_task(task_id)
        if not task:
            raise HTTPException(status_code=404, detail="任务不存在")
        
        # 获取已标记的问题列表
        problems = request_body.get("problems", [])
        analysis_type = request_body.get('analysis_type', 'function')
        
        if not problems:
            raise HTTPException(status_code=400, detail="没有已标记的问题")
        
        log_debug(f"开始导出已标记问题: task_id={task_id}, 问题数量={len(problems)}", "export_marked_problems")
        
        # 创建Excel工作簿
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "已标记问题"
        
        # 添加表头
        ws.append(FUNCTION_MARKED_PROBLEM_COLUMNS)
        
        # 添加数据行
        for problem in problems:
            description = "\n".join([
                f"需求描述：{str(problem.get('description_requirement', ''))}",
                f"问题描述：{str(problem.get('description_issue', ''))}",
                f"问题定位：{str(problem.get('description_location', ''))}",
                f"影响分析：{str(problem.get('description_impact', ''))}"
            ])

            row = [
                problem.get("source_code", ""),
                "代码审查",
                "程序问题",
                problem.get("issue_level", ""),
                1,
                problem.get("issue_name", ""),
                _extract_file_level_location(problem),
                description
            ]
            ws.append(row)
        
        # 设置格式
        _format_worksheet(ws)
        
        log_debug(f"已标记问题格式化完成: {len(problems)} 个", "export_marked_problems")
        
        # 导出为字节流
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # 生成文件名
        filename = f"marked_problems_task_{task_id}.xlsx"
        headers = {"Content-Disposition": f"attachment; filename={filename}"}
        
        log_debug(f"已标记问题导出成功: {filename}", "export_marked_problems")

        # 记录任务已导出过标记问题（目前针对函数审查）
        try:
            if analysis_type == 'function':
                task_repo.set_marked_problems_exported(task_id)
        except Exception:
            pass
        
        return StreamingResponse(
            output, 
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
            headers=headers
        )
        
    except Exception as e:
        log_error(f"导出已标记问题失败: task_id={task_id}", "export_marked_problems", e)
        raise HTTPException(status_code=500, detail=f"导出失败: {str(e)}")


@router.get("/download_commented_code/{task_id}")
async def download_commented_code(task_id: int, db: Session = Depends(get_db)):
    """
    下载函数审查生成的「注释后代码」，打包为 ZIP 文件（保留原目录结构）

    处理规则：
    - 针对任务涉及的所有源码文件，逐个生成注释版本；
    - 已生成注释的函数用 AI 注释后代码替换原始代码段；
    - 未生成注释的函数、函数之间的代码（全局变量、结构体等）、
      文件头部与尾部均保留原始代码。

    Args:
        task_id: 任务ID
        db: 数据库会话

    Returns:
        ZIP 文件下载响应
    """
    try:
        task_repo = TaskRepository(db)

        # 检查任务是否存在
        task = task_repo.get_task(task_id)
        if not task:
            raise HTTPException(status_code=404, detail="任务不存在")

        # 获取代码文件信息（解压路径）
        file_repo = FileRepository(db)
        file_info = file_repo.get_file(task.code_file_id)
        if not file_info or not file_info.extract_path:
            raise HTTPException(status_code=404, detail="代码文件不存在")
        if not os.path.exists(file_info.extract_path):
            raise HTTPException(status_code=404, detail="代码解压路径不存在")

        # 获取该任务下所有函数（含 AI 注释后代码），按文件路径分组
        function_repo = FunctionRepository(db)
        functions = function_repo.get_functions_by_task(task_id)

        funcs_by_file = {}
        funcs_by_basename = {}
        for func in functions:
            file_path = (func.file_path or "").replace("\\", "/").strip()
            if not file_path:
                continue
            item = {
                "function_name": func.function_name,
                "start_line": func.start_line,
                "end_line": func.end_line,
                "source_code": func.source_code,
                "commented_code": func.commented_code,
            }
            funcs_by_file.setdefault(file_path, []).append(item)
            funcs_by_basename.setdefault(os.path.basename(file_path), []).append(item)

        # 扫描项目中的所有源码文件
        project_root = os.path.abspath(file_info.extract_path)
        source_files = []
        seen_paths = set()
        for ext in ALL_SUPPORTED_EXTENSIONS:
            pattern = os.path.join(file_info.extract_path, "**", f"*{ext}")
            for full_path in glob.glob(pattern, recursive=True):
                if not os.path.isfile(full_path):
                    continue
                normalized = os.path.normcase(os.path.abspath(full_path))
                if normalized in seen_paths:
                    continue
                seen_paths.add(normalized)
                source_files.append(full_path)

        if not source_files:
            raise HTTPException(status_code=404, detail="未找到可下载的源代码文件")

        source_files.sort()

        commented_file_count = 0
        output = BytesIO()

        # 复用代码地图的多种编码读取逻辑，避免中文乱码
        from .codemap import read_file_with_multiple_encodings

        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zf:
            for full_path in source_files:
                rel_path = os.path.relpath(full_path, project_root).replace("\\", "/")

                # 读取原始文件内容（多编码兜底）
                try:
                    original_source = read_file_with_multiple_encodings(full_path)
                except Exception as read_err:
                    log_debug(f"跳过无法读取的文件: {full_path} ({read_err})", "download_commented_code")
                    continue

                # 优先按相对路径匹配函数，其次按文件名兜底
                file_funcs = funcs_by_file.get(rel_path) or funcs_by_basename.get(os.path.basename(rel_path), [])

                content = _build_commented_file_content(original_source, file_funcs)
                if file_funcs:
                    commented_file_count += 1

                zf.writestr(rel_path, content)

            # 附加说明文件，便于使用者了解文件内容
            readme = (
                "函数审查 - 注释后代码\n"
                "========================================\n"
                f"任务ID: {task_id}\n"
                f"包含源码文件数: {len(source_files)}\n"
                f"其中含AI注释的文件数: {commented_file_count}\n\n"
                "说明:\n"
                "1. 本压缩包中每个文件为对应源码文件的「注释后代码」版本，目录结构与原项目一致；\n"
                "2. 已生成注释的函数已替换为 AI 注释后的代码；\n"
                "3. 未生成注释的函数及函数之间的代码（全局变量、结构体、文件头尾等）保留原始代码。\n"
            )
            zf.writestr("注释代码说明.txt", readme)

        output.seek(0)

        # 生成 ZIP 包名：源码压缩包名-commented.zip（如 project.zip -> project-commented.zip）
        package_stem = os.path.splitext(os.path.basename(file_info.file_name or ""))[0].strip()
        if not package_stem:
            package_stem = f"task_{task_id}"
        filename = f"{package_stem}-commented.zip"

        # 兼容中文等非 ASCII 文件名：同时提供 ASCII 回退名与 RFC 5987 编码名
        ascii_stem = re.sub(r"[^0-9A-Za-z._-]", "_", package_stem) or f"task_{task_id}"
        ascii_fallback = f"{ascii_stem}-commented.zip"
        content_disposition = (
            f"attachment; filename=\"{ascii_fallback}\"; "
            f"filename*=UTF-8''{quote(filename)}"
        )
        headers = {
            "Content-Disposition": content_disposition,
            # 允许前端跨域/代理读取文件名
            "Access-Control-Expose-Headers": "Content-Disposition",
        }

        log_debug(
            f"注释后代码打包成功: 文件数={len(source_files)}, 含注释文件数={commented_file_count}, 包名={filename}",
            "download_commented_code",
        )

        return StreamingResponse(output, media_type="application/zip", headers=headers)

    except HTTPException:
        raise
    except Exception as e:
        log_error(f"下载注释代码失败: task_id={task_id}", "download_commented_code", e)
        raise HTTPException(status_code=500, detail=f"下载注释代码失败: {str(e)}")

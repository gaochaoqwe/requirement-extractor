"""
下载链路审查结果API
"""
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import io
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime

from ....db.database_chain import get_chain_db
from ....db.models_chain import ChainReviewProblem

router = APIRouter()


@router.get("/download/{task_id}")
async def download_chain_review_result(task_id: int, db: Session = Depends(get_chain_db)):
    """
    下载链路审查结果为Excel文件
    
    参考函数审查格式，但列名适应链路审查：
    - 链路名称（而非函数名称）
    - 链路功能描述
    - 调用链路
    """
    try:
        # 获取所有链路审查问题
        problems = db.query(ChainReviewProblem).filter(
            ChainReviewProblem.task_id == task_id
        ).all()
        
        if not problems:
            raise HTTPException(status_code=404, detail="没有找到链路审查结果")
        
        # 创建Excel工作簿
        wb = Workbook()
        wb.remove(wb.active)  # 删除默认sheet
        
        # 创建样式
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Sheet 1: 链路审查问题清单
        ws_problems = wb.create_sheet("链路审查问题清单")
        problem_headers = [
            "链路名称", "链路功能描述", "问题名称", "严重程度", "问题位置", 
            "问题原因", "需求描述", "问题描述", "问题定位", "影响分析", "改进建议"
        ]
        ws_problems.append(problem_headers)
        
        # 设置表头样式
        for cell in ws_problems[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # 写入问题数据
        for problem in problems:
            ws_problems.append([
                problem.chain_name,
                problem.chain_function_description or "",
                problem.issue_name,
                problem.issue_level,
                problem.issue_location,
                problem.problem_reason or "",
                problem.description_requirement or "",
                problem.description_issue or "",
                problem.description_location or "",
                problem.description_impact or "",
                problem.suggestion or ""
            ])
        
        # 调整列宽
        ws_problems.column_dimensions['A'].width = 25  # 链路名称
        ws_problems.column_dimensions['B'].width = 50  # 链路功能描述
        ws_problems.column_dimensions['C'].width = 25  # 问题名称
        ws_problems.column_dimensions['D'].width = 15  # 严重程度
        ws_problems.column_dimensions['E'].width = 30  # 问题位置
        ws_problems.column_dimensions['F'].width = 40  # 问题原因
        ws_problems.column_dimensions['G'].width = 40  # 需求描述
        ws_problems.column_dimensions['H'].width = 50  # 问题描述
        ws_problems.column_dimensions['I'].width = 50  # 问题定位
        ws_problems.column_dimensions['J'].width = 40  # 影响分析
        ws_problems.column_dimensions['K'].width = 40  # 改进建议
        
        # Sheet 2: 链路问题源代码
        ws_code = wb.create_sheet("问题源代码")
        code_headers = [
            "链路名称", "问题名称", "问题代码片段", "问题位置"
        ]
        ws_code.append(code_headers)
        
        # 设置表头样式
        for cell in ws_code[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # 写入源代码数据
        for problem in problems:
            if problem.source_code:
                ws_code.append([
                    problem.chain_name,
                    problem.issue_name,
                    problem.source_code,
                    problem.description_location or problem.issue_location
                ])
        
        # 调整列宽
        ws_code.column_dimensions['A'].width = 25
        ws_code.column_dimensions['B'].width = 25
        ws_code.column_dimensions['C'].width = 80
        ws_code.column_dimensions['D'].width = 50
        
        # Sheet 3: 调用链路详情
        ws_chain = wb.create_sheet("调用链路详情")
        chain_headers = [
            "链路名称", "调用链路", "链路功能描述", "入口函数", "涉及函数"
        ]
        ws_chain.append(chain_headers)
        
        # 设置表头样式
        for cell in ws_chain[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # 统计每个链路（去重）
        chains_dict = {}
        for problem in problems:
            chain_name = problem.chain_name
            if chain_name not in chains_dict:
                chains_dict[chain_name] = {
                    "call_chain": problem.call_chain or "",
                    "description": problem.chain_function_description or "",
                    "entry_function": problem.entry_function or "",
                    "affected_functions": problem.affected_functions or "[]"
                }
        
        # 写入链路数据
        for chain_name, chain_data in chains_dict.items():
            try:
                affected_functions = json.loads(chain_data["affected_functions"])
                affected_functions_str = ", ".join(affected_functions)
            except:
                affected_functions_str = chain_data["affected_functions"]
            
            ws_chain.append([
                chain_name,
                chain_data["call_chain"],
                chain_data["description"],
                chain_data["entry_function"],
                affected_functions_str
            ])
        
        # 调整列宽
        ws_chain.column_dimensions['A'].width = 25
        ws_chain.column_dimensions['B'].width = 60
        ws_chain.column_dimensions['C'].width = 50
        ws_chain.column_dimensions['D'].width = 20
        ws_chain.column_dimensions['E'].width = 60
        
        # 保存到内存
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        # 生成文件名
        filename = f"chain_review_task{task_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename={filename}"
            }
        )
        
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"下载失败: {str(e)}")


"""
下载分析结果的API端点 - Excel格式
"""
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from io import BytesIO
import openpyxl
from openpyxl.styles import Alignment
import json
from ...config_database import get_db  # 函数审查数据库切换模块
from ...config_database_chain import get_chain_db  # 链路审查数据库切换模块
from ...config_database_feature import get_feature_db  # 功能审查数据库切换模块
from ...db.deps import get_db
from ...db.repository.problem_repo import ProblemRepository
from ...db.repository.task_repo import TaskRepository
from ...db.repository.chain_problem_repo import ChainProblemRepository
from ...db.repository.file_problem_repo import FileProblemRepository
from ...db.repository.feature_repo import FeatureReviewProblemRepository

# 导入统一的日志系统
from ...utils.logger import log_debug, log_error

router = APIRouter()

# 函数级审查列
FUNCTION_COLUMNS = [
    "文件名称",
    "源代码",
    "违反规则",
    "问题等级",
    "问题名称",
    "问题位置",
    "问题描述"
]

# 业务级审查列
BUSINESS_COLUMNS = [
    "链路名称/文件名称",
    "链路功能/文件功能",
    "涉及函数/源代码",
    "问题原因",
    "问题等级", 
    "问题名称",
    "问题位置",
    "问题描述"
]


def _format_worksheet(ws):
    """格式化工作表：设置列宽和自动换行"""
    wrap_alignment = Alignment(wrapText=True, vertical="top")
    for column_cells in ws.columns:
        # 计算该列最大字符长度
        max_length = 0
        for cell in column_cells:
            cell.alignment = wrap_alignment  # 自动换行
            if cell.value:
                cell_length = len(str(cell.value))
                if cell_length > max_length:
                    max_length = cell_length
        # 基于最大长度设置列宽，限制在 20~60 之间
        adjusted_width = min(max(max_length + 2, 20), 60)
        column_letter = column_cells[0].column_letter
        ws.column_dimensions[column_letter].width = adjusted_width


@router.get("/download_result/{task_id}")
async def download_result(task_id: int, db: Session = Depends(get_db)):
    """
    下载分析结果为Excel文件
    
    Args:
        task_id: 任务ID
        db: 数据库会话
        
    Returns:
        Excel文件下载响应
    """
    try:
        task_repo = TaskRepository(db)
        
        # 检查任务是否存在
        task = task_repo.get_task(task_id)
        if not task:
            raise HTTPException(status_code=404, detail="任务不存在")
        
        # 判断审查类型
        analysis_type = 'function'  # 默认函数审查
        
        if task.task_name:
            if '链路审查' in task.task_name:
                analysis_type = 'chain'
            elif '功能审查' in task.task_name or '业务审查' in task.task_name:
                analysis_type = 'feature'
        
        # 或者根据task_description判断
        if task.task_description:
            if '链路审查' in task.task_description or '链路功能' in task.task_description:
                analysis_type = 'chain'
            elif '项目功能级别' in task.task_description:
                analysis_type = 'feature'
        
        log_debug(f"任务{task_id}的审查类型: {analysis_type}", "download_result")
        
        # 初始化变量
        function_problems = []
        chain_problems = []
        file_problems = []
        feature_problems = []
        
        # 根据审查类型获取对应的问题
        if analysis_type == 'function':
            # 函数审查：只查询函数级审查问题
            problem_repo = ProblemRepository(db)
            function_problems = problem_repo.get_problems_by_task(task_id)
        elif analysis_type == 'chain':
            # 链路审查：使用链路审查数据库
            chain_db = next(get_chain_db())
            try:
                chain_problem_repo = ChainProblemRepository(chain_db)
                chain_problems_raw = chain_problem_repo.get_problems_by_task(task_id)
                chain_problems = [problem.to_dict() for problem in chain_problems_raw]
            finally:
                chain_db.close()
        elif analysis_type == 'feature':
            # 功能审查：使用功能审查数据库
            feature_db = next(get_feature_db())
            try:
                feature_problem_repo = FeatureReviewProblemRepository(feature_db)
                feature_problems_raw = feature_problem_repo.get_problems_by_task(task_id)
                feature_problems = [problem.to_dict() for problem in feature_problems_raw]
            finally:
                feature_db.close()
        
        # 检查是否有任何问题
        if not function_problems and not chain_problems and not file_problems and not feature_problems:
            raise HTTPException(status_code=404, detail="该任务没有发现问题")
        
        # 创建Excel工作簿
        wb = openpyxl.Workbook()
        
        # 创建函数级审查sheet
        if function_problems:
            ws_function = wb.active
            ws_function.title = "函数级审查结果"
            
            # 添加表头
            ws_function.append(FUNCTION_COLUMNS)
            
            # 添加数据行
            for problem in function_problems:
                # 合并四段式描述
                description = "\n".join([
                    f"需求描述：{str(problem.get('description_requirement', ''))}",
                    f"问题描述：{str(problem.get('description_issue', ''))}",
                    f"问题定位：{str(problem.get('description_location', ''))}",
                    f"影响分析：{str(problem.get('description_impact', ''))}"
                ])
                
                row = [
                    problem.get("file_name", ""),
                    problem.get("source_code", ""),
                    problem.get("problem_reason", ""),
                    problem.get("issue_level", ""),
                    problem.get("issue_name", ""),
                    problem.get("issue_location", ""),
                    description
                ]
                ws_function.append(row)
            
            # 设置函数级审查sheet的格式
            _format_worksheet(ws_function)
            log_debug(f"添加函数级审查问题: {len(function_problems)} 个", "download_result")
        
        # 创建链路审查sheet
        if chain_problems:
            if function_problems:
                ws_chain = wb.create_sheet(title="链路审查结果")
            else:
                ws_chain = wb.active
                ws_chain.title = "链路审查结果"
            
            # 添加表头
            ws_chain.append(BUSINESS_COLUMNS)
            
            # 添加数据行
            for problem in chain_problems:
                try:
                    affected_functions = json.loads(problem.get("affected_functions", "[]"))
                    if isinstance(affected_functions, list):
                        affected_functions_str = ", ".join(affected_functions)
                    else:
                        affected_functions_str = str(affected_functions)
                except:
                    affected_functions_str = problem.get("affected_functions", "")
                
                # 合并四段式描述为详细描述
                detailed_description = "\n".join([
                    f"需求描述：{str(problem.get('description_requirement', ''))}",
                    f"问题描述：{str(problem.get('description_issue', ''))}",
                    f"问题定位：{str(problem.get('description_location', ''))}",
                    f"影响分析：{str(problem.get('description_impact', ''))}"
                ])
                
                row = [
                    problem.get("chain_name", ""),
                    problem.get("chain_function_description", ""),
                    affected_functions_str,
                    problem.get("problem_reason", ""),
                    problem.get("issue_level", ""),
                    problem.get("issue_name", ""),
                    problem.get("issue_location", ""),
                    detailed_description
                ]
                ws_chain.append(row)
            
            # 设置链路审查sheet的格式
            _format_worksheet(ws_chain)
            log_debug(f"添加链路审查问题: {len(chain_problems)} 个", "download_result")
        
        # 创建功能审查sheet
        if feature_problems:
            sheet_count = sum([bool(function_problems), bool(chain_problems)])
            if sheet_count > 0:
                ws_feature = wb.create_sheet(title="功能审查结果")
            else:
                ws_feature = wb.active
                ws_feature.title = "功能审查结果"
            
            # 添加表头
            ws_feature.append(BUSINESS_COLUMNS)
            
            # 添加数据行
            for problem in feature_problems:
                # 获取实现路径（相关函数列表）
                try:
                    implementation_path = json.loads(problem.get("implementation_path", "[]"))
                    if isinstance(implementation_path, list):
                        implementation_path_str = ", ".join(implementation_path)
                    else:
                        implementation_path_str = str(implementation_path)
                except:
                    implementation_path_str = problem.get("implementation_path", "")
                
                # 合并四段式描述为详细描述
                detailed_description = "\n".join([
                    f"需求描述：{str(problem.get('description_requirement', ''))}",
                    f"问题描述：{str(problem.get('description_issue', ''))}",
                    f"问题定位：{str(problem.get('description_location', ''))}",
                    f"影响分析：{str(problem.get('description_impact', ''))}"
                ])
                
                row = [
                    problem.get("feature_name", ""),  # 功能名称
                    problem.get("feature_description", ""),  # 功能描述
                    implementation_path_str,  # 实现路径（涉及函数）
                    problem.get("problem_reason", ""),
                    problem.get("issue_level", ""),
                    problem.get("issue_name", ""),
                    problem.get("issue_location", ""),
                    detailed_description
                ]
                ws_feature.append(row)
            
            # 设置功能审查sheet的格式
            _format_worksheet(ws_feature)
            log_debug(f"添加功能审查问题: {len(feature_problems)} 个", "download_result")
        
        # 导出为字节流
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # 生成文件名
        filename = f"code_review_result_task_{task_id}.xlsx"
        headers = {"Content-Disposition": f"attachment; filename={filename}"}
        
        return StreamingResponse(
            output, 
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
            headers=headers
        )
        
    except Exception as e:
        log_error(f"下载结果失败: task_id={task_id}", "download_result", e)
        raise HTTPException(status_code=500, detail=f"下载失败: {str(e)}")


@router.post("/export_marked_problems/{task_id}")
async def export_marked_problems(task_id: int, request_body: dict, db: Session = Depends(get_db)):
    """
    导出用户在代码视图中标记的问题为Excel文件
    
    Args:
        task_id: 任务ID
        request_body: 包含已标记问题列表的请求体 {"problems": [...]}
        db: 数据库会话
        
    Returns:
        Excel文件下载响应
    """
    try:
        task_repo = TaskRepository(db)
        
        # 检查任务是否存在
        task = task_repo.get_task(task_id)
        if not task:
            raise HTTPException(status_code=404, detail="任务不存在")
        
        # 获取已标记的问题列表
        problems = request_body.get("problems", [])
        analysis_type = request_body.get('analysis_type', 'function')
        
        if not problems:
            raise HTTPException(status_code=400, detail="没有已标记的问题")
        
        log_debug(f"开始导出已标记问题: task_id={task_id}, 问题数量={len(problems)}", "export_marked_problems")
        
        # 创建Excel工作簿
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "已标记问题"
        
        # 添加表头
        ws.append(FUNCTION_COLUMNS)
        
        # 添加数据行
        for problem in problems:
            # 合并四段式描述
            description = "\n".join([
                f"需求描述：{str(problem.get('description_requirement', ''))}",
                f"问题描述：{str(problem.get('description_issue', ''))}",
                f"问题定位：{str(problem.get('description_location', ''))}",
                f"影响分析：{str(problem.get('description_impact', ''))}"
            ])
            
            row = [
                problem.get("file_name", ""),
                problem.get("source_code", ""),
                problem.get("problem_reason", ""),
                problem.get("issue_level", ""),
                problem.get("issue_name", ""),
                problem.get("issue_location", ""),
                description
            ]
            ws.append(row)
        
        # 设置格式
        _format_worksheet(ws)
        
        log_debug(f"已标记问题格式化完成: {len(problems)} 个", "export_marked_problems")
        
        # 导出为字节流
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # 生成文件名
        filename = f"marked_problems_task_{task_id}.xlsx"
        headers = {"Content-Disposition": f"attachment; filename={filename}"}
        
        log_debug(f"已标记问题导出成功: {filename}", "export_marked_problems")

        # 记录任务已导出过标记问题（目前针对函数审查）
        try:
            if analysis_type == 'function':
                task_repo.set_marked_problems_exported(task_id)
        except Exception:
            pass
        
        return StreamingResponse(
            output, 
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 
            headers=headers
        )
        
    except Exception as e:
        log_error(f"导出已标记问题失败: task_id={task_id}", "export_marked_problems", e)
        raise HTTPException(status_code=500, detail=f"导出失败: {str(e)}")


@router.post("/functionReview/export_marked_problems/{task_id}")
async def export_marked_problems_function_review(task_id: int, request_body: dict, db: Session = Depends(get_db)):
    return await export_marked_problems(task_id=task_id, request_body=request_body, db=db)

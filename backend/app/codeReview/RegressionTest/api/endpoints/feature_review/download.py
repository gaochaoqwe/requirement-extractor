"""
下载功能点审查结果API
"""
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import io
import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime

from ....db.deps import get_db
from ....db.repository.feature_repo import FeatureRepository, FeatureReviewProblemRepository

router = APIRouter()


@router.get("/download/{task_id}")
async def download_feature_review_result(task_id: int, db: Session = Depends(get_db)):
    """
    下载功能点审查结果为Excel文件
    """
    try:
        feature_repo = FeatureRepository(db)
        problem_repo = FeatureReviewProblemRepository(db)
        
        # 获取所有功能点
        features = feature_repo.get_features_by_task(task_id)
        
        if not features:
            raise HTTPException(status_code=404, detail="没有找到功能点审查结果")
        
        # 创建Excel工作簿
        wb = Workbook()
        wb.remove(wb.active)  # 删除默认sheet
        
        # 创建样式
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Sheet 1: 功能点概览
        ws_overview = wb.create_sheet("功能点概览")
        overview_headers = ["功能点名称", "功能描述", "实现说明", "状态", "发现问题数"]
        ws_overview.append(overview_headers)
        
        # 设置表头样式
        for cell in ws_overview[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # 写入功能点数据
        for feature in features:
            problems = problem_repo.get_problems_by_feature(feature.id)
            ws_overview.append([
                feature.feature_name,
                feature.feature_description,
                feature.implementation_description or "暂无",
                feature.status,
                len(problems)
            ])
        
        # 调整列宽
        ws_overview.column_dimensions['A'].width = 20
        ws_overview.column_dimensions['B'].width = 40
        ws_overview.column_dimensions['C'].width = 50
        ws_overview.column_dimensions['D'].width = 15
        ws_overview.column_dimensions['E'].width = 15
        
        # Sheet 2: 功能点审查问题
        ws_problems = wb.create_sheet("功能点审查问题")
        problem_headers = [
            "功能点名称", "问题名称", "严重程度", "问题位置", 
            "文件名", "函数名", "问题原因", "问题描述", 
            "影响", "改进建议"
        ]
        ws_problems.append(problem_headers)
        
        # 设置表头样式
        for cell in ws_problems[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # 写入问题数据
        for feature in features:
            problems = problem_repo.get_problems_by_feature(feature.id)
            for problem in problems:
                ws_problems.append([
                    feature.feature_name,
                    problem.issue_name,
                    problem.issue_level,
                    problem.issue_location,
                    problem.file_name,
                    problem.function_name,
                    problem.problem_reason,
                    problem.description_issue,
                    problem.description_impact,
                    problem.suggestion
                ])
        
        # 调整列宽
        for col in ['A', 'B', 'C', 'D', 'E', 'F']:
            ws_problems.column_dimensions[col].width = 20
        for col in ['G', 'H', 'I', 'J']:
            ws_problems.column_dimensions[col].width = 40
        
        # Sheet 3: 功能点实现详情
        ws_impl = wb.create_sheet("功能点实现详情")
        impl_headers = ["功能点名称", "相关文件", "相关函数", "实现位置"]
        ws_impl.append(impl_headers)
        
        # 设置表头样式
        for cell in ws_impl[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # 写入实现详情
        for feature in features:
            try:
                related_files = json.loads(feature.related_files) if feature.related_files else []
                related_functions = json.loads(feature.related_functions) if feature.related_functions else []
                code_locations = json.loads(feature.code_locations) if feature.code_locations else []
                
                ws_impl.append([
                    feature.feature_name,
                    ", ".join(related_files),
                    ", ".join(related_functions),
                    "; ".join([f"{loc.get('file', '')}:{loc.get('function', '')}" for loc in code_locations])
                ])
            except Exception as e:
                ws_impl.append([
                    feature.feature_name,
                    "解析失败",
                    "解析失败",
                    str(e)
                ])
        
        # 调整列宽
        ws_impl.column_dimensions['A'].width = 20
        ws_impl.column_dimensions['B'].width = 50
        ws_impl.column_dimensions['C'].width = 50
        ws_impl.column_dimensions['D'].width = 60
        
        # 保存到内存
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        # 生成文件名
        filename = f"feature_review_task{task_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename={filename}"
            }
        )
        
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"下载失败: {str(e)}")


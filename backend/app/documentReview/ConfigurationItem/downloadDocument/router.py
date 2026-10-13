from fastapi import APIRouter, Query, Request, Response
from fastapi.responses import FileResponse, JSONResponse
import os
import sys
import sqlite3
import tempfile
import logging
import stat
import traceback
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment

# 配置日志
logger = logging.getLogger(__name__)
logger.setLevel(logging.DEBUG)

router = APIRouter()
from pathlib import Path

# 修正 BASE_DIR 计算 - 指向新的数据库位置
BASE_DIR = Path(__file__).resolve().parent.parent
DB_PATH = str(BASE_DIR / 'db' / 'documentReview.db')

# 确保临时目录存在 - 尝试多个位置
try:
    # 首选项目根目录
    TEMP_DIR = os.path.join(str(BASE_DIR.parent.parent.parent.parent), 'temp')
    os.makedirs(TEMP_DIR, exist_ok=True)
    
    # 测试写入权限
    test_file = os.path.join(TEMP_DIR, 'test_write.txt')
    with open(test_file, 'w') as f:
        f.write('test')
    os.remove(test_file)
    # print(f"[重要] 配置项模块 - 临时目录可写: {TEMP_DIR}")
except Exception as e:
    # 如果无法写入，尝试系统临时目录
    TEMP_DIR = tempfile.gettempdir()
    import sys
    sys.stderr.write(f"[警告] 配置项模块 - 项目目录无法写入，切换到系统临时目录: {TEMP_DIR}\n")
    sys.stderr.write(f"[警告] 原因: {str(e)}\n")

# 创建唯一的应用目录
APP_TEMP_DIR = os.path.join(TEMP_DIR, 'requirement_extractor')
os.makedirs(APP_TEMP_DIR, exist_ok=True)

# 打印关键路径用于调试
# print(f"[CRITICAL] 配置项下载模块 - 当前运行目录: {os.getcwd()}")
# print(f"[CRITICAL] 配置项下载模块 - Python路径: {sys.executable}")
# print(f"[CRITICAL] 配置项下载模块 - 数据库路径: {DB_PATH}")
# print(f"[CRITICAL] 配置项下载模块 - 临时目录路径: {APP_TEMP_DIR}")

# 同时支持带斜杠和不带斜杠的URL
@router.get("/api/download")
@router.get("/api/download/")
async def download(request: Request, session_id: int = Query(...)):
    # 详细记录请求信息
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")
    client_host = request.client.host if request.client else "unknown"
    log_prefix = f"[CONFIG-DOWNLOAD-{timestamp}-{client_host}-{session_id}]"
    
    sys.stderr.write(f"{log_prefix} 收到配置项下载请求 - URL: {request.url}\n")
    sys.stderr.write(f"{log_prefix} 客户端: {client_host}, headers: {dict(request.headers)}\n")
    sys.stderr.write(f"{log_prefix} 使用数据库: {DB_PATH}\n")
    logger.info(f"[下载] 收到配置项审查下载请求 - session_id: {session_id}")
    
    # 检查数据库文件是否存在
    if not os.path.exists(DB_PATH):
        error_msg = f"数据库文件不存在: {DB_PATH}"
        sys.stderr.write(f"{log_prefix} 错误: {error_msg}\n")
        return JSONResponse(status_code=500, content={"error": error_msg})
    
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    
    # 查询时包含问题等级字段
    cur.execute("SELECT * FROM review_results WHERE session_id=?", (session_id,))
    rows = cur.fetchall()
    columns = [d[0] for d in cur.description]
    conn.close()
    
    if not rows:
        return {"error": "无审查结果"}
    
    # 按 doc_type 拆分
    data_taskbook = [r for r in rows if r[columns.index('doc_type')] == 'taskBook']
    data_requirementspec = [r for r in rows if r[columns.index('doc_type')] == 'requirementSpec']
    
    # 创建Excel
    wb = Workbook()
    for sheet_name, data in [("研制任务书", data_taskbook), ("需求规格说明", data_requirementspec)]:
        ws = wb.create_sheet(title=sheet_name)
        # 调整表头顺序：将问题等级列放到问题名称前面
        ws.append([
            "原始需求内容", "违反规则（规则名）", "问题等级", "问题名称", "问题来源（文档名称+章节）", "问题描述"
        ])
        for r in data:
            ws.append([
                r[columns.index('requirement_content')],
                r[columns.index('violation_rules')],
                r[columns.index('issue_level')],  # 问题等级列移到第三列
                r[columns.index('issue_name')],
                r[columns.index('issue_location')],
                # 合并分析四段
                f"需求描述：{r[columns.index('description_requirement')] or ''}\n问题描述：{r[columns.index('description_issue')] or ''}\n问题定位：{r[columns.index('description_location')] or ''}\n影响分析：{r[columns.index('description_impact')] or ''}"
            ])
    
    # 删除默认sheet
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']
        
    # 生成带时间戳的唯一文件名
    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    file_name = f"config_report_{session_id}_{timestamp_str}.xlsx"
    tmp_file_path = os.path.join(APP_TEMP_DIR, file_name)
    
    sys.stderr.write(f"{log_prefix} 将生成文件: {tmp_file_path}\n")
    
    try:
        # 设置列宽和自动换行
        wrap_alignment = Alignment(wrapText=True, vertical="top")
        for worksheet in wb.worksheets:
            for col_cells in worksheet.columns:
                max_len = 0
                for cell in col_cells:
                    cell.alignment = wrap_alignment
                    if cell.value:
                        l = len(str(cell.value))
                        if l > max_len:
                            max_len = l
                adjusted_width = min(max(max_len + 2, 20), 60)
                column_letter = col_cells[0].column_letter
                worksheet.column_dimensions[column_letter].width = adjusted_width
        
        # 保存Excel文件
        sys.stderr.write(f"{log_prefix} [重要] 正在保存配置项Excel文件到: {tmp_file_path}\n")
        wb.save(tmp_file_path)
        
        # 确认文件存在并可读
        if os.path.exists(tmp_file_path):
            file_size = os.path.getsize(tmp_file_path)
            sys.stderr.write(f"{log_prefix} [重要] Excel文件保存成功，文件大小: {file_size} 字节\n")
            
            try:
                # 尝试将文件设置为只读，防止被意外删除
                os.chmod(tmp_file_path, stat.S_IREAD | stat.S_IRGRP | stat.S_IROTH)
                sys.stderr.write(f"{log_prefix} 文件已设置为只读保护\n")
            except Exception as chmod_err:
                sys.stderr.write(f"{log_prefix} 设置文件权限失败: {str(chmod_err)}\n")
            
            # 创建一个标记文件作为调试参考
            marker_file = os.path.join(APP_TEMP_DIR, f"config_marker_{session_id}_{timestamp_str}.txt")
            with open(marker_file, 'w', encoding='utf-8') as f:
                f.write(f"Config Excel file created at: {timestamp_str}\nPath: {tmp_file_path}\nSize: {file_size} bytes")
            
            sys.stderr.write(f"{log_prefix} [重要] 正在返回文件: {tmp_file_path}\n")
            
            # 使用FileResponse返回文件
            return FileResponse(
                path=tmp_file_path, 
                filename=f"配置项审查结果_{session_id}.xlsx",
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        else:
            error_msg = f"错误: 文件保存失败，文件不存在: {tmp_file_path}"
            sys.stderr.write(f"{log_prefix} {error_msg}\n")
            return JSONResponse(status_code=500, content={"error": error_msg})
    except Exception as e:
        error_trace = traceback.format_exc()
        error_msg = f"[下载] 配置项模块 - 生成或返回Excel文件时出错: {str(e)}"
        sys.stderr.write(f"{log_prefix} {error_msg}\n")
        sys.stderr.write(f"{log_prefix} \n{error_trace}\n")
        logger.error(error_msg)
        logger.error(error_trace)
        
        # 写入错误日志文件便于调试
        try:
            error_log_file = os.path.join(APP_TEMP_DIR, f"config_error_{session_id}_{timestamp_str}.log")
            with open(error_log_file, 'w', encoding='utf-8') as f:
                f.write(f"Error Time: {timestamp_str}\n")
                f.write(f"Error Message: {error_msg}\n")
                f.write(f"Stack Trace:\n{error_trace}\n")
            sys.stderr.write(f"{log_prefix} 错误日志已写入: {error_log_file}\n")
        except Exception as log_err:
            sys.stderr.write(f"{log_prefix} 无法写入错误日志: {str(log_err)}\n")
        
        return JSONResponse(
            status_code=500, 
            content={
                "error": f"配置项下载失败: {str(e)}", 
                "details": error_trace,
                "timestamp": timestamp_str
            }
        )


# ===== 新增：下载批注版研制任务书 =====
@router.get("/api/download/annotated/taskbook")
async def download_annotated_taskbook(session_id: int = Query(...)):
    """下载批注版研制任务书"""
    try:
        # 从数据库中获取批注文件路径
        conn = sqlite3.connect(DB_PATH)
        cur = conn.cursor()
        cur.execute("SELECT annotated_taskbook_path FROM review_sessions WHERE id = ?", (session_id,))
        result = cur.fetchone()
        conn.close()
        
        if result and result[0] and os.path.exists(result[0]):
            annotated_path = result[0]
            logger.info(f"从数据库找到研制任务书批注文件: {annotated_path}")
            return FileResponse(
                path=annotated_path,
                filename=f"研制任务书_批注版_{session_id}.docx",
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        
        # 备选方案：在临时目录查找对应session_id的批注文件
        temp_dir = os.path.join(os.path.dirname(BASE_DIR), 'temp', 'annotated_documents')
        if os.path.exists(temp_dir):
            pattern = f"研制任务书_批注版_session_{session_id}.docx"
            import glob
            candidates = glob.glob(os.path.join(temp_dir, pattern))
            
            if candidates:
                file_path = candidates[0]
                logger.info(f"在临时目录找到研制任务书批注文件: {file_path}")
                return FileResponse(
                    path=file_path,
                    filename=f"研制任务书_批注版_{session_id}.docx",
                    media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
        
        logger.error(f"未找到研制任务书批注版文档 - session_id: {session_id}")
        return JSONResponse(status_code=404, content={"error": "未找到研制任务书批注版文档"})
        
    except Exception as e:
        logger.error(f"下载研制任务书批注版失败: {e}", exc_info=True)
        return JSONResponse(status_code=500, content={"error": str(e)})


# ===== 新增：下载批注版需求规格说明 =====
@router.get("/api/download/annotated/requirementspec")
async def download_annotated_requirementspec(session_id: int = Query(...)):
    """下载批注版需求规格说明"""
    try:
        # 从数据库中获取批注文件路径
        conn = sqlite3.connect(DB_PATH)
        cur = conn.cursor()
        cur.execute("SELECT annotated_requirementspec_path FROM review_sessions WHERE id = ?", (session_id,))
        result = cur.fetchone()
        conn.close()
        
        if result and result[0] and os.path.exists(result[0]):
            annotated_path = result[0]
            logger.info(f"从数据库找到需求规格说明批注文件: {annotated_path}")
            return FileResponse(
                path=annotated_path,
                filename=f"需求规格说明_批注版_{session_id}.docx",
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        
        # 备选方案：在临时目录查找对应session_id的批注文件
        temp_dir = os.path.join(os.path.dirname(BASE_DIR), 'temp', 'annotated_documents')
        if os.path.exists(temp_dir):
            pattern = f"需求规格说明_批注版_session_{session_id}.docx"
            import glob
            candidates = glob.glob(os.path.join(temp_dir, pattern))
            
            if candidates:
                file_path = candidates[0]
                logger.info(f"在临时目录找到需求规格说明批注文件: {file_path}")
                return FileResponse(
                    path=file_path,
                    filename=f"需求规格说明_批注版_{session_id}.docx",
                    media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
        
        logger.error(f"未找到需求规格说明批注版文档 - session_id: {session_id}")
        return JSONResponse(status_code=404, content={"error": "未找到需求规格说明批注版文档"})
        
    except Exception as e:
        logger.error(f"下载需求规格说明批注版失败: {e}", exc_info=True)
        return JSONResponse(status_code=500, content={"error": str(e)})


# ===== 新增：下载所有批注版文档（打包） =====
@router.get("/api/download/annotated/all")
async def download_all_annotated(session_id: int = Query(...)):
    """下载所有批注版文档（打包为ZIP）"""
    try:
        import zipfile
        import tempfile
        
        # 从数据库中获取批注文件路径
        conn = sqlite3.connect(DB_PATH)
        cur = conn.cursor()
        cur.execute("SELECT annotated_taskbook_path, annotated_requirementspec_path FROM review_sessions WHERE id = ?", (session_id,))
        result = cur.fetchone()
        conn.close()
        
        if not result:
            return JSONResponse(status_code=404, content={"error": "未找到审查会话"})
        
        taskbook_path, requirementspec_path = result
        files_to_zip = []
        
        # 检查文件是否存在
        if taskbook_path and os.path.exists(taskbook_path):
            files_to_zip.append((taskbook_path, f"研制任务书_批注版_{session_id}.docx"))
        if requirementspec_path and os.path.exists(requirementspec_path):
            files_to_zip.append((requirementspec_path, f"需求规格说明_批注版_{session_id}.docx"))
        
        if not files_to_zip:
            return JSONResponse(status_code=404, content={"error": "未找到任何批注版文档"})
        
        # 创建临时ZIP文件
        temp_zip_path = os.path.join(APP_TEMP_DIR, f"批注版文档_{session_id}.zip")
        
        with zipfile.ZipFile(temp_zip_path, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            for file_path, archive_name in files_to_zip:
                zip_file.write(file_path, archive_name)
        
        logger.info(f"创建批注版文档压缩包: {temp_zip_path}, 包含 {len(files_to_zip)} 个文件")
        
        return FileResponse(
            path=temp_zip_path,
            filename=f"批注版文档_{session_id}.zip",
            media_type="application/zip"
        )
        
    except Exception as e:
        logger.error(f"下载所有批注版文档失败: {e}", exc_info=True)
        return JSONResponse(status_code=500, content={"error": str(e)})

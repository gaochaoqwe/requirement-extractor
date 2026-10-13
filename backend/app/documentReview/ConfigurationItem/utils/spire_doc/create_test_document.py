#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
创建测试用Word文档
"""

from docx import Document
from docx.shared import Inches
import os

def create_test_document(output_path="test_document.docx"):
    """创建一个包含典型内容的测试Word文档"""
    
    # 创建新文档
    doc = Document()
    
    # 添加标题
    doc.add_heading('软件需求规格说明书', 0)
    
    # 添加目录（模拟）
    doc.add_heading('目录', level=1)
    doc.add_paragraph('1. 引言')
    doc.add_paragraph('   1.1 目的')
    doc.add_paragraph('   1.2 范围')
    doc.add_paragraph('2. 系统概述')
    doc.add_paragraph('   2.1 系统功能')
    doc.add_paragraph('   2.2 系统性能')
    doc.add_paragraph('3. 功能需求')
    doc.add_paragraph('   3.1 用户管理')
    doc.add_paragraph('   3.2 数据处理')
    doc.add_paragraph('')
    
    # 添加正文内容
    doc.add_heading('引言', level=1)
    doc.add_paragraph('这是引言部分的内容，描述了文档的目的和范围。')
    
    doc.add_heading('目的', level=2) 
    doc.add_paragraph('本文档的目的是详细说明软件系统的需求规格。')
    
    doc.add_heading('范围', level=2)
    doc.add_paragraph('本文档涵盖了软件系统的所有功能和非功能需求。')
    
    doc.add_heading('系统概述', level=1)
    doc.add_paragraph('系统概述部分描述了整个系统的架构和主要组件。')
    
    doc.add_heading('系统功能', level=2)
    doc.add_paragraph('系统提供以下主要功能：')
    doc.add_paragraph('• 用户注册和登录')
    doc.add_paragraph('• 数据录入和处理') 
    doc.add_paragraph('• 报表生成和导出')
    
    doc.add_heading('系统性能', level=2)
    doc.add_paragraph('系统应满足以下性能要求：')
    doc.add_paragraph('• 响应时间不超过2秒')
    doc.add_paragraph('• 并发用户数不少于100人')
    
    doc.add_heading('功能需求', level=1)
    doc.add_paragraph('功能需求详细描述了系统应该具备的各项功能。')
    
    doc.add_heading('用户管理', level=2)
    doc.add_paragraph('用户管理模块负责用户的注册、登录、权限管理等功能。')
    
    doc.add_heading('数据处理', level=2)
    doc.add_paragraph('数据处理模块负责数据的录入、验证、存储和查询。')
    
    # 保存文档
    doc.save(output_path)
    print(f"✅ 测试文档已创建: {output_path}")
    
    return output_path

def create_simple_test_document(output_path="simple_test.docx"):
    """创建一个简单的测试文档"""
    
    doc = Document()
    
    # 添加简单内容
    doc.add_heading('测试文档', 0)
    doc.add_paragraph('这是一个用于测试文档标准化功能的简单文档。')
    
    doc.add_heading('概述', level=1)
    doc.add_paragraph('概述部分的内容。')
    
    doc.add_heading('功能说明', level=1)
    doc.add_paragraph('功能说明的内容。')
    
    doc.add_heading('系统要求', level=1)
    doc.add_paragraph('系统要求的内容。')
    
    doc.save(output_path)
    print(f"✅ 简单测试文档已创建: {output_path}")
    
    return output_path

if __name__ == "__main__":
    # 创建测试文档
    test_doc1 = create_test_document("test_requirement_spec.docx")
    test_doc2 = create_simple_test_document("simple_test.docx")
    
    print(f"\n📁 已创建测试文档:")
    print(f"- {test_doc1}")
    print(f"- {test_doc2}")
    print(f"\n💡 可以使用这些文档测试标准化功能") 
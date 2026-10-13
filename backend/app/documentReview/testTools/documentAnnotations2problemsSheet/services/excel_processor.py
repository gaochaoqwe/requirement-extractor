"""
测试用例总结Excel处理器
"""
import openpyxl
from pathlib import Path
from typing import Dict, List, Any, Tuple, Optional
import uuid
from datetime import datetime
import json
from collections import defaultdict

from ..config.config import REQUIRED_SHEET_NAME, REQUIRED_COLUMNS
from ..models.models import TestCaseSummarySession, TestItem
from ..database.database import get_db_session
from ..utils.logger import get_session_logger

class ExcelProcessor:
    def __init__(self, session_id: str):
        self.session_id = session_id
        self.logger = get_session_logger(session_id)
        self.required_sheet = REQUIRED_SHEET_NAME
        self.required_columns = REQUIRED_COLUMNS
        
    def process_uploaded_file(self, file_path: str, original_filename: str) -> Tuple[str, Dict[str, Any]]:
        """
        处理上传的测试用例Excel文件
        """
        try:
            self.logger.info(f"开始处理测试用例Excel文件: {original_filename}")
            
            # 验证Excel文件
            validation_result = self._validate_excel_file(file_path)
            if not validation_result["success"]:
                self.logger.error(f"Excel文件验证失败: {validation_result['error']}")
                return self.session_id, validation_result
            
            # 读取Excel数据
            read_result = self._read_test_case_data(file_path)
            if not read_result["success"]:
                self.logger.error(f"Excel数据读取失败: {read_result['error']}")
                return self.session_id, read_result
            
            # 保存到数据库
            save_result = self._save_to_database(
                self.session_id, 
                file_path, 
                original_filename, 
                read_result["test_items"]
            )
            
            if save_result["success"]:
                self.logger.info(f"测试用例文件处理完成 - 会话ID: {self.session_id}, 测试项数量: {len(read_result['test_items'])}")
            
            return self.session_id, save_result
            
        except Exception as e:
            error_msg = f"处理测试用例Excel文件时出错: {str(e)}"
            self.logger.error(error_msg)
            return self.session_id, {"success": False, "error": error_msg}
    
    def _validate_excel_file(self, file_path: str) -> Dict[str, Any]:
        """验证Excel文件格式"""
        try:
            self.logger.info("开始验证Excel文件")
            
            # 尝试打开Excel文件
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            
            # 检查是否存在所需的工作表
            if self.required_sheet not in workbook.sheetnames:
                available_sheets = ", ".join(workbook.sheetnames)
                return {
                    "success": False, 
                    "error": f"未找到工作表 '{self.required_sheet}'。可用工作表: {available_sheets}"
                }
            
            worksheet = workbook[self.required_sheet]
            
            # 检查文件是否为空
            if worksheet.max_row < 2:
                return {"success": False, "error": "Excel文件为空或只有表头"}
            
            # 检查必要的列是否存在
            max_col_index = max(self.required_columns.values())
            if worksheet.max_column <= max_col_index:
                return {
                    "success": False, 
                    "error": f"Excel文件列数不足，需要至少{max_col_index + 1}列"
                }
            
            workbook.close()
            
            self.logger.info("Excel文件验证成功")
            return {"success": True}
            
        except Exception as e:
            return {"success": False, "error": f"验证Excel文件时出错: {str(e)}"}
    
    def _read_test_case_data(self, file_path: str) -> Dict[str, Any]:
        """读取测试用例数据并按测试项分组"""
        try:
            self.logger.info("开始读取测试用例数据")
            
            # 使用openpyxl读取Excel文件
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            worksheet = workbook[self.required_sheet]
            
            # 按测试项名称分组存储测试用例
            test_items_dict = defaultdict(list)
            
            # 逐行处理数据（跳过表头行）
            for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    # 检查行是否有足够的列
                    max_col_index = max(self.required_columns.values())
                    if len(row) <= max_col_index:
                        continue
                    
                    # 提取测试用例信息
                    test_item_name = str(row[self.required_columns["测试项名称"]]) if row[self.required_columns["测试项名称"]] is not None else ""
                    input_condition = str(row[self.required_columns["输入条件"]]) if row[self.required_columns["输入条件"]] is not None else ""
                    test_result = str(row[self.required_columns["测试结果"]]) if row[self.required_columns["测试结果"]] is not None else ""
                    
                    # 跳过空的测试项名称
                    if not test_item_name or test_item_name.strip() == "" or test_item_name == "None":
                        continue
                    
                    # 将测试用例添加到对应的测试项
                    test_items_dict[test_item_name].append({
                        "row_index": row_idx,
                        "input_condition": input_condition,
                        "test_result": test_result
                    })
                    
                except Exception as e:
                    self.logger.warning(f"处理第{row_idx}行数据时出错: {str(e)}")
                    continue
            
            workbook.close()
            
            # 转换为最终格式
            test_items = []
            for test_item_name, test_cases in test_items_dict.items():
                test_items.append({
                    "test_item_name": test_item_name,
                    "test_cases": test_cases,
                    "test_cases_count": len(test_cases)
                })
            
            self.logger.info(f"成功读取测试用例数据，共{len(test_items)}个测试项，{sum(len(item['test_cases']) for item in test_items)}个测试用例")
            
            return {
                "success": True,
                "test_items": test_items,
                "total_test_items": len(test_items)
            }
            
        except Exception as e:
            error_msg = f"读取测试用例数据时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg}
    
    def _save_to_database(self, session_id: str, file_path: str, original_filename: str, test_items: List[Dict]) -> Dict[str, Any]:
        """保存数据到数据库"""
        try:
            self.logger.info(f"开始保存数据到数据库 - 会话ID: {session_id}")
            
            db = get_db_session()
            
            try:
                # 创建会话记录
                session = TestCaseSummarySession(
                    session_id=session_id,
                    original_filename=original_filename,
                    file_path=file_path,
                    status="uploaded",
                    total_test_items=len(test_items),
                    processed_test_items=0
                )
                db.add(session)
                db.flush()  # 获取ID
                
                # 批量创建测试项记录
                test_item_records = []
                for test_item_data in test_items:
                    test_item = TestItem(
                        session_id=session_id,
                        test_item_name=test_item_data["test_item_name"],
                        test_cases_data=json.dumps(test_item_data["test_cases"], ensure_ascii=False),
                        generation_status="pending"
                    )
                    test_item_records.append(test_item)
                
                db.add_all(test_item_records)
                db.commit()
                
                self.logger.info(f"数据保存成功 - 会话: {session_id}, 测试项: {len(test_item_records)}个")
                
                return {
                    "success": True,
                    "session_id": session_id,
                    "total_test_items": len(test_items),
                    "message": f"成功处理{len(test_items)}个测试项"
                }
                
            except Exception as e:
                db.rollback()
                raise e
            finally:
                db.close()
                
        except Exception as e:
            error_msg = f"保存数据到数据库时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg}
    
    def generate_output_excel(self, output_path: str) -> Dict[str, Any]:
        """
        生成输出Excel文件
        """
        try:
            self.logger.info(f"开始生成输出Excel文件: {output_path}")
            
            db = get_db_session()
            
            try:
                # 获取会话信息
                session = db.query(TestCaseSummarySession).filter(
                    TestCaseSummarySession.session_id == self.session_id
                ).first()
                
                if not session:
                    return {"success": False, "error": "会话不存在"}
                
                # 获取所有测试项
                test_items = db.query(TestItem).filter(
                    TestItem.session_id == self.session_id
                ).order_by(TestItem.id).all()
                
                if not test_items:
                    return {"success": False, "error": "没有找到测试项数据"}
                
                # 创建新的Excel文件
                workbook = openpyxl.Workbook()
                worksheet = workbook.active
                worksheet.title = "测试项总结"
                
                # 写入表头
                worksheet.cell(row=1, column=1, value="测试项名称")
                worksheet.cell(row=1, column=2, value="AI总结结果")
                
                # 写入数据
                updated_count = 0
                for idx, test_item in enumerate(test_items, start=2):
                    worksheet.cell(row=idx, column=1, value=test_item.test_item_name)
                    worksheet.cell(row=idx, column=2, value=test_item.ai_summary or "")
                    updated_count += 1
                
                # 保存文件
                workbook.save(output_path)
                workbook.close()
                
                self.logger.info(f"输出Excel文件生成成功: {output_path}, 更新了{updated_count}个测试项")
                
                return {
                    "success": True,
                    "output_path": output_path,
                    "total_items": len(test_items),
                    "updated_items": updated_count
                }
                
            finally:
                db.close()
                
        except Exception as e:
            error_msg = f"生成输出Excel文件时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg} 
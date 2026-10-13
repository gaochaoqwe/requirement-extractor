import os
import logging
import asyncio
from typing import List, Dict, Optional

from spire.doc import Document
import openpyxl
from openpyxl.styles import Alignment, PatternFill, Border, Side

from ..config.config import ANNOTATION_CONFIG, OUTPUT_COLUMNS, OUTPUT_DIR
from ..models.models import DocumentAnnotation, DocumentAnnotationSession
from ..database.database import get_db_session
from .ai_generator import AIGenerator


logger = logging.getLogger(__name__)


class DocumentProcessor:
    """文档批注处理器"""

    def __init__(self, session_id: Optional[str] = None) -> None:
        self.annotation_config = ANNOTATION_CONFIG
        self.session_id: Optional[str] = session_id
        self.ai_generator: Optional[AIGenerator] = AIGenerator(session_id) if session_id else None

    # ======================== 批注解析 ========================
    def process_document_annotations(self, doc_path: str, file_name: str) -> List[Dict]:
        """从Word文档中解析批注，抽取位置/说明/原文等基础信息。

        返回的每条记录结构：
        {
            "问题来源": "《文档名》3.2.1",
            "原文内容": "……",
            "问题描述": "说明文本",
            "位置": "3.2.1"
        }
        """
        try:
            logger.info(f"开始处理文档: {doc_path}")

            document = Document()
            document.LoadFromFile(doc_path)
            logger.info("文档加载成功，开始解析批注")

            annotations: List[Dict] = []

            try:
                comments_count = document.Comments.Count
                logger.info(f"文档中共有 {comments_count} 个批注")
            except Exception as e:  # pragma: no cover - spire 异常包装
                logger.error(f"无法获取批注数量: {e}")
                return []

            if comments_count == 0:
                logger.warning("文档中没有找到批注")
                return []

            for i in range(comments_count):
                try:
                    comment = document.Comments[i]
                    logger.debug(f"处理第 {i + 1} 个批注")

                    find_flag_weizhi = False
                    find_flag_wenti_miaoshu = False
                    weizhi = ""
                    wenti_miaoshu = ""

                    # 统计批注正文段落
                    try:
                        if hasattr(comment.Body, "Paragraphs"):
                            paras = comment.Body.Paragraphs
                            paragraphs_count = getattr(paras, "Count", len(paras))
                        else:
                            logger.warning(f"批注 {i + 1} 没有 Body.Paragraphs 属性")
                            continue
                    except Exception as e:
                        logger.warning(f"无法获取批注 {i + 1} 的段落信息: {e}")
                        continue

                    for j in range(paragraphs_count):
                        try:
                            paragraph = comment.Body.Paragraphs[j]
                            text = getattr(paragraph, "Text", "")
                            text = (text or "").strip()
                            if not text:
                                continue

                            logger.debug(f"批注 {i + 1} 段落 {j + 1} 内容: {text[:50]}…")

                            # 位置
                            if any(text.startswith(prefix) for prefix in self.annotation_config["位置标识"]):
                                weizhi = self._clean_annotation_text(text, self.annotation_config["位置标识"])
                                find_flag_weizhi = True
                                logger.debug(f"找到位置信息: {weizhi}")
                            # 说明
                            elif any(text.startswith(prefix) for prefix in self.annotation_config["说明标识"]):
                                wenti_miaoshu = self._clean_annotation_text(text, self.annotation_config["说明标识"])
                                find_flag_wenti_miaoshu = True
                                logger.debug(f"找到问题说明: {wenti_miaoshu}")
                        except Exception as e:
                            logger.warning(f"处理批注 {i + 1} 的段落 {j + 1} 时出错: {e}")
                            continue

                    # 只要位置或说明有一个存在就认为是有效批注
                    if not (find_flag_weizhi or find_flag_wenti_miaoshu):
                        logger.debug(f"批注 {i + 1} 未找到位置或说明标识，跳过")
                        continue

                    # 原文内容
                    original_text = ""
                    try:
                        if hasattr(comment, "OwnerParagraph") and comment.OwnerParagraph:
                            para = comment.OwnerParagraph
                            original_text = getattr(para, "Text", "") or ""
                    except Exception as e:
                        logger.warning(f"无法获取批注 {i + 1} 的原文: {e}")

                    if not find_flag_weizhi:
                        weizhi = "未指定"
                    if not find_flag_wenti_miaoshu:
                        wenti_miaoshu = "未提供详细说明"

                    wenti_laiyuan = f"{file_name}{weizhi}"
                    annotations.append(
                        {
                            "问题来源": wenti_laiyuan,
                            "原文内容": original_text,
                            "问题描述": wenti_miaoshu,
                            "位置": weizhi,
                        }
                    )
                    logger.info(f"成功解析批注 {i + 1}")
                except Exception as e:
                    logger.error(f"处理批注 {i + 1} 时出错: {e}")
                    continue

            logger.info(f"从文档 {file_name} 中提取到 {len(annotations)} 个有效批注")
            return annotations
        except Exception as e:
            logger.error(f"处理文档批注时出错: {e}")
            if "invalid start byte" in str(e):
                raise Exception("文档编码格式不支持，请确保文档是标准的.docx格式")
            if "LoadFromFile" in str(e):
                raise Exception("无法加载文档，请检查文件是否损坏或格式是否正确")
            raise

    def _clean_annotation_text(self, text: str, prefixes: List[str]) -> str:
        """清理批注文本，移除前缀"""
        for prefix in prefixes:
            if text.startswith(prefix):
                return text.replace(prefix, "").replace("。", "").strip()
        return text.strip()

    def _clean_filename(self, filename: str) -> str:
        """清理文件名"""
        # 去除扩展名
        name = filename.replace('.docx', '')
        
        # 去除数字前缀
        for i in range(1, 6):
            name = name.replace(f'{i}.', '')
        
        # 处理特定文档类型
        if '任务书' in name:
            name = name.split('任务书')[0] + '任务书'
        elif '规格说明' in name:
            name = name.split('规格说明')[0] + '规格说明'
        elif '设计说明' in name:
            name = name.split('设计说明')[0] + '设计说明'
        elif '接口说明' in name:
            name = name.split('接口说明')[0] + '接口说明'
        elif '报告' in name:
            name = name.split('报告')[0] + '报告'
        
        # 添加书名号
        return f'《{name}》'

    def analyze_problem_info(self, wenti_mingcheng: str, wenti_laiyuan: str, wenti_miaoshu: str) -> Dict:
        """
        分析问题信息，生成测试活动、问题类型等
        基于原有的wentidan_fexi函数逻辑
        """
        # 问题等级分析
        if any(keyword in wenti_mingcheng for keyword in ['越界', '章节缺失', '失效']):
            wentidengji = '重要问题'
        else:
            wentidengji = '一般问题'
        
        # 问题数目分析
        if '、' in wenti_laiyuan:
            comma_count = wenti_laiyuan.count('、') + 1
            wenti_shumu = str(comma_count)
        else:
            wenti_shumu = '1'
        
        # 根据问题来源判断问题类型和测试活动
        if any(ext in wenti_laiyuan for ext in ['.c', '.h', '.cpp']) or '源代码' in wenti_laiyuan:
            wentileixing = '程序问题'
            if '编码规范' in wenti_mingcheng:
                ceshi_huodong = '静态分析'
            else:
                ceshi_huodong = '代码审查'
            ceping_danwei_yijian = f"请开发人员针对{wenti_mingcheng}问题对代码进行修改，保证代码实现正确。"
            
        elif '文档' in wenti_laiyuan:
            ceshi_huodong = '文档审查'
            wentileixing = '软件需求问题'
            ceping_danwei_yijian = '请开发人员对文档问题进行修改,确保文档描述的正确且测试人员能够理解测试需求。'
            
        elif '任务书' in wenti_laiyuan:
            ceshi_huodong = '文档审查'
            wentileixing = '用户需求问题'
            ceping_danwei_yijian = '请开发人员对文档问题进行修改,确保文档描述的正确且测试人员能够理解测试需求。'
            
        elif any(keyword in wenti_laiyuan for keyword in ['规格说明', '设计说明', '影响域分析报告', '接口']):
            ceshi_huodong = '文档审查'
            wentileixing = '软件需求问题'
            ceping_danwei_yijian = '请开发人员对文档问题进行修改,确保文档描述的正确且测试人员能够理解测试需求。'
            
        else:
            # 默认值
            ceshi_huodong = '文档审查'
            wentileixing = '软件需求问题'
            ceping_danwei_yijian = '请开发人员对文档问题进行修改,确保文档描述的正确且测试人员能够理解测试需求。'
        
        # 生成问题描述（四段式）
        wenti_miaoshu_formatted = self._generate_problem_description(
            wenti_mingcheng, wenti_laiyuan, ceshi_huodong, wenti_miaoshu
        )
        
        return {
            "测试活动": ceshi_huodong,
            "问题类型": wentileixing,
            "问题等级": wentidengji,
            "问题数目": wenti_shumu,
            "问题描述": wenti_miaoshu_formatted,
            "测评单位意见": ceping_danwei_yijian
        }
    
    def _generate_problem_description(self, wenti_mingcheng: str, wenti_laiyuan: str, 
                                    ceshi_huodong: str, wenti_miaoshu: str) -> str:
        """
        生成四段式问题描述：需求描述、问题描述、问题定位、影响分析
        """
        # 需求描述 - 将问题转换为需求
        xuqiu_miaoshu = self._convert_problem_to_requirement(wenti_mingcheng)
        
        # 问题描述
        wenti_desc = f"在{ceshi_huodong}过程中发现：{wenti_miaoshu}"
        
        # 问题定位
        wenti_dingwei = f"问题位置：{wenti_laiyuan}"
        
        # 影响分析
        yingxiang_fenxi = self._generate_impact_analysis(wenti_mingcheng, ceshi_huodong)
        
        return f"需求描述：{xuqiu_miaoshu}\n问题描述：{wenti_desc}\n问题定位：{wenti_dingwei}\n影响分析：{yingxiang_fenxi}"
    
    def _convert_problem_to_requirement(self, wenti_mingcheng: str) -> str:
        """将问题描述转换为需求描述"""
        replacements = {
            '不规范': '规范',
            '不清晰': '清晰',
            '不符合': '符合',
            '不充分': '充分',
            '不一致': '一致',
            '混乱': '清晰',
            '不对': '正确',
            '缺失': '完整',
            '缺少': '需要补充',
            '错误': '正确',
            '没有': '需要补充',
            '不全': '需要进行详细补充',
            '同名': '不能同名，具有各自独立的名称',
            '存在未初始化变量': '需要对变量进行初始化',
            '存在死循环': '不存在死循环',
            '存在空语句': '不存在空语句',
            '数组边界没有界定': '需要对数组边界进行界定'
        }
        
        xuqiu_desc = wenti_mingcheng
        for problem, requirement in replacements.items():
            if problem in xuqiu_desc:
                xuqiu_desc = xuqiu_desc.replace(problem, requirement)
                break
        
        return xuqiu_desc
    
    def _generate_impact_analysis(self, wenti_mingcheng: str, ceshi_huodong: str) -> str:
        """生成影响分析"""
        if ceshi_huodong == '静态分析':
            return "该问题可能影响代码质量和可维护性，建议及时修复以避免潜在的运行时错误。"
        elif ceshi_huodong == '代码审查':
            return "该问题可能影响软件功能的正确实现，需要开发人员进行代码修改。"
        else:  # 文档审查
            return "该问题影响需求理解、文档维护以及测试工作的正常进行。"
    
    async def save_annotations_to_database_async(self, session_id: str, annotations: List[Dict]) -> None:
        """将批注信息保存到数据库（异步版本，包含AI生成）。"""
        try:
            if not self.ai_generator:
                self.ai_generator = AIGenerator(session_id)

            with get_db_session() as db:
                # 获取会话记录，用于更新已处理批注数
                session_row = db.query(DocumentAnnotationSession).filter(
                    DocumentAnnotationSession.session_id == session_id
                ).first()

                for idx, annotation_data in enumerate(annotations, start=1):
                    logger.info(f"处理批注 {idx}/{len(annotations)}")

                    # 1) AI 生成问题名称
                    problem_name_result = await self.ai_generator.generate_problem_name(
                        annotation_data["位置"],
                        annotation_data["问题描述"],
                        annotation_data["原文内容"],
                    )
                    if problem_name_result.get("success"):
                        wenti_mingcheng = problem_name_result["problem_name"]
                    else:
                        logger.warning(
                            "AI生成问题名称失败，使用默认值: %s",
                            problem_name_result.get("error"),
                        )
                        wenti_mingcheng = "文档问题"

                    # 2) 分析问题信息（测试活动/类型/等级/数目/默认测评意见）
                    analysis_result = self.analyze_problem_info(
                        wenti_mingcheng,
                        annotation_data["问题来源"],
                        annotation_data["问题描述"],
                    )

                    # 3) AI 生成需求描述（替换四段式中的第一段）
                    requirement_result = await self.ai_generator.generate_requirement_description(
                        wenti_mingcheng,
                        annotation_data["位置"],
                        annotation_data["问题描述"],
                    )
                    if requirement_result.get("success"):
                        xuqiu_miaoshu = requirement_result["requirement_description"]
                    else:
                        logger.warning(
                            "AI生成需求描述失败，使用规则转换: %s",
                            requirement_result.get("error"),
                        )
                        xuqiu_miaoshu = self._convert_problem_to_requirement(wenti_mingcheng)

                    # 4) AI 生成测评单位意见（可覆盖默认值）
                    opinion_result = await self.ai_generator.generate_evaluation_opinion(
                        wenti_mingcheng,
                        analysis_result["测试活动"],
                        analysis_result["问题类型"],
                    )
                    if opinion_result.get("success"):
                        ceping_danwei_yijian = opinion_result["evaluation_opinion"]
                    else:
                        logger.warning(
                            "AI生成测评单位意见失败，使用默认值: %s",
                            opinion_result.get("error"),
                        )
                        ceping_danwei_yijian = analysis_result["测评单位意见"]

                    # 5) 最终四段式问题描述
                    wenti_desc = (
                        f"在{analysis_result['测试活动']}过程中发现："
                        f"{annotation_data['问题描述']}"
                    )
                    wenti_dingwei = f"问题位置：{annotation_data['问题来源']}"
                    yingxiang_fenxi = self._generate_impact_analysis(
                        wenti_mingcheng,
                        analysis_result["测试活动"],
                    )
                    wenti_miaoshu_formatted = (
                        f"需求描述：{xuqiu_miaoshu}\n"
                        f"问题描述：{wenti_desc}\n"
                        f"问题定位：{wenti_dingwei}\n"
                        f"影响分析：{yingxiang_fenxi}"
                    )

                    # 6) 写入数据库
                    annotation = DocumentAnnotation(
                        session_id=session_id,
                        问题名称=wenti_mingcheng,
                        问题来源=annotation_data["问题来源"],
                        原文内容=annotation_data["原文内容"],
                        问题描述=annotation_data["问题描述"],
                        测试活动=analysis_result["测试活动"],
                        问题类型=analysis_result["问题类型"],
                        问题等级=analysis_result["问题等级"],
                        问题数目=analysis_result["问题数目"],
                        问题描述_formatted=wenti_miaoshu_formatted,
                        测评单位意见=ceping_danwei_yijian,
                    )
                    db.add(annotation)

                    # 更新会话的已处理批注数，并提交，让状态查询能实时看到进度
                    if session_row is not None:
                        current_processed = session_row.processed_annotations or 0
                        session_row.processed_annotations = current_processed + 1
                        db.commit()

                logger.info("成功保存 %d 个批注到数据库", len(annotations))
        except Exception as e:
            logger.error("保存批注到数据库时出错: %s", e)
            raise

    def save_annotations_to_database(self, session_id: str, annotations: List[Dict]) -> None:
        """同步包装（仅用于少量脚本场景，FastAPI 主流程建议直接用 async 版本）。"""
        loop = asyncio.new_event_loop()
        try:
            loop.run_until_complete(self.save_annotations_to_database_async(session_id, annotations))
        finally:
            loop.close()

    def generate_excel_output(self, session_id: str) -> str:
        """生成 Excel 问题报告单文件并返回路径。"""
        try:
            with get_db_session() as db:
                annotations = (
                    db.query(DocumentAnnotation)
                    .filter(DocumentAnnotation.session_id == session_id)
                    .all()
                )
                if not annotations:
                    raise ValueError("没有找到批注数据")

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "问题报告单"

                headers = OUTPUT_COLUMNS
                for col, header in enumerate(headers, start=1):
                    cell = ws.cell(row=1, column=col, value=header)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.fill = PatternFill(
                        start_color="CCCCCC",
                        end_color="CCCCCC",
                        fill_type="solid",
                    )

                for row, ann in enumerate(annotations, start=2):
                    ws.cell(row=row, column=1, value=ann.测试活动)
                    ws.cell(row=row, column=2, value=ann.问题类型)
                    ws.cell(row=row, column=3, value=ann.问题等级)
                    ws.cell(row=row, column=4, value=ann.问题数目)
                    ws.cell(row=row, column=5, value=ann.问题名称)
                    ws.cell(row=row, column=6, value=ann.问题来源)
                    ws.cell(row=row, column=7, value=ann.问题描述_formatted)
                    ws.cell(row=row, column=8, value=ann.测评单位意见)

                self._set_excel_formatting(ws, len(annotations) + 1)

                output_filename = f"问题报告单_{session_id}.xlsx"
                output_path = OUTPUT_DIR / output_filename
                wb.save(output_path)
                logger.info("成功生成Excel文件: %s", output_path)
                return str(output_path)
        except Exception as e:
            logger.error("生成Excel文件时出错: %s", e)
            raise

    def _set_excel_formatting(self, ws, max_row: int) -> None:
        """设置 Excel 基础格式。"""
        for row in range(1, max_row + 1):
            ws.row_dimensions[row].height = 40

        column_widths = [20, 20, 20, 20, 20, 20, 60, 60]
        for col, width in enumerate(column_widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for row in range(1, max_row + 1):
            for col in range(1, len(column_widths) + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal="left", vertical="top", wrap_text=True
                )
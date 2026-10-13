#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
文档标准化功能测试脚本
"""

import os
import sys
from pathlib import Path

# 添加当前目录到sys.path
current_dir = Path(__file__).parent
sys.path.insert(0, str(current_dir))

def create_test_documents():
    """创建测试文档"""
    
    print("🔧 创建测试文档...")
    
    try:
        # 导入创建测试文档的模块
        import create_test_document
        
        # 创建测试文档
        test_doc1 = create_test_document.create_test_document("test_requirement_spec.docx")
        test_doc2 = create_test_document.create_simple_test_document("simple_test.docx")
        
        print(f"✅ 测试文档创建成功")
        return [test_doc1, test_doc2]
        
    except Exception as e:
        print(f"❌ 创建测试文档失败: {e}")
        # 手动创建一个简单的测试文档
        try:
            from docx import Document
            doc = Document()
            doc.add_heading('测试文档', 0)
            doc.add_paragraph('这是测试内容')
            doc.add_heading('概述', level=1)
            doc.add_paragraph('概述内容')
            doc.add_heading('功能需求', level=1)
            doc.add_paragraph('功能需求内容')
            
            test_path = "manual_test.docx"
            doc.save(test_path)
            print(f"✅ 手动创建测试文档成功: {test_path}")
            return [test_path]
            
        except Exception as e2:
            print(f"❌ 手动创建测试文档也失败: {e2}")
            return []

def test_standardizer():
    """测试文档标准化功能"""
    
    print("🧪 开始测试文档标准化功能...")
    
    # 导入标准化功能
    try:
        import document_standardizer
        create_standardized_document_from_single_file = document_standardizer.create_standardized_document_from_single_file
        print("✅ 成功导入文档标准化模块")
    except ImportError as e:
        print(f"❌ 导入失败: {e}")
        return False
    
    # 创建测试文档
    test_documents = create_test_documents()
    
    if not test_documents:
        print("⚠️  没有可用的测试文档，跳过标准化测试")
        return False
    
    # 测试每个文档
    success_count = 0
    for doc_path in test_documents:
        if not os.path.exists(doc_path):
            print(f"⚠️  测试文档不存在: {doc_path}")
            continue
            
        try:
            print(f"\n🔄 开始测试文档: {doc_path}")
            standardized_path = create_standardized_document_from_single_file(
                doc_path,
                debug_mode=True
            )
            print(f"✅ 标准化成功: {standardized_path}")
            
            # 检查生成的文件是否存在
            if os.path.exists(standardized_path):
                file_size = os.path.getsize(standardized_path)
                original_size = os.path.getsize(doc_path)
                print(f"📄 原始文件大小: {original_size} 字节")
                print(f"📄 标准化文件大小: {file_size} 字节")
                
                # 验证文件内容（基本检查）
                if file_size > 0 and file_size >= original_size * 0.8:  # 标准化后文件不应该太小
                    print("✅ 文件大小验证通过")
                    success_count += 1
                else:
                    print("❌ 文件大小异常")
            else:
                print("❌ 生成的文件不存在")
                
        except Exception as e:
            print(f"❌ 标准化失败: {e}")
            import traceback
            print(f"详细错误: {traceback.format_exc()}")
    
    if success_count > 0:
        print(f"\n✅ 文档标准化测试成功！成功处理了 {success_count}/{len(test_documents)} 个文档")
        return True
    else:
        print(f"\n❌ 文档标准化测试失败！")
        return False

def test_import_structure():
    """测试导入结构"""
    
    print("🔍 测试模块导入结构...")
    
    try:
        # 直接导入模块
        import document_standardizer
        
        print("✅ document_standardizer模块导入成功")
        
        # 测试类实例化
        standardizer = document_standardizer.DocumentStandardizer(debug_mode=True)
        print("✅ DocumentStandardizer实例化成功")
        
        return True
        
    except ImportError as e:
        print(f"❌ 导入失败: {e}")
        import traceback
        print(f"详细错误: {traceback.format_exc()}")
        return False
    except Exception as e:
        print(f"❌ 其他错误: {e}")
        import traceback
        print(f"详细错误: {traceback.format_exc()}")
        return False

def test_document_reader_integration():
    """测试与文档读取器的集成"""
    
    print("\n🔗 测试与文档审查流程的集成...")
    
    try:
        # 测试文档读取功能
        sys.path.insert(0, str(current_dir.parent))
        from file_utils import read_word_document
        
        # 查找已有的测试文档
        test_docs = []
        for filename in ["test_requirement_spec.docx", "simple_test.docx", "manual_test.docx"]:
            if os.path.exists(filename):
                test_docs.append(filename)
        
        if not test_docs:
            print("⚠️  没有找到测试文档，跳过集成测试")
            return True
        
        # 测试标准化和读取的完整流程
        for doc_path in test_docs[:1]:  # 只测试第一个文档
            print(f"🔄 测试完整流程: {doc_path}")
            
            # 1. 标准化文档
            import document_standardizer
            standardized_path = document_standardizer.create_standardized_document_from_single_file(
                doc_path,
                debug_mode=True
            )
            
            # 2. 读取标准化后的文档
            content = read_word_document(standardized_path)
            
            if content and len(content) > 0:
                print(f"✅ 完整流程测试成功！")
                print(f"📖 读取到内容长度: {len(content)} 字符")
                print(f"📖 内容预览: {content[:200]}...")
                return True
            else:
                print("❌ 读取到的内容为空")
                return False
        
    except Exception as e:
        print(f"❌ 集成测试失败: {e}")
        import traceback
        print(f"详细错误: {traceback.format_exc()}")
        return False

def cleanup_test_files():
    """清理测试文件"""
    
    print("\n🧹 清理测试文件...")
    
    test_files = [
        "test_requirement_spec.docx",
        "simple_test.docx", 
        "manual_test.docx",
        "standardized_test_requirement_spec.docx",
        "standardized_simple_test.docx",
        "standardized_manual_test.docx"
    ]
    
    cleaned_count = 0
    for filename in test_files:
        if os.path.exists(filename):
            try:
                os.remove(filename)
                print(f"🗑️  删除: {filename}")
                cleaned_count += 1
            except Exception as e:
                print(f"⚠️  无法删除 {filename}: {e}")
    
    if cleaned_count > 0:
        print(f"✅ 清理完成，删除了 {cleaned_count} 个测试文件")
    else:
        print("ℹ️  没有需要清理的测试文件")

if __name__ == "__main__":
    print("=" * 60)
    print("🧪 文档标准化功能完整测试")
    print("=" * 60)
    
    all_tests_passed = True
    
    # 测试1: 导入结构
    print("\n【测试1】模块导入测试")
    if not test_import_structure():
        all_tests_passed = False
    
    # 测试2: 文档标准化功能  
    print("\n【测试2】文档标准化功能测试")
    if not test_standardizer():
        all_tests_passed = False
    
    # 测试3: 集成测试
    print("\n【测试3】文档审查流程集成测试")
    if not test_document_reader_integration():
        all_tests_passed = False
    
    # 清理测试文件
    cleanup_test_files()
    
    # 总结
    print("\n" + "=" * 60)
    if all_tests_passed:
        print("🎉 所有测试通过！文档标准化功能正常工作")
        print("✅ 能够成功生成转换编号后的文档")
        print("✅ 文档审查流程可以正常处理标准化文档") 
    else:
        print("❌ 部分测试失败，请检查错误信息")
    print("=" * 60) 
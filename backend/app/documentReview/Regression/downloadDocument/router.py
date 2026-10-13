# 回归测试下载模块
from fastapi import APIRouter, Query, Request, Response
from fastapi.responses import FileResponse, JSONResponse
import os
import sys
import sqlite3
import tempfile
import logging
import stat
import traceback
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment

# 配置日志
logger = logging.getLogger(__name__)
logger.setLevel(logging.DEBUG)

router = APIRouter()
from pathlib import Path

# 修正 BASE_DIR 计算 - 指向新的数据库位置
BASE_DIR = Path(__file__).resolve().parent.parent
DB_PATH = str(BASE_DIR / 'db' / 'regressionReview.db')

# 确保临时目录存在 - 尝试多个位置
try:
    # 首选项目根目录
    TEMP_DIR = os.path.join(str(BASE_DIR.parent.parent.parent.parent), 'temp')
    os.makedirs(TEMP_DIR, exist_ok=True)
    
    # 测试写入权限
    test_file = os.path.join(TEMP_DIR, 'test_write.txt')
    with open(test_file, 'w') as f:
        f.write('test')
    os.remove(test_file)
    # print(f"[重要] 临时目录可写: {TEMP_DIR}")
except Exception as e:
    # 如果无法写入，尝试系统临时目录
    TEMP_DIR = tempfile.gettempdir()
    import sys
    sys.stderr.write(f"[警告] 项目目录无法写入，切换到系统临时目录: {TEMP_DIR}\n")
    sys.stderr.write(f"[警告] 原因: {str(e)}\n")

# 创建唯一的应用目录
APP_TEMP_DIR = os.path.join(TEMP_DIR, 'requirement_extractor')
os.makedirs(APP_TEMP_DIR, exist_ok=True)

# 打印关键路径用于调试
# print(f"[CRITICAL] 下载模块 - 当前运行目录: {os.getcwd()}")
# print(f"[CRITICAL] 下载模块 - Python路径: {sys.executable}")
# print(f"[CRITICAL] 下载模块 - 数据库路径: {DB_PATH}")
# print(f"[CRITICAL] 下载模块 - 临时目录路径: {APP_TEMP_DIR}")

# 同时支持带斜杠和不带斜杠的URL
@router.get("/")
@router.get("")
async def download(request: Request, session_id: int = Query(...)):
    # 详细记录请求信息
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")
    client_host = request.client.host if request.client else "unknown"
    log_prefix = f"[REGRESSION-DOWNLOAD-{timestamp}-{client_host}-{session_id}]"
    
    sys.stderr.write(f"{log_prefix} 收到回归测试下载请求 - URL: {request.url}\n")
    sys.stderr.write(f"{log_prefix} 客户端: {client_host}, headers: {dict(request.headers)}\n")
    sys.stderr.write(f"{log_prefix} 使用数据库: {DB_PATH}\n")
    logger.info(f"[下载] 收到回归测试下载请求 - session_id: {session_id}")
    
    # 检查数据库文件是否存在
    if not os.path.exists(DB_PATH):
        error_msg = f"数据库文件不存在: {DB_PATH}"
        sys.stderr.write(f"{log_prefix} 错误: {error_msg}\n")
        return JSONResponse(status_code=500, content={"error": error_msg})
    
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    
    # 查询时包含问题等级字段
    cur.execute("SELECT * FROM regression_results WHERE session_id=?", (session_id,))
    rows = cur.fetchall()
    columns = [d[0] for d in cur.description]
    conn.close()
    
    if not rows:
        return {"error": "无审查结果"}
    
    # 只一个sheet：回归测试
    wb = Workbook()
    ws = wb.create_sheet(title="回归测试")
    # 添加问题等级列到表头
    ws.append([
        "问题等级", "原始需求内容", "违反规则（规则名）", "问题名称", "问题来源（文档名称+章节）", "问题描述"
    ])
    
    for r in rows:
        ws.append([
            r[columns.index('issue_level')],  # 新增：问题等级列
            r[columns.index('requirement_content')],
            r[columns.index('violation_rules')],
            r[columns.index('issue_name')],
            r[columns.index('issue_location')],
            f"需求描述：{r[columns.index('description_requirement')] or ''}\n问题描述：{r[columns.index('description_issue')] or ''}\n问题定位：{r[columns.index('description_location')] or ''}\n影响分析：{r[columns.index('description_impact')] or ''}"
        ])
    
    # 设置列宽和自动换行
    wrap_alignment = Alignment(wrapText=True, vertical="top")
    for worksheet in wb.worksheets:
        for col_cells in worksheet.columns:
            max_len = 0
            for cell in col_cells:
                cell.alignment = wrap_alignment
                if cell.value:
                    l = len(str(cell.value))
                    if l > max_len:
                        max_len = l
            adjusted_width = min(max(max_len + 2, 20), 60)
            column_letter = col_cells[0].column_letter
            worksheet.column_dimensions[column_letter].width = adjusted_width

    # 删除默认sheet
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']
        
    # 生成带时间戳的唯一文件名
    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    file_name = f"regression_report_{session_id}_{timestamp_str}.xlsx"
    tmp_file_path = os.path.join(APP_TEMP_DIR, file_name)
    
    sys.stderr.write(f"{log_prefix} 将生成文件: {tmp_file_path}\n")
    
    try:
        # 保存Excel文件
        sys.stderr.write(f"{log_prefix} [重要] 正在保存Excel文件到: {tmp_file_path}\n")
        wb.save(tmp_file_path)
        
        # 确认文件存在并可读
        if os.path.exists(tmp_file_path):
            file_size = os.path.getsize(tmp_file_path)
            sys.stderr.write(f"{log_prefix} [重要] Excel文件保存成功，文件大小: {file_size} 字节\n")
            
            try:
                # 尝试将文件设置为只读，防止被意外删除
                os.chmod(tmp_file_path, stat.S_IREAD | stat.S_IRGRP | stat.S_IROTH)
                sys.stderr.write(f"{log_prefix} 文件已设置为只读保护\n")
            except Exception as chmod_err:
                sys.stderr.write(f"{log_prefix} 设置文件权限失败: {str(chmod_err)}\n")
            
            # 创建一个标记文件作为调试参考
            marker_file = os.path.join(APP_TEMP_DIR, f"regression_marker_{session_id}_{timestamp_str}.txt")
            with open(marker_file, 'w', encoding='utf-8') as f:
                f.write(f"Regression Excel file created at: {timestamp_str}\nPath: {tmp_file_path}\nSize: {file_size} bytes")
            
            sys.stderr.write(f"{log_prefix} [重要] 正在返回文件: {tmp_file_path}\n")
            
            # 使用FileResponse返回文件
            return FileResponse(
                path=tmp_file_path, 
                filename=f"回归测试AI审查问题单_{session_id}.xlsx",
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        else:
            error_msg = f"错误: 文件保存失败，文件不存在: {tmp_file_path}"
            sys.stderr.write(f"{log_prefix} {error_msg}\n")
            return JSONResponse(status_code=500, content={"error": error_msg})
    except Exception as e:
        error_trace = traceback.format_exc()
        error_msg = f"[下载] 回归测试模块 - 生成或返回Excel文件时出错: {str(e)}"
        sys.stderr.write(f"{log_prefix} {error_msg}\n")
        sys.stderr.write(f"{log_prefix} \n{error_trace}\n")
        logger.error(error_msg)
        logger.error(error_trace)
        
        # 写入错误日志文件便于调试
        try:
            error_log_file = os.path.join(APP_TEMP_DIR, f"regression_error_{session_id}_{timestamp_str}.log")
            with open(error_log_file, 'w', encoding='utf-8') as f:
                f.write(f"Error at: {timestamp_str}\n")
                f.write(f"Error message: {str(e)}\n\n")
                f.write(error_trace)
            sys.stderr.write(f"{log_prefix} 错误日志已写入: {error_log_file}\n")
        except Exception as log_err:
            sys.stderr.write(f"{log_prefix} 无法写入错误日志: {str(log_err)}\n")
        
        return JSONResponse(
            status_code=500, 
            content={
                "error": f"回归测试下载失败: {str(e)}", 
                "details": error_trace,
                "timestamp": timestamp_str
            }
        )


# ===== 新增：下载批注版影响域分析报告 =====
@router.get("/annotated/impactreport")
async def download_annotated_impactreport(session_id: int = Query(...)):
    """下载批注版影响域分析报告"""
    try:
        # 从数据库中获取批注文件路径
        conn = sqlite3.connect(DB_PATH)
        cur = conn.cursor()
        cur.execute("SELECT annotated_impactreport_path FROM review_sessions WHERE id = ?", (session_id,))
        result = cur.fetchone()
        conn.close()
        
        if result and result[0] and os.path.exists(result[0]):
            annotated_path = result[0]
            logger.info(f"从数据库找到影响域分析报告批注文件: {annotated_path}")
            return FileResponse(
                path=annotated_path,
                filename=f"影响域分析报告_批注版_{session_id}.docx",
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        
        # 备选方案：在临时目录查找对应session_id的批注文件
        temp_dir = os.path.join(os.path.dirname(BASE_DIR.parent), 'temp', 'annotated_documents')
        if os.path.exists(temp_dir):
            pattern = f"影响域分析报告_批注版_session_{session_id}.docx"
            import glob
            candidates = glob.glob(os.path.join(temp_dir, pattern))
            
            if candidates:
                file_path = candidates[0]
                logger.info(f"在临时目录找到影响域分析报告批注文件: {file_path}")
                return FileResponse(
                    path=file_path,
                    filename=f"影响域分析报告_批注版_{session_id}.docx",
                    media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
        
        logger.error(f"未找到影响域分析报告批注版文档 - session_id: {session_id}")
        return JSONResponse(status_code=404, content={"error": "未找到影响域分析报告批注版文档"})
        
    except Exception as e:
        logger.error(f"下载影响域分析报告批注版失败: {e}", exc_info=True)
        return JSONResponse(status_code=500, content={"error": str(e)})

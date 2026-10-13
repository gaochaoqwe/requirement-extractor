"""
测试用例反向推导Excel处理器
"""
import openpyxl
from openpyxl.styles import Alignment
from pathlib import Path
from typing import Dict, List, Any, Tuple, Optional
import re
import uuid
from datetime import datetime
import json
from collections import defaultdict

from ..config.config import REQUIRED_SHEET_NAME, REQUIRED_COLUMNS, REQUIREMENT_SHEET_NAME, REQUIREMENT_COLUMNS, TEST_ROUND_COLUMN
from ..models.models import TestCaseReverseSession, TestItem
from ..database.database import get_db_session
from ..utils.logger import get_session_logger

class ExcelProcessor:
    def __init__(self, session_id: str):
        self.session_id = session_id
        self.logger = get_session_logger(session_id)

    def _normalize_multiline_text(self, text: Optional[str]) -> List[str]:
        """将多行文本清洗为非空行列表"""
        if not text:
            return []
        return [line.strip() for line in str(text).split('\n') if line and line.strip()]

    def _is_first_round(self, value: Any) -> bool:
        """判断测试轮次（E列）是否属于第一轮。

        仅保留第一轮测试用例，忽略第二轮、第三轮等。可识别“第一轮”“第1轮”“1”等写法；
        轮次列为空时按第一轮处理，避免漏取数据。
        """
        text = str(value or '').strip()
        if not text:
            return True
        if '一' in text:
            return True
        matched = re.search(r'\d+', text)
        return bool(matched and matched.group(0) == '1')

    def _build_case_name_description_sheet(self, workbook, test_items: List[TestItem]):
        """为普通模式生成测试用例名称和测试用例描述sheet"""
        worksheet = workbook.create_sheet(title="测试用例名称和测试用例描述", index=0)
        headers = ["测试项名称", "测试用例编号", "测试用例名称", "测试用例描述"]

        for col, header in enumerate(headers, 1):
            cell = worksheet.cell(row=1, column=col, value=header)
            cell.font = openpyxl.styles.Font(bold=True)
            cell.alignment = openpyxl.styles.Alignment(horizontal='center', vertical='center')

        row = 2
        for test_item in test_items:
            test_cases = json.loads(test_item.test_cases_data) if test_item.test_cases_data else []
            case_names = self._normalize_multiline_text(test_item.case_template_name)
            case_descriptions = self._normalize_multiline_text(test_item.case_template_expected)
            identifier_counters = defaultdict(int)

            for index, test_case in enumerate(test_cases):
                test_item_name = test_case.get("test_item_name") or test_item.test_item_name
                test_item_identifier = test_case.get("test_item_identifier") or test_item_name
                identifier_counters[test_item_identifier] += 1
                case_number = f"{test_item_identifier}-{identifier_counters[test_item_identifier]:02d}"

                case_name = case_names[index] if index < len(case_names) else test_item_name
                case_description = case_descriptions[index] if index < len(case_descriptions) else ""

                values = [test_item_name, case_number, case_name, case_description]
                for col, value in enumerate(values, 1):
                    cell = worksheet.cell(row=row, column=col, value=value)
                    cell.alignment = Alignment(wrap_text=True, vertical='top')

                worksheet.row_dimensions[row].height = 60
                row += 1

        worksheet.column_dimensions['A'].width = 30
        worksheet.column_dimensions['B'].width = 30
        worksheet.column_dimensions['C'].width = 50
        worksheet.column_dimensions['D'].width = 70

        return worksheet

    def _build_merged_normal_sheet(self, workbook, test_items: List[TestItem]):
        """为普通模式生成合并后的汇总sheet"""
        worksheet = workbook.create_sheet(title="测试方法（合并版）", index=1)
        headers = ["测试项名称", "测试内容", "测试方法"]

        for col, header in enumerate(headers, 1):
            cell = worksheet.cell(row=1, column=col, value=header)
            cell.font = openpyxl.styles.Font(bold=True)
            cell.alignment = openpyxl.styles.Alignment(horizontal='center', vertical='center')

        merged_data = defaultdict(lambda: {"contents": [], "methods": []})
        for test_item in test_items:
            item_data = merged_data[test_item.test_item_name]
            item_data["contents"].extend(self._normalize_multiline_text(test_item.test_content))
            item_data["methods"].extend(self._normalize_multiline_text(test_item.test_method))

        row = 2
        for test_item_name, item_data in merged_data.items():
            content_value = '\n'.join(item_data["contents"])
            method_value = '\n'.join(item_data["methods"])

            name_cell = worksheet.cell(row=row, column=1, value=test_item_name)
            name_cell.alignment = Alignment(wrap_text=True, vertical='top')

            content_cell = worksheet.cell(row=row, column=2, value=content_value)
            content_cell.alignment = Alignment(wrap_text=True, vertical='top')

            method_cell = worksheet.cell(row=row, column=3, value=method_value)
            method_cell.alignment = Alignment(wrap_text=True, vertical='top')

            worksheet.row_dimensions[row].height = 80
            row += 1

        worksheet.column_dimensions['A'].width = 30
        worksheet.column_dimensions['B'].width = 50
        worksheet.column_dimensions['C'].width = 50

        return worksheet

    def _read_requirement_descriptions(self, file_path: str) -> Dict[str, str]:
        """
        从需求分析表读取测试项名称与需求子项描述的映射

        Returns:
            Dict[str, str]: {测试项名称: 需求子项描述}，找不到需求分析表时返回空字典
        """
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True)

            if REQUIREMENT_SHEET_NAME not in workbook.sheetnames:
                self.logger.warning(f"Excel文件中未找到工作表 '{REQUIREMENT_SHEET_NAME}'，将不附加需求描述")
                workbook.close()
                return {}

            worksheet = workbook[REQUIREMENT_SHEET_NAME]
            requirement_map: Dict[str, str] = {}

            name_col = REQUIREMENT_COLUMNS["测试项名称"]
            desc_col = REQUIREMENT_COLUMNS["需求子项描述"]
            max_col = max(name_col, desc_col)

            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if not row or len(row) <= max_col:
                    continue

                test_item_name = row[name_col]
                requirement_desc = row[desc_col]

                if not test_item_name or not str(test_item_name).strip():
                    continue

                name_key = str(test_item_name).strip()
                desc_value = str(requirement_desc).strip() if requirement_desc else ""

                if name_key in requirement_map:
                    if desc_value and desc_value not in requirement_map[name_key]:
                        requirement_map[name_key] += "\n" + desc_value
                else:
                    requirement_map[name_key] = desc_value

            workbook.close()

            # 读取测试用例单的测试项名称，与需求分析表对比，打印未匹配项供排查
            try:
                wb2 = openpyxl.load_workbook(file_path, read_only=True)
                ws2 = wb2[REQUIRED_SHEET_NAME]
                case_item_names = set()
                for row in ws2.iter_rows(min_row=2, values_only=True):
                    if row and len(row) > REQUIRED_COLUMNS["测试项名称"]:
                        v = row[REQUIRED_COLUMNS["测试项名称"]]
                        if v and str(v).strip():
                            case_item_names.add(str(v).strip())
                wb2.close()

                unmatched = case_item_names - set(requirement_map.keys())
                if unmatched:
                    print("\n[需求匹配] 以下测试项在需求分析表中未找到对应需求描述，请检查两个sheet的测试项名称是否一致：")
                    for name in sorted(unmatched):
                        print(f"  未匹配: {repr(name)}")
                    print("[需求匹配] 需求分析表中已有的测试项名称：")
                    for name in sorted(requirement_map.keys()):
                        print(f"  已有:   {repr(name)}")
            except Exception as e:
                self.logger.warning(f"对比测试项名称时出错: {str(e)}")

            self.logger.info(f"从需求分析表读取到 {len(requirement_map)} 条需求描述")
            return requirement_map

        except Exception as e:
            self.logger.error(f"读取需求分析表时出错: {str(e)}")
            return {}

    def process_uploaded_file(self, file_path: str, original_filename: str) -> Tuple[str, Dict[str, Any]]:
        """
        处理上传的Excel文件
        
        Args:
            file_path: 文件路径
            original_filename: 原始文件名
            
        Returns:
            Tuple[str, Dict]: (session_id, result)
        """
        try:
            self.logger.info(f"开始处理Excel文件: {file_path}")
            
            # 验证文件
            validation_result = self._validate_excel_file(file_path)
            if not validation_result["success"]:
                return self.session_id, validation_result
            
            # 读取需求分析表（允许不存在，不影响主流程）
            requirement_map = self._read_requirement_descriptions(file_path)

            # 读取测试用例数据
            test_cases_data = self._read_test_case_data(file_path)
            if not test_cases_data["success"]:
                return self.session_id, test_cases_data

            # 将需求描述注入到每个测试项
            for item in test_cases_data["test_items"]:
                item["requirement_description"] = requirement_map.get(item["test_item_name"], "")

            # 保存到数据库
            save_result = self._save_to_database(original_filename, file_path, test_cases_data["test_items"])
            if not save_result["success"]:
                return self.session_id, save_result
            
            self.logger.info(f"Excel文件处理完成，共处理 {len(test_cases_data['test_items'])} 个测试项")
            
            return self.session_id, {
                "success": True,
                "message": "文件处理成功",
                "total_test_items": len(test_cases_data["test_items"])
            }
            
        except Exception as e:
            error_msg = f"处理Excel文件时出错: {str(e)}"
            self.logger.error(error_msg)
            return self.session_id, {"success": False, "error": error_msg}

    def _validate_excel_file(self, file_path: str) -> Dict[str, Any]:
        """验证Excel文件格式"""
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True)
            
            # 检查是否包含必需的工作表
            if REQUIRED_SHEET_NAME not in workbook.sheetnames:
                return {
                    "success": False,
                    "error": f"Excel文件中未找到工作表 '{REQUIRED_SHEET_NAME}'"
                }
            
            worksheet = workbook[REQUIRED_SHEET_NAME]
            
            # 检查列标题（检查第一行）
            header_row = list(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]
            
            missing_columns = []
            for col_name, col_index in REQUIRED_COLUMNS.items():
                if col_index >= len(header_row) or not header_row[col_index]:
                    missing_columns.append(f"{col_name}(第{col_index+1}列)")
            
            if missing_columns:
                return {
                    "success": False,
                    "error": f"Excel文件缺少必需的列: {', '.join(missing_columns)}"
                }
            
            workbook.close()
            return {"success": True}
            
        except Exception as e:
            return {"success": False, "error": f"验证Excel文件时出错: {str(e)}"}

    def _read_test_case_data(self, file_path: str) -> Dict[str, Any]:
        """读取测试用例数据"""
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True)
            worksheet = workbook[REQUIRED_SHEET_NAME]
            
            # 按测试项名称分组测试用例
            test_items = defaultdict(list)
            
            # 从第二行开始读取数据（跳过标题行）
            skipped_rounds = 0
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                if not row or len(row) <= max(REQUIRED_COLUMNS.values()):
                    continue

                # 只取第一轮测试用例（E列 测试轮次），忽略第二轮、第三轮
                if TEST_ROUND_COLUMN < len(row) and not self._is_first_round(row[TEST_ROUND_COLUMN]):
                    skipped_rounds += 1
                    continue

                # 提取必需的列数据
                test_item_name = row[REQUIRED_COLUMNS["测试项名称"]]
                test_item_identifier = row[REQUIRED_COLUMNS["测试项标识"]]
                input_condition = row[REQUIRED_COLUMNS["输入条件"]]
                expected_result = row[REQUIRED_COLUMNS["预期结果"]]
                
                # 跳过空行或无效数据
                if not test_item_name or not str(test_item_name).strip():
                    continue
                
                test_item_name = str(test_item_name).strip()
                test_item_identifier = str(test_item_identifier or "").strip()
                input_condition = str(input_condition or "").strip()
                expected_result = str(expected_result or "").strip()
                
                # 将测试用例添加到对应的测试项
                test_items[test_item_name].append({
                    "test_item_name": test_item_name,
                    "test_item_identifier": test_item_identifier,
                    "input_condition": input_condition,
                    "expected_result": expected_result
                })
            
            workbook.close()
            
            # 转换为列表格式
            test_items_list = []
            for test_item_name, test_cases in test_items.items():
                test_items_list.append({
                    "test_item_name": test_item_name,
                    "test_cases": test_cases
                })
            
            self.logger.info(f"成功读取 {len(test_items_list)} 个测试项，共 {sum(len(item['test_cases']) for item in test_items_list)} 个测试用例（已忽略 {skipped_rounds} 条非第一轮用例）")
            
            return {
                "success": True,
                "test_items": test_items_list
            }
            
        except Exception as e:
            error_msg = f"读取测试用例数据时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg}

    def _save_to_database(self, original_filename: str, file_path: str, test_items: List[Dict]) -> Dict[str, Any]:
        """保存数据到数据库"""
        try:
            with get_db_session() as db:
                # 创建会话记录
                session = TestCaseReverseSession(
                    session_id=self.session_id,
                    original_filename=original_filename,
                    file_path=file_path,
                    status="uploaded",
                    total_test_items=len(test_items),
                    processed_test_items=0
                )
                db.add(session)
                db.flush()  # 确保会话记录被保存
                
                # 保存测试项
                for test_item_data in test_items:
                    test_item = TestItem(
                        session_id=self.session_id,
                        test_item_name=test_item_data["test_item_name"],
                        test_cases_data=json.dumps(test_item_data["test_cases"], ensure_ascii=False),
                        requirement_description=test_item_data.get("requirement_description", ""),
                        generation_status="pending"
                    )
                    db.add(test_item)
                
                db.commit()
                
                self.logger.info(f"成功保存 {len(test_items)} 个测试项到数据库")
                
                return {"success": True}
                
        except Exception as e:
            error_msg = f"保存数据到数据库时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg}

    def generate_output_excel(self, session_id: str) -> Dict[str, Any]:
        """生成输出Excel文件"""
        try:
            self.logger.info(f"开始生成会话 {session_id} 的输出Excel文件")
            
            # 从数据库获取数据
            with get_db_session() as db:
                # 获取会话信息
                session = db.query(TestCaseReverseSession).filter(
                    TestCaseReverseSession.session_id == session_id
                ).first()
                
                if not session:
                    return {"success": False, "error": "会话不存在"}
                
                # 获取已完成的测试项
                test_items = db.query(TestItem).filter(
                    TestItem.session_id == session_id,
                    TestItem.generation_status == "completed"
                ).all()
                
                if not test_items:
                    return {"success": False, "error": "没有已完成的测试项"}
            
            # 检查处理模式
            processing_mode = session.processing_mode or "normal"
            
            # 创建新的Excel文件
            workbook = openpyxl.Workbook()
            if processing_mode == "normal":
                self._build_case_name_description_sheet(workbook, test_items)
                self._build_merged_normal_sheet(workbook, test_items)
                worksheet = workbook.worksheets[2]
            else:
                worksheet = workbook.active
            
            if processing_mode == "template":
                worksheet.title = "测试内容和测试方法及用例模板"
                # 用例模板模式的表头
                headers = [
                    "测试项名称", "测试内容", "测试方法", 
                    "用例模板（用例名称）", "用例模板（输入条件）", "用例模板（预期结果）"
                ]
            else:
                worksheet.title = "测试方法（拆分版）"
                # 普通模式的表头
                headers = ["测试项名称", "测试内容", "测试方法"]
            
            # 设置表头
            for col, header in enumerate(headers, 1):
                cell = worksheet.cell(row=1, column=col, value=header)
                cell.font = openpyxl.styles.Font(bold=True)
                cell.alignment = openpyxl.styles.Alignment(horizontal='center', vertical='center')
            
            # 写入数据
            row = 2
            
            for test_item in test_items:
                # 按\n分割测试内容和测试方法，同时处理模板字段
                content_lines = (test_item.test_content or "").split('\n') if test_item.test_content else [""]
                method_lines = (test_item.test_method or "").split('\n') if test_item.test_method else [""]
                
                template_name_lines = []
                template_input_lines = []
                template_expected_lines = []
                
                if processing_mode == "template":
                    template_name_raw = (test_item.case_template_name or "").split('\n') if test_item.case_template_name else [""]
                    template_input_raw = (test_item.case_template_input or "").split('\n') if test_item.case_template_input else [""]
                    template_expected_raw = (test_item.case_template_expected or "").split('\n') if test_item.case_template_expected else [""]
                    
                    # 模板字段直接使用，已经包含标准的\n换行符
                    template_name_lines = [line.strip() for line in template_name_raw if line.strip()]
                    template_input_lines = [line.strip() for line in template_input_raw if line.strip()]
                    template_expected_lines = [line.strip() for line in template_expected_raw if line.strip()]
                
                # 确定行数
                max_lines = max(len(content_lines), len(method_lines))
                if processing_mode == "template":
                    max_lines = max(max_lines, len(template_name_lines), len(template_input_lines), len(template_expected_lines))
                
                # 补齐列表
                while len(content_lines) < max_lines:
                    content_lines.append("")
                while len(method_lines) < max_lines:
                    method_lines.append("")
                
                if processing_mode == "template":
                    while len(template_name_lines) < max_lines:
                        template_name_lines.append("")
                    while len(template_input_lines) < max_lines:
                        template_input_lines.append("")
                    while len(template_expected_lines) < max_lines:
                        template_expected_lines.append("")
                
                # 创建行
                for i in range(max_lines):
                    # 测试项名称
                    name_cell = worksheet.cell(row=row, column=1, value=test_item.test_item_name)
                    name_cell.alignment = Alignment(wrap_text=True, vertical='top')
                    
                    # 测试内容
                    content_cell = worksheet.cell(row=row, column=2, value=content_lines[i].strip())
                    content_cell.alignment = Alignment(wrap_text=True, vertical='top')
                    
                    # 测试方法
                    method_cell = worksheet.cell(row=row, column=3, value=method_lines[i].strip())
                    method_cell.alignment = Alignment(wrap_text=True, vertical='top')
                    
                    # 模板字段（将内部的\n转换为Excel换行符）
                    if processing_mode == "template":
                        # 用例模板名称（通常不需要换行）
                        template_name_cell = worksheet.cell(row=row, column=4, value=template_name_lines[i])
                        template_name_cell.alignment = Alignment(wrap_text=True, vertical='top')
                        
                        # 用例模板输入条件（将\n转换为Excel内换行）
                        template_input_value = template_input_lines[i].replace('\\n', '\n') if template_input_lines[i] else ""
                        template_input_cell = worksheet.cell(row=row, column=5, value=template_input_value)
                        template_input_cell.alignment = Alignment(wrap_text=True, vertical='top')
                        
                        # 用例模板预期结果（将\n转换为Excel内换行）
                        template_expected_value = template_expected_lines[i].replace('\\n', '\n') if template_expected_lines[i] else ""
                        template_expected_cell = worksheet.cell(row=row, column=6, value=template_expected_value)
                        template_expected_cell.alignment = Alignment(wrap_text=True, vertical='top')
                    
                    worksheet.row_dimensions[row].height = 60
                    row += 1
            
            # 设置列宽
            worksheet.column_dimensions['A'].width = 30  # 测试项名称
            worksheet.column_dimensions['B'].width = 50  # 测试内容
            worksheet.column_dimensions['C'].width = 50  # 测试方法
            
            if processing_mode == "template":
                worksheet.column_dimensions['D'].width = 40  # 用例模板（用例名称）
                worksheet.column_dimensions['E'].width = 80  # 用例模板（输入条件）- 增加宽度以适应多步骤内容
                worksheet.column_dimensions['F'].width = 80  # 用例模板（预期结果）- 增加宽度以适应多步骤内容
            
            # 生成文件名和路径
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"test_case_reverse_result_{session_id}_{timestamp}.xlsx"
            
            # 确保输出目录存在
            from ..config.config import UPLOAD_DIR
            output_dir = UPLOAD_DIR / "outputs"
            output_dir.mkdir(parents=True, exist_ok=True)
            
            output_path = output_dir / filename
            
            # 保存文件
            workbook.save(output_path)
            workbook.close()
            
            self.logger.info(f"输出Excel文件生成成功: {output_path}")
            
            return {
                "success": True,
                "filename": filename,
                "file_path": str(output_path),
                "total_items": len(test_items)
            }
            
        except Exception as e:
            error_msg = f"生成输出Excel文件时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg} 
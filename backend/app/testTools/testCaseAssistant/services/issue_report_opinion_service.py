"""问题报告单归零说明生成服务。

上传包含“问题报告单”工作表的 Excel，读取 T 列“研制单位意见”，
调用 AI 生成“问题确认意见/归零说明”，写入同一行 AA 列，
内容格式为：<修改说明>\\n问题已归零。
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment

from ..config.config import MODULE_UPLOAD_DIR
from ..utils.logger import setup_session_logger
from .generation_progress import fail_progress, finish_progress, update_progress
from .test_case_service.test_case_design.case_design_ai_client import CaseDesignAIClient

# 问题报告单工作表名关键字
ISSUE_REPORT_SHEET_KEYWORD = '问题报告单'
# 表头关键字：研制单位意见（T列）/ 问题确认意见·归零说明（AA列）
OPINION_HEADER_KEYWORDS = ('研制单位意见',)
CONFIRMATION_HEADER_KEYWORDS = ('问题确认意见', '归零说明')
# 未匹配到表头时的默认列号（1-based）：T列=20，AA列=27
DEFAULT_OPINION_COLUMN = 20
DEFAULT_CONFIRMATION_COLUMN = 27
# 归零说明固定结尾（换行后追加）
CLOSED_SUFFIX = '问题已归零'

# 生成结果保存目录
OPINION_OUTPUT_DIR = MODULE_UPLOAD_DIR / 'issue_report_opinion'
OPINION_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


class IssueReportOpinionService:
    def __init__(self, session_id: str):
        self.session_id = session_id
        self.logger = setup_session_logger(session_id)
        self.ai_client = CaseDesignAIClient(session_id)

    def _find_sheet_name(self, workbook) -> Optional[str]:
        for name in workbook.sheetnames:
            if ISSUE_REPORT_SHEET_KEYWORD in str(name):
                return name
        return None

    def _locate_columns(self, worksheet) -> Tuple[int, int, int]:
        """定位表头行与意见/归零说明列，返回 (header_row, opinion_col, confirmation_col)，列号为 1-based。"""
        opinion_col = DEFAULT_OPINION_COLUMN
        confirmation_col = DEFAULT_CONFIRMATION_COLUMN
        header_row = 1

        scan_rows = min(worksheet.max_row or 1, 10)
        scan_cols = worksheet.max_column or 0
        for row in range(1, scan_rows + 1):
            found_opinion = None
            found_confirmation = None
            for col in range(1, scan_cols + 1):
                text = str(worksheet.cell(row=row, column=col).value or '')
                if not text:
                    continue
                if found_opinion is None and any(k in text for k in OPINION_HEADER_KEYWORDS):
                    found_opinion = col
                if found_confirmation is None and any(k in text for k in CONFIRMATION_HEADER_KEYWORDS):
                    found_confirmation = col
            if found_opinion:
                header_row = row
                opinion_col = found_opinion
                if found_confirmation:
                    confirmation_col = found_confirmation
                break

        return header_row, opinion_col, confirmation_col

    def _build_prompt(self, opinion_text: str) -> str:
        return (
            '你是测试工程师助手。\n'
            '任务：根据【研制单位意见】，生成该问题的【问题确认意见/归零说明】。\n'
            '要求：\n'
            '1) 你必须只输出 JSON，不能输出任何多余文字、不能使用 Markdown。\n'
            '2) 输出格式固定为：{"confirmation_opinion": "..."}\n'
            '3) 内容为一段连贯的中文说明，复述并书面化该问题是如何修改/处理并关闭的，需与研制单位意见保持一致。\n'
            '4) 不要输出“问题已归零”字样，不要输出编号、项目符号或换行。\n'
            '5) 为保证 JSON 合法：输出的字符串中禁止出现英文双引号(")，如需引用对象请用《》包裹。\n\n'
            f'研制单位意见：\n{opinion_text}\n'
        )

    async def _generate_opinion(self, opinion_text: str) -> Tuple[str, Optional[str]]:
        """调用 AI 生成归零说明；失败时返回 (空串, 错误信息)。"""
        prompt = self._build_prompt(opinion_text)
        try:
            ai_res = await self.ai_client._call_ai_with_retry_generic(
                prompt, required_keys=['confirmation_opinion']
            )
            if isinstance(ai_res, dict) and ai_res.get('success'):
                data = ai_res.get('data') if isinstance(ai_res.get('data'), dict) else {}
                text = str(data.get('confirmation_opinion') or '').strip()
                if text:
                    return text, None
                return '', 'AI未返回confirmation_opinion'
            return '', str((ai_res or {}).get('error') or 'AI返回失败')
        except Exception as e:
            return '', str(e)

    async def process_file(self, file_path: str, original_filename: str) -> Dict[str, Any]:
        """解析 Excel，按行生成归零说明并写入 AA 列，另存为新文件。"""
        update_progress(self.session_id, status='running', percent=3, message='正在打开并解析 Excel 文件…')
        try:
            workbook = openpyxl.load_workbook(file_path)
        except Exception as e:
            fail_progress(self.session_id, f'无法解析Excel文件: {e}')
            return {'success': False, 'error': f'无法解析Excel文件: {e}'}

        try:
            sheet_name = self._find_sheet_name(workbook)
            if not sheet_name:
                msg = f'未找到包含“{ISSUE_REPORT_SHEET_KEYWORD}”的工作表'
                fail_progress(self.session_id, msg)
                return {'success': False, 'error': msg}

            worksheet = workbook[sheet_name]
            update_progress(self.session_id, percent=6, message=f'正在定位工作表「{sheet_name}」的意见列…')
            header_row, opinion_col, confirmation_col = self._locate_columns(worksheet)
            self.logger.info(
                '问题报告单归零说明：sheet=%s header_row=%s opinion_col=%s confirmation_col=%s',
                sheet_name, header_row, opinion_col, confirmation_col,
            )

            # 第一遍：先扫描出所有「研制单位意见」非空、需要生成归零说明的行
            last_row = worksheet.max_row or 0
            pending_rows: List[int] = []
            skipped_rows = 0
            for row in range(header_row + 1, last_row + 1):
                opinion_text = str(worksheet.cell(row=row, column=opinion_col).value or '').strip()
                if opinion_text:
                    pending_rows.append(row)
                else:
                    skipped_rows += 1

            total = len(pending_rows)
            update_progress(
                self.session_id,
                status='running',
                percent=10,
                current=0,
                total=total,
                message=(
                    f'共 {total} 行需要生成归零说明（跳过 {skipped_rows} 行空白意见），开始逐行用 AI 生成…'
                    if total else '未找到可生成的行（“研制单位意见”列为空）'
                ),
            )

            updated_rows = 0
            errors: List[Dict[str, Any]] = []

            for index, row in enumerate(pending_rows, 1):
                opinion_text = str(worksheet.cell(row=row, column=opinion_col).value or '').strip()

                # 进度区间：10% -> 95%，按行均匀推进
                update_progress(
                    self.session_id,
                    status='running',
                    percent=10 + int(85 * (index - 1) / total),
                    current=index,
                    total=total,
                    current_label=f'第 {row} 行',
                    updated_rows=updated_rows,
                    message=f'正在生成第 {index}/{total} 行归零说明（工作表第 {row} 行）…',
                )

                ai_text, err = await self._generate_opinion(opinion_text)
                if not ai_text:
                    # AI 失败时回退为研制单位意见原文，避免该列留空
                    ai_text = opinion_text
                    errors.append({
                        'row': row,
                        'error': err or 'AI生成失败，已回退为研制单位意见原文',
                    })
                    self.logger.warning('问题报告单归零说明：第%s行生成失败 %s', row, err)

                cell = worksheet.cell(row=row, column=confirmation_col, value=f'{ai_text}\n{CLOSED_SUFFIX}')
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                updated_rows += 1

                update_progress(
                    self.session_id,
                    percent=10 + int(85 * index / total),
                    current=index,
                    total=total,
                    current_label=f'第 {row} 行',
                    updated_rows=updated_rows,
                    append_item={
                        'name': f'第 {row} 行',
                        'ok': not err,
                        'detail': '已生成归零说明' if not err
                        else f'AI 生成失败，已回退原文：{err or "未知原因"}',
                    },
                    message=f'已完成第 {index}/{total} 行（工作表第 {row} 行）',
                )

            update_progress(self.session_id, percent=96, message='正在保存生成结果 Excel…')
            output_filename = f'issue_report_opinion_{self.session_id}.xlsx'
            output_path = OPINION_OUTPUT_DIR / output_filename
            workbook.save(output_path)

            download_filename = f'{Path(original_filename).stem}_归零说明.xlsx'

            self.logger.info(
                '问题报告单归零说明：生成完成 updated=%s skipped=%s errors=%s -> %s',
                updated_rows, skipped_rows, len(errors), output_path,
            )
            finish_progress(
                self.session_id,
                message=(
                    f'生成完成：已生成 {updated_rows} 行归零说明，跳过 {skipped_rows} 行'
                    + (f'，其中 {len(errors)} 行 AI 生成失败已回退原文' if errors else '')
                ),
                current=updated_rows,
                total=total,
                updated_rows=updated_rows,
            )

            return {
                'success': True,
                'sheet_name': sheet_name,
                'header_row': header_row,
                'opinion_column': opinion_col,
                'confirmation_column': confirmation_col,
                'updated_rows': updated_rows,
                'skipped_rows': skipped_rows,
                'errors': errors,
                'output_filename': output_filename,
                'download_filename': download_filename,
            }
        except Exception as e:
            self.logger.exception('问题报告单归零说明：处理文件失败')
            fail_progress(self.session_id, f'处理文件时出错: {e}')
            return {'success': False, 'error': f'处理文件时出错: {e}'}
        finally:
            workbook.close()

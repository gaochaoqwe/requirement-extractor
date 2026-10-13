"""动态测试问题单生成服务。

上传包含“测试用例单”工作表的 Excel，找出测试结论（AA列）为“不通过”的用例，
按测试项（G列 测试项名称）分组，逐测试项调用 AI（提供 P列输入条件、Q列预期结果、Z列测试结果），
生成问题单，并以静态分析过滤模块的问题单 12 列格式汇总导出新的 Excel。
"""
from __future__ import annotations

from collections import OrderedDict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from ..config.config import MODULE_UPLOAD_DIR
from ..utils.logger import setup_session_logger
from .test_case_service.test_case_design.case_design_ai_client import CaseDesignAIClient

# 测试用例单 sheet 名关键字
CASE_SHEET_KEYWORD = '测试用例单'
# 测试用例单列索引（0-based）
COL_G_TEST_ITEM_NAME = 6      # G列 测试项名称
COL_P_INPUT_CONDITION = 15    # P列 输入条件
COL_Q_EXPECTED_RESULT = 16    # Q列 预期结果
COL_Z_TEST_RESULT = 25        # Z列 测试结果
COL_AA_TEST_CONCLUSION = 26   # AA列 测试结论

# 问题单固定取值
TEST_ACTIVITY = '动态测试'
ISSUE_TYPE = '程序问题'
ISSUE_LEVEL = '重要问题'
ISSUE_COUNT = '1'
ISSUE_SOURCE = '源程序'
DEFAULT_OPINION = '请开发人员确认是否是问题；若是问题请修改XXX；'

# 问题单输出表头（与静态分析过滤模块导出格式保持一致）
OUTPUT_HEADERS = [
    '问题名称', '问题来源', '问题复现', '问题总结', '测试活动', '问题类型',
    '问题等级', '问题个数', '问题名称', '问题来源', '问题描述', '测评单位意见',
]

OUTPUT_DIR = MODULE_UPLOAD_DIR / 'dynamic_issue_report'
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


class DynamicIssueReportService:
    def __init__(self, session_id: str):
        self.session_id = session_id
        self.logger = setup_session_logger(session_id)
        self.ai_client = CaseDesignAIClient(session_id)

    @staticmethod
    def _is_fail_status(value: Any) -> bool:
        text = str(value or '').strip()
        if not text:
            return False
        return '不通过' in text or '未通过' in text or text == '失败'

    def _find_case_sheet(self, workbook) -> Optional[str]:
        for name in workbook.sheetnames:
            if CASE_SHEET_KEYWORD in str(name):
                return name
        return None

    def _collect_failed_cases(self, worksheet) -> "OrderedDict[str, List[Dict[str, Any]]]":
        """收集测试结论为不通过的用例，按测试项名称分组（保持顺序）。"""
        groups: "OrderedDict[str, List[Dict[str, Any]]]" = OrderedDict()
        for row in range(2, (worksheet.max_row or 0) + 1):
            conclusion = worksheet.cell(row=row, column=COL_AA_TEST_CONCLUSION + 1).value
            if not self._is_fail_status(conclusion):
                continue

            test_item_name = str(worksheet.cell(row=row, column=COL_G_TEST_ITEM_NAME + 1).value or '').strip()
            if not test_item_name:
                continue

            groups.setdefault(test_item_name, []).append({
                'row': row,
                'input_condition': str(worksheet.cell(row=row, column=COL_P_INPUT_CONDITION + 1).value or '').strip(),
                'expected_result': str(worksheet.cell(row=row, column=COL_Q_EXPECTED_RESULT + 1).value or '').strip(),
                'test_result': str(worksheet.cell(row=row, column=COL_Z_TEST_RESULT + 1).value or '').strip(),
            })
        return groups

    def _build_prompt(self, test_item_name: str, cases: List[Dict[str, Any]]) -> str:
        case_parts = []
        for i, case in enumerate(cases, 1):
            case_parts.append(
                f'用例{i}：输入条件：{case["input_condition"] or "无"}；'
                f'预期结果：{case["expected_result"] or "无"}；'
                f'测试结果：{case["test_result"] or "无"}'
            )
        case_text = '\n'.join(case_parts) if case_parts else '(无)'

        return (
            '你是测试工程师助手。\n'
            '任务：根据动态测试中“测试结论=不通过”的测试用例，为同一测试项生成问题报告单。\n'
            '你必须只输出 JSON，不能输出任何多余文字、不能使用 Markdown。\n'
            '输出格式固定为：\n'
            '{"issues": [{"issue_name":"...","issue_reproduce":"...","issue_summary":"...",'
            '"requirement_desc":"...","problem_desc":"...","problem_location":"...",'
            '"impact_analysis":"...","test_unit_opinion":"..."}]}\n'
            '规则：\n'
            '1) issues 为该测试项下的问题单列表。通常同一类失效用例（即使有多条）只列 1 个问题单；'
            '仅当确实存在多类不同问题时才输出多条。\n'
            '2) issue_name：简洁的问题名称。\n'
            '3) issue_reproduce：问题复现，简述触发操作/输入与观察到的现象。\n'
            '4) issue_summary：一句话问题总结。\n'
            '5) 问题描述拆分为四个字段：\n'
            '   - requirement_desc：需求描述（该测试项应满足的需求/预期）\n'
            '   - problem_desc：问题描述（实际现象，与预期不符之处）\n'
            '   - problem_location：问题定位（结合输入条件与测试结果推断的可能的软件/功能位置）\n'
            '   - impact_analysis：影响分析\n'
            '6) test_unit_opinion：固定使用“请开发人员确认是否是问题；若是问题请修改XXX；”格式，'
            '其中 XXX 替换为具体要修改的内容。\n'
            '7) 为保证 JSON 合法：所有字符串中禁止出现英文双引号(")，如需引用对象请用《》包裹。\n\n'
            f'测试项名称：{test_item_name}\n\n'
            f'该测试项下不通过的用例：\n{case_text}\n'
        )

    def _build_fallback_issue(self, test_item_name: str, cases: List[Dict[str, Any]]) -> Dict[str, str]:
        """AI 失败时的回退问题单，避免整条测试项缺失。"""
        first = cases[0] if cases else {}
        return {
            'issue_name': f'{test_item_name}测试未通过',
            'issue_reproduce': first.get('test_result', '') or test_item_name,
            'issue_summary': f'{test_item_name}测试未通过',
            'issue_description': (
                f'需求描述：{first.get("expected_result", "") or "待补充"}\n'
                f'问题描述：{first.get("test_result", "") or "待补充"}\n'
                '问题定位：待人工确认\n'
                '影响分析：待人工确认'
            ),
            'test_unit_opinion': DEFAULT_OPINION,
        }

    @staticmethod
    def _assemble_description(issue: Dict[str, Any]) -> str:
        parts = [
            ('需求描述', issue.get('requirement_desc')),
            ('问题描述', issue.get('problem_desc')),
            ('问题定位', issue.get('problem_location')),
            ('影响分析', issue.get('impact_analysis')),
        ]
        lines = [f'{label}：{str(value or "").strip()}' for label, value in parts]
        return '\n'.join(lines)

    def _normalize_issue(self, raw: Dict[str, Any], test_item_name: str, cases: List[Dict[str, Any]]) -> Dict[str, str]:
        issue_name = str(raw.get('issue_name') or '').strip()
        desc_fields = ['requirement_desc', 'problem_desc', 'problem_location', 'impact_analysis']
        if not issue_name or not any(str(raw.get(k) or '').strip() for k in desc_fields):
            return self._build_fallback_issue(test_item_name, cases)
        description = self._assemble_description(raw).strip()

        reproduce = str(raw.get('issue_reproduce') or '').strip()
        if not reproduce:
            reproduce = '；'.join([c.get('test_result', '') for c in cases if c.get('test_result')])[:200]

        summary = str(raw.get('issue_summary') or '').strip() or issue_name
        opinion = str(raw.get('test_unit_opinion') or '').strip() or DEFAULT_OPINION

        return {
            'issue_name': issue_name,
            'issue_reproduce': reproduce,
            'issue_summary': summary,
            'issue_description': description,
            'test_unit_opinion': opinion,
        }

    async def _generate_issues(self, test_item_name: str, cases: List[Dict[str, Any]]) -> Tuple[List[Dict[str, str]], Optional[str]]:
        prompt = self._build_prompt(test_item_name, cases)
        try:
            ai_res = await self.ai_client._call_ai_with_retry_generic(prompt, required_keys=['issues'])
            if not isinstance(ai_res, dict) or not ai_res.get('success'):
                return [], str((ai_res or {}).get('error') or 'AI返回失败')
            data = ai_res.get('data') if isinstance(ai_res.get('data'), dict) else {}
            issues = data.get('issues')
            if not isinstance(issues, list) or not issues:
                return [], 'AI未返回issues'
            normalized = [self._normalize_issue(it, test_item_name, cases) for it in issues if isinstance(it, dict)]
            if not normalized:
                return [], 'AI返回issues结构不合法'
            return normalized, None
        except Exception as e:
            return [], str(e)

    def _write_output(self, issues: List[Dict[str, str]], output_path: Path) -> None:
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = '问题单'

        worksheet.append(OUTPUT_HEADERS)
        header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        normal_alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        for issue in issues:
            worksheet.append([
                issue['issue_name'],
                ISSUE_SOURCE,
                issue['issue_reproduce'],
                issue['issue_summary'],
                TEST_ACTIVITY,
                ISSUE_TYPE,
                ISSUE_LEVEL,
                ISSUE_COUNT,
                issue['issue_name'],
                ISSUE_SOURCE,
                issue['issue_description'],
                issue['test_unit_opinion'],
            ])
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            for cell in row:
                cell.alignment = normal_alignment

        widths = [30, 12, 40, 30, 12, 12, 12, 10, 30, 12, 60, 40]
        for idx, width in enumerate(widths):
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(idx + 1)].width = width
        worksheet.freeze_panes = 'A2'

        workbook.save(str(output_path))
        workbook.close()

    async def process_file(self, file_path: str, original_filename: str) -> Dict[str, Any]:
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
        except Exception as e:
            return {'success': False, 'error': f'无法解析Excel文件: {e}'}

        try:
            sheet_name = self._find_case_sheet(workbook)
            if not sheet_name:
                return {'success': False, 'error': f'未找到包含“{CASE_SHEET_KEYWORD}”的工作表'}

            worksheet = workbook[sheet_name]
            groups = self._collect_failed_cases(worksheet)

            if not groups:
                return {'success': False, 'error': '未找到测试结论为“不通过”的用例'}

            self.logger.info('动态测试问题单：共 %s 个测试项存在不通过用例', len(groups))

            all_issues: List[Dict[str, str]] = []
            errors: List[Dict[str, Any]] = []

            for test_item_name, cases in groups.items():
                generated, err = await self._generate_issues(test_item_name, cases)
                if generated:
                    all_issues.extend(generated)
                else:
                    # AI 失败时回退，保证该测试项仍出现在问题单中
                    all_issues.append(self._build_fallback_issue(test_item_name, cases))
                    errors.append({'test_item_name': test_item_name, 'error': err or 'AI生成失败，已回退'})
                    self.logger.warning('动态测试问题单：测试项 %s 生成失败 %s', test_item_name, err)

            output_filename = f'dynamic_issue_report_{self.session_id}.xlsx'
            self._write_output(all_issues, OUTPUT_DIR / output_filename)

            download_filename = f'{Path(original_filename).stem}_问题单.xlsx'
            self.logger.info(
                '动态测试问题单：生成完成 测试项=%s 问题单=%s errors=%s',
                len(groups), len(all_issues), len(errors),
            )

            return {
                'success': True,
                'sheet_name': sheet_name,
                'test_item_count': len(groups),
                'issue_count': len(all_issues),
                'errors': errors,
                'output_filename': output_filename,
                'download_filename': download_filename,
            }
        except Exception as e:
            self.logger.exception('动态测试问题单：处理文件失败')
            return {'success': False, 'error': f'处理文件时出错: {e}'}
        finally:
            workbook.close()

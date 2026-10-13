import openpyxl
from pathlib import Path
from typing import Dict, List, Any, Tuple, Optional
import uuid
from datetime import datetime

from ..config.config import REQUIRED_SHEET_NAME, REQUIRED_COLUMNS
from ..models.models import AssessmentSession, RequirementItem
from ..database.database import get_db_session
from ..utils.logger import get_session_logger

class ExcelProcessor:
    def __init__(self, session_id: str):
        self.session_id = session_id
        self.logger = get_session_logger(session_id)
        self.required_sheet = REQUIRED_SHEET_NAME
        self.required_columns = REQUIRED_COLUMNS
        
    def process_uploaded_file(self, file_path: str, original_filename: str) -> Tuple[str, Dict[str, Any]]:
        """
        处理上传的Excel文件
        """
        try:
            self.logger.info(f"开始处理Excel文件: {original_filename}")
            
            validation_result = self._validate_excel_file(file_path)
            if not validation_result["success"]:
                self.logger.error(f"Excel文件验证失败: {validation_result['error']}")
                return self.session_id, validation_result
            
            read_result = self._read_excel_data(file_path)
            if not read_result["success"]:
                self.logger.error(f"Excel数据读取失败: {read_result['error']}")
                return self.session_id, read_result
            
            save_result = self._save_to_database(
                self.session_id, 
                file_path, 
                original_filename, 
                read_result["data"]
            )
            
            if save_result["success"]:
                self.logger.info(f"Excel文件处理完成 - 会话ID: {self.session_id}, 数据行数: {len(read_result['data'])}")
            
            return self.session_id, save_result
            
        except Exception as e:
            error_msg = f"处理Excel文件时出错: {str(e)}"
            self.logger.error(error_msg)
            return self.session_id, {"success": False, "error": error_msg}
    
    def _validate_excel_file(self, file_path: str) -> Dict[str, Any]:
        """验证Excel文件格式和内容"""
        try:
            self.logger.info("开始验证Excel文件")
            
            if not Path(file_path).exists():
                return {"success": False, "error": "文件不存在"}
            
            try:
                workbook = openpyxl.load_workbook(file_path, read_only=True)
            except Exception as e:
                return {"success": False, "error": f"无法打开Excel文件: {str(e)}"}
            
            if self.required_sheet not in workbook.sheetnames:
                workbook.close()
                return {
                    "success": False, 
                    "error": f"未找到必需的工作表: {self.required_sheet}。可用工作表: {', '.join(workbook.sheetnames)}"
                }
            
            worksheet = workbook[self.required_sheet]
            
            columns_found = {}
            header_row_found = False
            
            for row_idx in range(1, min(6, worksheet.max_row + 1)):
                row_values = [cell.value for cell in worksheet[row_idx]]
                temp_columns = {}
                
                for required_col in self.required_columns.keys():
                    for col_idx, cell_value in enumerate(row_values):
                        if cell_value and required_col in str(cell_value):
                            temp_columns[required_col] = col_idx
                            break
                
                if len(temp_columns) >= 3:
                    columns_found = temp_columns
                    header_row_found = True
                    break
            
            workbook.close()
            
            if not header_row_found:
                missing_columns = [col for col in self.required_columns.keys() if col not in columns_found]
                return {
                    "success": False,
                    "error": f"未找到必需的列: {', '.join(missing_columns)}。请确保Excel文件包含这些列名。"
                }
            
            self.logger.info(f"Excel文件验证成功，找到列: {list(columns_found.keys())}")
            return {"success": True, "columns_found": columns_found}
            
        except Exception as e:
            error_msg = f"验证Excel文件时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg}
    
    def _read_excel_data(self, file_path: str) -> Dict[str, Any]:
        """读取Excel数据，使用openpyxl"""
        try:
            self.logger.info("开始读取Excel数据")
            
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            worksheet = workbook[self.required_sheet]
            
            header_row_idx = None
            columns_mapping = {}
            
            for row_idx in range(1, min(6, worksheet.max_row + 1)):
                row_values = [cell.value for cell in worksheet[row_idx]]
                temp_mapping = {}
                
                for required_col in self.required_columns.keys():
                    for val_idx, val in enumerate(row_values):
                        if val and required_col in str(val):
                            temp_mapping[required_col] = val_idx
                            break
                
                if len(temp_mapping) >= 3:
                    header_row_idx = row_idx
                    columns_mapping = temp_mapping
                    break
            
            if header_row_idx is None:
                workbook.close()
                return {"success": False, "error": "未找到包含必需列的表头行"}
            
            self.logger.info(f"找到表头行: 第{header_row_idx}行，列映射: {columns_mapping}")
            
            req_desc_col = columns_mapping.get("需求子项描述")
            test_item_col = columns_mapping.get("测试项描述") 
            test_method_col = columns_mapping.get("测试方法")
            adequacy_col = columns_mapping.get("充分性要求")
            pass_criteria_col = columns_mapping.get("通过准则")
            
            data_rows = []
            
            all_cols = [req_desc_col, test_item_col, test_method_col, adequacy_col, pass_criteria_col]
            valid_cols = [col for col in all_cols if col is not None]
            
            if not valid_cols:
                workbook.close()
                return {"success": False, "error": "未找到有效的列映射"}
            
            min_col = min(valid_cols) + 1
            max_col = max(valid_cols) + 1
            
            self.logger.info(f"开始批量读取数据，列范围: {min_col}-{max_col}, 行范围: {header_row_idx + 1}-{worksheet.max_row}")
            
            processed_count = 0
            for row_tuple in worksheet.iter_rows(min_row=header_row_idx + 1, max_row=worksheet.max_row, 
                                              min_col=min_col, max_col=max_col, values_only=True):
                processed_count += 1
                row_idx = header_row_idx + processed_count
                
                if processed_count % 1000 == 0:
                    self.logger.info(f"Excel读取进度: {processed_count}/{worksheet.max_row - header_row_idx}")
                
                test_item_val = ""
                test_method_val = ""
                
                if test_item_col is not None and test_item_col >= min_col - 1 and test_item_col <= max_col - 1:
                    col_idx_in_tuple = test_item_col - (min_col - 1)
                    if col_idx_in_tuple < len(row_tuple) and row_tuple[col_idx_in_tuple] is not None:
                        test_item_val = str(row_tuple[col_idx_in_tuple]).strip()
                
                if test_method_col is not None and test_method_col >= min_col - 1 and test_method_col <= max_col - 1:
                    col_idx_in_tuple = test_method_col - (min_col - 1)
                    if col_idx_in_tuple < len(row_tuple) and row_tuple[col_idx_in_tuple] is not None:
                        test_method_val = str(row_tuple[col_idx_in_tuple]).strip()
                
                if not test_item_val and not test_method_val:
                    continue
                
                def get_col_value(col_idx):
                    if col_idx is not None and col_idx >= min_col - 1 and col_idx <= max_col - 1:
                        col_pos = col_idx - (min_col - 1)
                        if col_pos < len(row_tuple) and row_tuple[col_pos] is not None:
                            return str(row_tuple[col_pos]).strip()
                    return ""
                
                row_data = {
                    "row_index": row_idx,
                    "original_requirement_description": get_col_value(req_desc_col),
                    "test_item_description": test_item_val,
                    "test_method": test_method_val,
                    "adequacy_requirement": get_col_value(adequacy_col),
                    "pass_criteria": get_col_value(pass_criteria_col)
                }
                
                data_rows.append(row_data)
            
            workbook.close()
            
            self.logger.info(f"成功读取Excel数据，共{len(data_rows)}行有效数据")
            return {
                "success": True, 
                "data": data_rows, 
                "columns_found": columns_mapping,
                "header_row": header_row_idx
            }
            
        except Exception as e:
            error_msg = f"读取Excel数据时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg}
    
    def _save_to_database(self, session_id: str, file_path: str, original_filename: str, data_rows: List[Dict]) -> Dict[str, Any]:
        """保存数据到数据库"""
        try:
            self.logger.info(f"开始保存数据到数据库 - 会话ID: {session_id}")
            
            db = get_db_session()
            
            try:
                session = AssessmentSession(
                    session_id=session_id,
                    original_filename=original_filename,
                    file_path=file_path,
                    status="uploaded",
                    total_rows=len(data_rows),
                    processed_rows=0
                )
                db.add(session)
                db.flush()
                
                requirement_items = []
                for row_data in data_rows:
                    item = RequirementItem(
                        session_id=session_id,
                        row_index=row_data["row_index"],
                        original_requirement_description=row_data["original_requirement_description"],
                        test_item_description=row_data["test_item_description"],
                        test_method=row_data["test_method"],
                        generation_status="pending"
                    )
                    requirement_items.append(item)
                
                db.add_all(requirement_items)
                db.commit()
                
                self.logger.info(f"数据保存成功 - 会话: {session_id}, 需求项: {len(requirement_items)}个")
                
                return {
                    "success": True,
                    "session_id": session_id,
                    "total_rows": len(data_rows),
                    "message": f"成功处理{len(data_rows)}行数据"
                }
                
            except Exception as e:
                db.rollback()
                raise e
            finally:
                db.close()
                
        except Exception as e:
            error_msg = f"保存数据到数据库时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg}
    
    def generate_output_excel(self, output_path: str) -> Dict[str, Any]:
        """
        生成输出Excel文件
        """
        try:
            self.logger.info(f"开始生成输出Excel文件: {output_path}")
            
            db = get_db_session()
            
            try:
                session = db.query(AssessmentSession).filter(
                    AssessmentSession.session_id == self.session_id
                ).first()
                
                if not session:
                    return {"success": False, "error": "会话不存在"}
                
                requirement_items = db.query(RequirementItem).filter(
                    RequirementItem.session_id == self.session_id
                ).order_by(RequirementItem.row_index).all()
                
                if not requirement_items:
                    return {"success": False, "error": "没有找到需求项数据"}
                
                workbook = openpyxl.load_workbook(session.file_path)
                worksheet = workbook[self.required_sheet]
                
                columns_mapping = {}
                header_row_idx = None
                
                for row_idx in range(1, min(6, worksheet.max_row + 1)):
                    row_values = [cell.value for cell in worksheet[row_idx]]
                    temp_mapping = {}
                    
                    for required_col in self.required_columns.keys():
                        for col_idx, val in enumerate(row_values):
                            if val and required_col in str(val):
                                temp_mapping[required_col] = col_idx
                                break
                    
                    if len(temp_mapping) >= 3:
                        header_row_idx = row_idx
                        columns_mapping = temp_mapping
                        break
                
                if header_row_idx is None:
                    workbook.close()
                    return {"success": False, "error": "无法找到表头行"}
                
                updated_count = 0
                for item in requirement_items:
                    excel_row_idx = item.row_index
                    
                    # 写入充分性要求
                    if item.adequacy_requirement:
                        adeq_col = columns_mapping.get("充分性要求")
                        if adeq_col is not None:
                            worksheet.cell(row=excel_row_idx, column=adeq_col + 1, 
                                        value=item.adequacy_requirement)
                    
                    # 写入通过准则
                    if item.pass_criteria:
                        pass_col = columns_mapping.get("通过准则")
                        if pass_col is not None:
                            worksheet.cell(row=excel_row_idx, column=pass_col + 1, 
                                        value=item.pass_criteria)
                    
                    updated_count += 1
                
                workbook.save(output_path)
                workbook.close()
                
                self.logger.info(f"输出Excel文件生成成功: {output_path}, 更新了{updated_count}行数据")
                
                return {
                    "success": True,
                    "output_path": output_path,
                    "total_items": len(requirement_items),
                    "updated_items": updated_count
                }
                
            finally:
                db.close()
                
        except Exception as e:
            error_msg = f"生成输出Excel文件时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg} 
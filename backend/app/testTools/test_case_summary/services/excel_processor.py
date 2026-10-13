"""
测试用例总结Excel处理器
"""
import openpyxl
from pathlib import Path
from typing import Dict, List, Any, Tuple, Optional
import uuid
from datetime import datetime
import json
from collections import defaultdict

from ..config.config import REQUIRED_SHEET_NAME, REQUIRED_COLUMNS
from ..models.models import TestCaseSummarySession, TestItem
from ..database.database import get_db_session
from ..utils.logger import get_session_logger

class ExcelProcessor:
    # E列“轮次”中用于识别第二轮用例的关键词
    ROUND2_KEYWORDS = ("二", "2")

    def __init__(self, session_id: str):
        self.session_id = session_id
        self.logger = get_session_logger(session_id)
        self.required_sheet = REQUIRED_SHEET_NAME
        self.required_columns = REQUIRED_COLUMNS
        
    def process_uploaded_file(self, file_path: str, original_filename: str, tester_name: str = '') -> Tuple[str, Dict[str, Any]]:
        """
        处理上传的测试用例Excel文件
        """
        try:
            self.logger.info(f"开始处理测试用例Excel文件: {original_filename}")
            
            # 验证Excel文件
            validation_result = self._validate_excel_file(file_path)
            if not validation_result["success"]:
                self.logger.error(f"Excel文件验证失败: {validation_result['error']}")
                return self.session_id, validation_result
            
            # 读取Excel数据
            read_result = self._read_test_case_data(file_path)
            if not read_result["success"]:
                self.logger.error(f"Excel数据读取失败: {read_result['error']}")
                return self.session_id, read_result
            
            # 保存到数据库
            save_result = self._save_to_database(
                self.session_id, 
                file_path, 
                original_filename, 
                read_result["test_items"],
                tester_name=tester_name
            )
            
            if save_result["success"]:
                self.logger.info(f"测试用例文件处理完成 - 会话ID: {self.session_id}, 测试项数量: {len(read_result['test_items'])}")
            
            return self.session_id, save_result
            
        except Exception as e:
            error_msg = f"处理测试用例Excel文件时出错: {str(e)}"
            self.logger.error(error_msg)
            return self.session_id, {"success": False, "error": error_msg}
    
    def _validate_excel_file(self, file_path: str) -> Dict[str, Any]:
        """验证Excel文件格式"""
        try:
            self.logger.info("开始验证Excel文件")
            
            # 尝试打开Excel文件
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            
            # 检查是否存在所需的工作表
            if self.required_sheet not in workbook.sheetnames:
                available_sheets = ", ".join(workbook.sheetnames)
                return {
                    "success": False, 
                    "error": f"未找到工作表 '{self.required_sheet}'。可用工作表: {available_sheets}"
                }
            
            worksheet = workbook[self.required_sheet]
            
            # 检查文件是否为空
            if worksheet.max_row < 2:
                return {"success": False, "error": "Excel文件为空或只有表头"}
            
            # 检查必要的列是否存在
            max_col_index = max(self.required_columns.values())
            if worksheet.max_column <= max_col_index:
                return {
                    "success": False, 
                    "error": f"Excel文件列数不足，需要至少{max_col_index + 1}列"
                }
            
            workbook.close()
            
            self.logger.info("Excel文件验证成功")
            return {"success": True}
            
        except Exception as e:
            return {"success": False, "error": f"验证Excel文件时出错: {str(e)}"}
    
    def _read_test_case_data(self, file_path: str) -> Dict[str, Any]:
        """读取测试用例数据，合并第二轮结果后按测试项分组"""
        try:
            self.logger.info("开始读取测试用例数据")

            # 使用openpyxl读取Excel文件
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            worksheet = workbook[self.required_sheet]

            max_col_index = max(self.required_columns.values())

            # 逐行读取原始用例记录（保持行序）
            records: List[Dict[str, Any]] = []
            for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    # 检查行是否有足够的列
                    if len(row) <= max_col_index:
                        continue

                    # 提取测试用例信息
                    test_item_name = self._cell_str(row[self.required_columns["测试项名称"]])

                    # 跳过空的测试项名称
                    if not test_item_name or test_item_name == "None":
                        continue

                    records.append({
                        "row_index": row_idx,
                        "test_item_name": test_item_name,
                        "case_no": self._normalize_key(row[self.required_columns["测试用例编号"]]),
                        "round_raw": self._cell_str(row[self.required_columns["轮次"]]),
                        "input_condition": self._cell_str(row[self.required_columns["输入条件"]]),
                        "test_result": self._cell_str(row[self.required_columns["测试结果"]]),
                    })

                except Exception as e:
                    self.logger.warning(f"处理第{row_idx}行数据时出错: {str(e)}")
                    continue

            workbook.close()

            # 合并第二轮结果：用第二轮测试结果覆盖第一轮，并删除第二轮用例
            merged_records = self._merge_round2_results(records)

            # 按测试项名称分组存储测试用例
            test_items_dict = defaultdict(list)
            for rec in merged_records:
                test_items_dict[rec["test_item_name"]].append({
                    "row_index": rec["row_index"],
                    "input_condition": rec["input_condition"],
                    "test_result": rec["test_result"]
                })

            # 转换为最终格式
            test_items = []
            for test_item_name, test_cases in test_items_dict.items():
                test_items.append({
                    "test_item_name": test_item_name,
                    "test_cases": test_cases,
                    "test_cases_count": len(test_cases)
                })

            total_cases = sum(len(item['test_cases']) for item in test_items)
            self.logger.info(f"成功读取测试用例数据，共{len(test_items)}个测试项，{total_cases}个测试用例")

            return {
                "success": True,
                "test_items": test_items,
                "total_test_items": len(test_items)
            }

        except Exception as e:
            error_msg = f"读取测试用例数据时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg}

    def _merge_round2_results(self, records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """
        合并第二轮测试结果

        规则：
        1. 第二轮用例由E列“轮次”标识（如“第二轮”）；
        2. 通过J列“测试用例编号”匹配对应的第一轮用例（编号相同）；
        3. 用第二轮用例Z列“测试结果”覆盖第一轮用例的测试结果；
        4. 匹配上的第二轮用例删除；未匹配上的第二轮用例保留，避免数据丢失。

        说明：用例不通过时结果体现在测试结果列，会在第二轮重新执行一遍，
        第二轮通过的结果与第一轮不同，故以第二轮结果为准生成符合性表。
        """
        # 第一轮（非第二轮）用例，按测试用例编号建立索引
        round1_by_case: Dict[str, Dict[str, Any]] = {}
        for rec in records:
            if not self._is_round2(rec["round_raw"]) and rec["case_no"]:
                round1_by_case.setdefault(rec["case_no"], rec)

        matched_case_nos = set()
        round2_count = 0
        for rec in records:
            if not self._is_round2(rec["round_raw"]):
                continue
            round2_count += 1
            base = round1_by_case.get(rec["case_no"]) if rec["case_no"] else None
            if base is not None:
                # 用第二轮的测试结果覆盖第一轮
                base["test_result"] = rec["test_result"]
                matched_case_nos.add(rec["case_no"])

        # 删除匹配上的第二轮用例，保序保留其余记录
        merged = [
            rec for rec in records
            if not (self._is_round2(rec["round_raw"]) and rec["case_no"] in matched_case_nos)
        ]

        if round2_count:
            self.logger.info(
                f"轮次合并完成 - 第二轮用例: {round2_count}条, 覆盖第一轮: {len(matched_case_nos)}条, "
                f"删除第二轮: {len(matched_case_nos)}条, 未匹配保留: {round2_count - len(matched_case_nos)}条, "
                f"合并前: {len(records)}条, 合并后: {len(merged)}条"
            )

        return merged

    def _is_round2(self, round_value: str) -> bool:
        """判断E列轮次是否为第二轮"""
        text = self._cell_str(round_value)
        if not text:
            return False
        return any(keyword in text for keyword in self.ROUND2_KEYWORDS)

    @staticmethod
    def _cell_str(value: Any) -> str:
        """单元格值转字符串（去除首尾空白）"""
        return "" if value is None else str(value).strip()

    @staticmethod
    def _normalize_key(value: Any) -> str:
        """测试用例编号归一化（避免 1.0 与 1 不匹配）"""
        if value is None:
            return ""
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()
    
    def _save_to_database(self, session_id: str, file_path: str, original_filename: str, test_items: List[Dict], tester_name: str = '') -> Dict[str, Any]:
        """保存数据到数据库"""
        try:
            self.logger.info(f"开始保存数据到数据库 - 会话ID: {session_id}")
            
            db = get_db_session()
            
            try:
                # 创建会话记录
                session = TestCaseSummarySession(
                    session_id=session_id,
                    original_filename=original_filename,
                    file_path=file_path,
                    status="uploaded",
                    total_test_items=len(test_items),
                    processed_test_items=0,
                    tester_name=tester_name if tester_name else None
                )
                db.add(session)
                db.flush()  # 获取ID
                
                # 批量创建测试项记录
                test_item_records = []
                for test_item_data in test_items:
                    test_item = TestItem(
                        session_id=session_id,
                        test_item_name=test_item_data["test_item_name"],
                        test_cases_data=json.dumps(test_item_data["test_cases"], ensure_ascii=False),
                        generation_status="pending"
                    )
                    test_item_records.append(test_item)
                
                db.add_all(test_item_records)
                db.commit()
                
                self.logger.info(f"数据保存成功 - 会话: {session_id}, 测试项: {len(test_item_records)}个")
                
                return {
                    "success": True,
                    "session_id": session_id,
                    "total_test_items": len(test_items),
                    "message": f"成功处理{len(test_items)}个测试项"
                }
                
            except Exception as e:
                db.rollback()
                raise e
            finally:
                db.close()
                
        except Exception as e:
            error_msg = f"保存数据到数据库时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg}
    
    def generate_output_excel(self, output_path: str) -> Dict[str, Any]:
        """
        生成输出Excel文件
        """
        try:
            self.logger.info(f"开始生成输出Excel文件: {output_path}")
            
            db = get_db_session()
            
            try:
                # 获取会话信息
                session = db.query(TestCaseSummarySession).filter(
                    TestCaseSummarySession.session_id == self.session_id
                ).first()
                
                if not session:
                    return {"success": False, "error": "会话不存在"}
                
                # 获取所有测试项
                test_items = db.query(TestItem).filter(
                    TestItem.session_id == self.session_id
                ).order_by(TestItem.id).all()
                
                if not test_items:
                    return {"success": False, "error": "没有找到测试项数据"}
                
                # 创建新的Excel文件
                workbook = openpyxl.Workbook()
                worksheet = workbook.active
                worksheet.title = "测试项总结"
                
                # 写入表头
                worksheet.cell(row=1, column=1, value="测试项名称")
                worksheet.cell(row=1, column=2, value="AI总结结果")
                
                # 写入数据
                updated_count = 0
                for idx, test_item in enumerate(test_items, start=2):
                    worksheet.cell(row=idx, column=1, value=test_item.test_item_name)
                    worksheet.cell(row=idx, column=2, value=test_item.ai_summary or "")
                    updated_count += 1
                
                # 保存文件
                workbook.save(output_path)
                workbook.close()
                
                self.logger.info(f"输出Excel文件生成成功: {output_path}, 更新了{updated_count}个测试项")
                
                return {
                    "success": True,
                    "output_path": output_path,
                    "total_items": len(test_items),
                    "updated_items": updated_count
                }
                
            finally:
                db.close()
                
        except Exception as e:
            error_msg = f"生成输出Excel文件时出错: {str(e)}"
            self.logger.error(error_msg)
            return {"success": False, "error": error_msg} 
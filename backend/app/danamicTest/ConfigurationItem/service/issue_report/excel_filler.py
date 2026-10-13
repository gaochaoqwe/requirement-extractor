"""
配置项测试问题报告单Excel填充器

负责将配置项测试问题报告单数据填充到Excel工作表中。
"""

import logging
import openpyxl
from typing import Dict, List, Optional, Any
from pathlib import Path

# 配置日志
logger = logging.getLogger(__name__)

class ConfigurationItemIssueReportExcelFiller:
    """配置项测试问题报告单Excel填充器"""
    
    def __init__(self):
        """初始化Excel填充器"""
        self.logger = logger
        self.column_mapping = self._get_column_mapping()
    
    def fill_issue_report_sheet(
        self,
        excel_file_path: str,
        sheet_name: str,
        issue_report_data: Dict[str, Any]
    ) -> bool:
        """
        填充配置项测试问题报告单工作表
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            issue_report_data: 问题报告数据字典
            
        Returns:
            bool: 填充是否成功
        """
        try:
            self.logger.debug(f"开始填充配置项测试问题报告单: {sheet_name}")
            
            # 打开Excel工作簿
            wb = openpyxl.load_workbook(excel_file_path)
            
            # 获取工作表
            if sheet_name not in wb.sheetnames:
                self.logger.error(f"工作表 '{sheet_name}' 不存在")
                return False
            
            sheet = wb[sheet_name]
            
            # 填充数据
            self._fill_issue_data(sheet, issue_report_data)
            
            # 保存文件
            wb.save(excel_file_path)
            
            self.logger.debug("配置项测试问题报告单填充完成")
            return True
            
        except Exception as e:
            self.logger.error(f"填充配置项测试问题报告单时发生错误: {str(e)}")
            return False
    
    def _fill_issue_data(self, sheet, issue_report_data: Dict[str, Any]) -> None:
        """填充问题数据"""
        basic_info = issue_report_data.get('basic_info', {})
        issue_numbering = issue_report_data.get('issue_numbering', {})
        issue_classification = issue_report_data.get('issue_classification', [])
        time_info = issue_report_data.get('time_info', {})
        description_templates = issue_report_data.get('description_templates', {})
        fixed_values = issue_report_data.get('fixed_values', {})
        issue_count = issue_report_data.get('issue_count', 16)
        
        # 获取问题编号列表
        issue_numbers = issue_numbering.get('issue_numbers', [])
        
        # 从第2行开始填充数据
        start_row = 2
        
        for i in range(issue_count):
            current_row = start_row + i
            classification = issue_classification[i] if i < len(issue_classification) else {}
            issue_number = issue_numbers[i] if i < len(issue_numbers) else f"Issue-{i+1}"
            
            # 填充基本信息
            self._fill_basic_issue_info(sheet, current_row, i+1, basic_info, issue_number)
            
            # 填充分类信息
            self._fill_classification_info(sheet, current_row, classification, description_templates)
            
            # 填充固定值
            self._fill_fixed_values(sheet, current_row, fixed_values)
            
            # 填充时间信息
            self._fill_time_info(sheet, current_row, time_info, basic_info)
            
            self.logger.debug(f"填充第{i+1}个问题完成，行号: {current_row}")
    
    def _fill_basic_issue_info(
        self, 
        sheet, 
        row: int, 
        issue_index: int, 
        basic_info: Dict[str, Any], 
        issue_number: str
    ) -> None:
        """填充基本问题信息"""
        # 1. 编号
        sheet.cell(row=row, column=1, value=issue_index)
        
        # 2. 软件配置项名称
        sheet.cell(row=row, column=2, value=basic_info.get('software_name', ''))
        
        # 3. 被测软件版本
        sheet.cell(row=row, column=3, value=basic_info.get('software_version', ''))
        
        # 4. 测试轮次
        sheet.cell(row=row, column=4, value='第一轮')
        
        # 5. 问题报告单编号
        sheet.cell(row=row, column=5, value=issue_number)
        
        self.logger.debug(f"基本问题信息填充完成: 编号={issue_index}, 软件名称={basic_info.get('software_name', '')}")
    
    def _fill_classification_info(
        self, 
        sheet, 
        row: int, 
        classification: Dict[str, str], 
        description_templates: Dict[str, str]
    ) -> None:
        """填充分类信息"""
        # 9. 测试活动
        test_activity = classification.get('test_activity', '')
        sheet.cell(row=row, column=9, value=test_activity)
        
        # 10. 问题类型
        issue_type = classification.get('issue_type', '')
        sheet.cell(row=row, column=10, value=issue_type)
        
        # 15. 问题描述（根据类型选择模板）
        description_type = classification.get('description_type', 'code')
        description_template = description_templates.get(description_type, description_templates.get('code', ''))
        sheet.cell(row=row, column=15, value=description_template)
        
        self.logger.debug(f"分类信息填充完成: 测试活动={test_activity}, 问题类型={issue_type}")
    
    def _fill_fixed_values(self, sheet, row: int, fixed_values: Dict[str, str]) -> None:
        """填充固定值"""
        # 6. 是否关联历史未归零问题
        sheet.cell(row=row, column=6, value=fixed_values.get('is_historical_issue', '否'))
        
        # 7. 关联版本
        sheet.cell(row=row, column=7, value=fixed_values.get('related_version', ''))
        
        # 8. 关联问题报告单编号
        sheet.cell(row=row, column=8, value=fixed_values.get('related_issue_number', ' '))
        
        # 11. 问题等级
        sheet.cell(row=row, column=11, value=fixed_values.get('issue_level', '一般问题'))
        
        # 12. 问题数目
        sheet.cell(row=row, column=12, value=fixed_values.get('issue_count', '1'))
        
        # 13. 问题名称
        sheet.cell(row=row, column=13, value=fixed_values.get('issue_name', ''))
        
        # 14. 问题来源
        sheet.cell(row=row, column=14, value=fixed_values.get('issue_source', ' '))
        
        # 16. 测评单位意见
        sheet.cell(row=row, column=16, value=fixed_values.get('evaluation_opinion', '请开发人员'))
        
        # 17. 问题修复优先级
        sheet.cell(row=row, column=17, value=fixed_values.get('fix_priority', '立即修复'))
        
        # 20. 研制单位意见
        sheet.cell(row=row, column=20, value=fixed_values.get('development_opinion', '确认是问题，已经在'))
        
        # 21. 问题修复时间
        sheet.cell(row=row, column=21, value=fixed_values.get('fix_time', '立即修复'))
        
        # 24. 主任设计师/总师意见
        sheet.cell(row=row, column=24, value=fixed_values.get('designer_opinion', ''))
        
        # 26. 问题处理结果
        sheet.cell(row=row, column=26, value=fixed_values.get('processing_result', '已归零'))
        
        # 27. 问题确认意见/归零说明
        sheet.cell(row=row, column=27, value=fixed_values.get('confirmation_opinion', '开发人员已经在，问题已归零'))
        
        # 31. 是否典型问题
        sheet.cell(row=row, column=31, value=fixed_values.get('is_typical_issue', '否'))
        
        # 32. 是否低级错误
        sheet.cell(row=row, column=32, value=fixed_values.get('is_low_level_error', '否'))
    
    def _fill_time_info(
        self, 
        sheet, 
        row: int, 
        time_info: Dict[str, str], 
        basic_info: Dict[str, Any]
    ) -> None:
        """填充时间信息"""
        ceshi_time = time_info.get('test_time', '')
        ceshi_nianyue = time_info.get('completion_year_month', '')
        tester_name = basic_info.get('tester_name', '')
        developer_name = basic_info.get('developer_name', '')
        software_version = basic_info.get('software_version', '')
        
        # 18. 测试人员
        sheet.cell(row=row, column=18, value=tester_name)
        
        # 19. 测试人员签名时间
        sheet.cell(row=row, column=19, value=ceshi_time)
        
        # 22. 研发人员
        sheet.cell(row=row, column=22, value=developer_name)
        
        # 23. 研发人员签名时间
        sheet.cell(row=row, column=23, value=ceshi_time)
        
        # 25. 完成年月
        sheet.cell(row=row, column=25, value=ceshi_nianyue)
        
        # 28. 确认人
        sheet.cell(row=row, column=28, value=tester_name)
        
        # 29. 确认时间
        sheet.cell(row=row, column=29, value=ceshi_time)
        
        # 30. 问题验证版本
        sheet.cell(row=row, column=30, value=software_version)
        
        self.logger.debug(f"时间信息填充完成: 测试时间={ceshi_time}, 完成年月={ceshi_nianyue}")
    
    def _get_column_mapping(self) -> Dict[str, int]:
        """获取列映射关系"""
        return {
            'issue_number': 1,                    # 编号
            'software_name': 2,                   # 软件配置项名称
            'software_version': 3,                # 被测软件版本
            'test_round': 4,                      # 测试轮次
            'report_number': 5,                   # 问题报告单编号
            'is_historical_issue': 6,             # 是否关联历史未归零问题
            'related_version': 7,                 # 关联版本
            'related_issue_number': 8,            # 关联问题报告单编号
            'test_activity': 9,                   # 测试活动
            'issue_type': 10,                     # 问题类型
            'issue_level': 11,                    # 问题等级
            'issue_count': 12,                    # 问题数目
            'issue_name': 13,                     # 问题名称
            'issue_source': 14,                   # 问题来源
            'issue_description': 15,              # 问题描述
            'evaluation_opinion': 16,             # 测评单位意见
            'fix_priority': 17,                   # 问题修复优先级
            'tester_name': 18,                    # 测试人员
            'tester_sign_time': 19,               # 测试人员签名时间
            'development_opinion': 20,            # 研制单位意见
            'fix_time': 21,                       # 问题修复时间
            'developer_name': 22,                 # 研发人员
            'developer_sign_time': 23,            # 研发人员签名时间
            'designer_opinion': 24,               # 主任设计师/总师意见
            'completion_year_month': 25,          # 完成年月
            'processing_result': 26,              # 问题处理结果
            'confirmation_opinion': 27,           # 问题确认意见/归零说明
            'confirmer': 28,                      # 确认人
            'confirmation_time': 29,              # 确认时间
            'verification_version': 30,           # 问题验证版本
            'is_typical_issue': 31,               # 是否典型问题
            'is_low_level_error': 32              # 是否低级错误
        }
    
    def get_filled_data_summary(self, excel_file_path: str, sheet_name: str, issue_count: int = 16) -> Dict[str, Any]:
        """获取已填充数据的摘要"""
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            summary = {
                'total_issues': issue_count,
                'filled_issues': 0,
                'issue_details': []
            }
            
            start_row = 2
            for i in range(issue_count):
                current_row = start_row + i
                issue_detail = {}
                
                for field, col in self.column_mapping.items():
                    cell_value = sheet.cell(row=current_row, column=col).value
                    issue_detail[field] = str(cell_value) if cell_value else ""
                
                # 检查是否已填充（至少有软件名称和问题编号）
                if issue_detail.get('software_name') and issue_detail.get('report_number'):
                    summary['filled_issues'] += 1
                
                summary['issue_details'].append(issue_detail)
            
            return summary
            
        except Exception as e:
            self.logger.error(f"获取填充数据摘要时发生错误: {str(e)}")
            return {'total_issues': 0, 'filled_issues': 0, 'issue_details': []}
    
    def update_single_issue(
        self,
        excel_file_path: str,
        sheet_name: str,
        issue_index: int,
        field_name: str,
        field_value: str
    ) -> bool:
        """
        更新单个问题的字段
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            issue_index: 问题索引（从1开始）
            field_name: 字段名称
            field_value: 字段值
            
        Returns:
            bool: 更新是否成功
        """
        try:
            if field_name not in self.column_mapping:
                self.logger.error(f"未知的字段名称: {field_name}")
                return False
            
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            row = 1 + issue_index  # 第1行是标题，数据从第2行开始
            col = self.column_mapping[field_name]
            
            sheet.cell(row=row, column=col, value=field_value)
            wb.save(excel_file_path)
            
            self.logger.info(f"更新问题{issue_index}的字段{field_name}完成")
            return True
            
        except Exception as e:
            self.logger.error(f"更新单个问题字段时发生错误: {str(e)}")
            return False

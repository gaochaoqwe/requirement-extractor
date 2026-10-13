#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
软件信息提取器

使用python-docx库进行Word文档表格识别和软件信息提取
简洁、可靠的表格识别方案
"""

import os
import re
from typing import Dict, Optional, Any

try:
    import docx
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

from ....utils.logging_utils.logger import log_info, log_error, log_warning, log_debug


class SoftwareInfoExtractor:
    """软件信息提取器 - 基于python-docx"""
    
    def __init__(self, debug_mode: bool = False):
        """
        初始化提取器
        
        Args:
            debug_mode: 是否启用调试模式
        """
        self.debug_mode = debug_mode
        
        # 检查python-docx是否可用
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx库未安装，请运行: pip install python-docx")
        
        # 表头关键词映射（按优先级排序，精确匹配优先）
        self.header_keywords = {
            'software_name': ['软件名称', '软件名', '名称', 'software name', 'name'],  # ✅ 移除'软件简称'
            'software_short_name': ['软件简称', 'short name'],  # ✅ 新增：软件简称单独识别
            'software_version': ['版本号', '版本', '发布号', '软件版本', 'version', 'software version', 'release'],
            'software_id': ['软件标识号', '软件标识', '标识号', '标识', 'software id', 'identifier', 'id']
        }
        
        self._debug_print("SoftwareInfoExtractor 初始化完成")
    
    def _debug_print(self, message: str, category: str = "SoftwareExtractor"):
        """调试输出"""
        if self.debug_mode:
            print(f"[{category}] {message}")
            log_debug(message, category)
    
    def extract_software_info(self, file_path: str) -> Dict[str, Any]:
        """
        从Word文档中提取软件信息
        
        Args:
            file_path: Word文档路径
            
        Returns:
            包含提取结果的字典
        """
        try:
            self._debug_print(f"开始处理文档: {file_path}")
            
            # 检查文件是否存在
            if not os.path.exists(file_path):
                return {
                    'found': False,
                    'message': f'文件不存在: {file_path}',
                    'data': {},
                    'table_index': None
                }
            
            # 打开文档
            try:
                doc = Document(file_path)
                self._debug_print(f"文档已打开，段落数: {len(doc.paragraphs)}")
            except Exception as e:
                self._debug_print(f"无法打开文档: {e}")
                return {
                    'found': False,
                    'message': f'无法打开文档: {e}',
                    'data': {},
                    'table_index': None
                }
            
            # 查找所有表格
            tables = doc.tables
            self._debug_print(f"找到 {len(tables)} 个表格")
            
            if not tables:
                return {
                    'found': False,
                    'message': '文档中未找到任何表格',
                    'data': {},
                    'table_index': None
                }
            
            # 遍历表格寻找软件标识表格
            for table_index, table in enumerate(tables):
                self._debug_print(f"分析表格 {table_index + 1}")
                
                # 检查表格是否为标识表格
                if self._is_identity_table(table, table_index + 1):
                    # 提取软件信息
                    software_info = self._extract_info_from_table(table, table_index + 1)
                    
                    if software_info:
                        self._debug_print(f"成功从表格 {table_index + 1} 提取软件信息")
                        return {
                            'found': True,
                            'message': f'成功从表格 {table_index + 1} 提取软件信息',
                            'data': software_info,
                            'table_index': table_index + 1
                        }
            
            return {
                'found': False,
                'message': '未找到包含软件标识信息的表格',
                'data': {},
                'table_index': None
            }
            
        except Exception as e:
            self._debug_print(f"处理文档时发生错误: {e}")
            log_error(f"文档处理错误: {e}", "SoftwareExtractor")
            return {
                'found': False,
                'message': f'处理文档时发生错误: {e}',
                'data': {},
                'table_index': None
            }
    
    def _is_identity_table(self, table, table_index: int) -> bool:
        """
        判断表格是否为软件标识表格
        
        Args:
            table: python-docx表格对象
            table_index: 表格索引
            
        Returns:
            是否为标识表格
        """
        try:
            if not table.rows:
                self._debug_print(f"表格 {table_index} 没有行")
                return False
            
            # 获取表头行
            header_row = table.rows[0]
            headers = [cell.text.strip() for cell in header_row.cells]
            
            self._debug_print(f"表格 {table_index} 表头: {headers}")
            
            # 检查关键字段
            found_fields = {
                'software_name': False,
                'software_version': False,
                'software_id': False
            }
            
            # 合并表头文本进行检查
            header_text = ' '.join(headers).lower()
            
            # ✅ 优先检查核心字段（software_name, software_version, software_id）
            core_fields = ['software_name', 'software_version', 'software_id']
            for field_type in core_fields:
                if field_type in self.header_keywords:
                    for keyword in self.header_keywords[field_type]:
                        if keyword.lower() in header_text:
                            found_fields[field_type] = True
                            self._debug_print(f"  找到 {field_type} 关键词: {keyword}")
                            break
            
            # ✅ 如果没有找到软件名称但找到软件简称，也认为有效
            if not found_fields['software_name']:
                for keyword in self.header_keywords.get('software_short_name', []):
                    if keyword.lower() in header_text:
                        found_fields['software_name'] = True
                        self._debug_print(f"  找到软件简称关键词（视为软件名称）: {keyword}")
                        break
            
            # 至少需要2个关键字段才认为是标识表格
            matched_fields = sum(found_fields.values())
            is_identity = matched_fields >= 2
            
            self._debug_print(f"表格 {table_index} 匹配字段数: {matched_fields}, 是否为标识表格: {is_identity}")
            
            return is_identity
            
        except Exception as e:
            self._debug_print(f"检查表格 {table_index} 时出错: {e}")
            return False
    
    def _extract_info_from_table(self, table, table_index: int) -> Optional[Dict[str, str]]:
        """
        从表格中提取软件信息
        
        Args:
            table: python-docx表格对象
            table_index: 表格索引
            
        Returns:
            提取的软件信息字典
        """
        try:
            if len(table.rows) < 2:
                self._debug_print(f"表格 {table_index} 数据行不足")
                return None
            
            # 获取表头和数据行
            header_row = table.rows[0]
            data_row = table.rows[1]
            
            headers = [cell.text.strip() for cell in header_row.cells]
            data = [cell.text.strip() for cell in data_row.cells]
            
            self._debug_print(f"表格 {table_index} 表头: {headers}")
            self._debug_print(f"表格 {table_index} 数据: {data}")
            
            software_info = {}
            
            # 查找每个字段对应的列（优先精确匹配）
            for i, header in enumerate(headers):
                if i < len(data) and data[i]:
                    value = data[i].strip()
                    
                    # ✅ 优先检查软件简称（精确匹配，避免与软件名称混淆）
                    if any(keyword in header for keyword in self.header_keywords.get('software_short_name', [])):
                        # 只在没有提取到软件名称时才用简称代替
                        if 'software_name' not in software_info:
                            software_info['software_short_name'] = value
                        self._debug_print(f"  发现软件简称列: {value}")
                        continue  # 跳过，避免被当作软件名称
                    
                    # 检查软件名称（排除了"软件简称"关键词）
                    if any(keyword in header for keyword in self.header_keywords['software_name']):
                        software_info['software_name'] = value
                        self._debug_print(f"  提取软件名称: {value}")
                    
                    # 检查版本号
                    elif any(keyword in header for keyword in self.header_keywords['software_version']):
                        software_info['software_version'] = value
                        self._debug_print(f"  提取版本号: {value}")
                    
                    # 检查标识号
                    elif any(keyword in header for keyword in self.header_keywords['software_id']):
                        software_info['software_id'] = value
                        self._debug_print(f"  提取标识号: {value}")
            
            # 验证提取结果
            if len(software_info) >= 2:
                # ✅ 提取系统名称和系统简称
                self._extract_system_names(software_info)
                
                self._debug_print(f"成功提取软件信息: {software_info}")
                return software_info
            else:
                self._debug_print(f"提取的信息不足: {software_info}")
                return None
                
        except Exception as e:
            self._debug_print(f"从表格 {table_index} 提取信息时出错: {e}")
            return None
    
    def _extract_system_names(self, software_info: Dict[str, str]) -> None:
        """
        提取系统名称和系统简称
        
        规则：
        1. 系统名称提取：
           - 优先：如果有软件简称，系统名称 = 软件名称 - 软件简称
           - 备用：查找关键字["设备", "系统", "吊舱", "搜索仪"]，从开头截取到关键字（包括）
        
        2. 系统简称提取：
           - 系统简称 = 系统名称 - 前面的数字字母部分
           - 前缀模式：数字、字母、"/"、"-" 的组合
        
        Args:
            software_info: 软件信息字典，会在原地修改
        """
        try:
            software_name = software_info.get('software_name', '')
            software_short_name = software_info.get('software_short_name', '')
            
            if not software_name:
                self._debug_print("软件名称为空，无法提取系统名称")
                return
            
            system_name = ''
            
            # 规则1：优先使用软件简称
            if software_short_name:
                # 系统名称 = 软件名称去掉软件简称
                if software_short_name in software_name:
                    system_name = software_name.replace(software_short_name, '').strip()
                    self._debug_print(f"使用软件简称提取系统名称: '{software_name}' - '{software_short_name}' = '{system_name}'")
            
            # 规则2：备用方案 - 使用关键字
            if not system_name:
                keywords = ['设备', '系统', '吊舱', '搜索仪']
                for keyword in keywords:
                    if keyword in software_name:
                        # 找到关键字位置，提取到关键字（包括）
                        keyword_pos = software_name.index(keyword)
                        system_name = software_name[:keyword_pos + len(keyword)]
                        self._debug_print(f"使用关键字 '{keyword}' 提取系统名称: '{system_name}'")
                        break
            
            # 如果两种方法都没有提取到，使用软件名称本身
            if not system_name:
                system_name = software_name
                self._debug_print(f"使用软件名称作为系统名称: '{system_name}'")
            
            # 提取系统简称：去掉前面的数字字母部分
            system_short_name = self._extract_system_short_name(system_name)
            
            # 更新软件信息
            software_info['system_name'] = system_name
            software_info['system_short_name'] = system_short_name
            
            self._debug_print(f"系统名称: '{system_name}'")
            self._debug_print(f"系统简称: '{system_short_name}'")
            
        except Exception as e:
            self._debug_print(f"提取系统名称时出错: {e}")
            log_error(f"提取系统名称失败: {e}", "SoftwareExtractor")
    
    def _extract_system_short_name(self, system_name: str) -> str:
        """
        从系统名称中提取系统简称
        
        规则：去掉前面由数字、字母、"/"、"-" 组成的前缀
        例如：
        - "D-3/BBTEEE超级搜索雷达" -> "超级搜索雷达"
        - "DDBBTEEE超级搜索雷达" -> "超级搜索雷达"
        - "超级搜索雷达" -> "超级搜索雷达"
        
        Args:
            system_name: 系统名称
            
        Returns:
            系统简称
        """
        if not system_name:
            return ''
        
        # 匹配前缀：由数字、字母、"/"、"-" 组成的部分
        # 使用正则表达式：^[0-9A-Za-z\-/]+
        match = re.match(r'^[0-9A-Za-z\-/]+', system_name)
        
        if match:
            prefix = match.group()
            # 去掉前缀后剩下的就是系统简称
            system_short_name = system_name[len(prefix):].strip()
            
            # 如果去掉前缀后为空，使用原系统名称
            if not system_short_name:
                system_short_name = system_name
            
            self._debug_print(f"系统名称 '{system_name}' 前缀 '{prefix}' -> 系统简称 '{system_short_name}'")
        else:
            # 没有匹配到前缀，直接使用系统名称
            system_short_name = system_name
            self._debug_print(f"系统名称 '{system_name}' 无前缀 -> 系统简称 '{system_short_name}'")
        
        return system_short_name


def extract_software_info_from_document(file_path: str, debug_mode: bool = False) -> Dict[str, Any]:
    """
    从Word文档中提取软件信息的便捷函数
    
    Args:
        file_path: Word文档路径
        debug_mode: 是否启用调试模式
        
    Returns:
        包含提取结果的字典
    """
    extractor = SoftwareInfoExtractor(debug_mode=debug_mode)
    return extractor.extract_software_info(file_path) 
"""
测试用例服务模块

负责测试用例数据的提取、存储和管理功能
"""

import openpyxl
from typing import List, Dict, Any, Optional
from pathlib import Path


class TestCaseService:
    """测试用例服务类"""
    
    def __init__(self):
        """初始化测试用例服务"""
        pass
    
    def extract_and_store_excel_data_to_db(
        self, 
        excel_file_path: str, 
        sheet_name: str = '测试用例单',
        task_id: str = None
    ) -> bool:
        """
        从测试用例Excel表中提取数据并存储到数据库
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            task_id: 任务ID
            
        Returns:
            bool: 是否成功
        """
        try:
            from ...db.optimized_models import OptimizedDynamicTestDatabase
            from ...utils.logging_utils.logger import log_info, log_warning, log_error
            
            log_info(f"开始从Excel提取测试用例数据: {excel_file_path}", "test_case_service")
            
            # 打开Excel文件
            wb = openpyxl.load_workbook(excel_file_path)
            if sheet_name not in wb.sheetnames:
                log_error(f"工作表 {sheet_name} 不存在", "test_case_service")
                return False
            
            sheet = wb[sheet_name]
            
            # 初始化数据库
            db = OptimizedDynamicTestDatabase()
            
            # 执行数据库迁移（如果需要）
            db.migrate_test_cases_table()
            
            # 生成任务ID（如果未提供）
            if not task_id:
                import uuid
                task_id = str(uuid.uuid4())
                log_info(f"生成新的任务ID: {task_id}", "test_case_service")
            
            # 定义Excel列映射（与test_case_form_handler.py中的write_test_cases_to_excel保持一致）
            # A列到AD列，共30列
            excel_column_mapping = {
                1: 'test_case_id',                    # A: 测试用例ID
                2: 'sequence_number',                 # B: 序号
                3: 'software_item_name',              # C: 软件配置项名称
                4: 'software_version',                # D: 被测软件版本
                5: 'test_round',                      # E: 测试轮次
                6: 'requirement_traceability',        # F: 需求追溯
                7: 'test_item_name',                  # G: 测试项名称
                8: 'test_item_identifier',            # H: 测试项标识
                9: 'test_case_name',                  # I: 测试用例名称
                10: 'test_case_number',               # J: 测试用例编号
                11: 'test_type',                      # K: 测试类型
                12: 'test_description',               # L: 测试描述
                13: 'preconditions',                  # M: 预置条件
                14: 'prerequisites_constraints',      # N: 前提与约束
                15: 'input_actions',                  # O: 输入/动作
                16: 'input_conditions',               # P: 输入条件
                17: 'expected_results',               # Q: 预期结果
                18: 'pass_criteria',                  # R: 通过准则
                19: 'case_designer',                  # S: 用例设计人
                20: 'reviewer',                       # T: 校对人
                21: 'design_time',                    # U: 设计时间
                22: 'review_time',                    # V: 校对时间
                23: 'test_personnel',                 # W: 测试人员
                24: 'supervisor',                     # X: 监督人员
                25: 'test_time',                      # Y: 测试时间
                26: 'test_result',                    # Z: 测试结果
                27: 'test_conclusion',                # AA: 测试结论
                28: 'issue_report_number',            # AB: 问题报告单编号
                29: 'execution_year_month',           # AC: 执行年月
                30: 'issues_found_count'              # AD: 用例发现问题数
            }
            
            # 提取数据
            extracted_data = []
            row_num = 2  # 从第2行开始（第1行是表头）
            
            while True:
                # 检查是否有数据（以第2列序号是否有值为准）
                sequence_cell = sheet.cell(row=row_num, column=2)
                if not sequence_cell.value:
                    break
                
                # 提取该行的所有数据
                row_data = {'requirement_id': f"test_req_{task_id}_{row_num}"}
                
                for col_num, field_name in excel_column_mapping.items():
                    cell_value = sheet.cell(row=row_num, column=col_num).value
                    
                    # 处理不同类型的值
                    if cell_value is not None:
                        if field_name in ['sequence_number', 'issues_found_count']:
                            # 数字字段转换为整数
                            try:
                                row_data[field_name] = int(cell_value)
                            except (ValueError, TypeError):
                                if field_name == 'sequence_number':
                                    row_data[field_name] = row_num - 1  # 使用行号作为序号
                                else:
                                    row_data[field_name] = 0
                        else:
                            # 其他字段转换为字符串
                            row_data[field_name] = str(cell_value).strip()
                    else:
                        row_data[field_name] = None
                
                extracted_data.append(row_data)
                row_num += 1
            
            log_info(f"从Excel提取到 {len(extracted_data)} 行测试用例数据", "test_case_service")
            
            # 存储到数据库
            if extracted_data:
                db.save_test_cases(task_id, extracted_data)
                log_info(f"成功存储 {len(extracted_data)} 条测试用例数据到数据库", "test_case_service")
                
                # 记录详细的存储信息
                for i, data in enumerate(extracted_data, 1):
                    log_info(f"  第{i}行: {data.get('test_case_name', 'N/A')} - {data.get('test_description', 'N/A')[:50]}...", "test_case_service")
            
            return True
            
        except Exception as e:
            log_error(f"从Excel提取并存储测试用例数据失败: {str(e)}", "test_case_service")
            import traceback
            traceback.print_exc()
            return False
    
    def get_test_cases_by_task(self, task_id: str) -> List[Dict[str, Any]]:
        """根据任务ID获取测试用例数据"""
        try:
            from ...db.optimized_models import OptimizedDynamicTestDatabase
            db = OptimizedDynamicTestDatabase()
            return db.get_test_cases_by_task(task_id)
        except Exception as e:
            print(f"获取测试用例数据失败: {e}")
            return []
    
    def get_all_test_cases(self) -> List[Dict[str, Any]]:
        """获取所有测试用例数据"""
        try:
            from ...db.optimized_models import OptimizedDynamicTestDatabase
            db = OptimizedDynamicTestDatabase()
            return db.get_all_test_cases()
        except Exception as e:
            print(f"获取所有测试用例数据失败: {e}")
            return []
    
    def update_test_case_field(self, test_case_id: int, field_name: str, field_value: Any) -> bool:
        """更新测试用例字段"""
        try:
            from ...db.optimized_models import OptimizedDynamicTestDatabase
            db = OptimizedDynamicTestDatabase()
            return db.update_test_case_field(test_case_id, field_name, field_value)
        except Exception as e:
            print(f"更新测试用例字段失败: {e}")
            return False 
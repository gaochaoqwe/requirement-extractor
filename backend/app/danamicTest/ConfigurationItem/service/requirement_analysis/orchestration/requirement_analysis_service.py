"""
配置项测试需求分析表服务

提供配置项测试需求分析表的完整服务接口，整合数据处理、Excel填充、格式设置和读取功能。
"""

import logging
from typing import Dict, List, Optional, Any, Tuple
from pathlib import Path

# 使用真实实现类
from ..test_generation.functional.functional_test_creator import (
    ConfigurationItemRequirementAnalysisDataProcessor,
)
from ..excel_operations.excel_main_filler import (
    ConfigurationItemRequirementAnalysisExcelFiller,
)
from ..excel_operations.excel_reader import (
    ConfigurationItemRequirementAnalysisReader,
)
from ..excel_operations.formatting.requirement_formatter import (
    RequirementFormatter as ConfigurationItemRequirementAnalysisFormatter,
)
from ..test_generation.content_generators.test_content_generator import (
    ConfigurationItemTestContentGenerator,
)
from ..content_analysis.interface_analyzer import InterfaceTestExtractor
from ..document_parsing.parsers.interface_table_parser import InterfaceTableParser
from ..test_generation.functional.interface_test_template_generator import InterfaceTestTemplateGenerator  # 🔧 新增：接口测试模板生成器

# 导入实时处理服务
try:
    from .realtime_processor import RealtimeRequirementAnalysisService
except ImportError:
    try:
        from ..service_realtime import RealtimeRequirementAnalysisService
    except ImportError:
        print("Warning: Could not import RealtimeRequirementAnalysisService")
        class RealtimeRequirementAnalysisService:
            def __init__(self): pass

# 导入任务日志函数
from ....utils.logging_utils.logger import log_info, log_debug, log_error, log_warning

# 配置日志
logger = logging.getLogger(__name__)

class ConfigurationItemRequirementAnalysisService:
    """配置项测试需求分析表服务"""
    
    def __init__(self):
        """初始化配置项测试需求分析表服务"""
        self.logger = logger
        self.data_processor = ConfigurationItemRequirementAnalysisDataProcessor()
        self.excel_filler = ConfigurationItemRequirementAnalysisExcelFiller()
        self.formatter = ConfigurationItemRequirementAnalysisFormatter()
        self.reader = ConfigurationItemRequirementAnalysisReader()
        self.test_content_generator = ConfigurationItemTestContentGenerator()
        self.interface_extractor = InterfaceTestExtractor()  # 🔧 新增：接口测试提取器
        self.interface_parser = InterfaceTableParser()  # 🔧 新增：接口表格解析器
        self.interface_test_template_generator = InterfaceTestTemplateGenerator()  # 🔧 新增：接口测试模板生成器
        self.realtime_service = RealtimeRequirementAnalysisService()  # 🔧 新增：实时处理服务
    
    def generate_requirement_analysis_table(
        self,
        excel_file_path: str,
        sheet_name: str,
        ruanjian_xinxi: Dict[str, List[str]],
        ruanjian_name: str,
        diama_biangeng_xuqiu: Dict[str, str],
        file_path: str,
        ceshileixing: str,
        neirong_dict: List[str],
        biaoshihao_ceshixianng: List[str],
        AI_fenx_flag_Anniu: int,
        jianceyi_flag: int,
        apply_formatting: bool = True,
        user_software_info: Optional[Dict[str, str]] = None,
        task_id: Optional[str] = None,  # 🔧 新增：任务ID参数用于数据库存储
        product_type: Optional[str] = None  # 🔧 新增：产品类型参数用于数据库检索
    ) -> bool:
        """
        生成配置项测试需求分析表
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            ruanjian_xinxi: 软件信息字典
            ruanjian_name: 软件名称
            diama_biangeng_xuqiu: 代码变更需求字典
            file_path: 文档文件路径
            ceshileixing: 测试类型
            neirong_dict: 内容描述列表
            biaoshihao_ceshixianng: 标识号测试项列表
            AI_fenx_flag_Anniu: AI分析标志
            jianceyi_flag: 检测仪标志
            apply_formatting: 是否应用格式设置
            user_software_info: 用户软件信息（包含前端提供的版本信息）
            task_id: 任务ID（用于数据库存储）
            
        Returns:
            bool: 生成是否成功
        """
        try:
            self.logger.info(f"开始生成配置项测试需求分析表: {sheet_name}")
            
            # Debug: 记录service入口参数
            log_info(f"[DEBUG] ConfigurationItemRequirementAnalysisService.generate_requirement_analysis_table 开始", "requirement_analysis_service")
            log_debug(f"   excel_file_path: {excel_file_path}", "requirement_analysis_service")
            log_debug(f"   sheet_name: {sheet_name}", "requirement_analysis_service")
            log_debug(f"   ruanjian_xinxi: {ruanjian_xinxi}", "requirement_analysis_service")
            log_debug(f"   ruanjian_name: {ruanjian_name}", "requirement_analysis_service")
            log_debug(f"   diama_biangeng_xuqiu: {diama_biangeng_xuqiu}", "requirement_analysis_service")
            log_debug(f"   file_path: {file_path}", "requirement_analysis_service")
            log_debug(f"   ceshileixing: {ceshileixing}", "requirement_analysis_service")
            log_debug(f"   neirong_dict: {neirong_dict}", "requirement_analysis_service")
            log_debug(f"   biaoshihao_ceshixianng: {biaoshihao_ceshixianng}", "requirement_analysis_service")
            log_debug(f"   AI_fenx_flag_Anniu: {AI_fenx_flag_Anniu}", "requirement_analysis_service")
            log_debug(f"   jianceyi_flag: {jianceyi_flag}", "requirement_analysis_service")
            log_debug(f"   apply_formatting: {apply_formatting}", "requirement_analysis_service")
            log_debug(f"   task_id: {task_id}", "requirement_analysis_service")
            
            # 第一步：组织数据
            log_info(f"[DEBUG] 准备调用 data_processor.organize_requirement_analysis_data", "requirement_analysis_service")
            requirement_analysis_data = self.data_processor.organize_requirement_analysis_data(
                ruanjian_xinxi=ruanjian_xinxi,
                ruanjian_name=ruanjian_name,
                diama_biangeng_xuqiu=diama_biangeng_xuqiu,
                file_path=file_path,
                ceshileixing=ceshileixing,
                neirong_dict=neirong_dict,
                biaoshihao_ceshixianng=biaoshihao_ceshixianng,
                AI_fenx_flag_Anniu=AI_fenx_flag_Anniu,
                jianceyi_flag=jianceyi_flag,
                user_software_info=user_software_info
            )
            log_debug(f"[DEBUG] data_processor.organize_requirement_analysis_data 返回: {requirement_analysis_data}", "requirement_analysis_service")
            
            # 验证数据
            if not self.data_processor.validate_requirement_analysis_data(requirement_analysis_data):
                self.logger.error("配置项测试需求分析表数据验证失败")
                return False
            
            # 第二步：生成测试内容
            test_items = requirement_analysis_data.get('test_items', [])
            ai_config = requirement_analysis_data.get('ai_config', {})
            
            # 从用户软件信息中提取产品类型，如果没有则使用传入的产品类型参数
            if user_software_info and 'product_type' in user_software_info:
                product_type_to_use = user_software_info['product_type']
            else:
                product_type_to_use = product_type
            
            log_info(f"[DEBUG] 使用产品类型进行数据库增强: {product_type_to_use}", "requirement_analysis_service")
            
            test_descriptions, test_methods = self.test_content_generator.generate_test_content(
                test_items, ai_config, product_type=product_type_to_use
            )
            
            # 更新测试内容到数据中
            self.data_processor.update_test_content(
                requirement_analysis_data, test_descriptions, test_methods
            )
            
            # 第三步：读取文档章节号
            document_info = requirement_analysis_data.get('document_info', {})
            file_path = document_info.get('document_path', '')
            ruanjian_name = document_info.get('software_name', '')
            
            if file_path:
                # 🔧 修复：使用中文需求名称而不是测试项标识符
                requirement_names = [item.get('pure_name', '') for item in test_items]
                self.logger.info(f"读取文档章节号，需求名称列表: {requirement_names}")
                chapter_numbers = self._read_document_chapters(file_path, requirement_names, ruanjian_name)
                self.logger.info(f"获取到的章节号: {chapter_numbers}")
                self.data_processor.update_chapter_numbers(requirement_analysis_data, chapter_numbers)
            
            # 第四步：强制要求 task_id，并先逐行写库，再从库渲染Excel
            if not task_id:
                self.logger.error("必须提供 task_id 才能生成需求分析表（启用行级持久化与渲染）")
                return False

            # 4.1 持久化到数据库（不写Excel）
            # log_info(f"[DEBUG] 逐行持久化到数据库(persist_only=True)，task_id={task_id}", "requirement_analysis_service")  # 注释掉详细DEBUG
            if not self.excel_filler.fill_requirement_analysis_sheet(
                excel_file_path=excel_file_path,
                sheet_name=sheet_name,
                requirement_analysis_data=requirement_analysis_data,
                task_id=task_id,
                persist_only=True
            ):
                self.logger.error("持久化到数据库失败")
                return False

            # 4.2 从数据库渲染Excel
            # log_info(f"[DEBUG] 从数据库渲染Excel: render_from_db(task_id={task_id})", "requirement_analysis_service")  # 注释掉详细DEBUG
            if not self.excel_filler.render_from_db(excel_file_path, sheet_name, task_id):
                self.logger.error("从数据库渲染Excel失败")
                return False
            
            # 第五步：应用格式（可选）
            if apply_formatting:
                test_count = requirement_analysis_data.get('test_count', 0)
                if not self.formatter.apply_formatting(
                    excel_file_path=excel_file_path,
                    sheet_name=sheet_name,
                    test_count=test_count,
                    requirement_analysis_data=requirement_analysis_data
                ):
                    self.logger.warning("格式设置失败，但数据填充成功")
                
                # 🔧 新增：应用组织资产库内容的颜色标记
                try:
                    test_items = requirement_analysis_data.get('test_items', [])
                    self.formatter.apply_database_source_highlighting(
                        excel_file_path=excel_file_path,
                        sheet_name=sheet_name,
                        test_items=test_items
                    )
                except Exception as color_error:
                    self.logger.warning(f"应用组织资产库颜色标记时出现警告: {color_error}")
            
            # 记录统计信息
            statistics = self.data_processor.get_requirement_analysis_statistics(requirement_analysis_data)
            log_debug(f"需求分析表统计: {statistics}", "requirement_analysis_service")
            
            self.logger.info("配置项测试需求分析表生成完成")
            return True
            
        except Exception as e:
            self.logger.error(f"生成配置项测试需求分析表时发生错误: {str(e)}")
            log_error(f"[DEBUG] 生成需求分析表异常: {str(e)}", "requirement_analysis_service", exception=e)
            return False
    
    def generate_requirement_analysis_table_realtime(
        self,
        excel_file_path: str,
        sheet_name: str,
        ruanjian_xinxi: Dict[str, List[str]],
        ruanjian_name: str,
        diama_biangeng_xuqiu: Dict[str, str],
        file_path: str,
        ceshileixing: str,
        neirong_dict: List[str],
        biaoshihao_ceshixianng: List[str],
        AI_fenx_flag_Anniu: int,
        jianceyi_flag: int,
        apply_formatting: bool = True,
        user_software_info: Optional[Dict[str, str]] = None,
        task_id: Optional[str] = None,
        product_type: Optional[str] = None
    ) -> bool:
        """
        使用实时处理模式生成配置项测试需求分析表
        
        特点：
        1. 英文标识最先生成
        2. 逐个处理测试项，生成一个存储一个
        3. 避免大数组存储所有AI生成内容
        4. 真正的流式处理
        """
        log_info("[REALTIME-MODE] 启用实时处理模式", "requirement_analysis_service")
        
        return self.realtime_service.generate_requirement_analysis_table_realtime(
            excel_file_path=excel_file_path,
            sheet_name=sheet_name,
            ruanjian_xinxi=ruanjian_xinxi,
            ruanjian_name=ruanjian_name,
            diama_biangeng_xuqiu=diama_biangeng_xuqiu,
            file_path=file_path,
            ceshileixing=ceshileixing,
            neirong_dict=neirong_dict,
            biaoshihao_ceshixianng=biaoshihao_ceshixianng,
            AI_fenx_flag_Anniu=AI_fenx_flag_Anniu,
            jianceyi_flag=jianceyi_flag,
            apply_formatting=apply_formatting,
            user_software_info=user_software_info,
            task_id=task_id,
            product_type=product_type
        )
    
    def generate_requirement_analysis_table_legacy(
        self,
        excel_file_path: str,
        sheet_name: str,
        ruanjian_xinxi: Dict[str, List[str]],
        ruanjian_name: str,
        diama_biangeng_xuqiu: Dict[str, str],
        file_path: str,
        ceshileixing: str,
        neirong_dict: List[str],
        biaoshihao_ceshixianng: List[str],
        AI_fenx_flag_Anniu: int,
        jianceyi_flag: int
    ) -> bool:
        """
        生成配置项测试需求分析表（兼容原有接口）
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            ruanjian_xinxi: 软件信息字典
            ruanjian_name: 软件名称
            diama_biangeng_xuqiu: 代码变更需求字典
            file_path: 文档文件路径
            ceshileixing: 测试类型
            neirong_dict: 内容描述列表
            biaoshihao_ceshixianng: 标识号测试项列表
            AI_fenx_flag_Anniu: AI分析标志
            jianceyi_flag: 检测仪标志
            
        Returns:
            bool: 生成是否成功
        """
        try:
            self.logger.info(f"开始生成配置项测试需求分析表（兼容模式）: {sheet_name}")
            
            # 使用兼容模式的填充方法
            log_warning(f"[DEBUG] 正在调用legacy方法: excel_filler.fill_requirement_analysis_sheet_legacy", "requirement_analysis_service")
            success = self.excel_filler.fill_requirement_analysis_sheet_legacy(
                excel_file_path=excel_file_path,
                sheet_name=sheet_name,
                ruanjian_xinxi=ruanjian_xinxi,
                ruanjian_name=ruanjian_name,
                diama_biangeng_xuqiu=diama_biangeng_xuqiu,
                file_path=file_path,
                ceshileixing=ceshileixing,
                neirong_dict=neirong_dict,
                biaoshihao_ceshixianng=biaoshihao_ceshixianng,
                AI_fenx_flag_Anniu=AI_fenx_flag_Anniu,
                jianceyi_flag=jianceyi_flag
            )
            
            # 应用兼容模式格式
            if success:
                test_count = len(diama_biangeng_xuqiu)
                # 读取章节号用于格式设置
                test_names = list(diama_biangeng_xuqiu.keys())
                chapter_numbers = self._read_document_chapters(file_path, test_names, ruanjian_name)
                
                self.formatter.apply_legacy_formatting(
                    excel_file_path, sheet_name, test_count, chapter_numbers
                )
            
            self.logger.info(f"配置项测试需求分析表生成完成（兼容模式）")
            return success
            
        except Exception as e:
            self.logger.error(f"生成配置项测试需求分析表时发生错误（兼容模式）: {str(e)}")
            return False
    
    def _read_document_chapters(self, file_path: str, test_names: List[str], ruanjian_name: str) -> List[str]:
        """读取文档章节号 - ConfigurationItem专用版本"""
        try:
            import docx
            
            self.logger.info(f"[INFO] 开始读取目录文档章节号")
            self.logger.info(f"[INFO] 目录文档路径: {file_path}")
            self.logger.info(f"[INFO] 需要匹配的需求名称: {test_names}")
            
            # 读取Word文档
            doc = docx.Document(file_path)
            chapter_numbers = []
            
            # 为每个需求名称查找对应的章节号
            for requirement_name in test_names:
                found_chapter = None
                self.logger.info(f"[SEARCH] 查找需求: '{requirement_name}'")
                
                # 遍历文档段落查找匹配
                for paragraph in doc.paragraphs:
                    paragraph_text = paragraph.text.strip()
                    if not paragraph_text:
                        continue
                    
                    self.logger.debug(f"   检查段落: '{paragraph_text}'")
                    
                    # 检查是否包含需求名称
                    if requirement_name in paragraph_text:
                        self.logger.info(f"   [SUCCESS] 找到匹配段落: '{paragraph_text}'")
                        
                        # 尝试提取章节编号 (例如: "3.2.1 用户身份验证功能")
                        import re
                        # 匹配章节号模式 (数字.数字.数字等格式)
                        chapter_pattern = r'^(\d+(?:\.\d+)*)\s+'
                        match = re.match(chapter_pattern, paragraph_text)
                        
                        if match:
                            found_chapter = match.group(1)
                            self.logger.info(f"   [EXTRACT] 提取到章节号: '{found_chapter}'")
                            break
                        else:
                            # 如果没有标准格式的章节号，尝试其他格式
                            # 查找是否有制表符分隔的章节号
                            if '\t' in paragraph_text:
                                parts = paragraph_text.split('\t')
                                if len(parts) > 0:
                                    potential_chapter = parts[0].strip()
                                    if re.match(r'^\d+(?:\.\d+)*$', potential_chapter):
                                        found_chapter = potential_chapter
                                        self.logger.info(f"   📋 从制表符分隔中提取章节号: '{found_chapter}'")
                                        break
                
                if found_chapter:
                    chapter_numbers.append(found_chapter)
                    self.logger.info(f"[SUCCESS] '{requirement_name}' -> '{found_chapter}'")
                else:
                    chapter_numbers.append('未找到具体章节号')
                    self.logger.warning(f"[FAILED] '{requirement_name}' -> 未找到章节号")
            
            self.logger.info(f"[INFO] 章节号匹配完成: {chapter_numbers}")
            return chapter_numbers
            
        except Exception as e:
            self.logger.error(f"读取文档章节号时发生错误: {str(e)}")
            import traceback
            self.logger.error(f"错误详情: {traceback.format_exc()}")
            return ['未找到具体章节号'] * len(test_names)
    
    def read_requirement_analysis_data(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表'
    ) -> Tuple[Dict[str, Any], str]:
        """
        读取配置项测试需求分析表数据
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            Tuple[Dict[str, Any], str]: (需求分析表数据, 状态信息)
        """
        try:
            return self.reader.read_requirement_analysis_data(excel_file_path, sheet_name)
        except Exception as e:
            self.logger.error(f"读取配置项测试需求分析表数据时发生错误: {str(e)}")
            return {}, f"读取失败: {e}"
    
    def update_test_item_field(
        self,
        excel_file_path: str,
        sheet_name: str,
        item_index: int,
        field_name: str,
        field_value: str,
        apply_formatting: bool = True
    ) -> bool:
        """
        更新配置项测试需求分析表中的单个字段
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            item_index: 测试项索引（从1开始）
            field_name: 字段名称
            field_value: 字段值
            apply_formatting: 是否重新应用格式设置
            
        Returns:
            bool: 更新是否成功
        """
        try:
            self.logger.info(f"更新配置项测试需求分析表字段: 测试项{item_index}, {field_name} = {field_value}")
            
            # 更新单个字段
            if not self.excel_filler.update_single_test_item(
                excel_file_path=excel_file_path,
                sheet_name=sheet_name,
                item_index=item_index,
                field_name=field_name,
                field_value=field_value
            ):
                return False
            
            # 重新应用格式（可选）
            if apply_formatting:
                # 需要知道测试项数量来应用格式
                summary = self.excel_filler.get_filled_data_summary(excel_file_path, sheet_name, 50)  # 假设最多50项
                test_count = summary.get('total_items', 0)
                self.formatter.apply_formatting(excel_file_path, sheet_name, test_count)
            
            self.logger.info(f"测试项{item_index}的字段{field_name}更新完成")
            return True
            
        except Exception as e:
            self.logger.error(f"更新配置项测试需求分析表字段时发生错误: {str(e)}")
            return False
    
    def get_requirement_analysis_status(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表'
    ) -> Dict[str, Any]:
        """
        获取配置项测试需求分析表的状态
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            Dict[str, Any]: 状态信息
        """
        try:
            return self.reader.get_requirement_analysis_summary(excel_file_path, sheet_name)
        except Exception as e:
            self.logger.error(f"获取配置项测试需求分析表状态时发生错误: {str(e)}")
            return {}
    
    def validate_requirement_analysis_table(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表',
        required_fields: Optional[List[str]] = None
    ) -> Dict[str, Any]:
        """
        验证配置项测试需求分析表的完整性
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            required_fields: 必需字段列表
            
        Returns:
            Dict[str, Any]: 验证结果
        """
        try:
            # 检查结构
            if not self.reader.validate_requirement_analysis_structure(excel_file_path, sheet_name):
                return {
                    'is_valid': False,
                    'error': '配置项测试需求分析表结构无效',
                    'table_type': '配置项测试需求分析表'
                }
            
            # 检查测试内容完整性
            content_check = self.reader.check_test_content_completeness(excel_file_path, sheet_name)
            
            validation_result = {
                'is_valid': content_check.get('is_complete', False),
                'missing_descriptions': content_check.get('missing_descriptions_indices', []),
                'missing_methods': content_check.get('missing_methods_indices', []),
                'missing_identifiers': content_check.get('missing_identifiers_indices', []),
                'total_test_items': content_check.get('total_test_items', 0),
                'completion_rate': 0,
                'validation_message': self._get_validation_message(content_check),
                'table_type': '配置项测试需求分析表'
            }
            
            # 计算完成率
            if validation_result['total_test_items'] > 0:
                missing_count = max(
                    len(validation_result['missing_descriptions']),
                    len(validation_result['missing_methods']),
                    len(validation_result['missing_identifiers'])
                )
                completed = validation_result['total_test_items'] - missing_count
                validation_result['completion_rate'] = round(completed / validation_result['total_test_items'] * 100, 2)
            
            self.logger.info(f"配置项测试需求分析表验证结果: {validation_result}")
            return validation_result
            
        except Exception as e:
            self.logger.error(f"验证配置项测试需求分析表时发生错误: {str(e)}")
            return {'is_valid': False, 'error': str(e)}
    
    def _get_validation_message(self, content_check: Dict[str, Any]) -> str:
        """
        生成验证消息
        
        Args:
            content_check: 内容检查结果
            
        Returns:
            str: 验证消息
        """
        if content_check.get('is_complete', False):
            return "配置项测试需求分析表验证通过，所有测试内容都已生成"
        
        messages = []
        
        missing_descriptions = content_check.get('missing_descriptions_count', 0)
        if missing_descriptions > 0:
            messages.append(f"有 {missing_descriptions} 个测试项缺失测试描述")
        
        missing_methods = content_check.get('missing_methods_count', 0)
        if missing_methods > 0:
            messages.append(f"有 {missing_methods} 个测试项缺失测试方法")
        
        missing_identifiers = content_check.get('missing_identifiers_count', 0)
        if missing_identifiers > 0:
            messages.append(f"有 {missing_identifiers} 个测试项缺失测试标识")
        
        return "; ".join(messages) if messages else "需求分析表状态未知"
    
    def generate_test_content_for_existing_table(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表',
        use_ai_analysis: bool = True,
        jianceyi_flag: int = 0
    ) -> bool:
        """
        为现有的配置项测试需求分析表生成测试内容
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            use_ai_analysis: 是否使用AI分析
            jianceyi_flag: 检测仪标志
            
        Returns:
            bool: 生成是否成功
        """
        try:
            self.logger.info("为现有配置项测试需求分析表生成测试内容")
            
            # 读取现有数据
            requirement_analysis_data, status_msg = self.reader.read_requirement_analysis_data(excel_file_path, sheet_name)
            test_items = requirement_analysis_data.get('test_items', [])
            
            if not test_items:
                self.logger.error("未找到测试项数据")
                return False
            
            # 构造AI配置
            ai_config = {
                'enable_ai_analysis': use_ai_analysis,
                'detector_flag': jianceyi_flag,
                'ai_analysis_attempts': 4,
                'use_knowledge_optimization': True
            }
            
            # 生成测试内容
            test_descriptions, test_methods = self.test_content_generator.generate_test_content(
                test_items, ai_config
            )
            
            # 更新Excel中的测试内容
            import openpyxl
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            start_row = 2  # 从第2行开始
            for i, (description, method) in enumerate(zip(test_descriptions, test_methods)):
                sheet.cell(row=start_row + i, column=19).value = description  # 测试项描述
                sheet.cell(row=start_row + i, column=20).value = method      # 测试方法
            
            wb.save(excel_file_path)
            
            self.logger.info(f"成功为{len(test_descriptions)}个测试项生成测试内容")
            return True
            
        except Exception as e:
            self.logger.error(f"为现有配置项测试需求分析表生成测试内容时发生错误: {str(e)}")
            return False
    
    def highlight_configuration_item_specific_content(
        self,
        excel_file_path: str,
        sheet_name: str,
        test_count: int
    ) -> bool:
        """
        高亮显示配置项测试特有的内容
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            test_count: 测试项数量
            
        Returns:
            bool: 高亮设置是否成功
        """
        try:
            return self.formatter.highlight_configuration_item_specific_fields(excel_file_path, sheet_name, test_count)
        except Exception as e:
            self.logger.error(f"高亮配置项测试特有内容时发生错误: {str(e)}")
            return False
    
    def highlight_auto_generated_content(
        self,
        excel_file_path: str,
        sheet_name: str,
        test_count: int
    ) -> bool:
        """
        高亮显示自动生成的内容
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            test_count: 测试项数量
            
        Returns:
            bool: 高亮设置是否成功
        """
        try:
            return self.formatter.highlight_auto_generated_fields(excel_file_path, sheet_name, test_count)
        except Exception as e:
            self.logger.error(f"高亮自动生成内容时发生错误: {str(e)}")
            return False
    
    def extract_test_requirements(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表'
    ) -> List[str]:
        """
        提取测试需求列表
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            List[str]: 测试需求列表
        """
        try:
            return self.reader.extract_test_requirements_only(excel_file_path, sheet_name)
        except Exception as e:
            self.logger.error(f"提取测试需求时发生错误: {str(e)}")
            return []
    
    def backup_requirement_analysis_table(
        self,
        excel_file_path: str,
        sheet_name: str,
        backup_file_path: str
    ) -> bool:
        """
        备份配置项测试需求分析表
        
        Args:
            excel_file_path: 源Excel文件路径
            sheet_name: 工作表名称
            backup_file_path: 备份文件路径
            
        Returns:
            bool: 备份是否成功
        """
        try:
            return self.excel_filler.backup_sheet_data(excel_file_path, sheet_name, backup_file_path)
        except Exception as e:
            self.logger.error(f"备份配置项测试需求分析表时发生错误: {str(e)}")
            return False

    def fill_requirement_analysis_table_for_form_builder(
        self, 
        excel_file_path: str, 
        sheet_name: str, 
        final_xuqiu_mioashu: Dict, 
        ruanjian_name: str, 
        total_lines: int, 
        AI_fenx_flag_Anniu: bool, 
        AI_deep_fenx_flag_Anniu: bool,
        file_path_mulu_word: str,
        user_software_info: Optional[Dict[str, str]] = None,
        task_id: Optional[str] = None,  # 🔧 新增：任务ID参数用于数据库存储
        product_type: Optional[str] = None  # 🔧 新增：产品类型参数用于数据库检索
    ) -> bool:
        """
        为表单构建器填充需求分析表
        
        这是从form_builder_service.py中提取的_fill_requirement_analysis_table方法
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            final_xuqiu_mioashu: 最终需求描述字典
            ruanjian_name: 软件名称
            total_lines: 总代码行数
            AI_fenx_flag_Anniu: AI分析标志
            AI_deep_fenx_flag_Anniu: 深度分析标志
            file_path_mulu_word: 目录文档路径
            user_software_info: 用户软件信息
            task_id: 任务ID（用于数据库存储）
            
        Returns:
            bool: 填充是否成功
        """
        try:
            log_info(f"开始填充需求分析表: {sheet_name}", "requirement_analysis_service")
            
            # 详细监控：输入参数
            log_info(f"📋 输入参数监控:", "requirement_analysis_service")
            log_info(f"   - Excel文件: {excel_file_path}", "requirement_analysis_service")
            log_info(f"   - 工作表名: {sheet_name}", "requirement_analysis_service")
            log_info(f"   - 软件名称: {ruanjian_name}", "requirement_analysis_service")
            log_info(f"   - final_xuqiu_mioashu数量: {len(final_xuqiu_mioashu)}", "requirement_analysis_service")
            log_info(f"   - final_xuqiu_mioashu内容: {final_xuqiu_mioashu}", "requirement_analysis_service")
            log_info(f"   - task_id: {task_id}", "requirement_analysis_service")  # 🔧 新增：task_id记录
            
            # 准备符合现有方法签名的参数
            # 优先使用前端提供的软件信息（软件标识和版本）
            software_version = '1.0.0'  # 默认版本
            software_id = ''  # 默认软件标识
            
            # 优先使用前端提供的信息
            if user_software_info:
                if user_software_info.get('software_version'):
                    software_version = user_software_info.get('software_version')
                    log_info(f"使用前端提供的软件版本: {software_version}", "requirement_analysis_service")
                # 修复：检查正确的字段名 software_identifier
                if user_software_info.get('software_identifier'):
                    software_id = user_software_info.get('software_identifier')
                    log_info(f"使用前端提供的软件标识: {software_id}", "requirement_analysis_service")
                elif user_software_info.get('software_id'):
                    software_id = user_software_info.get('software_id')
                    log_info(f"使用前端提供的软件标识(兼容字段): {software_id}", "requirement_analysis_service")
            
            # 如果前端没有提供软件标识，使用软件名称截取作为后备方案
            if not software_id:
                software_id = ruanjian_name[:8] if len(ruanjian_name) > 8 else ruanjian_name
                log_warning(f"未找到前端软件标识，使用软件名称截取: {software_id}", "requirement_analysis_service")
            
            ruanjian_xinxi = {
                '软件名称': [ruanjian_name],
                '软件标识': [software_id],  # 使用正确的软件标识而不是软件名称
                '被测软件版本': [software_version]  # 添加版本信息
            }
            diama_biangeng_xuqiu = {k: v.get('description', v) if isinstance(v, dict) else str(v) 
                                  for k, v in final_xuqiu_mioashu.items()}
            neirong_dict = list(diama_biangeng_xuqiu.values())
            
            # 关键修复：传递实际的需求名称而不是REQ标识符
            log_info(f"构建biaoshihao_ceshixianng参数:", "requirement_analysis_service")
            # 提取final_xuqiu_mioashu中的'name'字段作为实际需求名称
            biaoshihao_ceshixianng = []
            chapter_dict = {}  # 🔧 新增：构建章节编号字典
            for req_id, req_data in final_xuqiu_mioashu.items():
                if isinstance(req_data, dict):
                    req_name = req_data.get('name', req_id)
                    biaoshihao_ceshixianng.append(req_name)
                    # 🔧 提取章节编号信息
                    chapter_number = req_data.get('chapter_number', '')
                    if chapter_number:
                        chapter_dict[req_name] = chapter_number
                        log_debug(f"   添加章节映射: '{req_name}' -> '{chapter_number}'", "requirement_analysis_service")
                else:
                    biaoshihao_ceshixianng.append(str(req_data))
            
            log_info(f"   构建的biaoshihao_ceshixianng: {biaoshihao_ceshixianng}", "requirement_analysis_service")
            log_info(f"   构建的chapter_dict: {chapter_dict}", "requirement_analysis_service")
            
            log_debug(f"📊 最终构建的参数:", "requirement_analysis_service")
            log_debug(f"   - diama_biangeng_xuqiu.keys(): {list(diama_biangeng_xuqiu.keys())}", "requirement_analysis_service")
            log_debug(f"   - biaoshihao_ceshixianng: {biaoshihao_ceshixianng}", "requirement_analysis_service")
            log_debug(f"   - neirong_dict: {neirong_dict}", "requirement_analysis_service")
            
            # 调用现有的需求分析方法
            log_debug(f"准备调用 generate_requirement_analysis_table", "requirement_analysis_service")
            log_debug(f"   excel_file_path: {excel_file_path}", "requirement_analysis_service")
            log_debug(f"   sheet_name: {sheet_name}", "requirement_analysis_service")
            log_debug(f"   ruanjian_xinxi: {ruanjian_xinxi}", "requirement_analysis_service")
            log_debug(f"   ruanjian_name: {ruanjian_name}", "requirement_analysis_service")
            log_debug(f"   diama_biangeng_xuqiu: {diama_biangeng_xuqiu}", "requirement_analysis_service")
            log_debug(f"   neirong_dict: {neirong_dict}", "requirement_analysis_service")
            log_debug(f"   biaoshihao_ceshixianng: {biaoshihao_ceshixianng}", "requirement_analysis_service")
            log_debug(f"   AI_fenx_flag_Anniu: {1 if AI_fenx_flag_Anniu else 0}", "requirement_analysis_service")
            log_debug(f"   file_path (目录文档): {file_path_mulu_word}", "requirement_analysis_service")
            log_debug(f"   task_id: {task_id}", "requirement_analysis_service")  # 🔧 新增：task_id记录
            
            # 🚀 NEW: 使用实时处理模式
            use_realtime = True  # 启用实时处理模式
            
            if use_realtime and task_id:
                log_info("[REALTIME-SWITCH] 切换到实时处理模式", "requirement_analysis_service")
                # 调用实时处理方法
                success = self.generate_requirement_analysis_table_realtime(
                    excel_file_path=excel_file_path,
                    sheet_name=sheet_name,
                    ruanjian_xinxi=ruanjian_xinxi,
                    ruanjian_name=ruanjian_name,
                    diama_biangeng_xuqiu=diama_biangeng_xuqiu,
                    file_path=file_path_mulu_word,
                    ceshileixing="配置项测试",
                    neirong_dict=neirong_dict,
                    biaoshihao_ceshixianng=biaoshihao_ceshixianng,
                    AI_fenx_flag_Anniu=1 if AI_fenx_flag_Anniu else 0,
                    jianceyi_flag=0,
                    apply_formatting=True,
                    user_software_info=user_software_info,
                    task_id=task_id,
                    product_type=product_type
                )
            else:
                log_info("[LEGACY-MODE] 使用传统处理模式", "requirement_analysis_service")
                # 调用传统需求分析表生成方法
                success = self.generate_requirement_analysis_table(
                    excel_file_path=excel_file_path,
                    sheet_name=sheet_name,
                    ruanjian_xinxi=ruanjian_xinxi,
                    ruanjian_name=ruanjian_name,
                    diama_biangeng_xuqiu=diama_biangeng_xuqiu,
                    file_path=file_path_mulu_word,  # 传递目录文档路径，用于提取章节编号
                    ceshileixing="配置项测试",
                    neirong_dict=neirong_dict,
                    biaoshihao_ceshixianng=biaoshihao_ceshixianng,
                    AI_fenx_flag_Anniu=1 if AI_fenx_flag_Anniu else 0,
                    jianceyi_flag=0,  # 默认值
                    apply_formatting=True,
                    user_software_info=user_software_info,
                    task_id=task_id,  # 🔧 新增：传递task_id给数据库存储
                    product_type=product_type
                )
            
            log_debug(f"generate_requirement_analysis_table 返回: {success}", "requirement_analysis_service")
            
            if success:
                log_info(f"需求分析表填充成功: {sheet_name}", "requirement_analysis_service")
            else:
                log_error(f"需求分析表填充失败: {sheet_name}", "requirement_analysis_service")
                
            return success
            
        except Exception as e:
            log_error(f"填充需求分析表时发生错误", "requirement_analysis_service", e)
            return False

    def extract_and_store_excel_data_to_db(
        self, 
        excel_file_path: str, 
        sheet_name: str = '需求分析表',
        task_id: str = None
    ) -> bool:
        """
        从需求分析Excel表中提取数据并存储到数据库
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            task_id: 任务ID
            
        Returns:
            bool: 是否成功
        """
        try:
            import openpyxl
            from ...db.optimized_models import OptimizedDynamicTestDatabase
            from ...utils.logging_utils.logger import log_info, log_warning, log_error
            
            log_info(f"开始从Excel提取需求分析数据: {excel_file_path}", "requirement_analysis_service")
            
            # 打开Excel文件
            wb = openpyxl.load_workbook(excel_file_path)
            if sheet_name not in wb.sheetnames:
                log_error(f"工作表 {sheet_name} 不存在", "requirement_analysis_service")
                return False
            
            sheet = wb[sheet_name]
            
            # 初始化数据库
            db = OptimizedDynamicTestDatabase()
            
            # 执行数据库迁移（如果需要）
            db.migrate_requirement_analysis_table()
            
            # 生成任务ID（如果未提供）
            if not task_id:
                import uuid
                task_id = str(uuid.uuid4())
                log_info(f"生成新的任务ID: {task_id}", "requirement_analysis_service")
            
            # 定义Excel列映射（与excel_filler.py中的field_mapping保持一致）
            excel_column_mapping = {
                2: 'index_number',                      # 序号
                3: 'software_name',                     # 软件名称
                4: 'software_version',                  # 软件版本
                5: 'round_info',                        # 轮次信息
                6: 'standard_sub_requirement_name',     # 标准子需求名称
                7: 'standard_sub_requirement_identifier', # 标准子需求标识
                8: 'standard_sub_requirement_unit',     # 标准子需求测试单元
                9: 'task_requirement_trace',            # 任务需求追踪
                10: 'software_requirement_trace',       # 软件需求规格说明追踪
                11: 'software_requirement_name',        # 软件需求规格说明需求名称
                12: 'software_requirement_identifier',  # 软件需求规格说明需求标识
                13: 'requirement_description',          # 需求描述
                14: 'test_type',                        # 测试类型
                15: 'test_item_name',                   # 测试项名称
                16: 'test_item_identifier',             # 测试项标识
                17: 'process_identifier',               # 过程标识
                18: 'requirement_sub_description',      # 需求子项描述
                19: 'test_description',                 # 测试项描述
                20: 'test_method',                      # 测试方法
                21: 'adequacy_requirement',             # 充分性要求
                22: 'pass_criteria'                     # 通过准则
            }
            
            # 提取数据
            extracted_data = []
            row_num = 2  # 从第2行开始（第1行是表头）
            
            while True:
                # 检查是否有数据（以第2列序号是否有值为准）
                index_cell = sheet.cell(row=row_num, column=2)
                if not index_cell.value:
                    break
                
                # 提取该行的所有数据
                row_data = {'requirement_id': f"req_{task_id}_{row_num}"}
                
                for col_num, field_name in excel_column_mapping.items():
                    cell_value = sheet.cell(row=row_num, column=col_num).value
                    
                    # 处理不同类型的值
                    if cell_value is not None:
                        if field_name == 'index_number':
                            # 序号转换为整数
                            try:
                                row_data[field_name] = int(cell_value)
                            except (ValueError, TypeError):
                                row_data[field_name] = row_num - 1  # 使用行号作为序号
                        else:
                            # 其他字段转换为字符串
                            row_data[field_name] = str(cell_value).strip()
                    else:
                        row_data[field_name] = None
                
                extracted_data.append(row_data)
                row_num += 1
            
            log_info(f"从Excel提取到 {len(extracted_data)} 行数据", "requirement_analysis_service")
            
            # 存储到数据库
            if extracted_data:
                db.save_requirement_analysis(task_id, extracted_data)
                log_info(f"成功存储 {len(extracted_data)} 条需求分析数据到数据库", "requirement_analysis_service")
                
                # 记录详细的存储信息
                for i, data in enumerate(extracted_data, 1):
                    log_info(f"  第{i}行: {data.get('test_item_name', 'N/A')} - {data.get('requirement_description', 'N/A')[:50]}...", "requirement_analysis_service")
            
            return True
            
        except Exception as e:
            log_error(f"从Excel提取并存储需求分析数据失败: {str(e)}", "requirement_analysis_service")
            import traceback
            traceback.print_exc()
            return False
    
    def generate_interface_tests_from_docx(self, 
                                         docx_path: str,
                                         excel_file_path: str,
                                         sheet_name: str = "需求分析表",
                                         project_name: str = "软件需求规格说明",
                                         base_identifier: str = "CT-388.222.JJK1-2.00") -> bool:
        """
        从Word文档中提取接口测试项并生成到Excel
        
        Args:
            docx_path: Word文档路径
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            project_name: 项目名称
            base_identifier: 基础标识符
            
        Returns:
            bool: 是否生成成功
        """
        try:
            log_info("=" * 80, "requirement_analysis_service")
            log_info("🔗 开始生成接口测试项", "requirement_analysis_service")
            log_info("=" * 80, "requirement_analysis_service")
            
            # 1. 从Word文档提取接口表格数据
            log_info(f"📖 从Word文档提取接口表格: {docx_path}", "requirement_analysis_service")
            interface_data = self.interface_parser.extract_interface_table_from_docx(docx_path)
            
            if not interface_data:
                log_warning("未找到接口表格数据", "requirement_analysis_service")
                return False
            
            log_info(f"✅ 成功提取 {len(interface_data)} 个接口", "requirement_analysis_service")
            
            # 2. 转换为接口测试项
            log_info("🔄 转换接口数据为测试项", "requirement_analysis_service")
            interface_tests = self.interface_extractor.extract_interface_tests(
                interface_data=interface_data,
                project_name=project_name,
                base_identifier=base_identifier
            )
            
            if not interface_tests:
                log_warning("未生成任何接口测试项", "requirement_analysis_service")
                return False
            
            log_info(f"✅ 成功生成 {len(interface_tests)} 个接口测试项", "requirement_analysis_service")
            
            # 3. 生成测试内容（测试项描述和测试方法）
            log_info("🤖 生成接口测试内容", "requirement_analysis_service")
            
            # 准备AI配置（使用默认配置）
            ai_config = {
                'enable_test_description': True,
                'enable_test_method': True,
                'use_knowledge_optimization': True
            }
            
            # 生成测试内容 - 🔧 修改：接口测试项使用固定模板
            test_descriptions = []
            test_methods = []
            
            for interface_test in interface_tests:
                interface_name = interface_test.get('name', '')
                interface_desc = interface_test.get('requirement_description', '')
                
                # 使用接口测试模板生成器
                description, method = self.interface_test_template_generator.generate_interface_test_content(
                    interface_name=interface_name,
                    interface_desc=interface_desc
                )
                
                test_descriptions.append(description)
                test_methods.append(method)
            
            log_info(f"📝 使用固定模板为 {len(interface_tests)} 个接口测试项生成内容", "requirement_analysis_service")
            
            self.data_processor.update_test_content(
                test_items=interface_tests,
                test_descriptions=test_descriptions,
                test_methods=test_methods
            )
            
            log_info(f"✅ 成功生成 {len(test_descriptions)} 个测试描述和 {len(test_methods)} 个测试方法", "requirement_analysis_service")
            
            # 4. 更新接口测试项的测试内容
            self.data_processor.update_test_content(
                test_items=interface_tests,
                test_descriptions=test_descriptions,
                test_methods=test_methods
            )
            
            log_info(f"✅ 成功生成 {len(test_descriptions)} 个测试描述和 {len(test_methods)} 个测试方法", "requirement_analysis_service")
            
            # 5. 写入Excel文件
            log_info(f"📝 写入Excel文件: {excel_file_path}", "requirement_analysis_service")
            
            success = self.excel_filler.fill_excel_with_test_items(
                test_items=interface_tests,
                excel_file_path=excel_file_path,
                sheet_name=sheet_name
            )
            
            if success:
                log_info("✅ 接口测试项Excel生成完成", "requirement_analysis_service")
                
                # 6. 应用格式设置
                log_info("🎨 应用Excel格式设置", "requirement_analysis_service")
                self.formatter.format_requirement_analysis_excel(
                    excel_file_path=excel_file_path,
                    sheet_name=sheet_name
                )
                
                log_info("=" * 80, "requirement_analysis_service")
                log_info("🎉 接口测试项生成完成", "requirement_analysis_service")
                log_info("=" * 80, "requirement_analysis_service")
                
                return True
            else:
                log_error("Excel写入失败", "requirement_analysis_service")
                return False
                
        except Exception as e:
            log_error(f"生成接口测试项时发生错误: {e}", "requirement_analysis_service")
            import traceback
            traceback.print_exc()
            return False
    
    def add_interface_tests_to_existing_excel(self,
                                            docx_path: str,
                                            excel_file_path: str,
                                            sheet_name: str = "需求分析表",
                                            project_name: str = "软件需求规格说明",
                                            base_identifier: str = "CT-388.222.JJK1-2.00") -> bool:
        """
        将接口测试项添加到现有的Excel文件中
        
        Args:
            docx_path: Word文档路径
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            project_name: 项目名称
            base_identifier: 基础标识符
            
        Returns:
            bool: 是否添加成功
        """
        try:
            log_info("🔗 开始将接口测试项添加到现有Excel", "requirement_analysis_service")
            
            # 1. 读取现有的测试项
            existing_tests = self.reader.read_test_items_from_excel(
                excel_file_path=excel_file_path,
                sheet_name=sheet_name
            )
            
            log_info(f"📖 读取到现有测试项: {len(existing_tests)} 个", "requirement_analysis_service")
            
            # 2. 提取接口测试项
            interface_data = self.interface_parser.extract_interface_table_from_docx(docx_path)
            
            if not interface_data:
                log_warning("未找到接口表格数据", "requirement_analysis_service")
                return False
            
            # 3. 生成接口测试项，调整索引以避免冲突
            start_index = len(existing_tests) + 1
            interface_tests = self.interface_extractor.extract_interface_tests(
                interface_data=interface_data,
                project_name=project_name,
                base_identifier=base_identifier
            )
            
            # 调整接口测试项的索引
            for i, test in enumerate(interface_tests):
                test['index'] = start_index + i
            
            # 4. 生成测试内容 - 🔧 修改：接口测试项使用固定模板
            test_descriptions = []
            test_methods = []
            
            for interface_test in interface_tests:
                interface_name = interface_test.get('name', '')
                interface_desc = interface_test.get('requirement_description', '')
                
                # 使用接口测试模板生成器
                description, method = self.interface_test_template_generator.generate_interface_test_content(
                    interface_name=interface_name,
                    interface_desc=interface_desc
                )
                
                test_descriptions.append(description)
                test_methods.append(method)
            
            log_info(f"📝 使用固定模板为 {len(interface_tests)} 个接口测试项生成内容", "requirement_analysis_service")
            
            self.data_processor.update_test_content(
                test_items=interface_tests,
                test_descriptions=test_descriptions,
                test_methods=test_methods
            )
            
            # 5. 合并测试项
            all_tests = existing_tests + interface_tests
            
            log_info(f"📋 合并后总测试项数量: {len(all_tests)}", "requirement_analysis_service")
            
            # 6. 重新写入Excel
            success = self.excel_filler.fill_excel_with_test_items(
                test_items=all_tests,
                excel_file_path=excel_file_path,
                sheet_name=sheet_name
            )
            
            if success:
                # 应用格式设置
                self.formatter.format_requirement_analysis_excel(
                    excel_file_path=excel_file_path,
                    sheet_name=sheet_name
                )
                
                log_info(f"✅ 成功添加 {len(interface_tests)} 个接口测试项到Excel", "requirement_analysis_service")
                return True
            else:
                log_error("Excel写入失败", "requirement_analysis_service")
                return False
                
        except Exception as e:
            log_error(f"添加接口测试项到Excel时发生错误: {e}", "requirement_analysis_service")
            return False

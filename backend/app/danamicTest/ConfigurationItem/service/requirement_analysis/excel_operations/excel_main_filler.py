"""
"""
# 导入任务日志函数
# 配置日志

import logging
import openpyxl
import os
from typing import Dict, List, Optional, Any, Tuple
from pathlib import Path
from ..document_parsing.utils.text_utils import extract_requirement_name

from ....utils.logging_utils.logger import log_info, log_debug, log_error, log_warning

logger = logging.getLogger(__name__)


from .data_filling import BasicInfoFiller, FunctionalTestFiller, DatabasePersistence, ExcelDataFiller
from .formatting.requirement_formatter import RequirementFormatter

class ConfigurationItemRequirementAnalysisExcelFiller:
    """配置项测试需求分析表Excel填充器（主协调器）"""

    def __init__(self):
        """初始化Excel填充器"""
        self.logger = logging.getLogger(__name__)
        self.basic_info_filler = BasicInfoFiller()
        self.functional_test_filler = FunctionalTestFiller()
        self.database_persistence = DatabasePersistence()
        self.requirement_formatter = RequirementFormatter()
        self.excel_data_filler = ExcelDataFiller()

    def fill_requirement_analysis_sheet(
        self,
        excel_file_path: str,
        sheet_name: str,
        requirement_analysis_data: Dict[str, Any],
        task_id: str,
        persist_only: bool = True
    ) -> bool:
        """
        主填充方法 - 协调所有子模块

        两种模式：
        - persist_only=True: 仅将 test_items 入库，不写Excel。
        - persist_only=False: 直接将 test_items 写入Excel（不经过DB），用于调试/回退。
        """
        try:
            log_info(f"[EXCEL_FILLER] [DEBUG] fill_requirement_analysis_sheet 开始", "excel_main_filler")
            basic_info: Dict[str, Any] = requirement_analysis_data.get('basic_info', {}) or {}
            test_items: List[Dict[str, Any]] = requirement_analysis_data.get('test_items', []) or []
            if not isinstance(test_items, list):
                test_items = []

            log_info(f"[EXCEL_FILLER] [DEBUG] 测试项数量: {len(test_items)}", "excel_main_filler")
            log_info(f"[EXCEL_FILLER] [DEBUG] persist_only={persist_only}", "excel_main_filler")

            if persist_only:
                # 仅入库，后续由 render_from_db 统一渲染
                self.database_persistence._persist_rows_to_db(basic_info, test_items, task_id)
                return True
            else:
                # 直接写入Excel（调试/回退路径）
                if not os.path.exists(excel_file_path):
                    # 如果目标工作簿不存在，则创建空工作簿
                    wb = openpyxl.Workbook()
                    # 默认Sheet可能与目标sheet同名/不同名，统一处理
                else:
                    wb = openpyxl.load_workbook(excel_file_path)

                if sheet_name in wb.sheetnames:
                    sheet = wb[sheet_name]
                else:
                    sheet = wb.create_sheet(sheet_name)

                # 使用新版数据填充逻辑
                self.excel_data_filler.fill_excel_data_new(sheet, basic_info, test_items)
                wb.save(excel_file_path)
                return True

        except Exception as e:
            import traceback
            log_error(f"[EXCEL_FILLER] [DEBUG] fill_requirement_analysis_sheet 失败: {e}", "excel_main_filler")
            log_error(f"[EXCEL_FILLER] [DEBUG] 异常堆栈: {traceback.format_exc()}", "excel_main_filler")
            return False
    
    def render_from_db(self, excel_file_path: str, sheet_name: str, task_id: str) -> bool:
        """从数据库渲染Excel表格"""
        log_info(f"[EXCEL_FILLER] [DEBUG] 开始调用render_from_db", "excel_main_filler")
        log_info(f"[EXCEL_FILLER] [DEBUG] 参数: file={excel_file_path}, sheet={sheet_name}, task={task_id}", "excel_main_filler")
        
        result = self.database_persistence.render_from_db(excel_file_path, sheet_name, task_id)
        log_info(f"[EXCEL_FILLER] [DEBUG] render_from_db返回结果: {result}", "excel_main_filler")
        
        return result

    def get_filled_data_summary(self, excel_file_path: str, sheet_name: str, test_count: int) -> Dict[str, Any]:
        """代理到 ExcelDataFiller 获取摘要"""
        return self.excel_data_filler.get_filled_data_summary(excel_file_path, sheet_name, test_count)

    def backup_sheet_data(self, excel_file_path: str, sheet_name: str, backup_file_path: str) -> bool:
        """备份工作表到指定路径（与 DataFiller 的简化备份区分）"""
        try:
            import shutil
            # 直接复制整个文件到备份路径
            shutil.copy2(excel_file_path, backup_file_path)
            log_info(f"[EXCEL_FILLER] 备份完成: {backup_file_path}", "excel_main_filler")
            return True
        except Exception as e:
            log_error(f"[EXCEL_FILLER] 备份失败: {e}", "excel_main_filler")
            return False

    def update_single_test_item(self, excel_file_path: str, sheet_name: str, item_index: int, field_name: str, field_value: str) -> bool:
        """
        更新单个测试项的指定字段。
        item_index 从1开始；数据区从第2行开始。
        支持字段: name(15), test_description(19), test_method(20), chapter_number(10),
                 test_type(14), test_item_identifier(16), software_requirement_identifier(12), software_requirement_name(11)
        """
        try:
            column_map = {
                'name': 15,
                'test_item_name': 15,
                'test_description': 19,
                'test_method': 20,
                'chapter_number': 10,
                'test_type': 14,
                'test_item_identifier': 16,
                'software_requirement_identifier': 12,
                'software_requirement_name': 11,
            }
            col = column_map.get(field_name)
            if not col:
                log_warning(f"[EXCEL_FILLER] 不支持的字段: {field_name}", "excel_main_filler")
                return False

            wb = openpyxl.load_workbook(excel_file_path)
            if sheet_name not in wb.sheetnames:
                log_error(f"[EXCEL_FILLER] 工作表不存在: {sheet_name}", "excel_main_filler")
                return False
            sheet = wb[sheet_name]
            row = 1 + item_index  # 数据从第2行开始
            sheet.cell(row=row, column=col, value=field_value)
            wb.save(excel_file_path)
            return True
        except Exception as e:
            log_error(f"[EXCEL_FILLER] 更新单项失败: {e}", "excel_main_filler")
            return False

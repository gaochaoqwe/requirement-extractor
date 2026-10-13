"""
"""
# 导入任务日志函数
# 配置日志

import logging
import openpyxl
import os
from typing import Dict, List, Optional, Any, Tuple
from pathlib import Path
from ...document_parsing.utils.text_utils import extract_requirement_name

from .....utils.logging_utils.logger import log_info, log_debug, log_error, log_warning

logger = logging.getLogger(__name__)


class DatabasePersistence:
    """数据库持久化"""

    def __init__(self):
        self.logger = logging.getLogger(__name__)

    def _persist_rows_to_db(self, basic_info: Dict[str, Any], test_items: List[Dict[str, Any]], task_id: str) -> None:
        """将 test_items 逐行计算并保存到数据库（不写Excel）。"""
        log_info(f"[DB_PERSIST] [DEBUG] 开始持久化到数据库，Task ID: {task_id}", "database_persistence")
        log_info(f"[DB_PERSIST] [DEBUG] 测试项数量: {len(test_items)}", "database_persistence")
        
        try:
            log_info(f"[DB_PERSIST] [DEBUG] 导入必要模块...", "database_persistence")
            from ....content_analysis.test_type_classifier import ceshi_leixing_biaoshi_choose
            from ....test_generation.ai_content.ai_content_optimizer import translate_chinese_to_english, batch_translate_requirements
            from .....db.optimized_models import OptimizedDynamicTestDatabase
            log_info(f"[DB_PERSIST] [DEBUG] 模块导入成功", "database_persistence")
        except Exception as e:
            log_error(f"[DB_PERSIST] [DEBUG] 模块导入失败: {e}", "database_persistence")
            raise
        
        # 统一使用日志工具写入任务日志文件

        db = OptimizedDynamicTestDatabase()
        software_id = basic_info.get('software_identifier', '')
        software_version = basic_info.get('software_version', '')
        
        # 统计信息
        total_items = len(test_items)
        success_count = 0
        error_count = 0
        error_details = []
        
        log_info(f"[DB_BATCH_START] 开始批量持久化需求分析数据 - total_items: {total_items}, task_id: {task_id}", "database_persistence")

        # 预先计算第23列英文标识
        last_col_values: List[str] = [""] * len(test_items)
        translate_indices: List[int] = []
        translate_names: List[str] = []
        for i, test_item in enumerate(test_items):
            candidate = None
            # 0) 优先使用已有的 english_identifier（与Excel填充逻辑保持一致）
            eng_id = test_item.get('english_identifier', '')
            if isinstance(eng_id, str) and eng_id.strip():
                candidate = eng_id.strip()
            # 1) 接口/需求追踪字段中的标识
            if not candidate:
                trace_val = test_item.get('software_requirement_trace', '')
                if isinstance(trace_val, str) and trace_val.strip():
                    candidate = trace_val.strip()
            # 2) 需求正文结构化中的标识号
            if not candidate:
                desc_val = test_item.get('requirement_description', '')
                if isinstance(desc_val, dict):
                    candidate = desc_val.get('identifier') or desc_val.get('标识号')
            if candidate and str(candidate).strip():
                last_col_values[i] = str(candidate).strip()
            else:
                name_for_trans = test_item.get('pure_name') or test_item.get('name') or f'测试项{i+1}'
                translate_indices.append(i)
                translate_names.append(str(name_for_trans).strip())
        if translate_names:
            try:
                translated = batch_translate_requirements(translate_names)
                for idx, val in zip(translate_indices, translated):
                    last_col_values[idx] = val
            except Exception as e:
                log_warning(f"[DB_TRANSLATION_WARNING] 英文翻译失败，使用默认值 - error: {str(e)}", "database_persistence")

        for i, test_item in enumerate(test_items):
            # 计算纯名称(base_name)与全名(full_name)
            raw_name = test_item.get('name') or f'测试项{i+1}'
            pure_name = test_item.get('pure_name')
            if not pure_name:
                # 去除性能/边界测试后缀（兼容中英文括号）
                try:
                    name_tmp = str(raw_name)
                    for suf in ['（性能测试）', '（边界测试）', '(性能测试)', '(边界测试)']:
                        if name_tmp.endswith(suf):
                            name_tmp = name_tmp[: -len(suf)]
                    pure_name = name_tmp.strip()
                except Exception:
                    pure_name = str(raw_name)
            
            # 🔧 修复：确保需求名称不带任何测试类型后缀
            def extract_pure_requirement_name(name: str) -> str:
                """提取纯净的需求名称，去除所有测试类型后缀"""
                if not name:
                    return ''
                
                # 去除各种测试类型后缀（支持中英文括号）
                suffixes_to_remove = [
                    '（性能测试）', '（边界测试）', '（余量测试）', '（强度测试）',
                    '(性能测试)', '(边界测试)', '(余量测试)', '(强度测试)'
                ]
                
                clean_name = str(name).strip()
                for suffix in suffixes_to_remove:
                    if clean_name.endswith(suffix):
                        clean_name = clean_name[:-len(suffix)].strip()
                        break
                
                return clean_name
            
            # 获取纯净的需求名称（不带后缀）
            pure_requirement_name = extract_pure_requirement_name(raw_name)
            
            base_name = pure_requirement_name  # 🔧 使用纯净名称作为base_name
            full_name = str(raw_name)
            requirement_name = base_name  # 兼容后续旧变量名使用
            # 追踪名称规范化全过程
            try:
                log_info(
                    f"[DB_NAME_TRACE] idx={i+1} raw_name='{raw_name}' -> pure_name='{pure_name}' | base_name='{base_name}' | full_name='{full_name}'",
                    "database_persistence",
                )
            except Exception:
                pass
            try:
                test_name = full_name
                _, test_type_label = ceshi_leixing_biaoshi_choose(test_name)
                try:
                    function_identifier = translate_chinese_to_english(str(base_name))
                except Exception:
                    function_identifier = f"FUNC{i+1:03d}"

                final_last_col = last_col_values[i] if last_col_values[i] else function_identifier

                # 构建追踪信息（第10/17列格式）
                software_name = basic_info.get('software_name', '软件')
                chapter_number = test_item.get('chapter_number', '')
                if chapter_number and chapter_number.strip() and chapter_number != '未找到具体章节号':
                    req_text = f"{chapter_number} {base_name}"
                else:
                    req_text = str(base_name)
                tracking_info = f"《{software_name}软件需求规格说明》 {req_text}"

                # 计算子描述（仅写入对应小节）
                full_desc = test_item.get('requirement_description', '')
                sub_desc = test_item.get('requirement_sub_description')
                if not sub_desc:
                    try:
                        if isinstance(full_desc, dict):
                            full_text = str(full_desc.get('text') or '')
                        else:
                            full_text = str(full_desc or '')
                        lines = [ln.strip() for ln in full_text.splitlines() if ln and ln.strip()]
                        selected: List[str] = []
                        if '性能' in test_type_label:
                            for ln in lines:
                                if ln.startswith('g)') or '性能' in ln or '性能：' in ln:
                                    selected.append(ln)
                        elif '边界' in test_type_label:
                            for ln in lines:
                                if ln.startswith('h)') or '边界' in ln or '边界：' in ln:
                                    selected.append(ln)
                        sub_desc = '\n'.join(selected).strip() if selected else full_text
                    except Exception:
                        sub_desc = full_desc

                # 🔧 修复：根据测试类型生成正确的标识符后缀
                if test_type_label == '性能测试':
                    type_suffix = 'PF01'
                elif test_type_label == '边界测试':
                    type_suffix = 'BD01'
                else:
                    type_suffix = 'FN01'  # 功能测试或其他
                
                row_index = test_item.get('index', i + 1)
                row_data = {
                    'requirement_id': f"REQ_{task_id}_{row_index:03d}",
                    'software_name': software_name,
                    'software_version': software_version,
                    'round_info': '第一轮',
                    'standard_sub_requirement_name': base_name,
                    'standard_sub_requirement_identifier': f"CT-{software_id}-{software_version}-{final_last_col}-{type_suffix}",
                    'standard_sub_requirement_unit': test_item.get('standard_sub_requirement_unit', '0.2'),
                    'task_requirement_trace': test_item.get('task_requirement_trace', '无'),
                    'software_requirement_trace': tracking_info,
                    'software_requirement_name': base_name,
                    'software_requirement_identifier': final_last_col,
                    'requirement_description': full_desc,
                    'test_type': test_type_label,
                    'test_item_name': full_name,
                    'test_item_identifier': f"CT-{software_id}-{software_version}-{final_last_col}-{type_suffix}",
                    'process_identifier': tracking_info,
                    'requirement_sub_description': sub_desc,
                    'test_description': test_item.get('test_description', ''),
                    'test_method': test_item.get('test_method', ''),
                    'adequacy_requirement': test_item.get('adequacy_requirement', ''),
                    'pass_criteria': test_item.get('pass_criteria', ''),
                    'english_identifier': final_last_col,
                    'software_identifier': software_id
                }
                # 追踪即将入库的两列值
                try:
                    log_info(
                        f"[DB_NAME_TRACE] SAVE idx={i+1} standard_sub_requirement_name(col6)='{row_data['standard_sub_requirement_name']}', "
                        f"software_requirement_name(col11)='{row_data['software_requirement_name']}'",
                        "database_persistence",
                    )
                except Exception:
                    pass
                
                # 尝试保存到数据库
                db.save_requirement_row(task_id, row_index, row_data)
                success_count += 1
                
            except Exception as e:
                error_count += 1
                error_msg = f"保存第{i+1}项失败: {requirement_name}, error: {str(e)}"
                error_details.append(error_msg)
                log_error(f"[DB_ITEM_ERROR] X 单项持久化失败 - item: {i+1}/{total_items}, name: {requirement_name}, error: {str(e)}", "database_persistence")
                # 不中断处理，继续处理下一项
                continue
        
        # 记录批量处理结果
        log_info(f"[DB_BATCH_COMPLETE] 批量持久化完成 - total: {total_items}, success: {success_count}, errors: {error_count}", "database_persistence")
        
        if error_count > 0:
            log_warning(f"[DB_BATCH_ERRORS] 持久化过程中出现 {error_count} 个错误，详情: {'; '.join(error_details[:5])}{'...' if len(error_details) > 5 else ''}", "database_persistence")
        
        if success_count == 0:
            raise Exception(f"所有 {total_items} 项数据都保存失败")
        elif error_count > 0:
            self.logger.warning(f"部分数据保存失败: {error_count}/{total_items}")
        else:
            log_info(f"[DB_BATCH_SUCCESS] √ 所有数据保存成功: {success_count}/{total_items}", "database_persistence")


    def render_from_db(self, excel_file_path: str, sheet_name: str, task_id: str) -> bool:
        """从数据库读取行并渲染到Excel（按index_number顺序）。"""
        try:
            log_info(f"[RENDER_DB] [DEBUG] 开始从数据库渲染Excel", "database_persistence")
            log_info(f"[RENDER_DB] [DEBUG] Excel文件: {excel_file_path}", "database_persistence")
            log_info(f"[RENDER_DB] [DEBUG] 工作表: {sheet_name}", "database_persistence")
            log_info(f"[RENDER_DB] [DEBUG] Task ID: {task_id}", "database_persistence")
            
            from .....db.optimized_models import OptimizedDynamicTestDatabase
            log_info(f"[RENDER_DB] [DEBUG] 数据库模块导入成功", "database_persistence")
            
            db = OptimizedDynamicTestDatabase()
            log_info(f"[RENDER_DB] [DEBUG] 数据库实例创建成功", "database_persistence")
            
            rows = db.get_requirement_rows(task_id)
            log_info(f"[RENDER_DB] [DEBUG] 从数据库获取到 {len(rows)} 行数据", "database_persistence")

            wb = openpyxl.load_workbook(excel_file_path)
            log_info(f"[RENDER_DB] [DEBUG] Excel文件加载成功", "database_persistence")
            log_info(f"[RENDER_DB] [DEBUG] 可用工作表: {wb.sheetnames}", "database_persistence")
            
            if sheet_name not in wb.sheetnames:
                log_error(f"[RENDER_DB] [DEBUG] 工作表 '{sheet_name}' 不存在", "database_persistence")
                return False
            sheet = wb[sheet_name]
            log_info(f"[RENDER_DB] [DEBUG] 工作表获取成功", "database_persistence")

            start_row = 2
            log_info(f"[RENDER_DB] [DEBUG] 开始填充数据，从第{start_row}行开始", "database_persistence")
            
            for i, r in enumerate(rows):
                row = start_row + i
                if i < 3:  # 只记录前3行的详细信息
                    log_info(f"[RENDER_DB] [DEBUG] 填充第{i+1}行数据到Excel行{row}", "database_persistence")
                english_id = r.get('english_identifier') or r.get('software_requirement_identifier') or ''
                software_id = r.get('software_identifier') or ''
                software_version = r.get('software_version') or ''

                # 写入列
                sheet.cell(row=row, column=2, value=i + 1)  # 修复：序号从1开始连续递增
                sheet.cell(row=row, column=3, value=r.get('software_name'))
                sheet.cell(row=row, column=4, value=software_version)
                sheet.cell(row=row, column=5, value=r.get('round_info'))
                sheet.cell(row=row, column=6, value=r.get('standard_sub_requirement_name'))
                # 列7/16依赖列23(english_id)
                sheet.cell(row=row, column=7, value=r.get('test_item_identifier') or f"CT-{software_id}-{software_version}-{english_id}-FN01")
                sheet.cell(row=row, column=8, value=r.get('standard_sub_requirement_unit'))
                sheet.cell(row=row, column=9, value=r.get('task_requirement_trace'))
                sheet.cell(row=row, column=10, value=r.get('software_requirement_trace'))
                sheet.cell(row=row, column=11, value=r.get('software_requirement_name'))
                # 列12严格等于列23
                sheet.cell(row=row, column=12, value=english_id)
                sheet.cell(row=row, column=13, value=r.get('requirement_description'))
                sheet.cell(row=row, column=14, value=r.get('test_type'))
                sheet.cell(row=row, column=15, value=r.get('test_item_name'))
                sheet.cell(row=row, column=16, value=r.get('test_item_identifier') or f"CT-{software_id}-{software_version}-{english_id}-FN01")
                sheet.cell(row=row, column=17, value=r.get('software_requirement_trace'))
                sheet.cell(row=row, column=18, value=r.get('requirement_sub_description'))
                sheet.cell(row=row, column=19, value=r.get('test_description'))
                sheet.cell(row=row, column=20, value=r.get('test_method'))
                sheet.cell(row=row, column=21, value=r.get('adequacy_requirement'))
                sheet.cell(row=row, column=22, value=r.get('pass_criteria'))
                sheet.cell(row=row, column=23, value=english_id)
                # 追踪写入到Excel的两列值
                try:
                    col6_val = "" if r.get('standard_sub_requirement_name') is None else str(r.get('standard_sub_requirement_name'))
                    col11_val = "" if r.get('software_requirement_name') is None else str(r.get('software_requirement_name'))
                    log_info(
                        f"[RENDER_NAME_TRACE] row={row} col6='{col6_val}', col11='{col11_val}'",
                        "database_persistence",
                    )
                except Exception:
                    pass

            log_info(f"[RENDER_DB] [DEBUG] 数据填充完成，开始保存Excel文件", "database_persistence")
            wb.save(excel_file_path)
            log_info(f"[RENDER_DB] [DEBUG] Excel文件保存成功", "database_persistence")
            
            self.logger.info(f"从数据库渲染完成，共写入{len(rows)}行")
            return True
        except Exception as e:
            import traceback
            log_error(f"[RENDER_DB] [DEBUG] 从数据库渲染Excel失败: {e}", "database_persistence")
            log_error(f"[RENDER_DB] [DEBUG] 异常堆栈: {traceback.format_exc()}", "database_persistence")
            return False
    


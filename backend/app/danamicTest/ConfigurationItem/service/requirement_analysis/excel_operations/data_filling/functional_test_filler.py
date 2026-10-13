"""
"""
# 导入任务日志函数
# 配置日志

import logging
import openpyxl
import os
from typing import Dict, List, Optional, Any, Tuple
from pathlib import Path
from ...document_parsing.utils.text_utils import extract_requirement_name

from .....utils.logging_utils.logger import log_info, log_debug, log_error, log_warning
from .excel_data_filler import ExcelDataFiller

logger = logging.getLogger(__name__)


class FunctionalTestFiller:
    """功能测试填充器"""

    def __init__(self):
        self.logger = logging.getLogger(__name__)

    def fill_requirement_analysis_sheet_legacy(
        self,
        excel_file_path: str,
        sheet_name: str,
        ruanjian_xinxi: Dict[str, List[str]],
        ruanjian_name: str,
        diama_biangeng_xuqiu: Dict[str, str],
        file_path: str,
        ceshileixing: str,
        neirong_dict: List[str],
        biaoshihao_ceshixianng: List[str],
        AI_fenx_flag_Anniu: int,
        jianceyi_flag: int
    ) -> bool:
        """
        填充配置项测试需求分析表内容（兼容原有接口）
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            ruanjian_xinxi: 软件信息字典
            ruanjian_name: 软件名称
            diama_biangeng_xuqiu: 代码变更需求字典
            file_path: 文档文件路径
            ceshileixing: 测试类型
            neirong_dict: 内容描述列表
            biaoshihao_ceshixianng: 标识号测试项列表
            AI_fenx_flag_Anniu: AI分析标志
            jianceyi_flag: 检测仪标志
            
        Returns:
            bool: 填充是否成功
        """
        try:
            self.logger.info(f"[FIX] 开始填充配置项测试需求分析表（兼容模式）: {sheet_name}")
            
            # [ALERT] 详细日志：输入参数监控
            self.logger.info(f"📋 输入参数监控:")
            self.logger.info(f"   - Excel文件: {excel_file_path}")
            self.logger.info(f"   - 工作表名: {sheet_name}")
            self.logger.info(f"   - 软件名称: {ruanjian_name}")
            self.logger.info(f"   - 测试类型: {ceshileixing}")
            self.logger.info(f"   - diama_biangeng_xuqiu数量: {len(diama_biangeng_xuqiu)}")
            self.logger.info(f"   - neirong_dict数量: {len(neirong_dict)}")
            self.logger.info(f"   - biaoshihao_ceshixianng数量: {len(biaoshihao_ceshixianng)}")
            
            # [ALERT] 关键数据内容监控
            self.logger.info(f"[SEARCH] 关键数据内容:")
            self.logger.info(f"   - diama_biangeng_xuqiu.keys(): {list(diama_biangeng_xuqiu.keys())}")
            self.logger.info(f"   - biaoshihao_ceshixianng: {biaoshihao_ceshixianng}")
            
            # 打开Excel工作簿
            wb = openpyxl.load_workbook(excel_file_path)
            
            # 获取工作表
            if sheet_name not in wb.sheetnames:
                self.logger.error(f"[ERROR] 工作表 '{sheet_name}' 不存在，可用工作表: {wb.sheetnames}")
                return False
            
            sheet = wb[sheet_name]
            self.logger.info(f"[SUCCESS] 成功打开工作表: {sheet_name} ({sheet.max_row}行 x {sheet.max_column}列)")
            
            # 基本信息
            len_dict = len(diama_biangeng_xuqiu)
            allkeys = list(diama_biangeng_xuqiu.keys())
            
            # [ALERT] 数据一致性检查
            self.logger.info(f"📊 数据一致性检查:")
            self.logger.info(f"   - allkeys: {allkeys}")
            self.logger.info(f"   - len_dict: {len_dict}")
            self.logger.info(f"   - len(allkeys): {len(allkeys)}")
            self.logger.info(f"   - len(biaoshihao_ceshixianng): {len(biaoshihao_ceshixianng)}")
            self.logger.info(f"   - len(neirong_dict): {len(neirong_dict)}")
            
            # 读取文档获取章节号
            zhangjiehao, zhangjiehao_du = self._read_document_chapters(file_path, allkeys, ruanjian_name)
            
            # 生成测试需求标识前缀
            ceshi_xuqiu_biaoshi = self._generate_test_requirement_prefix(ruanjian_xinxi, ceshileixing)
            
            # 数据验证
            if len(neirong_dict) != len_dict:
                self.logger.warning('测试项数量与文档数量不一致！')
                neirong_dict = ['' for _ in range(len_dict)]
            if len(zhangjiehao) != len_dict:
                self.logger.warning('章节号数量与文档数量不一致！')
            
            # 生成测试内容
            test_descriptions, test_methods = self._generate_test_content_legacy(
                allkeys, neirong_dict, AI_fenx_flag_Anniu, jianceyi_flag
            )
            
            # 填充Excel数据
            # log_info(f"[FIX] [DEBUG] 开始调用 _fill_excel_data_legacy", "excel_filler")
            self._fill_excel_data_legacy(
                sheet, len_dict, allkeys, ruanjian_xinxi, ruanjian_name,
                zhangjiehao, neirong_dict, biaoshihao_ceshixianng,
                ceshi_xuqiu_biaoshi, test_descriptions, test_methods
            )
            # log_info(f"[SUCCESS] [DEBUG] _fill_excel_data_legacy 调用完成", "excel_filler")
            
            # [SEARCH] 在调用legacy方法后，检查第10列的值
            # log_info(f"[DEBUG] 调用 _fill_excel_data_legacy 后，检查第10列的值:", "excel_filler")
            for row in range(2, min(6, len_dict + 2)):  # 检查前几行
                current_value = sheet.cell(row=row, column=10).value
                # log_info(f"         - 行{row}列10当前值: '{current_value}'", "excel_filler")
            
            # 保存文件
            wb.save(excel_file_path)
            
            self.logger.info("配置项测试需求分析表填充完成（兼容模式）")
            return True
            
        except Exception as e:
            self.logger.error(f"填充配置项测试需求分析表时发生错误: {str(e)}")
            return False
    

    def fill_requirement_analysis_sheet(
        self,
        excel_file_path: str,
        sheet_name: str,
        requirement_analysis_data: Dict[str, Any],
        task_id: Optional[str] = None,
        persist_only: bool = False
    ) -> bool:
        """
        填充需求分析表工作表
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            requirement_analysis_data: 需求分析表数据
            
        Returns:
            bool: 填充是否成功
        """
        try:
            self.logger.info(f"开始填充配置项测试需求分析表: {sheet_name}")
            
            # Debug: 记录excel_filler入口参数 - 注释掉详细的DEBUG信息
            # log_info(f"[DEBUG] ConfigurationItemRequirementAnalysisExcelFiller.fill_requirement_analysis_sheet 开始", "excel_filler")
            # log_debug(f"   excel_file_path: {excel_file_path}", "excel_filler")
            # log_debug(f"   sheet_name: {sheet_name}", "excel_filler")
            # log_debug(f"   requirement_analysis_data: {requirement_analysis_data}", "excel_filler")
            
            # 若只入库，不写Excel
            if persist_only:
                if not task_id:
                    self.logger.error("persist_only=True 时必须提供 task_id")
                    return False
                basic_info = requirement_analysis_data.get('basic_info', {})
                test_items = requirement_analysis_data.get('test_items', [])
                self._persist_rows_to_db(basic_info, test_items, task_id)
                self.logger.info("仅持久化到数据库完成(persist_only=True)")
                return True

            # 打开Excel工作簿
            wb = openpyxl.load_workbook(excel_file_path)

            # 获取工作表
            if sheet_name not in wb.sheetnames:
                self.logger.error(f"工作表 '{sheet_name}' 不存在")
                return False

            sheet = wb[sheet_name]
            
            # 获取数据
            basic_info = requirement_analysis_data.get('basic_info', {})
            test_items = requirement_analysis_data.get('test_items', [])
            
            # [WARNING] 问题发现：这里只调用了 _fill_excel_data_new，没有调用 _fill_excel_data_legacy
            # log_warning(f"[ALERT] [DEBUG] 仅调用 _fill_excel_data_new 方法!", "excel_filler")
            # log_info(f"[DEBUG] 准备调用 _fill_excel_data_new", "excel_filler")
            self.logger.debug(f"   basic_info: {basic_info}")
            self.logger.debug(f"   test_items: {test_items}")
            
            # [SEARCH] 在调用新方法前，先检查第10列的当前值
            # log_info(f"[DEBUG] 调用 _fill_excel_data_new 前，检查第10列的值:", "excel_filler")
            for row in range(2, min(6, len(test_items) + 2)):  # 检查前几行
                current_value = sheet.cell(row=row, column=10).value
                # log_info(f"         - 行{row}列10当前值: '{current_value}'", "excel_filler")
            
            # 填充数据
            self._fill_excel_data_new(sheet, basic_info, test_items, task_id=task_id)
            
            # [SEARCH] 在调用新方法后，再次检查第10列的值
            # log_info(f"[DEBUG] 调用 _fill_excel_data_new 后，检查第10列的值:", "excel_filler")
            for row in range(2, min(6, len(test_items) + 2)):  # 检查前几行
                current_value = sheet.cell(row=row, column=10).value
                # log_info(f"         - 行{row}列10当前值: '{current_value}'", "excel_filler")
            
            # 保存文件
            wb.save(excel_file_path)
            
            self.logger.info("配置项测试需求分析表填充完成")
            return True
            
        except Exception as e:
            import traceback
            error_details = traceback.format_exc()
            self.logger.error(f"填充配置项测试需求分析表时发生错误: {str(e)}")
            self.logger.error(f"错误详情: {error_details}")
            # log_error(f"[ALERT] [ERROR] fill_requirement_analysis_sheet异常: {str(e)}", "excel_filler")
            # log_error(f"[ALERT] [ERROR] 异常堆栈: {error_details}", "excel_filler")
            return False
    

    def _fill_excel_data_new(
        self,
        sheet,
        basic_info: Dict[str, Any],
        test_items: List[Dict[str, Any]],
        task_id: Optional[str] = None
    ) -> None:
        """新版填充实现的委托方法，交由 ExcelDataFiller 完成具体写入"""
        try:
            filler = ExcelDataFiller()
            filler.fill_excel_data_new(sheet, basic_info, test_items)
        except Exception as e:
            self.logger.error(f"_fill_excel_data_new 执行失败: {e}")


    def _generate_test_requirement_prefix(
        self,
        ruanjian_xinxi: Dict[str, List[str]],
        ceshileixing: str
    ) -> str:
        """生成测试需求标识前缀"""
        ruanjian_biaoshi = ruanjian_xinxi.get('软件标识', [''])[0]
        ruanjian_version = ruanjian_xinxi.get('被测软件版本', [''])[0]
        
        if ceshileixing == '回归测试':
            return 'CT_R-' + ruanjian_biaoshi + '_' + ruanjian_version + '_'
        elif ceshileixing == '配置项测试':
            return 'CT-' + ruanjian_biaoshi + '_' + ruanjian_version + '_'
        else:  # 系统测试
            return 'ST-' + ruanjian_biaoshi + '_' + ruanjian_version + '_'
    

    def _generate_test_content_legacy(
        self,
        allkeys: List[str],
        neirong_dict: List[str],
        AI_fenx_flag_Anniu: int,
        jianceyi_flag: int
    ) -> Tuple[List[str], List[str]]:
        """生成测试内容（兼容模式）"""
        try:
            from ...test_generation.content_generators.test_content_generator import ConfigurationItemTestContentGenerator
            
            generator = ConfigurationItemTestContentGenerator()
            
            # 构造测试项数据
            test_items = []
            for i, test_name in enumerate(allkeys):
                test_item = {
                    'name': test_name,
                    'requirement_description': neirong_dict[i] if i < len(neirong_dict) else ''
                }
                test_items.append(test_item)
            
            # 构造AI配置 - 使用AI开关控制器
            from .....utils.ai_switch_controller.ai_switch_controller import get_ai_controller, AIModule
            ai_controller = get_ai_controller()
            ai_config = ai_controller.get_ai_config(AIModule.REQUIREMENT_ANALYSIS)
            
            # 添加额外配置
            ai_config.update({
                'detector_flag': jianceyi_flag,
                'use_knowledge_optimization': True
            })
            
            # 生成测试内容
            test_descriptions, test_methods = generator.generate_test_content(test_items, ai_config)
            
            return test_descriptions, test_methods
            
        except Exception as e:
            self.logger.error(f"生成测试内容时发生错误: {str(e)}")
            # 返回默认内容
            default_descriptions = [f'测试软件{name}功能实现是否正常' for name in allkeys]
            default_methods = [f'通过检测仪发送测试指令,通过检测仪状态监控界面查看测试软件{name}功能实现是否正常' for name in allkeys]
            return default_descriptions, default_methods
    

    def update_single_test_item(
        self,
        excel_file_path: str,
        sheet_name: str,
        item_index: int,
        field_name: str,
        field_value: str
    ) -> bool:
        """
        更新单个测试项的字段
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            item_index: 测试项索引（从1开始）
            field_name: 字段名称
            field_value: 字段值
            
        Returns:
            bool: 更新是否成功
        """
        try:
            self.logger.info(f"更新测试项{item_index}的字段{field_name}")
            
            # 打开Excel工作簿
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            # 计算行号（测试项从第2行开始）
            row = 1 + item_index
            
            # 字段映射
            field_mapping = {
                'index': 2,                       # 序号
                'software_name': 3,               # 软件名称
                'software_version': 4,            # 软件版本
                'round_info': 5,                  # 轮次信息
                'standard_sub_requirement_name': 6, # 标准子需求名称
                'standard_sub_requirement_identifier': 7, # 标准子需求标识
                'standard_sub_requirement_unit': 8, # 标准子需求测试单元
                'task_requirement_trace': 9,     # 任务需求追踪
                'software_requirement_trace': 10, # 软件需求规格说明追踪
                'software_requirement_name': 11,  # 软件需求规格说明需求名称
                'software_requirement_identifier': 12, # 软件需求规格说明需求标识
                'requirement_description': 13,   # 需求描述
                'test_type': 14,                  # 测试类型
                'test_item_name': 15,             # 测试项名称
                'test_item_identifier': 16,       # 测试项标识
                'process_identifier': 17,         # 过程标识
                'requirement_sub_description': 18, # 需求子项描述
                'test_description': 19,           # 测试项描述
                'test_method': 20,                # 测试方法
                'adequacy_requirement': 21,       # 充分性要求
                'pass_criteria': 22               # 通过准则
            }
            
            # 获取列号
            if field_name not in field_mapping:
                self.logger.error(f"未知的字段名称: {field_name}")
                return False
            
            col = field_mapping[field_name]
            
            # 更新单元格
            sheet.cell(row=row, column=col).value = field_value
            
            # 保存文件
            wb.save(excel_file_path)
            
            self.logger.info(f"测试项{item_index}的字段{field_name}更新完成")
            return True
            
        except Exception as e:
            self.logger.error(f"更新单个测试项字段时发生错误: {str(e)}")
            return False
    


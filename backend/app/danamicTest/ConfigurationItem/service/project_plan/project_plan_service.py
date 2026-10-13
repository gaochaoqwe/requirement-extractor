"""
项目策划服务模块

负责项目策划数据的提取、存储和管理功能
"""

import openpyxl
from typing import List, Dict, Any, Optional
from pathlib import Path


class ProjectPlanService:
    """项目策划服务类"""
    
    def __init__(self):
        """初始化项目策划服务"""
        pass
    
    def extract_and_store_excel_data_to_db(
        self, 
        excel_file_path: str, 
        sheet_name: str = '项目策划表',
        task_id: str = None
    ) -> bool:
        """
        从项目策划Excel表中提取数据并存储到数据库
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            task_id: 任务ID
            
        Returns:
            bool: 是否成功
        """
        try:
            from ...db.optimized_models import OptimizedDynamicTestDatabase
            from ...utils.logging_utils.logger import log_info, log_warning, log_error
            
            log_info(f"开始从Excel提取项目策划数据: {excel_file_path}", "project_plan_service")
            
            # 打开Excel文件
            wb = openpyxl.load_workbook(excel_file_path)
            if sheet_name not in wb.sheetnames:
                log_error(f"工作表 {sheet_name} 不存在", "project_plan_service")
                return False
            
            sheet = wb[sheet_name]
            
            # 初始化数据库
            db = OptimizedDynamicTestDatabase()
            
            # 执行数据库迁移（如果需要）
            db.migrate_project_plan_table()
            
            # 生成任务ID（如果未提供）
            if not task_id:
                import uuid
                task_id = str(uuid.uuid4())
                log_info(f"生成新的任务ID: {task_id}", "project_plan_service")
            
            # 定义Excel列映射（与project_plan/excel_filler.py中的填充逻辑保持一致）
            # 从第18行开始的数据区域，A到M列
            excel_column_mapping = {
                1: 'item_id',                           # A: ID（通常保持空白）
                2: 'sequence_number',                   # B: 序号
                3: 'software_requirement',              # C: 软件需求
                4: 'standard_sub_requirement_name',     # D: 标准子需求名称
                5: 'standard_sub_requirement_unit',     # E: 标准子需求单位
                6: 'requirement_tracking',              # F: 软件需求规格说明/影响域分析报告追踪
                7: 'importance',                        # G: 重要性
                8: 'test_completed',                    # H: 是否完成测试
                9: 'requirement_analysis_personnel',    # I: 需求分析（测试主管）
                10: 'document_review_personnel',        # J: 文档审查（测试人员）
                11: 'static_analysis_personnel',        # K: 静态分析（测试人员）
                12: 'code_review_personnel',            # L: 代码审查（测试人员）
                13: 'dynamic_test_personnel'            # M: 动态测试（测试人员）
            }
            
            # 提取数据 - 从第18行开始
            extracted_data = []
            row_num = 18  # 从第18行开始（项目策划表的数据区域起始行）
            
            while True:
                # 检查是否有数据（以第2列序号是否有值为准）
                sequence_cell = sheet.cell(row=row_num, column=2)
                if not sequence_cell.value:
                    break
                
                # 提取该行的所有数据
                row_data = {'plan_item_id': f"plan_{task_id}_{row_num}"}
                
                for col_num, field_name in excel_column_mapping.items():
                    cell_value = sheet.cell(row=row_num, column=col_num).value
                    
                    # 处理不同类型的值
                    if cell_value is not None:
                        if field_name == 'sequence_number':
                            # 序号转换为整数
                            try:
                                row_data[field_name] = int(cell_value)
                            except (ValueError, TypeError):
                                row_data[field_name] = row_num - 17  # 使用行号作为序号（从18行开始，所以减17）
                        else:
                            # 其他字段转换为字符串
                            row_data[field_name] = str(cell_value).strip()
                    else:
                        row_data[field_name] = None
                
                extracted_data.append(row_data)
                row_num += 1
            
            log_info(f"从Excel提取到 {len(extracted_data)} 行项目策划数据", "project_plan_service")
            
            # 存储到数据库
            if extracted_data:
                db.save_project_plan(task_id, extracted_data)
                log_info(f"成功存储 {len(extracted_data)} 条项目策划数据到数据库", "project_plan_service")
                
                # 记录详细的存储信息
                for i, data in enumerate(extracted_data, 1):
                    log_info(f"  第{i}行: {data.get('software_requirement', 'N/A')} - {data.get('importance', 'N/A')}", "project_plan_service")
            
            return True
            
        except Exception as e:
            log_error(f"从Excel提取并存储项目策划数据失败: {str(e)}", "project_plan_service")
            import traceback
            traceback.print_exc()
            return False
    
    def get_project_plan_by_task(self, task_id: str) -> List[Dict[str, Any]]:
        """根据任务ID获取项目策划数据"""
        try:
            from ...db.optimized_models import OptimizedDynamicTestDatabase
            db = OptimizedDynamicTestDatabase()
            return db.get_project_plan_by_task(task_id)
        except Exception as e:
            print(f"获取项目策划数据失败: {e}")
            return []
    
    def get_all_project_plans(self) -> List[Dict[str, Any]]:
        """获取所有项目策划数据"""
        try:
            from ...db.optimized_models import OptimizedDynamicTestDatabase
            db = OptimizedDynamicTestDatabase()
            return db.get_all_project_plans()
        except Exception as e:
            print(f"获取所有项目策划数据失败: {e}")
            return []
    
    def update_project_plan_field(self, plan_id: int, field_name: str, field_value: Any) -> bool:
        """更新项目策划字段"""
        try:
            from ...db.optimized_models import OptimizedDynamicTestDatabase
            db = OptimizedDynamicTestDatabase()
            return db.update_project_plan_field(plan_id, field_name, field_value)
        except Exception as e:
            print(f"更新项目策划字段失败: {e}")
            return False 
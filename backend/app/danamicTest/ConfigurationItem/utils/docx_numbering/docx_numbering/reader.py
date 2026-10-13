"""
Main reader class for extracting numbered content from DOCX files.
"""

import re
from typing import Dict, List, Optional, Tuple, Union
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from .formatters import NumberFormatConverter
from .exceptions import InvalidDocxFileError, NumberingNotFoundError, NumberFormatError


class NumberedDocxReader:
    """
    A class for reading DOCX files and extracting automatic numbering information.
    
    This class can parse DOCX files and extract the automatic numbering that appears
    in the document, converting it to the appropriate text representation.
    """
    
    def __init__(self, docx_path: Union[str, Path], gap_text: str = "\t"):
        """
        Initialize the NumberedDocxReader.
        
        Args:
            docx_path: Path to the DOCX file
            gap_text: Text to use as separator between numbering and content (default: tab)
        """
        self.docx_path = Path(docx_path)
        self.gap_text = gap_text
        self.formatter = NumberFormatConverter()
        
        # Validate file path
        if not self.docx_path.exists():
            raise InvalidDocxFileError(f"DOCX file not found: {self.docx_path}")
        
        if not self.docx_path.is_file():
            raise InvalidDocxFileError(f"Path is not a file: {self.docx_path}")
        
        if self.docx_path.suffix.lower() != '.docx':
            raise InvalidDocxFileError(f"File must have .docx extension: {self.docx_path}")
        
        # Load the document
        try:
            self.document = Document(str(self.docx_path))
        except Exception as e:
            raise InvalidDocxFileError(f"Failed to load DOCX file: {e}")
        
        # Parse numbering information
        self.numbering_styles = self._parse_numbering_styles()
        
        # Initialize counters for tracking numbering state
        self.counters = {}
        self.number_cache = {}
    
    def _parse_numbering_styles(self) -> Dict[Tuple[int, int], Dict[str, str]]:
        """
        Parse numbering styles from the document's numbering.xml.
        
        Returns:
            Dictionary mapping (numId, level) to style information
        """
        numbering_styles = {}
        
        # Get numbering part from document
        numbering_part = self.document.part.numbering_part
        if numbering_part is None:
            return numbering_styles
        
        try:
            numbering_element = numbering_part._element
        except Exception as e:
            # If we can't access numbering element, return empty styles
            return numbering_styles
        
        # Create mapping from abstractNumId to numId
        abstract_to_num = {}
        try:
            for num_element in numbering_element.findall(qn("w:num")):
                num_id = int(num_element.get(qn("w:numId")))
                abstract_num_id_element = num_element.find(qn("w:abstractNumId"))
                if abstract_num_id_element is not None:
                    abstract_num_id = int(abstract_num_id_element.get(qn("w:val")))
                    abstract_to_num[abstract_num_id] = num_id
        except (ValueError, AttributeError, TypeError) as e:
            # If there's an error parsing the numbering structure, return empty styles
            return numbering_styles
        
        # Parse abstract numbering definitions
        try:
            for abstract_num in numbering_element.findall(qn("w:abstractNum")):
                abstract_num_id = int(abstract_num.get(qn("w:abstractNumId")))
                
                if abstract_num_id not in abstract_to_num:
                    continue
                    
                num_id = abstract_to_num[abstract_num_id]
                
                # Parse each level
                for level in abstract_num.findall(qn("w:lvl")):
                    level_id = int(level.get(qn("w:ilvl")))
                    
                    # Extract style attributes
                    style = {}
                    for child in level:
                        try:
                            tag_name = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                            val_attr = child.get(qn("w:val"))
                            if val_attr:
                                style[tag_name] = val_attr
                        except (AttributeError, IndexError):
                            continue
                    
                    # Handle special case for custom decimal formats
                    if "numFmt" not in style:
                        try:
                            # Look for custom format in AlternateContent
                            alt_content = level.find(qn("mc:AlternateContent"))
                            if alt_content is not None:
                                fallback = alt_content.find(qn("mc:Fallback"))
                                if fallback is not None:
                                    num_fmt = fallback.find(qn("w:numFmt"))
                                    if num_fmt is not None and num_fmt.get(qn("w:val")) == "decimal":
                                        choice = alt_content.find(qn("mc:Choice"))
                                        if choice is not None:
                                            custom_fmt = choice.find(qn("w:numFmt"))
                                            if custom_fmt is not None:
                                                format_attr = custom_fmt.get(qn("w:format"))
                                                if format_attr:
                                                    style["numFmt"] = "decimal" + format_attr.split(",")[0]
                        except (AttributeError, IndexError):
                            pass
                    
                    # Handle decimalZero format
                    if style.get("numFmt") == "decimalZero":
                        style["numFmt"] = "decimal01"
                    
                    # Set default start value if not present
                    if "start" not in style:
                        style["start"] = "1"
                    
                    numbering_styles[(num_id, level_id)] = style
        except (ValueError, AttributeError, TypeError) as e:
            # If there's an error parsing abstract numbering, continue with what we have
            pass
        
        return numbering_styles
    
    def _get_paragraph_numbering(self, paragraph) -> Optional[str]:
        """
        Get the numbering text for a paragraph.
        
        Args:
            paragraph: The paragraph element to check
            
        Returns:
            The numbering text if the paragraph has numbering, None otherwise
        """
        # Check if paragraph has numbering properties
        try:
            num_pr = paragraph._element.pPr
            if num_pr is None:
                return None
            
            num_pr_element = num_pr.find(qn("w:numPr"))
            if num_pr_element is None:
                return None
            
            # Get numbering ID and level
            num_id_element = num_pr_element.find(qn("w:numId"))
            ilvl_element = num_pr_element.find(qn("w:ilvl"))
            
            if num_id_element is None or ilvl_element is None:
                return None
            
            num_id = int(num_id_element.get(qn("w:val")))
            level = int(ilvl_element.get(qn("w:val")))
            
            # Skip if numId is 0 (no numbering)
            if num_id == 0:
                return None
                
        except (AttributeError, ValueError, TypeError):
            return None
        
        # Get style information
        style_key = (num_id, level)
        if style_key not in self.numbering_styles:
            return None
        
        style = self.numbering_styles[style_key]
        
        # Try to get the actual numbering value from the document first
        try:
            # First, try to extract the numbering from the paragraph text itself
            para_text = paragraph.text.strip()
            manual_num = self._extract_displayed_numbering(para_text)
            if manual_num:
                # Use the displayed numbering directly
                return manual_num + self.gap_text
            
            # Fallback: Update counter (but reset appropriately)
            if style_key not in self.counters:
                self.counters[style_key] = int(style.get("start", "1"))
            else:
                # Reset lower level counters when a higher level changes
                num_id, current_level = style_key
                for (counter_num_id, counter_level) in list(self.counters.keys()):
                    if counter_num_id == num_id and counter_level > current_level:
                        del self.counters[(counter_num_id, counter_level)]
                
                self.counters[style_key] += 1
            
            current_num = self.counters[style_key]
            
            # Format the number according to the style
            formatted_number = self._format_number(current_num, style.get("numFmt", "decimal"))
            
            # Cache the formatted number for this level
            self.number_cache[style_key] = formatted_number
            
            # Build the level text
            level_text = style.get("lvlText", "%1")
            
            # Replace placeholders with actual numbers
            for i in range(level + 1):
                placeholder = f"%{i + 1}"
                level_key = (num_id, i)
                if level_key in self.number_cache:
                    level_text = level_text.replace(placeholder, self.number_cache[level_key])
                else:
                    level_text = level_text.replace(placeholder, "")
            
            # Add suffix
            suffix = style.get("suff", "tab")
            if suffix == "space":
                level_text += " "
            elif suffix == "nothing":
                pass
            else:  # default to tab
                level_text += self.gap_text
            
            return level_text
        except (ValueError, TypeError) as e:
            # If there's an error processing the numbering, return empty string
            return ""
    
    def _format_number(self, num: int, format_type: str) -> str:
        """
        Format a number according to the specified format type.
        
        Args:
            num: The number to format
            format_type: The format type (e.g., 'decimal', 'upperRoman', etc.)
            
        Returns:
            The formatted number as a string
        """
        try:
            if format_type.startswith("decimal"):
                # Handle zero-padded decimal formats
                if format_type == "decimal":
                    return str(num)
                else:
                    # Extract padding length from format like "decimal001"
                    padding = len(format_type) - len("decimal")
                    return str(num).zfill(padding)
            
            elif format_type == "upperRoman":
                return self.formatter.int_to_upper_roman(num)
            
            elif format_type == "lowerRoman":
                return self.formatter.int_to_upper_roman(num).lower()
            
            elif format_type == "upperLetter":
                return self.formatter.int_to_upper_letter(num)
            
            elif format_type == "lowerLetter":
                return self.formatter.int_to_upper_letter(num).lower()
            
            elif format_type == "ordinal":
                # Handle ordinal numbers (1st, 2nd, 3rd, etc.)
                if 11 <= num <= 13:
                    return f"{num}th"
                else:
                    endings = {1: "st", 2: "nd", 3: "rd"}
                    return f"{num}{endings.get(num % 10, 'th')}"
            
            elif format_type == "cardinalText":
                return self.formatter.int_to_cardinal_text(num)
            
            elif format_type == "ordinalText":
                return self.formatter.int_to_ordinal_text(num)
            
            elif format_type == "chineseCounting":
                return self.formatter.int_to_chinese_counting(num)
            
            elif format_type == "chineseLegalSimplified":
                return self.formatter.int_to_chinese_legal_simplified(num)
            
            elif format_type == "ideographTraditional":
                if 1 <= num <= len(self.formatter.IDEOGRAPH_TRADITIONAL):
                    return self.formatter.IDEOGRAPH_TRADITIONAL[num - 1]
                else:
                    return str(num)
            
            elif format_type == "ideographZodiac":
                if 1 <= num <= len(self.formatter.IDEOGRAPH_ZODIAC):
                    return self.formatter.IDEOGRAPH_ZODIAC[num - 1]
                else:
                    return str(num)
            
            elif format_type == "decimalEnclosedCircleChinese":
                # For circled numbers, just return the number
                return str(num)
            
            elif format_type == "bullet":
                # For bullet points, return a bullet character
                return "•"
            
            else:
                # Default to decimal format
                return str(num)
                
        except Exception:
            # If formatting fails, return decimal format
            return str(num)
    
    def get_paragraphs_with_numbering(self) -> List[str]:
        """
        Get all paragraphs with their numbering text.
        
        ⚠️ DEPRECATED: This method is deprecated. Use get_full_document_content() instead.

        Returns:
            List of strings, each containing the numbering text followed by paragraph content
        """
        import warnings
        warnings.warn("get_paragraphs_with_numbering() is deprecated. Use get_full_document_content() instead.", 
                     DeprecationWarning, stacklevel=2)
        # Reset counters for fresh read
        self.counters.clear()
        self.number_cache.clear()

        result = []

        for paragraph in self.document.paragraphs:
            numbering_text = self._get_paragraph_numbering(paragraph)
            if numbering_text is None:
                # 如果没有自动编号，尝试检测手动编号
                manual_numbering = self._extract_manual_numbering(paragraph.text)
                if manual_numbering:
                    numbering_text = manual_numbering + self.gap_text
                else:
                    numbering_text = ""

            result.append(numbering_text + paragraph.text)

        return result

    def _extract_displayed_numbering(self, text: str) -> Optional[str]:
        """
        从段落文本开头提取显示的编号
        
        Args:
            text: 段落文本
            
        Returns:
            提取到的编号，如果没有则返回None
        """
        if not text:
            return None
            
        text = text.strip()
        
        # 匹配各种编号格式：1. 1.1 1.2.1 (1) a) A) 等
        numbering_patterns = [
            r'^(\d+(?:\.\d+)*\.?)\s+',           # 1. 1.1 1.2.1 等
            r'^(\d+(?:\.\d+)*)\s+',              # 1 1.1 1.2 等（无点号）
            r'^(\(\d+\))\s+',                    # (1) (2) 等
            r'^([a-z]\))\s+',                    # a) b) c) 等
            r'^([A-Z]\))\s+',                    # A) B) C) 等
            r'^([IVXLCDM]+\.?)\s+',              # I. II. III. 等罗马数字
            r'^([一二三四五六七八九十百千万]+[、．.])\s+',  # 中文数字
        ]
        
        for pattern in numbering_patterns:
            match = re.match(pattern, text)
            if match:
                numbering = match.group(1)
                # 确保编号后面不是纯数字（避免误匹配）
                remaining_text = text[len(match.group(0)):].strip()
                if remaining_text and not remaining_text[0].isdigit():
                    return numbering
        
        return None
    
    def _extract_manual_numbering(self, text: str) -> Optional[str]:
        """
        尝试从段落文本中提取手动编号。

        Args:
            text: 段落文本

        Returns:
            提取的编号文本，如果没有找到则返回None
        """
        if not text or not text.strip():
            return None

        # 匹配各种编号格式的正则表达式（按优先级排序）
        patterns = [
            # 多级数字编号：1.1.1.1, 1.1.1, 1.1, 1 等（优先匹配更长的）
            r'^(\d+\.\d+\.\d+\.\d+)\s+',
            r'^(\d+\.\d+\.\d+)\s+',
            r'^(\d+\.\d+)\s+',
            # 单级数字编号：1, 2, 3 等（后面必须跟空格和非数字字符）
            r'^(\d+)\s+(?!\d)',
            # 带点的数字编号：1., 2., 3. 等
            r'^(\d+(?:\.\d+)*)\.\s+',
            # 中文编号：一、二、三、等
            r'^([一二三四五六七八九十百千万]+)、\s*',
            # 中文数字编号：1、2、3、等
            r'^(\d+)、\s*',
            # 罗马数字编号：I, II, III, IV 等（完整的罗马数字）
            r'^([IVXLCDM]+)\s+',
            r'^([ivxlcdm]+)\s+',
            # 字母编号：A, B, C 等（单个或多个字母）
            r'^([A-Z]+)\s+',
            r'^([a-z]+)\s+',
            # 带括号的编号：(1), (2), (a), (A), (一), (二) 等
            r'^\(([^)]+)\)\s*',
            # 带方括号的编号：[1], [2], [a], [A] 等
            r'^\[([^\]]+)\]\s*',
            # 特殊格式：检测"标题1"、"标题1.1"等格式
            r'^标题(\d+(?:\.\d+)*)',
            # 检测"第X章"、"第X节"等格式
            r'^第([一二三四五六七八九十百千万\d]+)[章节条款项部分]\s*',
            # 检测"Chapter X"、"Section X"等英文格式
            r'^(?:Chapter|Section|Part)\s+(\d+(?:\.\d+)*)\s+',
            # 检测项目符号后的编号
            r'^[•·▪▫◦‣⁃]\s*(\d+(?:\.\d+)*)\s+',
        ]

        for pattern in patterns:
            match = re.match(pattern, text)
            if match:
                numbering = match.group(1)
                # 验证提取的编号是否合理
                if self._is_valid_numbering(numbering, text):
                    return numbering

        return None

    def _is_valid_numbering(self, numbering: str, full_text: str) -> bool:
        """
        验证提取的编号是否合理。

        Args:
            numbering: 提取的编号
            full_text: 完整的段落文本

        Returns:
            如果编号合理则返回True
        """
        # 排除一些明显不是编号的情况

        # 1. 如果文本太短，可能不是标题
        if len(full_text.strip()) < 3:
            return False

        # 2. 如果文本包含明显的正文特征，则不是标题
        content_indicators = [
            '这是', '内容', '正文', '描述', '说明', '详细',
            '具体', '包括', '例如', '比如', '如下', '以下',
            '本文', '文章', '段落', '句子', '文字', '文档'
        ]

        for indicator in content_indicators:
            if indicator in full_text:
                return False

        # 3. 如果编号后面的文本看起来像标题
        remaining_text = full_text[len(numbering):].strip()
        if remaining_text.startswith(('标题', '章', '节', '条', '款', '项', '第', 'Chapter', 'Section', 'Part')):
            return True

        # 4. 如果文本较短且包含标题关键词，可能是标题
        title_keywords = [
            '标题', '章节', '部分', '概述', '介绍', '总结', '结论',
            '背景', '目标', '方法', '结果', '讨论', '参考',
            '摘要', '关键词', '引言', '实验', '分析', '建议'
        ]

        if len(remaining_text) < 30:  # 短文本更可能是标题
            for keyword in title_keywords:
                if keyword in remaining_text:
                    return True

        # 5. 检查是否包含常见的标题模式
        title_patterns = [
            r'^[A-Z][a-z]+',  # 首字母大写的英文单词
            r'^[一二三四五六七八九十百千万]+',  # 中文数字
            r'^\w+$',  # 单个词
        ]

        for pattern in title_patterns:
            if re.match(pattern, remaining_text) and len(remaining_text) < 20:
                return True

        # 6. 默认情况：如果编号格式正确且文本不太长，认为是有效的
        return len(remaining_text) < 40  # 标题通常不会太长

    def get_numbered_paragraphs_only(self) -> List[str]:
        """
        Get only paragraphs that have numbering.
        
        ⚠️ DEPRECATED: This method is deprecated. Use get_full_document_content() with filtering instead.

        Returns:
            List of strings containing only numbered paragraphs
        """
        import warnings
        warnings.warn("get_numbered_paragraphs_only() is deprecated. Use get_full_document_content() with filtering instead.", 
                     DeprecationWarning, stacklevel=2)
        numbered_paragraphs = []

        for paragraph in self.document.paragraphs:
            # 首先检查自动编号
            numbering_text = self._get_paragraph_numbering(paragraph)
            if numbering_text is not None:
                numbered_paragraphs.append(numbering_text + paragraph.text)
            else:
                # 检查手动编号
                manual_numbering = self._extract_manual_numbering(paragraph.text)
                if manual_numbering:
                    numbered_paragraphs.append(manual_numbering + self.gap_text + paragraph.text)

        return numbered_paragraphs
    
    def get_numbering_structure(self) -> List[Dict[str, Union[str, int]]]:
        """
        Get the numbering structure with level information.

        Returns:
            List of dictionaries with numbering information
        """
        # Reset counters for fresh read
        self.counters.clear()
        self.number_cache.clear()

        result = []

        for paragraph in self.document.paragraphs:
            # 🔧 使用修复后的编号提取逻辑（与get_full_document_content一致）
            
            # 首先尝试从段落文本中提取显示的编号
            displayed_numbering = self._extract_displayed_numbering(paragraph.text.strip())
            
            if displayed_numbering:
                # 使用显示的编号
                numbering_text = displayed_numbering
                level = self._infer_numbering_level_from_display(displayed_numbering)
            else:
                # 回退到自动编号检查
                numbering_text = self._get_paragraph_numbering(paragraph)
                level = None

                if numbering_text is not None:
                    # 获取自动编号的级别
                    try:
                        num_pr = paragraph._element.pPr
                        if num_pr is not None:
                            num_pr_element = num_pr.find(qn("w:numPr"))
                            if num_pr_element is not None:
                                ilvl_element = num_pr_element.find(qn("w:ilvl"))
                                if ilvl_element is not None:
                                    level = int(ilvl_element.get(qn("w:val")))
                    except:
                        level = 0
                else:
                    # 检查手动编号
                    manual_numbering = self._extract_manual_numbering(paragraph.text)
                    if manual_numbering:
                        numbering_text = manual_numbering
                        # 根据编号格式推断级别
                        level = self._infer_numbering_level(manual_numbering)

            if numbering_text is not None:
                result.append({
                    "level": level if level is not None else 0,
                    "numbering": numbering_text.rstrip(self.gap_text + " "),
                    "text": paragraph.text,
                    "full_text": numbering_text + self.gap_text + paragraph.text
                })

        return result
    
    def get_word_like_content(self, indent_style: str = "space", numbering_separator: str = " ") -> List[str]:
        """
        获取像Word显示效果一样的完整内容（编号+文本组合）
        
        Args:
            indent_style: 缩进样式 ("space", "tab", "none")
            numbering_separator: 编号和文本之间的分隔符
            
        Returns:
            格式化后的完整内容列表
        """
        structure = self.get_numbering_structure()
        word_like_content = []
        
        for item in structure:
            level = item.get('level', 0)
            numbering = item.get('numbering', '')
            text = item.get('text', '')
            
            # 处理缩进
            if indent_style == "space":
                indent = "    " * level  # 4个空格per级别
            elif indent_style == "tab":
                indent = "\t" * level
            else:
                indent = ""
            
            # 格式化编号（确保编号后面有适当的分隔符）
            if numbering:
                # 如果编号本身没有结尾的点或括号，根据内容添加
                formatted_numbering = numbering
                if not any(numbering.endswith(suffix) for suffix in ['.', ')', '】', '、']):
                    # 如果是纯数字或字母，添加点号
                    if numbering.replace('.', '').replace('-', '').isalnum():
                        if not numbering.endswith('.'):
                            formatted_numbering = numbering + "."
                
                # 组合内容
                combined_text = f"{indent}{formatted_numbering}{numbering_separator}{text}"
            else:
                # 没有编号的情况
                combined_text = f"{indent}{text}"
            
            word_like_content.append(combined_text)
        
        return word_like_content
    
    def get_formatted_document(self, format_style: str = "standard") -> str:
        """
        获取格式化的完整文档内容
        
        Args:
            format_style: 格式样式 ("standard", "compact", "detailed")
            
        Returns:
            格式化后的完整文档字符串
        """
        if format_style == "standard":
            content_lines = self.get_word_like_content(indent_style="space", numbering_separator=" ")
        elif format_style == "compact":
            content_lines = self.get_word_like_content(indent_style="none", numbering_separator=" ")
        else:  # detailed
            content_lines = self.get_word_like_content(indent_style="space", numbering_separator="  ")
        
        return '\n'.join(content_lines)
    
    def get_enhanced_word_like_content(self) -> List[Dict[str, Union[str, int]]]:
        """
        获取增强的Word样式内容，包含更多元数据
        
        Returns:
            包含格式化内容和元数据的字典列表
        """
        structure = self.get_numbering_structure()
        enhanced_content = []
        
        for item in structure:
            level = item.get('level', 0)
            numbering = item.get('numbering', '')
            text = item.get('text', '')
            
            # 生成缩进
            indent = "    " * level
            
            # 格式化编号
            formatted_numbering = numbering
            if numbering and not any(numbering.endswith(suffix) for suffix in ['.', ')', '】', '、']):
                if numbering.replace('.', '').replace('-', '').isalnum():
                    if not numbering.endswith('.'):
                        formatted_numbering = numbering + "."
            
            # 组合内容 - 避免编号重复
            if formatted_numbering:
                # 🔧 检查文本是否已经包含编号，避免重复
                text_has_numbering = self._extract_displayed_numbering(text.strip()) is not None
                if text_has_numbering:
                    # 文本已包含编号，直接使用（只添加缩进）
                    word_like_text = f"{indent}{text}"
                else:
                    # 文本无编号，添加格式化编号
                    word_like_text = f"{indent}{formatted_numbering} {text}"
            else:
                word_like_text = f"{indent}{text}"
            
            enhanced_content.append({
                "level": level,
                "numbering": numbering,
                "formatted_numbering": formatted_numbering,
                "text": text,
                "word_like_text": word_like_text,
                "indent": indent,
                "raw_full_text": item.get('full_text', ''),
                "character_count": len(text),
                "has_numbering": bool(numbering)
            })
        
        return enhanced_content
    
    def get_full_document_content(self) -> List[Dict[str, Union[str, int]]]:
        """
        获取文档的所有内容，包括有编号和无编号的段落
        
        Returns:
            包含所有段落信息的字典列表
        """
        # 🔧 重置缓存以确保使用最新的编号逻辑
        self.counters.clear()
        self.number_cache.clear()
        
        all_content = []
        
        for paragraph in self.document.paragraphs:
            # 跳过空段落
            if not paragraph.text.strip():
                continue
                
            # 优先检查显示的编号
            displayed_numbering = self._extract_displayed_numbering(paragraph.text)
            level = None
            formatted_numbering = ""
            has_numbering = False
            original_numbering = ""
            
            if displayed_numbering:
                # 有显示的编号（最准确的方式）
                level = self._infer_numbering_level_from_display(displayed_numbering)
                formatted_numbering = displayed_numbering
                original_numbering = displayed_numbering
                has_numbering = True
            else:
                # 检查Word的自动编号 - 但使用修复的逻辑
                numbering_text = self._get_paragraph_numbering(paragraph)
                if numbering_text is not None:
                    # 有自动编号 - 需要从编号中提取显示编号而不是使用原始结果
                    level = self._get_paragraph_level(paragraph)
                    
                    # 🔧 尝试从自动编号文本中提取显示的编号
                    display_num = self._extract_displayed_numbering(numbering_text)
                    if display_num:
                        # 如果能提取到显示编号，使用它
                        formatted_numbering = display_num
                        original_numbering = display_num
                    else:
                        # 否则直接使用编号文本（但这可能仍然是错误的）
                        formatted_numbering = numbering_text.rstrip(self.gap_text + " ")
                        original_numbering = formatted_numbering
                    # 自动添加点号（如果需要）
                    if not any(formatted_numbering.endswith(suffix) for suffix in ['.', ')', '】', '、']):
                        if formatted_numbering.replace('.', '').replace('-', '').isalnum():
                            if not formatted_numbering.endswith('.'):
                                formatted_numbering = formatted_numbering + "."
                    has_numbering = True
                else:
                    # 检查手动编号
                    manual_numbering = self._extract_manual_numbering(paragraph.text)
                    if manual_numbering:
                        # 有手动编号
                        level = self._infer_numbering_level(manual_numbering)
                        formatted_numbering = manual_numbering
                        original_numbering = manual_numbering
                        # 处理手动编号格式
                        if not any(formatted_numbering.endswith(suffix) for suffix in ['.', ')', '】', '、']):
                            if formatted_numbering.replace('.', '').replace('-', '').isalnum():
                                if not formatted_numbering.endswith('.'):
                                    formatted_numbering = formatted_numbering + "."
                        has_numbering = True
                    else:
                        # 普通段落，无编号
                        level = 0
                        formatted_numbering = ""
                        original_numbering = ""
                        has_numbering = False
            
            # 生成缩进
            if level is None:
                level = 0
            indent = "    " * level
            
            # 组合Word样式文本 - 避免编号重复
            if has_numbering:
                # 🔧 检查段落文本是否已经包含编号，避免重复
                text_has_numbering = self._extract_displayed_numbering(paragraph.text.strip()) is not None
                if text_has_numbering:
                    # 段落文本已包含编号，直接使用（只添加缩进）
                    word_like_text = f"{indent}{paragraph.text}"
                else:
                    # 段落文本无编号，添加自动编号
                    word_like_text = f"{indent}{formatted_numbering} {paragraph.text}"
            else:
                word_like_text = f"{paragraph.text}"  # 普通段落不缩进
            
            all_content.append({
                "level": level,
                "numbering": original_numbering,
                "formatted_numbering": formatted_numbering,
                "text": paragraph.text,
                "word_like_text": word_like_text,
                "indent": indent,
                "character_count": len(paragraph.text),
                "has_numbering": has_numbering,
                "is_title": has_numbering,  # 暂时认为有编号的是标题
                "paragraph_style": paragraph.style.name if paragraph.style else "Normal"
            })
        
        return all_content
    
    def get_complete_formatted_document(self, format_style: str = "standard") -> str:
        """
        获取完整的格式化文档内容（包含所有段落）
        
        Args:
            format_style: 格式样式 ("standard", "compact", "detailed")
            
        Returns:
            格式化后的完整文档字符串
        """
        all_content = self.get_full_document_content()
        
        content_lines = []
        for item in all_content:
            if format_style == "compact":
                # 紧凑模式：无缩进
                if item['has_numbering']:
                    content_lines.append(f"{item['formatted_numbering']} {item['text']}")
                else:
                    content_lines.append(item['text'])
            elif format_style == "detailed":
                # 详细模式：双空格分隔
                if item['has_numbering']:
                    content_lines.append(f"{item['indent']}{item['formatted_numbering']}  {item['text']}")
                else:
                    content_lines.append(item['text'])
            else:  # standard
                # 标准模式：单空格分隔
                content_lines.append(item['word_like_text'])
        
        return '\n'.join(content_lines)

    def _infer_numbering_level_from_display(self, numbering: str) -> int:
        """
        从显示的编号推断级别
        
        Args:
            numbering: 显示的编号文本，如 "1.2.3."
            
        Returns:
            编号级别（0为顶级）
        """
        if not numbering:
            return 0
        
        # 移除末尾的点号和空格
        clean_numbering = numbering.rstrip('.、 ')
        
        # 计算点号数量来确定级别
        if '.' in clean_numbering:
            # 数字格式：1.2.3 -> 级别2（因为有2个点号）
            level = clean_numbering.count('.')
        else:
            # 单级编号：1, a), (1) 等 -> 级别0
            level = 0
        
        return level
    
    def _infer_numbering_level(self, numbering: str) -> int:
        """
        根据编号格式推断编号级别。

        Args:
            numbering: 编号文本

        Returns:
            推断的级别（0开始）
        """
        # 计算点号的数量来推断级别
        if '.' in numbering:
            # 对于 1.1.1 这样的格式，级别 = 点号数量
            return numbering.count('.')
        else:
            # 对于单级编号，级别为0
            return 0

    def get_enhanced_numbering_info(self) -> List[Dict[str, Union[str, int, bool]]]:
        """
        获取增强的编号信息，包括样式信息和编号类型。
        
        ⚠️ DEPRECATED: This method is deprecated. Use get_full_document_content() instead.

        Returns:
            包含详细编号信息的字典列表
        """
        import warnings
        warnings.warn("get_enhanced_numbering_info() is deprecated. Use get_full_document_content() instead.", 
                     DeprecationWarning, stacklevel=2)
        result = []

        for paragraph in self.document.paragraphs:
            # 检查自动编号
            auto_numbering = self._get_paragraph_numbering(paragraph)
            manual_numbering = None
            numbering_type = "none"

            if auto_numbering is not None:
                numbering_type = "auto"
                numbering_text = auto_numbering.rstrip(self.gap_text + " ")
                level = self._get_paragraph_level(paragraph)
            else:
                # 检查手动编号
                manual_numbering = self._extract_manual_numbering(paragraph.text)
                if manual_numbering:
                    numbering_type = "manual"
                    numbering_text = manual_numbering
                    level = self._infer_numbering_level(manual_numbering)
                else:
                    continue

            # 获取段落样式信息
            style_name = paragraph.style.name if paragraph.style else "Normal"

            # 检查是否是标题样式
            is_heading = style_name.startswith(('Heading', 'Title', '标题'))

            result.append({
                "text": paragraph.text,
                "numbering": numbering_text,
                "level": level if level is not None else 0,
                "numbering_type": numbering_type,
                "style_name": style_name,
                "is_heading": is_heading,
                "full_text": (numbering_text + self.gap_text + paragraph.text) if numbering_text else paragraph.text
            })

        return result

    def _get_paragraph_level(self, paragraph) -> Optional[int]:
        """
        获取段落的编号级别。

        Args:
            paragraph: 段落对象

        Returns:
            编号级别，如果没有则返回None
        """
        try:
            num_pr = paragraph._element.pPr
            if num_pr is not None:
                num_pr_element = num_pr.find(qn("w:numPr"))
                if num_pr_element is not None:
                    ilvl_element = num_pr_element.find(qn("w:ilvl"))
                    if ilvl_element is not None:
                        return int(ilvl_element.get(qn("w:val")))
        except:
            pass
        return None
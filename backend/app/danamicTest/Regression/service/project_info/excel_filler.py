"""
回归测试项目基本信息Excel填充器

负责将回归测试项目基本信息数据填充到Excel工作表中。
"""

import logging
import openpyxl
from typing import Dict, List, Optional, Any
from pathlib import Path

# 配置日志
logger = logging.getLogger(__name__)

class RegressionProjectInfoExcelFiller:
    """回归测试项目基本信息Excel填充器"""
    
    def __init__(self):
        """初始化Excel填充器"""
        self.logger = logger
        self.cell_mapping = self._get_cell_mapping()
    
    def fill_project_info_sheet(
        self,
        excel_file_path: str,
        sheet_name: str,
        project_data: Dict[str, Any]
    ) -> bool:
        """
        填充回归测试项目基本信息工作表
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            project_data: 项目数据字典
            
        Returns:
            bool: 填充是否成功
        """
        try:
            self.logger.info(f"开始填充回归测试项目基本信息表: {sheet_name}")
            
            # 打开Excel工作簿
            wb = openpyxl.load_workbook(excel_file_path)
            
            # 获取工作表
            if sheet_name not in wb.sheetnames:
                self.logger.error(f"工作表 '{sheet_name}' 不存在")
                return False
            
            sheet = wb[sheet_name]
            
            # 填充数据
            self._fill_data_to_sheet(sheet, project_data)
            
            # 保存文件
            wb.save(excel_file_path)
            
            self.logger.info("回归测试项目基本信息表填充完成")
            return True
            
        except Exception as e:
            self.logger.error(f"填充回归测试项目基本信息表时发生错误: {str(e)}")
            return False
    
    def _fill_data_to_sheet(self, sheet, project_data: Dict[str, Any]) -> None:
        """
        将数据填充到工作表
        
        Args:
            sheet: openpyxl工作表对象
            project_data: 项目数据字典
        """
        user_info = project_data.get('user_software_info', {})
        xinghao_list = project_data.get('xinghao_xinxi_list', [])
        suoshulingyu = project_data.get('suoshulingyu', '')
        fixed_values = project_data.get('fixed_values', {})
        
        # 填充可选信息（优先使用用户输入）
        if user_info and any(user_info.values()):
            self._fill_user_software_info(sheet, user_info, suoshulingyu)
        elif xinghao_list:
            self._fill_xinghao_info(sheet, xinghao_list)
        else:
            self._fill_default_info(sheet, suoshulingyu)
        
        # 填充人员信息
        self._fill_personnel_info(sheet, project_data)
        
        # 填充固定值（回归测试专用）
        self._fill_fixed_values(sheet, fixed_values)
    
    def _fill_user_software_info(self, sheet, user_info: Dict[str, str], default_domain: str) -> None:
        """
        填充用户软件信息
        
        Args:
            sheet: openpyxl工作表对象
            user_info: 用户软件信息字典
            default_domain: 默认领域
        """
        # 飞机型号
        sheet.cell(row=4, column=3, value=user_info.get('aircraft_model', ''))
        
        # 软件项目代号
        sheet.cell(row=5, column=3, value=user_info.get('software_id', ''))
        
        # 所属领域
        domain_value = user_info.get('domain', default_domain)
        sheet.cell(row=6, column=3, value=domain_value)
        
        # 项目总师
        sheet.cell(row=7, column=3, value=user_info.get('project_director', ''))
        
        # 主任设计师
        sheet.cell(row=9, column=3, value=user_info.get('chief_designer', ''))
        
        # 研制部门（优先使用用户输入）
        development_dept = user_info.get('development_department', '')
        if development_dept:
            sheet.cell(row=8, column=3, value=development_dept)
    
    def _fill_xinghao_info(self, sheet, xinghao_list: List[str]) -> None:
        """
        填充型号信息（备用数据源）
        
        Args:
            sheet: openpyxl工作表对象
            xinghao_list: 型号信息列表
        """
        if len(xinghao_list) >= 6:
            sheet.cell(row=4, column=3, value=xinghao_list[0])  # 飞机型号
            sheet.cell(row=5, column=3, value=xinghao_list[1])  # 软件项目代号
            sheet.cell(row=6, column=3, value=xinghao_list[2])  # 所属领域
            sheet.cell(row=7, column=3, value=xinghao_list[3])  # 项目总师
            sheet.cell(row=9, column=3, value=xinghao_list[5])  # 主任设计师
    
    def _fill_default_info(self, sheet, default_domain: str) -> None:
        """
        填充默认信息
        
        Args:
            sheet: openpyxl工作表对象
            default_domain: 默认领域
        """
        sheet.cell(row=4, column=3, value='')  # 飞机型号
        sheet.cell(row=5, column=3, value='')  # 软件项目代号
        sheet.cell(row=6, column=3, value=default_domain)  # 所属领域
        sheet.cell(row=7, column=3, value='')  # 项目总师
        sheet.cell(row=9, column=3, value='')  # 主任设计师
    
    def _fill_personnel_info(self, sheet, project_data: Dict[str, Any]) -> None:
        """
        填充人员信息
        
        Args:
            sheet: openpyxl工作表对象
            project_data: 项目数据字典
        """
        # 研制部门（如果用户信息中没有，使用开发人员信息中的）
        user_info = project_data.get('user_software_info', {})
        if not user_info.get('development_department'):
            yanzhibumen = project_data.get('yanzhibumen', '')
            if yanzhibumen:
                sheet.cell(row=8, column=3, value=yanzhibumen)
        
        # 开发人员
        kaifaren = project_data.get('kaifaren', '')
        sheet.cell(row=10, column=3, value=kaifaren)
        
        # 测试主管
        zhuguan_name = project_data.get('zhuguan_name', '')
        sheet.cell(row=11, column=3, value=zhuguan_name)
    
    def _fill_fixed_values(self, sheet, fixed_values: Dict[str, str]) -> None:
        """
        填充固定值（回归测试专用）
        
        Args:
            sheet: openpyxl工作表对象
            fixed_values: 固定值字典
        """
        # 测试级别（回归测试）
        sheet.cell(row=12, column=3, value=fixed_values.get('test_level', '回归测试'))
        
        # 测试类别
        sheet.cell(row=13, column=3, value=fixed_values.get('test_category', '半实物环境'))
        
        # 测试环境
        sheet.cell(row=14, column=3, value=fixed_values.get('test_environment', '半实物环境'))
    
    def _get_cell_mapping(self) -> Dict[str, tuple]:
        """
        获取单元格映射关系
        
        Returns:
            Dict[str, tuple]: 字段名到单元格位置的映射
        """
        return {
            'aircraft_model': (4, 3),        # 飞机型号
            'software_id': (5, 3),           # 软件项目代号
            'domain': (6, 3),                # 所属领域
            'project_director': (7, 3),      # 项目总师
            'development_department': (8, 3), # 研制部门
            'chief_designer': (9, 3),        # 主任设计师
            'developer': (10, 3),            # 开发人员
            'test_manager': (11, 3),         # 测试主管
            'test_level': (12, 3),           # 测试级别
            'test_category': (13, 3),        # 测试类别
            'test_environment': (14, 3)      # 测试环境
        }
    
    def get_filled_data_summary(self, excel_file_path: str, sheet_name: str) -> Dict[str, str]:
        """
        获取已填充数据的摘要（用于验证和调试）
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            Dict[str, str]: 已填充的数据摘要
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            summary = {}
            for field, (row, col) in self.cell_mapping.items():
                cell_value = sheet.cell(row=row, column=col).value
                summary[field] = str(cell_value) if cell_value else ""
            
            return summary
            
        except Exception as e:
            self.logger.error(f"获取填充数据摘要时发生错误: {str(e)}")
            return {}
    
    def fill_regression_specific_data(self, sheet, regression_data: Dict[str, Any]) -> None:
        """
        填充回归测试特有的数据
        
        Args:
            sheet: openpyxl工作表对象
            regression_data: 回归测试特有数据
        """
        # 可以在这里添加回归测试特有的数据填充逻辑
        # 例如：版本信息、变更记录等
        pass

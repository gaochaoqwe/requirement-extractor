"""
回归测试项目策划表格式设置器

负责设置回归测试项目策划表的格式、颜色和样式。
"""

import logging
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment
from typing import Dict, List, Optional, Any
from pathlib import Path

# 配置日志
logger = logging.getLogger(__name__)

class RegressionProjectPlanFormatter:
    """回归测试项目策划表格式设置器"""
    
    def __init__(self):
        """初始化格式设置器"""
        self.logger = logger
        self.color_fills = self._get_color_fills()
        self.border_style = self._get_border_style()
        self.alignment_style = self._get_alignment_style()
    
    def apply_formatting(
        self,
        excel_file_path: str,
        sheet_name: str,
        test_count: int,
        project_plan_data: Optional[Dict[str, Any]] = None
    ) -> bool:
        """
        应用回归测试项目策划表格式设置
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            test_count: 测试项数量
            project_plan_data: 项目策划表数据（可选，用于动态格式设置）
            
        Returns:
            bool: 格式设置是否成功
        """
        try:
            self.logger.info(f"开始应用回归测试项目策划表格式: {sheet_name}")
            
            # 打开Excel工作簿
            wb = openpyxl.load_workbook(excel_file_path)
            
            # 获取工作表
            if sheet_name not in wb.sheetnames:
                self.logger.error(f"工作表 '{sheet_name}' 不存在")
                return False
            
            sheet = wb[sheet_name]
            
            # 应用颜色填充
            self._apply_color_fills(sheet, test_count)
            
            # 应用边框样式
            self._apply_borders(sheet, test_count)
            
            # 应用对齐方式
            self._apply_alignment(sheet, test_count)
            
            # 应用特殊格式（如章节号缺失标记）
            if project_plan_data:
                self._apply_special_formatting(sheet, test_count, project_plan_data)
            
            # 保存文件
            wb.save(excel_file_path)
            
            self.logger.info("回归测试项目策划表格式设置完成")
            return True
            
        except Exception as e:
            self.logger.error(f"应用回归测试项目策划表格式时发生错误: {str(e)}")
            return False
    
    def _apply_color_fills(self, sheet, test_count: int) -> None:
        """应用颜色填充"""
        start_row = 18  # 测试项数据从第18行开始
        
        for i in range(test_count):
            current_row = start_row + i
            
            # 回归测试项目策划表使用浅黄色填充
            self._set_row_colors(sheet, current_row)
            
        self.logger.debug(f"颜色填充完成，共处理{test_count}行")
    
    def _set_row_colors(self, sheet, row: int) -> None:
        """设置单行颜色"""
        light_yellow_fill = self.color_fills['light_yellow']
        
        # 回归测试项目策划表的测试项行使用浅黄色
        for col in range(1, 15):  # 1-14列
            sheet.cell(row=row, column=col).fill = light_yellow_fill
    
    def _apply_borders(self, sheet, test_count: int) -> None:
        """应用边框样式"""
        # 应用边框到测试项数据行
        start_row = 18
        max_row = start_row + test_count - 1
        max_col = 14  # 回归测试项目策划表有14列
        
        for row in sheet.iter_rows(min_row=start_row, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                cell.border = self.border_style
        
        self.logger.debug(f"边框样式应用完成，范围: {start_row}-{max_row}行, 1-{max_col}列")
    
    def _apply_alignment(self, sheet, test_count: int) -> None:
        """应用对齐方式"""
        # 应用对齐方式到测试项数据行
        start_row = 18
        max_row = start_row + test_count - 1
        max_col = 14  # 回归测试项目策划表有14列
        
        for row in sheet.iter_rows(min_row=start_row, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                cell.alignment = self.alignment_style
        
        self.logger.debug(f"对齐方式应用完成，范围: {start_row}-{max_row}行, 1-{max_col}列")
    
    def _apply_special_formatting(self, sheet, test_count: int, project_plan_data: Dict[str, Any]) -> None:
        """应用特殊格式设置"""
        try:
            test_items = project_plan_data.get('test_items', [])
            red_fill = self.color_fills['red']
            
            start_row = 18
            for i, test_item in enumerate(test_items):
                current_row = start_row + i
                chapter_number = test_item.get('chapter_number', '')
                
                # 如果章节号未找到，使用红色标记
                if chapter_number == '未找到具体章节号':
                    for col in range(1, 15):
                        cell = sheet.cell(row=current_row, column=col)
                        cell.fill = red_fill
                        cell.border = self.border_style
                        cell.alignment = self.alignment_style
            
            self.logger.debug("特殊格式设置完成")
            
        except Exception as e:
            self.logger.error(f"应用特殊格式设置时发生错误: {str(e)}")
    
    def _get_color_fills(self) -> Dict[str, PatternFill]:
        """获取颜色填充样式"""
        return {
            'light_yellow': PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid"),
            'light_green': PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid"),
            'light_blue': PatternFill(start_color="99CCFF", end_color="99CCFF", fill_type="solid"),
            'light_orange': PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid"),
            'light_pink': PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid"),
            'red': PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        }
    
    def _get_border_style(self) -> Border:
        """获取边框样式"""
        return Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000')
        )
    
    def _get_alignment_style(self) -> Alignment:
        """获取对齐样式"""
        return Alignment(wrap_text=True, vertical='center')
    
    def highlight_regression_specific_fields(
        self,
        excel_file_path: str,
        sheet_name: str,
        test_count: int
    ) -> bool:
        """
        高亮显示回归测试特有的字段
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            test_count: 测试项数量
            
        Returns:
            bool: 高亮设置是否成功
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            # 回归测试特有字段使用橙色高亮
            orange_fill = self.color_fills['light_orange']
            
            start_row = 18
            for i in range(test_count):
                current_row = start_row + i
                
                # 测试标识（第14列）- 回归测试特有的标识生成逻辑
                sheet.cell(row=current_row, column=14).fill = orange_fill
                
                # 章节号（第6列）- 回归测试的文档读取逻辑
                sheet.cell(row=current_row, column=6).fill = orange_fill
            
            wb.save(excel_file_path)
            self.logger.info("回归测试特有字段高亮完成")
            return True
            
        except Exception as e:
            self.logger.error(f"高亮回归测试特有字段时发生错误: {str(e)}")
            return False
    
    def highlight_auto_generated_fields(
        self,
        excel_file_path: str,
        sheet_name: str,
        test_count: int
    ) -> bool:
        """
        高亮显示自动生成的字段
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            test_count: 测试项数量
            
        Returns:
            bool: 高亮设置是否成功
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            # 自动生成字段使用粉色高亮
            pink_fill = self.color_fills['light_pink']
            
            start_row = 18
            for i in range(test_count):
                current_row = start_row + i
                
                # 自动生成的字段
                auto_generated_columns = [
                    2,   # 序号
                    5,   # 标准子需求单位
                    6,   # 章节号（从文档读取）
                    7,   # 重要程度
                    8,   # 是否充分测试
                    9,   # 需求分析人员
                    10,  # 文档审查人员
                    11,  # 静态分析人员
                    12,  # 代码审查人员
                    13,  # 动态测试人员
                    14   # 测试标识
                ]
                
                for col in auto_generated_columns:
                    sheet.cell(row=current_row, column=col).fill = pink_fill
            
            wb.save(excel_file_path)
            self.logger.info("自动生成字段高亮完成")
            return True
            
        except Exception as e:
            self.logger.error(f"高亮自动生成字段时发生错误: {str(e)}")
            return False
    
    def highlight_missing_chapters(
        self,
        excel_file_path: str,
        sheet_name: str,
        missing_indices: List[int]
    ) -> bool:
        """
        高亮显示缺失章节号的行
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            missing_indices: 缺失章节号的行索引列表（从0开始）
            
        Returns:
            bool: 高亮设置是否成功
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            # 缺失章节号使用红色高亮
            red_fill = self.color_fills['red']
            
            start_row = 18
            for index in missing_indices:
                current_row = start_row + index
                
                # 高亮整行
                for col in range(1, 15):  # 回归测试项目策划表有14列
                    cell = sheet.cell(row=current_row, column=col)
                    cell.fill = red_fill
                    cell.border = self.border_style
                    cell.alignment = self.alignment_style
            
            wb.save(excel_file_path)
            self.logger.info(f"缺失章节号高亮完成，共{len(missing_indices)}行")
            return True
            
        except Exception as e:
            self.logger.error(f"高亮缺失章节号时发生错误: {str(e)}")
            return False
    
    def apply_conditional_formatting(
        self,
        excel_file_path: str,
        sheet_name: str,
        project_plan_data: Dict[str, Any]
    ) -> bool:
        """
        应用条件格式设置
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            project_plan_data: 项目策划表数据
            
        Returns:
            bool: 条件格式设置是否成功
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            test_items = project_plan_data.get('test_items', [])
            start_row = 18
            
            for i, test_item in enumerate(test_items):
                current_row = start_row + i
                importance = test_item.get('importance', '重要')
                
                # 根据重要程度设置特殊颜色
                if importance == '非常重要':
                    # 非常重要的测试项使用特殊标记
                    sheet.cell(row=current_row, column=7).fill = self.color_fills['light_orange']
                elif importance == '一般':
                    # 一般重要的测试项使用特殊标记
                    sheet.cell(row=current_row, column=7).fill = self.color_fills['light_blue']
            
            wb.save(excel_file_path)
            self.logger.info("条件格式设置完成")
            return True
            
        except Exception as e:
            self.logger.error(f"应用条件格式设置时发生错误: {str(e)}")
            return False
    
    def reset_formatting(
        self,
        excel_file_path: str,
        sheet_name: str,
        test_count: int
    ) -> bool:
        """
        重置格式设置
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            test_count: 测试项数量
            
        Returns:
            bool: 重置是否成功
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            # 清除所有格式
            start_row = 18
            max_row = start_row + test_count - 1
            max_col = 14
            
            for row in sheet.iter_rows(min_row=start_row, max_row=max_row, min_col=1, max_col=max_col):
                for cell in row:
                    cell.fill = PatternFill()  # 清除填充
                    cell.border = Border()     # 清除边框
                    cell.alignment = Alignment()  # 清除对齐
            
            wb.save(excel_file_path)
            self.logger.info("格式重置完成")
            return True
            
        except Exception as e:
            self.logger.error(f"重置格式时发生错误: {str(e)}")
            return False
    
    def apply_header_formatting(
        self,
        excel_file_path: str,
        sheet_name: str
    ) -> bool:
        """
        应用标题行格式
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            bool: 标题格式设置是否成功
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            # 标题行使用特殊格式
            header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
            bold_font = openpyxl.styles.Font(bold=True)
            
            # 应用到第17行（测试项标题行）
            for col in range(1, 15):  # 回归测试项目策划表有14列
                cell = sheet.cell(row=17, column=col)
                cell.fill = header_fill
                cell.font = bold_font
                cell.border = self.border_style
                cell.alignment = Alignment(horizontal='center', vertical='center')
            
            wb.save(excel_file_path)
            self.logger.info("标题行格式设置完成")
            return True
            
        except Exception as e:
            self.logger.error(f"应用标题行格式时发生错误: {str(e)}")
            return False

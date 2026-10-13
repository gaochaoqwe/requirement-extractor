"""
回归测试需求分析表服务

提供回归测试需求分析表的完整服务接口，整合数据处理、Excel填充、格式设置和读取功能。
"""

import logging
from typing import Dict, List, Optional, Any, Tuple
from pathlib import Path

from .core import (
    RegressionRequirementAnalysisDataProcessor,
    RegressionRequirementAnalysisExcelFiller,
    RegressionRequirementAnalysisFormatter,
    RegressionRequirementAnalysisReader
)
from .content_generation import RegressionTestContentGenerator

# 配置日志
logger = logging.getLogger(__name__)

class RegressionRequirementAnalysisService:
    """回归测试需求分析表服务"""
    
    def __init__(self):
        """初始化回归测试需求分析表服务"""
        self.logger = logger
        self.data_processor = RegressionRequirementAnalysisDataProcessor()
        self.excel_filler = RegressionRequirementAnalysisExcelFiller()
        self.formatter = RegressionRequirementAnalysisFormatter()
        self.reader = RegressionRequirementAnalysisReader()
        self.test_content_generator = RegressionTestContentGenerator()
    
    def generate_requirement_analysis_table(
        self,
        excel_file_path: str,
        sheet_name: str,
        ruanjian_xinxi: Dict[str, List[str]],
        ruanjian_name: str,
        diama_biangeng_xuqiu: Dict[str, str],
        file_path: str,
        ceshileixing: str,
        neirong_dict: List[str],
        biaoshihao_ceshixianng: List[str],
        AI_fenx_flag_Anniu: int,
        jianceyi_flag: int,
        apply_formatting: bool = True
    ) -> bool:
        """
        生成回归测试需求分析表
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            ruanjian_xinxi: 软件信息字典
            ruanjian_name: 软件名称
            diama_biangeng_xuqiu: 代码变更需求字典
            file_path: 文档文件路径
            ceshileixing: 测试类型
            neirong_dict: 内容描述列表
            biaoshihao_ceshixianng: 标识号测试项列表
            AI_fenx_flag_Anniu: AI分析标志
            jianceyi_flag: 检测仪标志
            apply_formatting: 是否应用格式设置
            
        Returns:
            bool: 生成是否成功
        """
        try:
            self.logger.info(f"开始生成回归测试需求分析表: {sheet_name}")
            
            # 第一步：组织数据
            requirement_analysis_data = self.data_processor.organize_requirement_analysis_data(
                ruanjian_xinxi=ruanjian_xinxi,
                ruanjian_name=ruanjian_name,
                diama_biangeng_xuqiu=diama_biangeng_xuqiu,
                file_path=file_path,
                ceshileixing=ceshileixing,
                neirong_dict=neirong_dict,
                biaoshihao_ceshixianng=biaoshihao_ceshixianng,
                AI_fenx_flag_Anniu=AI_fenx_flag_Anniu,
                jianceyi_flag=jianceyi_flag
            )
            
            # 验证数据
            if not self.data_processor.validate_requirement_analysis_data(requirement_analysis_data):
                self.logger.error("回归测试需求分析表数据验证失败")
                return False
            
            # 第二步：生成测试内容
            test_items = requirement_analysis_data.get('test_items', [])
            ai_config = requirement_analysis_data.get('ai_config', {})
            
            test_descriptions, test_methods = self.test_content_generator.generate_test_content(
                test_items, ai_config
            )
            
            # 更新测试内容到数据中
            self.data_processor.update_test_content(
                requirement_analysis_data, test_descriptions, test_methods
            )
            
            # 第三步：生成测试标识
            requirement_analysis_data = self.data_processor.generate_test_identifiers(requirement_analysis_data)
            
            # 第四步：读取文档章节号
            document_info = requirement_analysis_data.get('document_info', {})
            file_path = document_info.get('document_path', '')
            ruanjian_name = document_info.get('software_name', '')
            
            if file_path:
                test_names = [item.get('name', '') for item in test_items]
                chapter_numbers = self._read_document_chapters(file_path, test_names, ruanjian_name)
                self.data_processor.update_chapter_numbers(requirement_analysis_data, chapter_numbers)
            
            # 第五步：填充Excel
            if not self.excel_filler.fill_requirement_analysis_sheet(
                excel_file_path=excel_file_path,
                sheet_name=sheet_name,
                requirement_analysis_data=requirement_analysis_data
            ):
                self.logger.error("Excel填充失败")
                return False
            
            # 第六步：应用格式（可选）
            if apply_formatting:
                test_count = requirement_analysis_data.get('test_count', 0)
                if not self.formatter.apply_formatting(
                    excel_file_path=excel_file_path,
                    sheet_name=sheet_name,
                    test_count=test_count,
                    requirement_analysis_data=requirement_analysis_data
                ):
                    self.logger.warning("格式设置失败，但数据填充成功")
            
            # 记录统计信息
            statistics = self.data_processor.get_requirement_analysis_statistics(requirement_analysis_data)
            self.logger.debug(f"需求分析表统计: {statistics}")  # 改为DEBUG级别
            
            self.logger.info("回归测试需求分析表生成完成")
            return True
            
        except Exception as e:
            self.logger.error(f"生成回归测试需求分析表时发生错误: {str(e)}")
            return False
    
    def generate_requirement_analysis_table_legacy(
        self,
        excel_file_path: str,
        sheet_name: str,
        ruanjian_xinxi: Dict[str, List[str]],
        ruanjian_name: str,
        diama_biangeng_xuqiu: Dict[str, str],
        file_path: str,
        ceshileixing: str,
        neirong_dict: List[str],
        biaoshihao_ceshixianng: List[str],
        AI_fenx_flag_Anniu: int,
        jianceyi_flag: int
    ) -> bool:
        """
        生成回归测试需求分析表（兼容原有接口）
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            ruanjian_xinxi: 软件信息字典
            ruanjian_name: 软件名称
            diama_biangeng_xuqiu: 代码变更需求字典
            file_path: 文档文件路径
            ceshileixing: 测试类型
            neirong_dict: 内容描述列表
            biaoshihao_ceshixianng: 标识号测试项列表
            AI_fenx_flag_Anniu: AI分析标志
            jianceyi_flag: 检测仪标志
            
        Returns:
            bool: 生成是否成功
        """
        try:
            self.logger.info(f"开始生成回归测试需求分析表（兼容模式）: {sheet_name}")
            
            # 使用兼容模式的填充方法
            success = self.excel_filler.fill_requirement_analysis_sheet_legacy(
                excel_file_path=excel_file_path,
                sheet_name=sheet_name,
                ruanjian_xinxi=ruanjian_xinxi,
                ruanjian_name=ruanjian_name,
                diama_biangeng_xuqiu=diama_biangeng_xuqiu,
                file_path=file_path,
                ceshileixing=ceshileixing,
                neirong_dict=neirong_dict,
                biaoshihao_ceshixianng=biaoshihao_ceshixianng,
                AI_fenx_flag_Anniu=AI_fenx_flag_Anniu,
                jianceyi_flag=jianceyi_flag
            )
            
            # 应用兼容模式格式
            if success:
                test_count = len(diama_biangeng_xuqiu)
                # 读取章节号用于格式设置
                test_names = list(diama_biangeng_xuqiu.keys())
                chapter_numbers = self._read_document_chapters(file_path, test_names, ruanjian_name)
                
                self.formatter.apply_legacy_formatting(
                    excel_file_path, sheet_name, test_count, chapter_numbers
                )
            
            self.logger.info(f"回归测试需求分析表生成完成（兼容模式）")
            return success
            
        except Exception as e:
            self.logger.error(f"生成回归测试需求分析表时发生错误（兼容模式）: {str(e)}")
            return False
    
    def _read_document_chapters(self, file_path: str, test_names: List[str], ruanjian_name: str) -> List[str]:
        """读取文档章节号"""
        try:
            from ...forms.formFilter.document_reader import read_word_document
            chapter_numbers, _ = read_word_document(file_path, test_names, ruanjian_name)
            return chapter_numbers
        except Exception as e:
            self.logger.error(f"读取文档章节号时发生错误: {str(e)}")
            return ['未找到具体章节号'] * len(test_names)
    
    def read_requirement_analysis_data(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表'
    ) -> Tuple[Dict[str, Any], str]:
        """
        读取回归测试需求分析表数据
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            Tuple[Dict[str, Any], str]: (需求分析表数据, 状态信息)
        """
        try:
            return self.reader.read_requirement_analysis_data(excel_file_path, sheet_name)
        except Exception as e:
            self.logger.error(f"读取回归测试需求分析表数据时发生错误: {str(e)}")
            return {}, f"读取失败: {e}"
    
    def update_test_item_field(
        self,
        excel_file_path: str,
        sheet_name: str,
        item_index: int,
        field_name: str,
        field_value: str,
        apply_formatting: bool = True
    ) -> bool:
        """
        更新回归测试需求分析表中的单个字段
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            item_index: 测试项索引（从1开始）
            field_name: 字段名称
            field_value: 字段值
            apply_formatting: 是否重新应用格式设置
            
        Returns:
            bool: 更新是否成功
        """
        try:
            self.logger.info(f"更新回归测试需求分析表字段: 测试项{item_index}, {field_name} = {field_value}")
            
            # 更新单个字段
            if not self.excel_filler.update_single_test_item(
                excel_file_path=excel_file_path,
                sheet_name=sheet_name,
                item_index=item_index,
                field_name=field_name,
                field_value=field_value
            ):
                return False
            
            # 重新应用格式（可选）
            if apply_formatting:
                # 需要知道测试项数量来应用格式
                summary = self.excel_filler.get_filled_data_summary(excel_file_path, sheet_name, 50)  # 假设最多50项
                test_count = summary.get('total_items', 0)
                self.formatter.apply_formatting(excel_file_path, sheet_name, test_count)
            
            self.logger.info(f"测试项{item_index}的字段{field_name}更新完成")
            return True
            
        except Exception as e:
            self.logger.error(f"更新回归测试需求分析表字段时发生错误: {str(e)}")
            return False
    
    def get_requirement_analysis_status(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表'
    ) -> Dict[str, Any]:
        """
        获取回归测试需求分析表的状态
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            Dict[str, Any]: 状态信息
        """
        try:
            return self.reader.get_requirement_analysis_summary(excel_file_path, sheet_name)
        except Exception as e:
            self.logger.error(f"获取回归测试需求分析表状态时发生错误: {str(e)}")
            return {}
    
    def validate_requirement_analysis_table(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表',
        required_fields: Optional[List[str]] = None
    ) -> Dict[str, Any]:
        """
        验证回归测试需求分析表的完整性
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            required_fields: 必需字段列表
            
        Returns:
            Dict[str, Any]: 验证结果
        """
        try:
            # 检查结构
            if not self.reader.validate_requirement_analysis_structure(excel_file_path, sheet_name):
                return {
                    'is_valid': False,
                    'error': '回归测试需求分析表结构无效',
                    'table_type': '回归测试需求分析表'
                }
            
            # 检查测试内容完整性
            content_check = self.reader.check_test_content_completeness(excel_file_path, sheet_name)
            
            validation_result = {
                'is_valid': content_check.get('is_complete', False),
                'missing_descriptions': content_check.get('missing_descriptions_indices', []),
                'missing_methods': content_check.get('missing_methods_indices', []),
                'missing_identifiers': content_check.get('missing_identifiers_indices', []),
                'total_test_items': content_check.get('total_test_items', 0),
                'completion_rate': 0,
                'validation_message': self._get_validation_message(content_check),
                'table_type': '回归测试需求分析表'
            }
            
            # 计算完成率
            if validation_result['total_test_items'] > 0:
                missing_count = max(
                    len(validation_result['missing_descriptions']),
                    len(validation_result['missing_methods']),
                    len(validation_result['missing_identifiers'])
                )
                completed = validation_result['total_test_items'] - missing_count
                validation_result['completion_rate'] = round(completed / validation_result['total_test_items'] * 100, 2)
            
            self.logger.info(f"回归测试需求分析表验证结果: {validation_result}")
            return validation_result
            
        except Exception as e:
            self.logger.error(f"验证回归测试需求分析表时发生错误: {str(e)}")
            return {'is_valid': False, 'error': str(e)}
    
    def _get_validation_message(self, content_check: Dict[str, Any]) -> str:
        """
        生成验证消息
        
        Args:
            content_check: 内容检查结果
            
        Returns:
            str: 验证消息
        """
        if content_check.get('is_complete', False):
            return "回归测试需求分析表验证通过，所有测试内容都已生成"
        
        messages = []
        
        missing_descriptions = content_check.get('missing_descriptions_count', 0)
        if missing_descriptions > 0:
            messages.append(f"有 {missing_descriptions} 个测试项缺失测试描述")
        
        missing_methods = content_check.get('missing_methods_count', 0)
        if missing_methods > 0:
            messages.append(f"有 {missing_methods} 个测试项缺失测试方法")
        
        missing_identifiers = content_check.get('missing_identifiers_count', 0)
        if missing_identifiers > 0:
            messages.append(f"有 {missing_identifiers} 个测试项缺失测试标识")
        
        return "; ".join(messages) if messages else "需求分析表状态未知"
    
    def generate_test_content_for_existing_table(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表',
        use_ai_analysis: bool = True,
        jianceyi_flag: int = 0
    ) -> bool:
        """
        为现有的回归测试需求分析表生成测试内容
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            use_ai_analysis: 是否使用AI分析
            jianceyi_flag: 检测仪标志
            
        Returns:
            bool: 生成是否成功
        """
        try:
            self.logger.info("为现有回归测试需求分析表生成测试内容")
            
            # 读取现有数据
            requirement_analysis_data, status_msg = self.reader.read_requirement_analysis_data(excel_file_path, sheet_name)
            test_items = requirement_analysis_data.get('test_items', [])
            
            if not test_items:
                self.logger.error("未找到测试项数据")
                return False
            
            # 构造AI配置
            ai_config = {
                'enable_ai_analysis': use_ai_analysis,
                'detector_flag': jianceyi_flag,
                'ai_analysis_attempts': 4,
                'use_knowledge_optimization': True
            }
            
            # 生成测试内容
            test_descriptions, test_methods = self.test_content_generator.generate_test_content(
                test_items, ai_config
            )
            
            # 更新Excel中的测试内容
            import openpyxl
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            start_row = 2  # 从第2行开始
            for i, (description, method) in enumerate(zip(test_descriptions, test_methods)):
                sheet.cell(row=start_row + i, column=20).value = description  # 测试项描述
                sheet.cell(row=start_row + i, column=21).value = method      # 测试方法
            
            wb.save(excel_file_path)
            
            self.logger.info(f"成功为{len(test_descriptions)}个测试项生成测试内容")
            return True
            
        except Exception as e:
            self.logger.error(f"为现有回归测试需求分析表生成测试内容时发生错误: {str(e)}")
            return False
    
    def highlight_regression_specific_content(
        self,
        excel_file_path: str,
        sheet_name: str,
        test_count: int
    ) -> bool:
        """
        高亮显示回归测试特有的内容
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            test_count: 测试项数量
            
        Returns:
            bool: 高亮设置是否成功
        """
        try:
            return self.formatter.highlight_regression_specific_content(excel_file_path, sheet_name, test_count)
        except Exception as e:
            self.logger.error(f"高亮回归测试特有内容时发生错误: {str(e)}")
            return False
    
    def highlight_auto_generated_content(
        self,
        excel_file_path: str,
        sheet_name: str,
        test_count: int
    ) -> bool:
        """
        高亮显示自动生成的内容
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            test_count: 测试项数量
            
        Returns:
            bool: 高亮设置是否成功
        """
        try:
            return self.formatter.highlight_auto_generated_content(excel_file_path, sheet_name, test_count)
        except Exception as e:
            self.logger.error(f"高亮自动生成内容时发生错误: {str(e)}")
            return False
    
    def extract_test_requirements(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表'
    ) -> List[str]:
        """
        提取测试需求列表
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            List[str]: 测试需求列表
        """
        try:
            return self.reader.extract_test_requirements_only(excel_file_path, sheet_name)
        except Exception as e:
            self.logger.error(f"提取测试需求时发生错误: {str(e)}")
            return []
    
    def backup_requirement_analysis_table(
        self,
        excel_file_path: str,
        sheet_name: str,
        backup_file_path: str
    ) -> bool:
        """
        备份回归测试需求分析表
        
        Args:
            excel_file_path: 源Excel文件路径
            sheet_name: 工作表名称
            backup_file_path: 备份文件路径
            
        Returns:
            bool: 备份是否成功
        """
        try:
            return self.excel_filler.backup_sheet_data(excel_file_path, sheet_name, backup_file_path)
        except Exception as e:
            self.logger.error(f"备份回归测试需求分析表时发生错误: {str(e)}")
            return False

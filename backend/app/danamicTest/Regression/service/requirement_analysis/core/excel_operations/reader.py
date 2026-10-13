"""
回归测试需求分析表读取器

负责回归测试需求分析表的读取功能，包括：
- 需求分析表数据读取
- 基本信息提取
- 测试项数据提取
- 结构验证
- 内容完整性检查
- 测试类型统计
"""

import logging
import openpyxl
from typing import Dict, List, Optional, Any, Tuple
from pathlib import Path

# 配置日志
logger = logging.getLogger(__name__)

class RegressionRequirementAnalysisReader:
    """回归测试需求分析表读取器"""
    
    def __init__(self):
        """初始化回归测试需求分析表读取器"""
        self.logger = logger
        
        # 回归测试需求分析表列映射
        self.column_mapping = {
            'serial_number': 2,                           # 编号
            'software_name': 3,                           # 软件配置项名称
            'software_version': 4,                        # 被测软件版本
            'test_round': 5,                              # 轮次
            'standard_sub_requirement_name': 6,           # 标准子需求名称
            'standard_sub_requirement_identifier': 7,     # 标准子需求标识
            'standard_sub_requirement_unit': 8,           # 标准子需求单位
            'task_requirement_trace': 9,                  # 任务书需求追溯
            'software_requirement_spec_trace': 10,        # 软件需求规格说明追踪
            'software_requirement_name': 11,              # 软件需求规格说明需求名称
            'software_requirement_identifier': 12,        # 软件需求规格说明需求标识
            'software_requirement_description': 13,       # 软件需求描述
            'test_type': 14,                              # 测试类型
            'test_item_name': 15,                         # 测试项名称
            'test_item_identifier': 16,                   # 测试项标识
            'software_requirement_trace_relation': 17,    # 软件需求规格说明条目追踪关系
            'design_constraint': 18,                      # 设计约束
            'requirement_sub_item_description': 19,       # 需求子项描述
            'test_description': 20,                       # 测试项描述
            'test_method': 21,                            # 测试方法
            'adequacy_requirement': 22,                   # 充分性要求
            'pass_criteria': 23                           # 通过准则
        }
    
    def read_requirement_analysis_data(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表'
    ) -> Tuple[Dict[str, Any], str]:
        """
        读取回归测试需求分析表数据
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            Tuple[Dict[str, Any], str]: (需求分析表数据, 状态信息)
        """
        try:
            self.logger.info(f"开始读取回归测试需求分析表: {sheet_name}")
            
            # 检查文件是否存在
            if not Path(excel_file_path).exists():
                return {}, f"文件不存在: {excel_file_path}"
            
            # 打开Excel文件
            wb = openpyxl.load_workbook(excel_file_path)
            
            # 检查工作表是否存在
            if sheet_name not in wb.sheetnames:
                return {}, f"工作表不存在: {sheet_name}"
            
            sheet = wb[sheet_name]
            
            # 提取基本信息
            basic_info = self._extract_basic_info(sheet)
            
            # 提取测试项数据
            test_items = self._extract_test_items(sheet)
            
            # 构建需求分析表数据
            requirement_analysis_data = {
                'basic_info': basic_info,
                'test_items': test_items,
                'test_count': len(test_items),
                'table_type': '回归测试需求分析表',
                'column_mapping': self.column_mapping,
                'file_path': excel_file_path,
                'sheet_name': sheet_name
            }
            
            status_msg = f"成功读取回归测试需求分析表，共{len(test_items)}个测试项"
            self.logger.info(status_msg)
            
            return requirement_analysis_data, status_msg
            
        except Exception as e:
            error_msg = f"读取回归测试需求分析表时发生错误: {str(e)}"
            self.logger.error(error_msg)
            return {}, error_msg
    
    def _extract_basic_info(self, sheet) -> Dict[str, Any]:
        """
        提取基本信息
        
        Args:
            sheet: Excel工作表
            
        Returns:
            Dict[str, Any]: 基本信息
        """
        try:
            # 从第一行数据中提取基本信息
            first_row = 2  # 第一行数据
            
            software_name = sheet.cell(row=first_row, column=self.column_mapping['software_name']).value or ''
            software_version = sheet.cell(row=first_row, column=self.column_mapping['software_version']).value or ''
            test_round = sheet.cell(row=first_row, column=self.column_mapping['test_round']).value or '第一轮'
            unit = sheet.cell(row=first_row, column=self.column_mapping['standard_sub_requirement_unit']).value or '0.2'
            task_trace = sheet.cell(row=first_row, column=self.column_mapping['task_requirement_trace']).value or '无'
            design_constraint = sheet.cell(row=first_row, column=self.column_mapping['design_constraint']).value or '无'
            adequacy_requirement = sheet.cell(row=first_row, column=self.column_mapping['adequacy_requirement']).value or '所有用例执行完毕，对于未执行的用例说明未执行原因。'
            pass_criteria = sheet.cell(row=first_row, column=self.column_mapping['pass_criteria']).value or '各测试步骤、测试用例执行结果与预期一致，功能实现正确'
            
            return {
                'software_name': software_name,
                'software_version': software_version,
                'test_round': test_round,
                'unit': unit,
                'task_trace': task_trace,
                'design_constraint': design_constraint,
                'adequacy_requirement': adequacy_requirement,
                'pass_criteria': pass_criteria,
                'test_type': '回归测试'  # 默认为回归测试
            }
            
        except Exception as e:
            self.logger.error(f"提取基本信息时发生错误: {str(e)}")
            return {}
    
    def _extract_test_items(self, sheet) -> List[Dict[str, Any]]:
        """
        提取测试项数据
        
        Args:
            sheet: Excel工作表
            
        Returns:
            List[Dict[str, Any]]: 测试项列表
        """
        try:
            test_items = []
            row = 2  # 从第2行开始
            
            while True:
                # 检查编号列是否有值
                serial_number = sheet.cell(row=row, column=self.column_mapping['serial_number']).value
                if not serial_number:
                    break
                
                # 提取测试项数据
                test_item = {
                    'serial_number': serial_number,
                    'name': sheet.cell(row=row, column=self.column_mapping['test_item_name']).value or '',
                    'standard_sub_requirement_name': sheet.cell(row=row, column=self.column_mapping['standard_sub_requirement_name']).value or '',
                    'standard_sub_requirement_identifier': sheet.cell(row=row, column=self.column_mapping['standard_sub_requirement_identifier']).value or '',
                    'software_requirement_name': sheet.cell(row=row, column=self.column_mapping['software_requirement_name']).value or '',
                    'software_requirement_identifier': sheet.cell(row=row, column=self.column_mapping['software_requirement_identifier']).value or '',
                    'software_requirement_description': sheet.cell(row=row, column=self.column_mapping['software_requirement_description']).value or '',
                    'test_type': sheet.cell(row=row, column=self.column_mapping['test_type']).value or '',
                    'test_identifier': sheet.cell(row=row, column=self.column_mapping['test_item_identifier']).value or '',
                    'chapter_number': sheet.cell(row=row, column=self.column_mapping['software_requirement_spec_trace']).value or '未找到具体章节号',
                    'requirement_sub_item_description': sheet.cell(row=row, column=self.column_mapping['requirement_sub_item_description']).value or '',
                    'test_description': sheet.cell(row=row, column=self.column_mapping['test_description']).value or '',
                    'test_method': sheet.cell(row=row, column=self.column_mapping['test_method']).value or '',
                    'trace_relation': sheet.cell(row=row, column=self.column_mapping['software_requirement_trace_relation']).value or '未找到具体章节号'
                }
                
                test_items.append(test_item)
                row += 1
                
                # 防止无限循环
                if row > 1000:
                    break
            
            self.logger.info(f"成功提取{len(test_items)}个测试项")
            return test_items
            
        except Exception as e:
            self.logger.error(f"提取测试项数据时发生错误: {str(e)}")
            return []
    
    def validate_requirement_analysis_structure(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表'
    ) -> bool:
        """
        验证回归测试需求分析表结构
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            bool: 结构是否有效
        """
        try:
            # 检查文件是否存在
            if not Path(excel_file_path).exists():
                self.logger.error(f"文件不存在: {excel_file_path}")
                return False
            
            # 打开Excel文件
            wb = openpyxl.load_workbook(excel_file_path)
            
            # 检查工作表是否存在
            if sheet_name not in wb.sheetnames:
                self.logger.error(f"工作表不存在: {sheet_name}")
                return False
            
            sheet = wb[sheet_name]
            
            # 检查是否有数据
            if not sheet.cell(row=2, column=2).value:  # 检查第一个测试项的编号
                self.logger.warning("需求分析表中没有数据")
                return False
            
            # 检查关键列是否存在数据
            required_columns = [2, 3, 15, 16]  # 编号、软件名称、测试项名称、测试项标识
            for col in required_columns:
                if not sheet.cell(row=2, column=col).value:
                    self.logger.error(f"关键列{col}缺少数据")
                    return False
            
            self.logger.info("回归测试需求分析表结构验证通过")
            return True
            
        except Exception as e:
            self.logger.error(f"验证需求分析表结构时发生错误: {str(e)}")
            return False
    
    def check_test_content_completeness(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表'
    ) -> Dict[str, Any]:
        """
        检查测试内容完整性
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            Dict[str, Any]: 完整性检查结果
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            missing_descriptions_indices = []
            missing_methods_indices = []
            missing_identifiers_indices = []
            total_test_items = 0
            
            row = 2  # 从第2行开始
            while True:
                # 检查编号列是否有值
                serial_number = sheet.cell(row=row, column=self.column_mapping['serial_number']).value
                if not serial_number:
                    break
                
                total_test_items += 1
                
                # 检查测试项描述
                test_description = sheet.cell(row=row, column=self.column_mapping['test_description']).value
                if not test_description or test_description.strip() == '':
                    missing_descriptions_indices.append(total_test_items)
                
                # 检查测试方法
                test_method = sheet.cell(row=row, column=self.column_mapping['test_method']).value
                if not test_method or test_method.strip() == '':
                    missing_methods_indices.append(total_test_items)
                
                # 检查测试项标识
                test_identifier = sheet.cell(row=row, column=self.column_mapping['test_item_identifier']).value
                if not test_identifier or test_identifier.strip() == '':
                    missing_identifiers_indices.append(total_test_items)
                
                row += 1
                
                # 防止无限循环
                if row > 1000:
                    break
            
            is_complete = (len(missing_descriptions_indices) == 0 and 
                          len(missing_methods_indices) == 0 and 
                          len(missing_identifiers_indices) == 0)
            
            result = {
                'is_complete': is_complete,
                'total_test_items': total_test_items,
                'missing_descriptions_indices': missing_descriptions_indices,
                'missing_methods_indices': missing_methods_indices,
                'missing_identifiers_indices': missing_identifiers_indices,
                'missing_descriptions_count': len(missing_descriptions_indices),
                'missing_methods_count': len(missing_methods_indices),
                'missing_identifiers_count': len(missing_identifiers_indices)
            }
            
            self.logger.info(f"测试内容完整性检查完成: 总计{total_test_items}项，缺失描述{len(missing_descriptions_indices)}项，缺失方法{len(missing_methods_indices)}项")
            return result
            
        except Exception as e:
            self.logger.error(f"检查测试内容完整性时发生错误: {str(e)}")
            return {
                'is_complete': False,
                'total_test_items': 0,
                'missing_descriptions_indices': [],
                'missing_methods_indices': [],
                'missing_identifiers_indices': [],
                'error': str(e)
            }
    
    def get_requirement_analysis_summary(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表'
    ) -> Dict[str, Any]:
        """
        获取回归测试需求分析表摘要
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            Dict[str, Any]: 摘要信息
        """
        try:
            requirement_analysis_data, status_msg = self.read_requirement_analysis_data(excel_file_path, sheet_name)
            
            if not requirement_analysis_data:
                return {'error': status_msg}
            
            basic_info = requirement_analysis_data.get('basic_info', {})
            test_items = requirement_analysis_data.get('test_items', [])
            
            # 统计测试类型分布
            test_type_distribution = {}
            for test_item in test_items:
                test_type = test_item.get('test_type', '未知')
                test_type_distribution[test_type] = test_type_distribution.get(test_type, 0) + 1
            
            # 统计内容完整性
            complete_descriptions = sum(1 for item in test_items if item.get('test_description'))
            complete_methods = sum(1 for item in test_items if item.get('test_method'))
            complete_identifiers = sum(1 for item in test_items if item.get('test_identifier'))
            
            # 统计章节号情况
            missing_chapters = sum(1 for item in test_items if item.get('chapter_number') == '未找到具体章节号')
            
            summary = {
                'total_test_items': len(test_items),
                'software_name': basic_info.get('software_name', ''),
                'software_version': basic_info.get('software_version', ''),
                'test_type': basic_info.get('test_type', ''),
                'test_type_distribution': test_type_distribution,
                'completion_status': {
                    'descriptions': complete_descriptions,
                    'methods': complete_methods,
                    'identifiers': complete_identifiers,
                    'missing_chapters': missing_chapters
                },
                'completion_rate': {
                    'descriptions': round(complete_descriptions / len(test_items) * 100, 2) if test_items else 0,
                    'methods': round(complete_methods / len(test_items) * 100, 2) if test_items else 0,
                    'identifiers': round(complete_identifiers / len(test_items) * 100, 2) if test_items else 0
                },
                'table_type': '回归测试需求分析表',
                'file_path': excel_file_path,
                'sheet_name': sheet_name
            }
            
            self.logger.info(f"回归测试需求分析表摘要生成完成")
            return summary
            
        except Exception as e:
            self.logger.error(f"获取需求分析表摘要时发生错误: {str(e)}")
            return {'error': str(e)}
    
    def extract_test_requirements_only(
        self,
        excel_file_path: str,
        sheet_name: str = '需求分析表'
    ) -> List[str]:
        """
        仅提取测试需求列表
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            
        Returns:
            List[str]: 测试需求列表
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            test_requirements = []
            row = 2  # 从第2行开始
            
            while True:
                # 检查编号列是否有值
                serial_number = sheet.cell(row=row, column=self.column_mapping['serial_number']).value
                if not serial_number:
                    break
                
                # 提取测试项名称
                test_name = sheet.cell(row=row, column=self.column_mapping['test_item_name']).value
                if test_name:
                    test_requirements.append(test_name)
                
                row += 1
                
                # 防止无限循环
                if row > 1000:
                    break
            
            self.logger.info(f"成功提取{len(test_requirements)}个测试需求")
            return test_requirements
            
        except Exception as e:
            self.logger.error(f"提取测试需求时发生错误: {str(e)}")
            return []
    
    def get_test_items_by_type(
        self,
        excel_file_path: str,
        test_type: str,
        sheet_name: str = '需求分析表'
    ) -> List[Dict[str, Any]]:
        """
        根据测试类型获取测试项
        
        Args:
            excel_file_path: Excel文件路径
            test_type: 测试类型
            sheet_name: 工作表名称
            
        Returns:
            List[Dict[str, Any]]: 指定类型的测试项列表
        """
        try:
            requirement_analysis_data, _ = self.read_requirement_analysis_data(excel_file_path, sheet_name)
            test_items = requirement_analysis_data.get('test_items', [])
            
            filtered_items = [
                item for item in test_items 
                if item.get('test_type', '') == test_type
            ]
            
            self.logger.info(f"找到{len(filtered_items)}个{test_type}类型的测试项")
            return filtered_items
            
        except Exception as e:
            self.logger.error(f"根据测试类型获取测试项时发生错误: {str(e)}")
            return []

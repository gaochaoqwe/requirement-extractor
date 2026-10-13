"""
回归测试需求分析表Excel填充器

负责回归测试需求分析表的Excel填充功能，包括：
- 22列完整的回归测试需求分析表填充
- 兼容模式支持原有接口
- 测试类型分析和标识生成
- 文档章节号集成
- 数据验证和错误处理
"""

import logging
import openpyxl
from typing import Dict, List, Optional, Any, Tuple
from pathlib import Path

# 导入文本处理工具
from ...text_processing.text_utils import extract_requirement_name, extract_requirement_names_batch

# 配置日志
logger = logging.getLogger(__name__)

class RegressionRequirementAnalysisExcelFiller:
    """回归测试需求分析表Excel填充器"""
    
    def __init__(self):
        """初始化回归测试需求分析表Excel填充器"""
        self.logger = logger
        
        # 回归测试需求分析表列映射
        self.column_mapping = {
            'serial_number': 2,                           # 编号
            'software_name': 3,                           # 软件配置项名称
            'software_version': 4,                        # 被测软件版本
            'test_round': 5,                              # 轮次
            'standard_sub_requirement_name': 6,           # 标准子需求名称
            'standard_sub_requirement_identifier': 7,     # 标准子需求标识
            'standard_sub_requirement_unit': 8,           # 标准子需求单位
            'task_requirement_trace': 9,                  # 任务书需求追溯
            'software_requirement_spec_trace': 10,        # 软件需求规格说明追踪
            'software_requirement_name': 11,              # 软件需求规格说明需求名称
            'software_requirement_identifier': 12,        # 软件需求规格说明需求标识
            'software_requirement_description': 13,       # 软件需求描述
            'test_type': 14,                              # 测试类型
            'test_item_name': 15,                         # 测试项名称
            'test_item_identifier': 16,                   # 测试项标识
            'software_requirement_trace_relation': 17,    # 软件需求规格说明条目追踪关系
            'design_constraint': 18,                      # 设计约束
            'requirement_sub_item_description': 19,       # 需求子项描述
            'test_description': 20,                       # 测试项描述
            'test_method': 21,                            # 测试方法
            'adequacy_requirement': 22,                   # 充分性要求
            'pass_criteria': 23                           # 通过准则
        }
    
    def fill_requirement_analysis_sheet(
        self,
        excel_file_path: str,
        sheet_name: str,
        requirement_analysis_data: Dict[str, Any]
    ) -> bool:
        """
        填充回归测试需求分析表
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            requirement_analysis_data: 需求分析表数据
            
        Returns:
            bool: 填充是否成功
        """
        try:
            self.logger.info(f"开始填充回归测试需求分析表: {sheet_name}")
            
            # 打开Excel文件
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            # 清空现有数据
            self._clear_existing_data(sheet)
            
            # 提取数据
            basic_info = requirement_analysis_data.get('basic_info', {})
            test_items = requirement_analysis_data.get('test_items', [])
            
            # 填充测试项数据
            for index, test_item in enumerate(test_items):
                row_number = index + 2  # 从第2行开始
                self._fill_single_test_item(sheet, row_number, test_item, basic_info)
            
            # 保存文件
            wb.save(excel_file_path)
            
            self.logger.info(f"回归测试需求分析表填充完成，共填充{len(test_items)}个测试项")
            return True
            
        except Exception as e:
            self.logger.error(f"填充回归测试需求分析表时发生错误: {str(e)}")
            return False
    
    def _clear_existing_data(self, sheet):
        """清空现有数据"""
        try:
            # 清空从第2行开始的数据（保留表头）
            for row in range(2, 300):
                for cell in sheet[row]:
                    cell.value = None
            self.logger.debug("现有数据清空完成")
        except Exception as e:
            self.logger.warning(f"清空现有数据时发生错误: {str(e)}")
    
    def _fill_single_test_item(
        self,
        sheet,
        row_number: int,
        test_item: Dict[str, Any],
        basic_info: Dict[str, Any]
    ):
        """
        填充单个测试项
        
        Args:
            sheet: Excel工作表
            row_number: 行号
            test_item: 测试项数据
            basic_info: 基本信息
        """
        try:
            # 基本信息列
            sheet.cell(row=row_number, column=self.column_mapping['serial_number']).value = test_item.get('serial_number', '')
            sheet.cell(row=row_number, column=self.column_mapping['software_name']).value = basic_info.get('software_name', '')
            sheet.cell(row=row_number, column=self.column_mapping['software_version']).value = basic_info.get('software_version', '')
            sheet.cell(row=row_number, column=self.column_mapping['test_round']).value = basic_info.get('test_round', '第一轮')
            
            # 标准子需求信息
            sheet.cell(row=row_number, column=self.column_mapping['standard_sub_requirement_name']).value = test_item.get('name', '')
            sheet.cell(row=row_number, column=self.column_mapping['standard_sub_requirement_identifier']).value = test_item.get('standard_sub_requirement_identifier', '')
            sheet.cell(row=row_number, column=self.column_mapping['standard_sub_requirement_unit']).value = basic_info.get('unit', '0.2')
            
            # 追溯信息
            sheet.cell(row=row_number, column=self.column_mapping['task_requirement_trace']).value = basic_info.get('task_trace', '无')
            sheet.cell(row=row_number, column=self.column_mapping['software_requirement_spec_trace']).value = test_item.get('chapter_number', '未找到具体章节号')
            sheet.cell(row=row_number, column=self.column_mapping['software_requirement_name']).value = test_item.get('name', '')
            sheet.cell(row=row_number, column=self.column_mapping['software_requirement_identifier']).value = test_item.get('software_requirement_identifier', '')
            
            # 需求描述
            sheet.cell(row=row_number, column=self.column_mapping['software_requirement_description']).value = test_item.get('requirement_description', '')
            
            # 测试信息
            sheet.cell(row=row_number, column=self.column_mapping['test_type']).value = test_item.get('test_type', '')
            sheet.cell(row=row_number, column=self.column_mapping['test_item_name']).value = test_item.get('name', '')
            sheet.cell(row=row_number, column=self.column_mapping['test_item_identifier']).value = test_item.get('test_identifier', '')
            
            # 其他信息
            sheet.cell(row=row_number, column=self.column_mapping['software_requirement_trace_relation']).value = test_item.get('chapter_number', '未找到具体章节号')
            sheet.cell(row=row_number, column=self.column_mapping['design_constraint']).value = basic_info.get('design_constraint', '无')
            sheet.cell(row=row_number, column=self.column_mapping['requirement_sub_item_description']).value = test_item.get('requirement_description', '')
            
            # 测试内容
            sheet.cell(row=row_number, column=self.column_mapping['test_description']).value = test_item.get('test_description', '')
            sheet.cell(row=row_number, column=self.column_mapping['test_method']).value = test_item.get('test_method', '')
            
            # 固定内容
            sheet.cell(row=row_number, column=self.column_mapping['adequacy_requirement']).value = basic_info.get('adequacy_requirement', '所有用例执行完毕，对于未执行的用例说明未执行原因。')
            sheet.cell(row=row_number, column=self.column_mapping['pass_criteria']).value = basic_info.get('pass_criteria', '各测试步骤、测试用例执行结果与预期一致，功能实现正确')
            
        except Exception as e:
            self.logger.error(f"填充测试项{row_number}时发生错误: {str(e)}")
    
    def fill_requirement_analysis_sheet_legacy(
        self,
        excel_file_path: str,
        sheet_name: str,
        ruanjian_xinxi: Dict[str, List[str]],
        ruanjian_name: str,
        diama_biangeng_xuqiu: Dict[str, str],
        file_path: str,
        ceshileixing: str,
        neirong_dict: List[str],
        biaoshihao_ceshixianng: List[str],
        AI_fenx_flag_Anniu: int,
        jianceyi_flag: int
    ) -> bool:
        """
        填充回归测试需求分析表（兼容模式）
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            ruanjian_xinxi: 软件信息字典
            ruanjian_name: 软件名称
            diama_biangeng_xuqiu: 代码变更需求字典
            file_path: 文档文件路径
            ceshileixing: 测试类型
            neirong_dict: 内容描述列表
            biaoshihao_ceshixianng: 标识号测试项列表
            AI_fenx_flag_Anniu: AI分析标志
            jianceyi_flag: 检测仪标志
            
        Returns:
            bool: 填充是否成功
        """
        try:
            self.logger.info(f"🔧 开始填充回归测试需求分析表（兼容模式）: {sheet_name}")
            
            # 🚨 详细日志：输入参数监控
            self.logger.info(f"📋 输入参数监控:")
            self.logger.info(f"   - Excel文件: {excel_file_path}")
            self.logger.info(f"   - 工作表名: {sheet_name}")
            self.logger.info(f"   - 软件名称: {ruanjian_name}")
            self.logger.info(f"   - 测试类型: {ceshileixing}")
            self.logger.info(f"   - diama_biangeng_xuqiu数量: {len(diama_biangeng_xuqiu)}")
            self.logger.info(f"   - neirong_dict数量: {len(neirong_dict)}")
            self.logger.info(f"   - biaoshihao_ceshixianng数量: {len(biaoshihao_ceshixianng)}")
            
            # 🚨 关键数据内容监控
            self.logger.info(f"🔍 关键数据内容:")
            self.logger.info(f"   - diama_biangeng_xuqiu.keys(): {list(diama_biangeng_xuqiu.keys())}")
            self.logger.info(f"   - biaoshihao_ceshixianng: {biaoshihao_ceshixianng}")
            
            # 生成测试内容
            test_descriptions, test_methods = self._generate_test_content_legacy(
                diama_biangeng_xuqiu, neirong_dict, AI_fenx_flag_Anniu, jianceyi_flag
            )
            
            # 读取章节号
            chapter_numbers = self._read_document_chapters(file_path, list(diama_biangeng_xuqiu.keys()), ruanjian_name)
            
            # 填充Excel数据
            success = self._fill_excel_data_legacy(
                excel_file_path, sheet_name, ruanjian_xinxi, ruanjian_name,
                diama_biangeng_xuqiu, ceshileixing, neirong_dict, biaoshihao_ceshixianng,
                test_descriptions, test_methods, chapter_numbers
            )
            
            self.logger.info("回归测试需求分析表填充完成（兼容模式）")
            return success
            
        except Exception as e:
            self.logger.error(f"填充回归测试需求分析表时发生错误（兼容模式）: {str(e)}")
            return False
    
    def _generate_test_content_legacy(
        self,
        diama_biangeng_xuqiu: Dict[str, str],
        neirong_dict: List[str],
        AI_fenx_flag_Anniu: int,
        jianceyi_flag: int
    ) -> Tuple[List[str], List[str]]:
        """生成测试内容（兼容模式）"""
        try:
            from .test_content_generator import RegressionTestContentGenerator
            
            # 构建测试项列表
            test_items = []
            # 🚨 修复：提取实际的需求名称而不是REQ标识符
            # 原始错误：test_names = list(diama_biangeng_xuqiu.keys())  # ['REQ-001', 'REQ-002', ...]
            # 修复后：如果diama_biangeng_xuqiu的值包含name字段，则提取name；否则直接使用键名
            test_names = []
            for req_id, req_data in diama_biangeng_xuqiu.items():
                if isinstance(req_data, dict) and 'name' in req_data:
                    test_names.append(req_data['name'])
                else:
                    # 如果没有name字段，使用键名（保持向下兼容）
                    test_names.append(req_id)
            
            for index, test_name in enumerate(test_names):
                requirement_description = neirong_dict[index] if index < len(neirong_dict) else ''
                test_item = {
                    'name': test_name,
                    'requirement_description': requirement_description,
                    'test_type': ''  # 将在生成过程中确定
                }
                test_items.append(test_item)
            
            # 构建AI配置
            ai_config = {
                'enable_ai_analysis': AI_fenx_flag_Anniu == 1,
                'detector_flag': jianceyi_flag,
                'ai_analysis_attempts': 4,
                'use_knowledge_optimization': True
            }
            
            # 生成测试内容
            generator = RegressionTestContentGenerator()
            test_descriptions, test_methods = generator.generate_test_content(test_items, ai_config)
            
            return test_descriptions, test_methods
            
        except Exception as e:
            self.logger.error(f"生成测试内容时发生错误（兼容模式）: {str(e)}")
            return [], []
    
    def _read_document_chapters(self, file_path: str, test_names: List[str], ruanjian_name: str) -> List[str]:
        """读取文档章节号"""
        try:
            from ...forms.formFilter.document_reader import read_word_document
            chapter_numbers, _ = read_word_document(file_path, test_names, ruanjian_name)
            return chapter_numbers
        except Exception as e:
            self.logger.error(f"读取文档章节号时发生错误: {str(e)}")
            return ['未找到具体章节号'] * len(test_names)
    
    def _fill_excel_data_legacy(
        self,
        excel_file_path: str,
        sheet_name: str,
        ruanjian_xinxi: Dict[str, List[str]],
        ruanjian_name: str,
        diama_biangeng_xuqiu: Dict[str, str],
        ceshileixing: str,
        neirong_dict: List[str],
        biaoshihao_ceshixianng: List[str],
        test_descriptions: List[str],
        test_methods: List[str],
        chapter_numbers: List[str]
    ) -> bool:
        """填充Excel数据（兼容模式）"""
        try:
            # 打开Excel文件
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            # 清空现有数据
            self._clear_existing_data(sheet)
            
            # 提取基本信息
            software_identifier = ruanjian_xinxi.get('软件标识', [''])[0] if isinstance(ruanjian_xinxi.get('软件标识', ['']), list) else ruanjian_xinxi.get('软件标识', '')
            software_version = ruanjian_xinxi.get('被测软件版本', [''])[0] if isinstance(ruanjian_xinxi.get('被测软件版本', ['']), list) else ruanjian_xinxi.get('被测软件版本', '')
            
            # 生成测试需求标识前缀
            if ceshileixing == '回归测试':
                test_requirement_prefix = f'CT_R-{software_identifier}_{software_version}_'
            elif ceshileixing == '配置项测试':
                test_requirement_prefix = f'CT-{software_identifier}_{software_version}_'
            else:  # 系统测试
                test_requirement_prefix = f'ST-{software_identifier}_{software_version}_'
            
            # 填充每个测试项
            # 🚨 修复：提取实际的需求名称而不是REQ标识符 
            # 原始错误：test_names = list(diama_biangeng_xuqiu.keys())  # ['REQ-001', 'REQ-002', ...]
            # 修复后：如果diama_biangeng_xuqiu的值包含name字段，则提取name；否则直接使用键名
            test_names = []
            for req_id, req_data in diama_biangeng_xuqiu.items():
                if isinstance(req_data, dict) and 'name' in req_data:
                    test_names.append(req_data['name'])
                else:
                    # 如果没有name字段，使用键名（保持向下兼容）
                    test_names.append(req_id)
            
            self.logger.info(f"🔧 开始Excel数据填充（回归测试）")
            self.logger.info(f"📊 填充参数:")
            self.logger.info(f"   - test_names: {test_names}")
            self.logger.info(f"   - biaoshihao_ceshixianng: {biaoshihao_ceshixianng}")
            
            for index, test_name in enumerate(test_names):
                row_number = index + 2  # 从第2行开始
                self.logger.info(f"📝 正在填充第{index+1}行数据 (Excel行{row_number}):")
                
                # 获取测试类型和标识
                try:
                    from ...content_generation.test_type_classifier import ceshi_leixing_biaoshi_choose
                    type_identifier, test_type = ceshi_leixing_biaoshi_choose(test_name)
                except:
                    type_identifier, test_type = 'FN01', '功能测试'
                
                # 🚨 修复：正确处理需求名称和标识符
                # biaoshihao_ceshixianng 实际包含的是需求名称，不是英文标识符
                requirement_name = biaoshihao_ceshixianng[index] if index < len(biaoshihao_ceshixianng) else f'默认需求{index+1}'
                
                # 生成正确的英文标识符（基于test_names中的REG-XXX格式）
                if test_name.startswith('REG-'):
                    english_identifier = test_name  # 直接使用REG-001格式
                else:
                    english_identifier = f'REG-{index+1:03d}'  # 生成REG-001格式
                
                # 🚨 关键数据处理监控
                self.logger.info(f"   🔍 第{index}项数据处理:")
                self.logger.info(f"      - 原始test_name: '{test_name}'")
                self.logger.info(f"      - biaoshihao_ceshixianng[{index}]: '{biaoshihao_ceshixianng[index] if index < len(biaoshihao_ceshixianng) else 'N/A'}'")
                self.logger.info(f"      - requirement_name（需求名称）: '{requirement_name}'")
                self.logger.info(f"      - english_identifier（英文标识符）: '{english_identifier}'")
                
                # 获取需求描述
                requirement_description = neirong_dict[index] if index < len(neirong_dict) else ''
                
                # 获取测试内容
                test_description = test_descriptions[index] if index < len(test_descriptions) else ''
                test_method = test_methods[index] if index < len(test_methods) else ''
                
                # 获取章节号
                chapter_number = chapter_numbers[index] if index < len(chapter_numbers) else '未找到具体章节号'
                
                # 🚨 关键步骤：使用正确的需求名称（已经是纯名称，无需再提取）
                pure_requirement_name = requirement_name  # 直接使用传入的纯需求名称
                self.logger.info(f"      - 使用需求名称: '{pure_requirement_name}'")
                
                # 填充数据
                sheet.cell(row=row_number, column=2).value = index + 1  # 编号
                sheet.cell(row=row_number, column=3).value = ruanjian_name  # 软件配置项名称
                sheet.cell(row=row_number, column=4).value = software_version  # 被测软件版本
                sheet.cell(row=row_number, column=5).value = '第一轮'  # 轮次
                
                # 🚨 关键列监控：列6(F) 标准子需求名称
                sheet.cell(row=row_number, column=6).value = pure_requirement_name  # 标准子需求名称
                self.logger.info(f"      ✅ 列6(F) 标准子需求名称: '{pure_requirement_name}'")
                
                sheet.cell(row=row_number, column=7).value = f'CT_R-{english_identifier}'  # 标准子需求标识
                sheet.cell(row=row_number, column=8).value = '0.2'  # 标准子需求单位
                sheet.cell(row=row_number, column=9).value = '无'  # 任务书需求追溯
                sheet.cell(row=row_number, column=10).value = chapter_number  # 软件需求规格说明追踪
                
                # 🚨 关键列监控：列11(K) 软件需求规格说明需求名称
                sheet.cell(row=row_number, column=11).value = pure_requirement_name  # 软件需求规格说明需求名称
                self.logger.info(f"      ✅ 列11(K) 软件需求规格说明需求名称: '{pure_requirement_name}'")
                
                # 🚨 关键列监控：列12(L) 软件需求规格说明需求标识
                sheet.cell(row=row_number, column=12).value = english_identifier  # 软件需求规格说明需求标识
                self.logger.info(f"      ✅ 列12(L) 软件需求规格说明需求标识: '{english_identifier}'")
                
                sheet.cell(row=row_number, column=13).value = requirement_description  # 软件需求描述
                sheet.cell(row=row_number, column=14).value = test_type  # 测试类型
                
                # 🚨 关键列监控：列15(O) 测试项名称
                sheet.cell(row=row_number, column=15).value = pure_requirement_name  # 测试项名称
                self.logger.info(f"      ✅ 列15(O) 测试项名称: '{pure_requirement_name}'")
                
                # 🚨 关键列监控：列16(P) 测试项标识
                test_item_identifier = f'{test_requirement_prefix}{english_identifier}{type_identifier}'
                sheet.cell(row=row_number, column=16).value = test_item_identifier  # 测试项标识
                self.logger.info(f"      ✅ 列16(P) 测试项标识: '{test_item_identifier}'")
                sheet.cell(row=row_number, column=17).value = chapter_number  # 软件需求规格说明条目追踪关系
                sheet.cell(row=row_number, column=18).value = '无'  # 设计约束
                
                # 修复：需求子项描述应该与软件需求描述内容一致
                requirement_sub_desc = requirement_description
                sheet.cell(row=row_number, column=19).value = requirement_sub_desc  # 需求子项描述
                self.logger.info(f"      列19 需求子项描述: '{requirement_sub_desc[:50]}...' (长度: {len(requirement_sub_desc)})")
                
                # 修复：测试项描述和测试方法 - 确保使用AI生成的内容或合理的默认值
                final_test_description = test_description.strip()
                final_test_method = test_method.strip()
                
                self.logger.info(f"      原始AI生成内容 - test_description: '{final_test_description[:100]}...'")
                self.logger.info(f"      原始AI生成内容 - test_method: '{final_test_method[:100]}...'")
                
                # 如果是默认值模式，生成更合适的描述
                if not final_test_description or final_test_description.startswith('测试软件REQ-') or final_test_description == f'测试软件{test_name}功能实现是否正常':
                    # 使用更智能的默认描述
                    if pure_requirement_name:
                        final_test_description = f'验证{pure_requirement_name}的实现是否符合需求规格说明中的功能要求'
                    else:
                        final_test_description = f'验证软件功能实现是否符合需求规格说明中的要求'
                    self.logger.info(f"      使用默认测试项描述: '{final_test_description}'")
                
                if not final_test_method or final_test_method.startswith('通过检测仪发送测试指令'):
                    # 使用更智能的默认方法
                    if pure_requirement_name:
                        final_test_method = f'1. 设置测试环境和条件；2. 执行{pure_requirement_name}相关的功能操作；3. 检查功能执行结果；4. 验证结果是否符合预期'
                    else:
                        final_test_method = f'1. 设置测试环境和条件；2. 执行相关的功能操作；3. 检查功能执行结果；4. 验证结果是否符合预期'
                    self.logger.info(f"      使用默认测试方法: '{final_test_method[:50]}...'")
                
                sheet.cell(row=row_number, column=20).value = final_test_description  # 测试项描述
                sheet.cell(row=row_number, column=21).value = final_test_method  # 测试方法
                self.logger.info(f"      列20 最终测试项描述: '{final_test_description[:50]}...' (长度: {len(final_test_description)})")
                self.logger.info(f"      列21 最终测试方法: '{final_test_method[:50]}...' (长度: {len(final_test_method)})")
                
                sheet.cell(row=row_number, column=22).value = '所有用例执行完毕，对于未执行的用例说明未执行原因。'  # 充分性要求
                sheet.cell(row=row_number, column=23).value = '各测试步骤、测试用例执行结果与预期一致，功能实现正确'  # 通过准则
            
            # 保存文件
            wb.save(excel_file_path)
            
            self.logger.info(f"Excel数据填充完成（兼容模式），共填充{len(test_names)}个测试项")
            return True
            
        except Exception as e:
            self.logger.error(f"填充Excel数据时发生错误（兼容模式）: {str(e)}")
            return False
    
    def update_single_test_item(
        self,
        excel_file_path: str,
        sheet_name: str,
        item_index: int,
        field_name: str,
        field_value: str
    ) -> bool:
        """
        更新单个测试项的字段
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            item_index: 测试项索引（从1开始）
            field_name: 字段名称
            field_value: 字段值
            
        Returns:
            bool: 更新是否成功
        """
        try:
            if field_name not in self.column_mapping:
                self.logger.error(f"未知的字段名称: {field_name}")
                return False
            
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            row_number = item_index + 1  # 转换为Excel行号
            column_number = self.column_mapping[field_name]
            
            sheet.cell(row=row_number, column=column_number).value = field_value
            wb.save(excel_file_path)
            
            self.logger.info(f"测试项{item_index}的字段{field_name}更新成功")
            return True
            
        except Exception as e:
            self.logger.error(f"更新测试项字段时发生错误: {str(e)}")
            return False
    
    def get_filled_data_summary(
        self,
        excel_file_path: str,
        sheet_name: str,
        max_rows: int = 100
    ) -> Dict[str, Any]:
        """
        获取已填充数据的摘要
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            max_rows: 最大检查行数
            
        Returns:
            Dict[str, Any]: 数据摘要
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            filled_rows = 0
            for row in range(2, max_rows + 2):  # 从第2行开始检查
                if sheet.cell(row=row, column=2).value:  # 检查编号列
                    filled_rows += 1
                else:
                    break
            
            return {
                'total_items': filled_rows,
                'has_data': filled_rows > 0,
                'last_filled_row': filled_rows + 1 if filled_rows > 0 else 1
            }
            
        except Exception as e:
            self.logger.error(f"获取填充数据摘要时发生错误: {str(e)}")
            return {'total_items': 0, 'has_data': False, 'last_filled_row': 1}
    
    def backup_sheet_data(
        self,
        excel_file_path: str,
        sheet_name: str,
        backup_file_path: str
    ) -> bool:
        """
        备份工作表数据
        
        Args:
            excel_file_path: 源Excel文件路径
            sheet_name: 工作表名称
            backup_file_path: 备份文件路径
            
        Returns:
            bool: 备份是否成功
        """
        try:
            import shutil
            shutil.copy2(excel_file_path, backup_file_path)
            self.logger.info(f"工作表数据备份成功: {backup_file_path}")
            return True
        except Exception as e:
            self.logger.error(f"备份工作表数据时发生错误: {str(e)}")
            return False

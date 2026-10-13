"""
回归测试问题报告单格式设置器

负责设置回归测试问题报告单的格式、颜色和样式。
"""

import logging
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment
from typing import Dict, List, Optional, Any
from pathlib import Path

# 配置日志
logger = logging.getLogger(__name__)

class RegressionIssueReportFormatter:
    """回归测试问题报告单格式设置器"""
    
    def __init__(self):
        """初始化格式设置器"""
        self.logger = logger
        self.color_fills = self._get_color_fills()
        self.border_style = self._get_border_style()
    
    def apply_formatting(
        self,
        excel_file_path: str,
        sheet_name: str,
        issue_report_data: Optional[Dict[str, Any]] = None
    ) -> bool:
        """
        应用回归测试问题报告单格式设置
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            issue_report_data: 问题报告数据（可选，用于动态格式设置）
            
        Returns:
            bool: 格式设置是否成功
        """
        try:
            self.logger.debug(f"开始应用回归测试问题报告单格式: {sheet_name}")
            
            # 打开Excel工作簿
            wb = openpyxl.load_workbook(excel_file_path)
            
            # 获取工作表
            if sheet_name not in wb.sheetnames:
                self.logger.error(f"工作表 '{sheet_name}' 不存在")
                return False
            
            sheet = wb[sheet_name]
            
            # 获取问题数量
            issue_count = issue_report_data.get('issue_count', 16) if issue_report_data else 16
            
            # 应用颜色填充
            self._apply_color_fills(sheet, issue_count)
            
            # 应用边框样式
            self._apply_borders(sheet, issue_count)
            
            # 应用对齐方式
            self._apply_alignment(sheet, issue_count)
            
            # 保存文件
            wb.save(excel_file_path)
            
            self.logger.debug("回归测试问题报告单格式设置完成")
            return True
            
        except Exception as e:
            self.logger.error(f"应用回归测试问题报告单格式时发生错误: {str(e)}")
            return False
    
    def _apply_color_fills(self, sheet, issue_count: int) -> None:
        """应用颜色填充"""
        start_row = 2  # 数据从第2行开始
        
        for i in range(issue_count):
            current_row = start_row + i
            
            # 根据原有逻辑设置颜色
            self._set_row_colors(sheet, current_row, i + 1)
            
        self.logger.debug(f"颜色填充完成，共处理{issue_count}行")
    
    def _set_row_colors(self, sheet, row: int, item_index: int) -> None:
        """设置单行颜色"""
        light_yellow_fill = self.color_fills['light_yellow']
        light_green_fill = self.color_fills['light_green']
        light_blue_fill = self.color_fills['light_blue']
        
        # 按照原有逻辑设置颜色
        # 1-3列：浅黄色（编号、软件名称、版本）
        for col in [1, 2, 3]:
            sheet.cell(row=row, column=col).fill = light_yellow_fill
        
        # 4列：浅绿色（测试轮次）
        sheet.cell(row=row, column=4).fill = light_green_fill
        
        # 5列：浅黄色（问题报告单编号）
        sheet.cell(row=row, column=5).fill = light_yellow_fill
        
        # 6列：浅绿色（是否关联历史问题）
        sheet.cell(row=row, column=6).fill = light_green_fill
        
        # 7-8列：浅黄色（关联版本、关联问题编号）
        for col in [7, 8]:
            sheet.cell(row=row, column=col).fill = light_yellow_fill
        
        # 9-11列：浅绿色（测试活动、问题类型、问题等级）
        for col in [9, 10, 11]:
            sheet.cell(row=row, column=col).fill = light_green_fill
        
        # 12-16列：浅蓝色（问题数目、名称、来源、描述、测评意见）
        for col in range(12, 17):
            sheet.cell(row=row, column=col).fill = light_blue_fill
        
        # 17-19列：浅黄色（修复优先级、测试人员、签名时间）
        for col in [17, 18, 19]:
            sheet.cell(row=row, column=col).fill = light_yellow_fill
        
        # 20列：浅蓝色（研制单位意见）
        sheet.cell(row=row, column=20).fill = light_blue_fill
        
        # 21列：浅绿色（问题修复时间）
        sheet.cell(row=row, column=21).fill = light_green_fill
        
        # 22-24列：浅蓝色（研发人员、签名时间、总师意见）
        for col in [22, 23, 24]:
            sheet.cell(row=row, column=col).fill = light_blue_fill
        
        # 25-26列：浅绿色（完成年月、处理结果）
        for col in [25, 26]:
            sheet.cell(row=row, column=col).fill = light_green_fill
        
        # 27列：浅蓝色（确认意见）
        sheet.cell(row=row, column=27).fill = light_blue_fill
        
        # 28列：浅绿色（确认人）
        sheet.cell(row=row, column=28).fill = light_green_fill
        
        # 29-30列：浅黄色（确认时间、验证版本）
        for col in [29, 30]:
            sheet.cell(row=row, column=col).fill = light_yellow_fill
        
        # 31-32列：浅绿色（是否典型问题、是否低级错误）
        for col in [31, 32]:
            sheet.cell(row=row, column=col).fill = light_green_fill
    
    def _apply_borders(self, sheet, issue_count: int) -> None:
        """应用边框样式"""
        # 应用边框到所有数据行和列
        max_row = 1 + issue_count  # 标题行 + 数据行
        max_col = 32  # 32列
        
        for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                cell.border = self.border_style
        
        self.logger.debug(f"边框样式应用完成，范围: 1-{max_row}行, 1-{max_col}列")
    
    def _apply_alignment(self, sheet, issue_count: int) -> None:
        """应用对齐方式"""
        # 应用对齐方式到所有数据行和列
        max_row = 1 + issue_count  # 标题行 + 数据行
        max_col = 32  # 32列
        
        alignment = Alignment(wrap_text=True, vertical='center')
        
        for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                cell.alignment = alignment
        
        self.logger.debug(f"对齐方式应用完成，范围: 1-{max_row}行, 1-{max_col}列")
    
    def _get_color_fills(self) -> Dict[str, PatternFill]:
        """获取颜色填充样式"""
        return {
            'light_yellow': PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid"),
            'light_green': PatternFill(start_color="CCFFCC", end_color="CCFFCC", fill_type="solid"),
            'light_blue': PatternFill(start_color="99CCFF", end_color="99CCFF", fill_type="solid"),
            'light_orange': PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid"),
            'light_pink': PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid")
        }
    
    def _get_border_style(self) -> Border:
        """获取边框样式"""
        return Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000')
        )
    
    def highlight_regression_specific_fields(
        self,
        excel_file_path: str,
        sheet_name: str,
        issue_count: int = 16
    ) -> bool:
        """
        高亮显示回归测试特有的字段
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            issue_count: 问题数量
            
        Returns:
            bool: 高亮设置是否成功
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            # 回归测试特有字段使用橙色高亮
            orange_fill = self.color_fills['light_orange']
            
            start_row = 2
            for i in range(issue_count):
                current_row = start_row + i
                
                # 问题报告单编号（第5列）- 回归测试特有的编号格式
                sheet.cell(row=current_row, column=5).fill = orange_fill
                
                # 测试活动（第9列）- 回归测试的分类逻辑
                sheet.cell(row=current_row, column=9).fill = orange_fill
                
                # 问题类型（第10列）- 回归测试的分类逻辑
                sheet.cell(row=current_row, column=10).fill = orange_fill
            
            wb.save(excel_file_path)
            self.logger.info("回归测试特有字段高亮完成")
            return True
            
        except Exception as e:
            self.logger.error(f"高亮回归测试特有字段时发生错误: {str(e)}")
            return False
    
    def highlight_auto_generated_fields(
        self,
        excel_file_path: str,
        sheet_name: str,
        issue_count: int = 16
    ) -> bool:
        """
        高亮显示自动生成的字段
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            issue_count: 问题数量
            
        Returns:
            bool: 高亮设置是否成功
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            # 自动生成字段使用粉色高亮
            pink_fill = self.color_fills['light_pink']
            
            start_row = 2
            for i in range(issue_count):
                current_row = start_row + i
                
                # 自动生成的字段
                auto_generated_columns = [
                    1,   # 编号
                    5,   # 问题报告单编号
                    11,  # 问题等级
                    12,  # 问题数目
                    16,  # 测评单位意见
                    17,  # 问题修复优先级
                    20,  # 研制单位意见
                    21,  # 问题修复时间
                    26,  # 问题处理结果
                    27,  # 问题确认意见
                    31,  # 是否典型问题
                    32   # 是否低级错误
                ]
                
                for col in auto_generated_columns:
                    sheet.cell(row=current_row, column=col).fill = pink_fill
            
            wb.save(excel_file_path)
            self.logger.info("自动生成字段高亮完成")
            return True
            
        except Exception as e:
            self.logger.error(f"高亮自动生成字段时发生错误: {str(e)}")
            return False
    
    def apply_conditional_formatting(
        self,
        excel_file_path: str,
        sheet_name: str,
        issue_report_data: Dict[str, Any]
    ) -> bool:
        """
        应用条件格式设置
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            issue_report_data: 问题报告数据
            
        Returns:
            bool: 条件格式设置是否成功
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            issue_classification = issue_report_data.get('issue_classification', [])
            issue_count = len(issue_classification)
            
            start_row = 2
            for i, classification in enumerate(issue_classification):
                current_row = start_row + i
                test_activity = classification.get('test_activity', '')
                
                # 根据测试活动类型设置特殊颜色
                if test_activity == '文档审查':
                    # 文档审查问题使用特殊标记
                    sheet.cell(row=current_row, column=9).fill = self.color_fills['light_blue']
                elif test_activity == '代码审查':
                    # 代码审查问题使用特殊标记
                    sheet.cell(row=current_row, column=9).fill = self.color_fills['light_green']
                elif test_activity == '动态测试':
                    # 动态测试问题使用特殊标记
                    sheet.cell(row=current_row, column=9).fill = self.color_fills['light_yellow']
            
            wb.save(excel_file_path)
            self.logger.info("条件格式设置完成")
            return True
            
        except Exception as e:
            self.logger.error(f"应用条件格式设置时发生错误: {str(e)}")
            return False
    
    def reset_formatting(
        self,
        excel_file_path: str,
        sheet_name: str,
        issue_count: int = 16
    ) -> bool:
        """
        重置格式设置
        
        Args:
            excel_file_path: Excel文件路径
            sheet_name: 工作表名称
            issue_count: 问题数量
            
        Returns:
            bool: 重置是否成功
        """
        try:
            wb = openpyxl.load_workbook(excel_file_path)
            sheet = wb[sheet_name]
            
            # 清除所有格式
            max_row = 1 + issue_count
            max_col = 32
            
            for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
                for cell in row:
                    cell.fill = PatternFill()  # 清除填充
                    cell.border = Border()     # 清除边框
                    cell.alignment = Alignment()  # 清除对齐
            
            wb.save(excel_file_path)
            self.logger.info("格式重置完成")
            return True
            
        except Exception as e:
            self.logger.error(f"重置格式时发生错误: {str(e)}")
            return False

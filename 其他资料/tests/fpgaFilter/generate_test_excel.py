"""
生成FPGA T-Linter测试Excel文件（不使用pandas）
覆盖多种规则类型
"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from pathlib import Path


def create_fpga_test_excel(output_path: str):
    """创建FPGA测试Excel文件"""
    
    # 创建工作簿
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "结果1"
    
    # 定义表头 - 8列格式匹配FPGA处理器
    # A列：序号, B列：检测项(规则编号), C列：问题描述, D列：文件名
    # E列：行号, F列：列号, G列：严重程度, H列：建议
    headers = [
        "序号", "检测项", "问题描述", "文件名", "行号", "列号", "严重程度", "建议"
    ]
    
    # 写入表头
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    
    # 测试数据 - 覆盖多个规则
    test_data = [
        # A类规则 - 建议准则（不确定规则）
        {
            "序号": 1,
            "检测项": "5.1.1.5 建议准则A-1-5",
            "问题描述": "时钟信号命名不规范，建议使用clk前缀",
            "文件名": "clock_divider.v",
            "行号": 15,
            "列号": 1,
            "严重程度": "建议",
            "建议": "将信号名改为clk_div"
        },
        {
            "序号": 2,
            "检测项": "5.1.2.3 建议准则A-2-3",
            "问题描述": "复位信号命名不规范，建议使用rst后缀",
            "文件名": "counter.v",
            "行号": 23,
            "列号": 1,
            "严重程度": "建议",
            "建议": "将信号名改为sys_rst"
        },
        {
            "序号": 3,
            "检测项": "5.1.3.1 建议准则A-3-1",
            "问题描述": "状态机状态命名不清晰",
            "文件名": "state_machine.v",
            "行号": 45,
            "列号": 1,
            "严重程度": "建议",
            "建议": "使用更具描述性的状态名称"
        },
        
        # R类规则 - 强制规则（确定规则）
        {
            "序号": 4,
            "检测项": "5.2.1.1 强制准则R-1-1",
            "问题描述": "存在组合逻辑环路",
            "文件名": "alu.v",
            "行号": 30,
            "列号": 1,
            "严重程度": "错误",
            "建议": "消除组合逻辑环路，添加寄存器打断"
        },
        {
            "序号": 5,
            "检测项": "5.2.2.5 强制准则R-2-5",
            "问题描述": "时钟域交叉未做同步处理",
            "文件名": "fifo.v",
            "行号": 40,
            "列号": 1,
            "严重程度": "错误",
            "建议": "添加双触发器同步器"
        },
        {
            "序号": 6,
            "检测项": "5.2.3.2 强制准则R-3-2",
            "问题描述": "复位信号未正确连接",
            "文件名": "memory_ctrl.v",
            "行号": 20,
            "列号": 1,
            "严重程度": "错误",
            "建议": "确保所有寄存器都连接复位信号"
        },
        {
            "序号": 7,
            "检测项": "5.2.4.1 强制准则R-4-1",
            "问题描述": "使用了阻塞赋值在时序逻辑中",
            "文件名": "data_path.v",
            "行号": 24,
            "列号": 1,
            "严重程度": "错误",
            "建议": "时序逻辑应使用非阻塞赋值(<=)"
        },
        
        # 更多A类规则
        {
            "序号": 8,
            "检测项": "5.1.4.2 建议准则A-4-2",
            "问题描述": "模块端口顺序不规范",
            "文件名": "uart_tx.v",
            "行号": 6,
            "列号": 1,
            "严重程度": "建议",
            "建议": "按照输入、输出、双向的顺序排列端口"
        },
        {
            "序号": 9,
            "检测项": "5.1.5.3 建议准则A-5-3",
            "问题描述": "缺少必要的注释说明",
            "文件名": "spi_master.v",
            "行号": 22,
            "列号": 1,
            "严重程度": "建议",
            "建议": "添加模块功能说明和端口描述"
        },
        
        # 更多R类规则
        {
            "序号": 10,
            "检测项": "5.2.5.4 强制准则R-5-4",
            "问题描述": "敏感列表不完整",
            "文件名": "i2c_slave.v",
            "行号": 22,
            "列号": 1,
            "严重程度": "错误",
            "建议": "在always块中包含所有读取的信号"
        },
        {
            "序号": 11,
            "检测项": "5.2.1.3 强制准则R-1-3",
            "问题描述": "存在锁存器推断",
            "文件名": "pipeline.v",
            "行号": 27,
            "列号": 1,
            "严重程度": "错误",
            "建议": "在组合逻辑中为所有分支赋值"
        },
        {
            "序号": 12,
            "检测项": "5.2.2.1 强制准则R-2-1",
            "问题描述": "信号位宽不匹配",
            "文件名": "arbiter.v",
            "行号": 25,
            "列号": 1,
            "严重程度": "错误",
            "建议": "确保赋值两侧位宽一致"
        },
        
        # A类规则
        {
            "序号": 13,
            "检测项": "5.1.6.1 建议准则A-6-1",
            "问题描述": "case语句缺少default分支",
            "文件名": "decoder.v",
            "行号": 13,
            "列号": 1,
            "严重程度": "建议",
            "建议": "添加default分支提高代码健壮性"
        },
        {
            "序号": 14,
            "检测项": "5.1.7.2 建议准则A-7-2",
            "问题描述": "参数定义未使用parameter关键字",
            "文件名": "encoder.v",
            "行号": 11,
            "列号": 1,
            "严重程度": "建议",
            "建议": "使用parameter或localparam定义常量"
        },
        
        # R类规则
        {
            "序号": 15,
            "检测项": "5.2.3.5 强制准则R-3-5",
            "问题描述": "未考虑溢出情况",
            "文件名": "multiplier.v",
            "行号": 31,
            "列号": 1,
            "严重程度": "错误",
            "建议": "添加溢出检测和处理逻辑"
        },
        {
            "序号": 16,
            "检测项": "5.2.4.3 强制准则R-4-3",
            "问题描述": "除零未做保护",
            "文件名": "divider.v",
            "行号": 35,
            "列号": 1,
            "严重程度": "错误",
            "建议": "添加除数为零的检查"
        },
        
        # VHDL文件测试
        {
            "序号": 17,
            "检测项": "5.1.8.1 建议准则A-8-1",
            "问题描述": "实体名称与文件名不一致",
            "文件名": "controller.vhd",
            "行号": 45,
            "列号": 1,
            "严重程度": "建议",
            "建议": "保持实体名与文件名一致"
        },
        {
            "序号": 18,
            "检测项": "5.2.5.2 强制准则R-5-2",
            "问题描述": "进程敏感列表缺少信号",
            "文件名": "processor.vhd",
            "行号": 26,
            "列号": 1,
            "严重程度": "错误",
            "建议": "将所有读取的信号添加到敏感列表"
        },
        
        # 更多测试数据
        {
            "序号": 19,
            "检测项": "5.1.9.3 建议准则A-9-3",
            "问题描述": "总线信号命名不规范",
            "文件名": "bus_interface.v",
            "行号": 10,
            "列号": 1,
            "严重程度": "建议",
            "建议": "使用统一的总线命名约定"
        },
        {
            "序号": 20,
            "检测项": "5.2.6.1 强制准则R-6-1",
            "问题描述": "存在竞争冒险",
            "文件名": "dma_ctrl.v",
            "行号": 29,
            "列号": 1,
            "严重程度": "错误",
            "建议": "消除竞争条件，使用同步设计"
        },
    ]
    
    # 写入数据 - 8列顺序
    for row_idx, data in enumerate(test_data, start=2):
        ws.cell(row=row_idx, column=1, value=data["序号"])         # A列
        ws.cell(row=row_idx, column=2, value=data["检测项"])       # B列
        ws.cell(row=row_idx, column=3, value=data["问题描述"])     # C列
        ws.cell(row=row_idx, column=4, value=data["文件名"])       # D列
        ws.cell(row=row_idx, column=5, value=data["行号"])         # E列
        ws.cell(row=row_idx, column=6, value=data["列号"])         # F列
        ws.cell(row=row_idx, column=7, value=data["严重程度"])     # G列
        ws.cell(row=row_idx, column=8, value=data["建议"])         # H列
    
    # 设置列宽 - 8列
    ws.column_dimensions['A'].width = 8    # 序号
    ws.column_dimensions['B'].width = 30   # 检测项
    ws.column_dimensions['C'].width = 40   # 问题描述
    ws.column_dimensions['D'].width = 20   # 文件名
    ws.column_dimensions['E'].width = 8    # 行号
    ws.column_dimensions['F'].width = 8    # 列号
    ws.column_dimensions['G'].width = 12   # 严重程度
    ws.column_dimensions['H'].width = 40   # 建议
    
    # 设置行高
    for row in range(1, len(test_data) + 2):
        ws.row_dimensions[row].height = 30
    
    # 保存文件
    wb.save(output_path)
    print(f"✅ FPGA测试Excel文件已生成: {output_path}")
    print(f"   - 共 {len(test_data)} 条测试数据")
    print(f"   - 覆盖 A 类规则（建议）和 R 类规则（强制）")
    print(f"   - 包含 Verilog 和 VHDL 文件")


if __name__ == "__main__":
    # 输出路径
    output_path = Path(__file__).parent.parent.parent.parent / "T-Linter导出表单 copy.xlsx"
    create_fpga_test_excel(str(output_path))
